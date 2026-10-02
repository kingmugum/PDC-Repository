from pathlib import Path

from openpyxl import load_workbook

from core.cross_document_semantics import build_semantic_source_units, attach_semantic_traceability, _structured_source_facts
from core.quality_audit import (
    apply_quality_audits,
    build_swe6_export_preservation_audit,
    build_unsupported_generation_report,
    normalize_requirement_extensions,
)
from core.swe6_exporter import build_swe6_cases, SWE6Exporter, _source_table_fact_assignment


def _req(text, source_text=None, srs='SRS_001'):
    r={
        'candidate_id':'C1','srs_id':srs,'scenario_candidate_id':'SCN-CAND-001','category':'기능','function_name':'Feature',
        'requirement':text,'user_input':'','system_input_preconditions':'','processing_action':text,'output':text,
        'acceptance_criteria':text,'failure_situations':[],'user_intervention_points':[],'derivation_type':'explicit',
        'derivation_reason':'','clarification_needed':[],'source_evidence':[{'document':'x.docx','location':'Paragraph 1','text':source_text or text}],
        'confidence':0.95,'classification_basis':'source','activation_trigger':'','preconditions':'','behavior_flows':[],
        'evaluation_method':'','exception_conditions':[],'related_artifacts':[],'source_requirement_ids':[],
    }
    normalize_requirement_extensions(r)
    return r


def _data(reqs):
    if not isinstance(reqs,list): reqs=[reqs]
    return {'schema_version':'REQ-STUDIO-CANONICAL-REQ-2.0','source_document':'x.docx','scenario_candidates':[{
        'scenario_candidate_id':'SCN-CAND-001','scenario_name':'x','user_goal_context':'','scenario_flow':[],
        'expected_outcome':'','source_evidence':[]}], 'requirements':reqs,'gaps':[]}


def test_result_columns_are_blank_and_na_not_a_passfail_option(tmp_path):
    req=_req('입력을 확인한다.')
    req.update({'requirement_level':'Software','allocation_status':'SW_IMPLEMENTATION_REQUIREMENT','swe1_eligibility':'Eligible','swe6_eligibility':'Eligible','verification_domain':'SWE.6 Software Qualification','source_semantic_unit_ids':['S1'],'source_backed_atomic_behaviors':[{'source_semantic_unit_id':'S1','behavior_text':'입력을 확인한다.','knowledge_state':'KNOWN'}],'semantic_provenance_status':'COMPLETE'})
    out=SWE6Exporter(tmp_path).export_excel(tmp_path/'x.docx',_data(req),source_text='입력을 확인한다.')
    wb=load_workbook(out,data_only=False)
    try:
        ws=wb['2_테스트 케이스']
        assert [ws.cell(4,c).value for c in (18,19,20,21)] == [None,None,None,None]
        formulas=[getattr(dv,'formula1','') for dv in ws.data_validations.dataValidation]
        assert any('PASS,FAIL' in str(f) for f in formulas)
        assert not any('N/A' in str(f) for f in formulas)
    finally:
        wb.close()


def test_table_fact_value_role_blocks_enum_range_and_encoding_auto_use():
    facts=[]
    facts += _structured_source_facts('Output_FadeInTime: 254 (Value x 100ms, 255 Invalid)',unit_id='U1',location='Table 1 / Row 1')
    facts += _structured_source_facts('Output_MLMnSOH: 0x00=NOK, 0x01=OK',unit_id='U2',location='Table 1 / Row 2')
    facts += _structured_source_facts('Par_NotPConfirmationTime = 600 ms',unit_id='U3',location='Table 1 / Row 3')
    by={f['source_semantic_unit_id']:f for f in facts}
    assert by['U1']['value_role'] in {'ENCODING_RULE','INVALID_MARKER'}
    assert by['U2']['value_role']=='ENUM_MAPPING'
    assert by['U3']['value_role']=='TIMING_CRITERION'
    req={'source_table_fact_matches':[dict(f,join_confidence='HIGH') for f in facts]}
    v,c,val,audit=_source_table_fact_assignment(req,'expected','Output_FadeInTime Output_MLMnSOH')
    assert (v,c,val)==('','','')
    assert audit
    req2={'source_table_fact_matches':[dict(by['U3'],join_confidence='HIGH')]}
    v,c,val,audit=_source_table_fact_assignment(req2,'execution','Par_NotPConfirmationTime')
    assert v=='Par_NotPConfirmationTime' and '600' in val


def test_trailing_page_number_heading_is_review_context_and_bitrate_is_allocation_review_context():
    compact='''[DOCUMENT] x.docx
[SRC C1 | Paragraph 1 | text]
5.4.2.1 커넥터 타입 53
[SRC C2 | Table 1 / Row 2 | table]
CAN 500 Kbit/s, LIN 19.2 Kbit/s
'''
    units=build_semantic_source_units(compact)
    by={u['source_location']:u for u in units}
    assert by['Paragraph 1']['coverage_eligibility']=='review_context'
    assert by['Paragraph 1']['source_unit_type']=='heading_or_document_context'
    assert by['Table 1 / Row 2']['source_unit_type']=='interface_fact'
    assert by['Table 1 / Row 2']['coverage_eligibility']=='allocation_review_context'


def test_structured_provenance_collector_does_not_false_flag_100ms():
    req=_req('LIN wake-up 후 60ms 대기하고 100ms 이내 schedule table을 시작한다.')
    req.update({
        'requirement_level':'Software','allocation_status':'SW_IMPLEMENTATION_REQUIREMENT','swe1_eligibility':'Eligible','swe6_eligibility':'Eligible','verification_domain':'SWE.6 Software Qualification',
        'source_semantic_unit_ids':['U1'],'source_backed_atomic_behaviors':[{'source_semantic_unit_id':'U1','behavior_text':'LIN wake-up 후 60ms 대기하고 100ms 이내 schedule table을 시작한다.','knowledge_state':'KNOWN'}],
        'semantic_provenance_status':'COMPLETE',
        'fact_level_allocations':[{
            'source_fact':'LIN wake-up 후 60ms 대기하고 100ms 이내 schedule table을 시작한다.',
            'structured_source_facts':[{
                'source_literal':'60ms 대기 / 100ms 이내 시작',
                'numeric_values':['60ms','100ms'],'timing_values':['60ms','100ms'],'range_values':[],
                'explicit_relations':['100ms 이내'],'enum_mappings':[],'identifiers':[]
            }],
            'verification_domain':'SWE.6 Software Qualification','allocation_status':'SW_IMPLEMENTATION_REQUIREMENT'
        }]
    })
    audit=build_swe6_export_preservation_audit(_data(req),[{'srs_id':'SRS_001','tc_id':'TC_001','description':'60ms / 100ms','prep_desc':'','exec_desc':'60ms 대기','expected_desc':'100ms 이내 시작'}])
    row=audit['rows'][0]
    assert '100ms' not in row['missing_source_fact_tokens']


def test_unsupported_generation_audit_ignores_schema_metadata_keys():
    req=_req('출력은 100ms 이내 활성화된다.')
    req['applicability']={'vehicle_lines':[],'baseline_versions':[],'feature_variants':[],'enable_conditions':[],'exclusion_conditions':[],'knowledge_state':'UNKNOWN','knowledge_state_reason':'','source_evidence':[]}
    req['external_dependencies']=[{'type':'external_document','name':'ES95400-30','purpose':'','knowledge_state':'KNOWN','source_evidence':[],'required_for':[]}]
    report=build_unsupported_generation_report(_data(req),'출력은 100ms 이내 활성화된다. ES95400-30')
    assertions=[x['generated_assertion'] for x in report['findings']]
    assert 'baseline_versions' not in assertions
    assert 'knowledge_state' not in assertions
    assert report['uncertain_assertion_count'] < 10


def test_same_semantic_unit_is_fragmented_for_different_srs_ownership():
    compact='''[DOCUMENT] x.docx
[SRC C1 | Paragraph 82 | text]
사용자 설정을 CAN으로 수신하여 NVM에 저장한다. 차량 On/Off 정보를 수신하여 Slave를 제어한다. 밝기는 USM 설정으로만 변경하고 Rheostat과 연동하지 않는다.
'''
    r1=_req('사용자 설정을 NVM에 저장한다.','사용자 설정을 CAN으로 수신하여 NVM에 저장한다.',srs='SRS_001')
    r2=_req('차량 On/Off 정보에 따라 Slave를 제어한다.','차량 On/Off 정보를 수신하여 Slave를 제어한다.',srs='SRS_002')
    d=_data([r1,r2])
    attach_semantic_traceability(d,compact)
    f1={x['source_fact_fragment_id'] for x in d['requirements'][0].get('source_fact_fragments',[])}
    f2={x['source_fact_fragment_id'] for x in d['requirements'][1].get('source_fact_fragments',[])}
    assert f1 and f2
    assert f1 != f2
    assert d.get('source_fact_multi_srs_allocation_conflicts',[]) == []


def test_high_confidence_unlinked_software_structural_constraint_is_recovered_as_review_candidate():
    compact='''[DOCUMENT] x.docx
[SRC C1 | Paragraph 10 | text]
Watchdog Timer를 사용할 경우 Watchdog Timer는 가장 낮은 Task 이후 작동해야 한다. Watchdog Timer는 ISR 내부에 있을 수 없다.
'''
    d={'source_document':'x.docx','requirements':[],'gaps':[],'scenario_candidates':[]}
    out=apply_quality_audits(d,compact)
    assert out.get('semantic_missing_recovery',{}).get('recovered_candidate_count',0) >= 1
    recovered=[r for r in out['requirements'] if 'Watchdog' in r.get('requirement','')]
    assert recovered
    assert recovered[0]['knowledge_state']=='KNOWN'
    assert recovered[0]['swe6_eligibility']=='Eligible'

def test_child_intent_independent_coverage_requires_relevant_structured_observation():
    from core.quality_audit import _child_intent_disposition
    req=_req('Input_A를 수신하고 Output_B를 송신한다.')
    child={'source_backed_behavior':'Input_A를 수신한다.'}
    cases=[{'tc_id':'TC_001','exec_desc':'Input_A를 수신한다.','expected_desc':'정상 동작한다.',
            'prep_var':'Unrelated_X','prep_compare':'=','prep_value':'1',
            'exec_var':'','exec_compare':'','exec_value':'','expected_var':'','expected_compare':'','expected_value':''}]
    disp=_child_intent_disposition(req,child,cases)
    assert disp['intent_status']=='REPRESENTED_IN_GENERIC_TC'
    assert disp['independent_tc_coverage']=='NOT_ESTABLISHED'
