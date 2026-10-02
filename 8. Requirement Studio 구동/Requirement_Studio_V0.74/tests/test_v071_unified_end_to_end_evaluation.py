import json
from pathlib import Path
from types import SimpleNamespace

from docx import Document
from openpyxl import Workbook

from core.review_exchange import ReviewExchangeBuilder
from core.evaluation_orchestrator import MultiModelEvaluationOrchestrator


def _meta():
    return SimpleNamespace(
        display_name="H-Chat / GPT", provider_id="hchat_gpt", model="gpt-5.6-terra",
        endpoint_family="hchat", auth_mode="key", project_header_enabled=False,
        response_parsing_mode="text",
    )


def _xlsx(path: Path, sheet="01_SWE1"):
    wb = Workbook(); ws = wb.active; ws.title = sheet
    if sheet == "2_테스트 케이스":
        ws.append(["Requirement Studio SWE.6"]); ws.append([])
        headers=[""]*21
        headers[0]="TC ID"; headers[1]="SRS ID"
        headers[17]="Output Value"; headers[18]="PASS / FAIL"; headers[19]="Comment"; headers[20]="Capture CANoe (Optional)"
        ws.append(headers)
        for name in ["3_Deferred Intent", "4_Source_Intent_Audit", "5_Source_Fact_Audit"]:
            w = wb.create_sheet(name); w.append(["ID", "Value"]); w.append(["X", "Y"])
    else:
        ws.append(["ID", "Value"]); ws.append(["SRS_001", "Source-backed value"])
    wb.save(path)


def test_unified_review_exchange_copies_all_available_outputs(tmp_path):
    root = tmp_path / "root"; root.mkdir()
    (root / "04_CHANGE_DECISION_V0.70_to_V0.71.txt").write_text("decision", encoding="utf-8")
    source = tmp_path / "source.docx"; Document().save(source)
    swe1w = tmp_path / "swe1.docx"; Document().save(swe1w)
    swe1x = tmp_path / "swe1.xlsx"; _xlsx(swe1x)
    swe6x = tmp_path / "swe6.xlsx"; _xlsx(swe6x, "2_테스트 케이스")
    result = ReviewExchangeBuilder(root, app_version="v0.71").build(
        source_document=source, provider_metadata=_meta(), analysis_text="",
        requirement_data={"requirements": [], "gaps": []}, evaluation={}, normalized_summary={},
        quality_review_config={"enabled": True, "evaluation_artifact": "unified", "selected_outputs": {"swe1_word": True, "swe1_excel": True, "swe6_excel": True}},
        artifact_paths={"swe1_word": swe1w, "swe1_excel": swe1x, "swe6_excel": swe6x},
        evaluation_mode="initial",
    )
    run = Path(result["run_dir"])
    payload = json.loads((run / "02_REQUIREMENT_STUDIO_REVIEW_PACKAGE.json").read_text(encoding="utf-8"))
    assert payload["review_handoff"]["evaluation_artifact_focus"] == "unified"
    assert set(payload["review_handoff"]["unified_artifacts"]) == {"swe1_word", "swe1_excel", "swe6_excel"}
    assert len(list(run.glob("03_SWE1_WORD_*"))) == 1
    assert len(list(run.glob("03_SWE1_EXCEL_*"))) == 1
    assert len(list(run.glob("03_SWE6_EXCEL_*"))) == 1


def test_unified_evidence_reads_docx_tables_and_multiple_artifacts(tmp_path):
    root = tmp_path / "root"; root.mkdir(); (root / "config").mkdir()
    run = root / "review_exchange" / "run"; run.mkdir(parents=True)
    source = run / "01_SOURCE_source.docx"; d=Document(); d.add_paragraph("source"); d.save(source)
    w = run / "03_SWE1_WORD_swe1.docx"; d=Document(); t=d.add_table(rows=1, cols=2); t.cell(0,0).text="SRS_001"; t.cell(0,1).text="Annex requirement detail"; d.save(w)
    x1=run/"03_SWE1_EXCEL_swe1.xlsx"; _xlsx(x1)
    x6=run/"03_SWE6_EXCEL_swe6.xlsx"; _xlsx(x6, "2_테스트 케이스")
    review={
        "review_handoff": {"evaluation_artifact_focus":"unified", "file_3_swe1_word":w.name, "file_3_swe1_excel":x1.name, "file_3_swe6_excel":x6.name},
        "quality_review_config":{"selected_outputs":{"swe1_word":True,"swe1_excel":True,"swe6_excel":True}},
        "canonical_requirement":{"requirements":[]}, "calculated_review_evidence":{}, "local_qa":{}, "evaluation_request":{}
    }
    (run/"02_REQUIREMENT_STUDIO_REVIEW_PACKAGE.json").write_text(json.dumps(review,ensure_ascii=False),encoding="utf-8")
    (run/"05_REQUIREMENT_STUDIO_REVIEW_SUMMARY.txt").write_text("summary",encoding="utf-8")
    orch=MultiModelEvaluationOrchestrator(root, app_version="v0.71", provider_config={})
    evidence,_=orch.build_evidence_package(run)
    assert evidence["evaluation_mode"] == "unified"
    assert set(evidence["actual_artifacts"]) == {"swe1_word","swe1_excel","swe6_excel"}
    assert evidence["actual_artifacts"]["swe1_word"]["tables"][0]["rows"][0][1] == "Annex requirement detail"
    assert len(evidence["evidence_package_sha256"]) == 64


def test_final_result_completeness_audit_and_110_line_split(tmp_path):
    provider_results={
        "gpt":{"findings":[{"finding_key":"A"}],"next_version_recommendations":[],"verdict":"PASS","tool_quality_gate":"PASS","artifact_readiness_gate":"PASS","official_release":"PASS"},
        "gemini":{"findings":[{"finding_key":"A"},{"finding_key":"B"}],"next_version_recommendations":[],"verdict":"PASS","tool_quality_gate":"PASS","artifact_readiness_gate":"PASS","official_release":"PASS"},
        "claude":{"findings":[],"next_version_recommendations":[],"verdict":"PASS","tool_quality_gate":"PASS","artifact_readiness_gate":"PASS","official_release":"PASS"},
    }
    # Use manually shaped consensus to validate occurrence accounting.
    consensus={"evaluation_completeness":"COMPLETE","tool_quality_gate":"PASS","artifact_readiness_gate":"PASS","official_release":"PASS","consensus_findings":[{"reviewer_count":2},{"reviewer_count":1}]}
    final=MultiModelEvaluationOrchestrator._final_result(provider_results,consensus,{},Path("06_EVALUATION_EVIDENCE.json"))
    assert final["completeness_audit"]["provider_finding_occurrences"] == 3
    assert final["completeness_audit"]["mapped_finding_occurrences"] == 3
    assert final["completeness_audit"]["unmapped_finding_occurrences"] == 0
    assert final["completeness_audit"]["passed"] is True
    text="\n".join(f"line {i}" for i in range(225))
    parts=MultiModelEvaluationOrchestrator.split_final_result(text,tmp_path,lines_per_file=110)
    assert [len(p.read_text(encoding="utf-8").splitlines()) for p in parts] == [110,110,5]
