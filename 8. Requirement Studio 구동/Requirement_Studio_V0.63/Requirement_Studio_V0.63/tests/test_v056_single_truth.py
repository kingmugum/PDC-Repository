import json
from pathlib import Path
from types import SimpleNamespace

from docx import Document

from core.gold_source_registry import GoldSourceRegistry
from core.quality_audit import apply_quality_audits, build_conflict_register, build_testability_result, finalize_test_intent_coverage
from core.review_exchange import ReviewExchangeBuilder
from core.swe1_exporter import SWE1Exporter


def test_repeated_equivalent_final_matrix_is_single_truth():
    compact = "\n".join([
        "[SRC DOC-C1 | Paragraph 10 | paragraph]\n[REQ20_01] Input_FootLpOn == On(0x1) 이면 Foot Lamp를 On 제어한다.",
        "[SRC DOC-C2 | Paragraph 40 | paragraph]\n[REQ20_01] Input_FootLpOn == On(0x1)이면 Foot Lamp를 On 제어한다",
    ])
    data = {"requirements": [{
        "candidate_id": "C1", "srs_id": "S1", "derivation_type": "explicit",
        "requirement": "Input_FootLpOn On 시 Foot Lamp On 제어",
        "processing_action": "Foot Lamp를 On 제어한다.",
        "source_requirement_ids": ["REQ20_01"],
        "source_evidence": [{"document":"a.docx", "location": "Paragraph 10", "text": "[REQ20_01] Input_FootLpOn == On(0x1) 이면 Foot Lamp를 On 제어한다."}],
    }], "gaps": []}
    out = apply_quality_audits(data, compact)
    cov = out["source_coverage"]
    assert set(cov["covered"]) == {"SRC-OCC-0001", "SRC-OCC-0002"}
    repair = cov["traceability_repair"]
    assert repair["repeated_equivalent_link_count"] >= 1
    repeated = [r for r in repair["linkage_audit_records"] if r.get("link_type") == "repeated_equivalent"]
    assert repeated and all(r["accepted_final"] for r in repeated)
    assert all(r["final_coverage_status"] == "Covered" for r in repeated)


def test_grouped_swe1_word_exposes_source_backed_atomic_behavior(tmp_path):
    req_data = {
        "requirements": [{
            "candidate_id":"C1", "srs_id":"S1", "function_name":"Driving dimming", "category":"기능",
            "requirement":"주행 중 감광 조건을 만족하면 감광 제어한다.",
            "activation_trigger":"주행 조건 충족", "processing_action":"감광 제어", "output":"감광 밝기",
            "acceptance_criteria":"감광 밝기 적용", "derivation_type":"explicit", "source_requirement_ids":["REQ02_16","REQ02_17","REQ02_18"],
            "source_evidence":[{"document":"gold.docx","location":"Paragraph 139-143","text":"주행 감광 및 P->not-P 판단"}],
            "source_backed_atomic_behaviors": [],
        }],
        "source_coverage": {"source_requirement_occurrences": [
            {"occurrence_type":"declaration","coverage_status":"Covered","source_req_id":"REQ02_16","source_location":"Paragraph 139","source_excerpt":"[REQ02_16] 주행 중 감광은 USM 설정 가능","linked_candidate_ids":["C1"],"linked_srs_ids":["S1"]},
            {"occurrence_type":"declaration","coverage_status":"Covered","source_req_id":"REQ02_18","source_location":"Paragraph 143","source_excerpt":"[REQ02_18] P -> not P 이동 시 Par_NotPConfirmationTime 판단 시간을 적용한다.","linked_candidate_ids":["C1"],"linked_srs_ids":["S1"]},
        ]},
    }
    exporter = SWE1Exporter(tmp_path)
    src = tmp_path / "gold.docx"
    Document().save(src)
    out = exporter.export_word(src, req_data)
    doc = Document(out)
    all_text = "\n".join(p.text for p in doc.paragraphs) + "\n" + "\n".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    assert "Source-backed Atomic Behavior" in all_text
    assert "REQ02_18" in all_text
    assert "Par_NotPConfirmationTime" in all_text


def test_review_exchange_uses_actual_swe1_word_as_file3(tmp_path):
    root = tmp_path / "Requirement_Studio_V0.56"
    root.mkdir()
    (root / "04_CHANGE_DECISION_V0.55_to_V0.56.txt").write_text("decision", encoding="utf-8")
    source = tmp_path / "source.docx"
    Document().save(source)
    swe1 = tmp_path / "SWE.1 요구사항 정리_source_260930_V0.0.docx"
    Document().save(swe1)
    builder = ReviewExchangeBuilder(root, app_version="v0.56")
    meta = SimpleNamespace(display_name="H-Chat / GPT", provider_id="hchat_gpt", model="gpt-5.6-terra", endpoint_family="x", auth_mode="key", project_header_enabled=False, response_parsing_mode="text")
    result = builder.build(
        source_document=source, provider_metadata=meta, analysis_text="", requirement_data={"requirements":[],"gaps":[]}, evaluation={}, normalized_summary={}, quality_review_config={}, artifact_paths={"swe1_word":str(swe1)}, evaluation_mode="initial"
    )
    run = Path(result["run_dir"])
    assert not (run / "03_REQUIREMENT_STUDIO_OUTPUT_REVIEW.xlsx").exists()
    file3 = list(run.glob("03_SWE.1 요구사항 정리_*.docx"))
    assert len(file3) == 1
    payload = json.loads((run / "02_REQUIREMENT_STUDIO_REVIEW_PACKAGE.json").read_text(encoding="utf-8"))
    assert payload["review_handoff"]["file_3_swe1_word"] == file3[0].name


def test_bundled_gold_contract_splits_req0301_and_req0307():
    root = Path(__file__).resolve().parents[1]
    reg = GoldSourceRegistry(root)
    contract, _ = reg.find_contract("b9d58a87cc72b66919c6af7220b215b9e66e06ab426eae20e04a8642dbffb15c")
    payload = reg.contract_to_review_payload(contract)
    rows = payload["canonical_requirement"]["requirements"]
    singleton_sets = [tuple(r.get("source_requirement_ids") or []) for r in rows]
    assert ("REQ03_01",) in singleton_sets
    assert ("REQ03_07",) in singleton_sets


def test_parameter_direction_conflict_has_correct_evidence():
    compact = "\n".join([
        "[SRC DOC-C1 | Table 3 / Row 1 | table]\nPar_TimerMoodWait | 무드램프 On 대기 시간 (On -> Off 시)",
        "[SRC DOC-C2 | Table 17 / Row 5 | table]\nPar_TimerMoodWait | 무드램프 On 대기 시간 (Off -> On 시) | 3 s",
    ])
    conflicts = build_conflict_register({"requirements":[],"gaps":[],"source_coverage":{"source_requirement_occurrences":[]}}, compact)
    rows = [c for c in conflicts if c.get("classification") == "Direction / Definition Contradiction"]
    assert len(rows) == 1
    assert "On -> Off" in rows[0]["source_a"]["text"] or "On -> Off" in rows[0]["source_b"]["text"]
    assert "Off -> On" in rows[0]["source_a"]["text"] or "Off -> On" in rows[0]["source_b"]["text"]


def test_testability_invariant_concrete_complete_implies_feasible():
    data = {"requirements":[{
        "srs_id":"S1","requirement":"Controller shall output value.","category":"기능","processing_action":"output value","output":"value","acceptance_criteria":"value output",
        "external_dependencies":["CAN DB"],
    }], "conflict_register":[]}
    data["testability_and_decomposition_result"] = build_testability_result(data)
    finalize_test_intent_coverage(data, [{"srs_id":"S1","tc_id":"TC_001"}])
    row = data["testability_and_decomposition_result"]["by_srs"][0]
    # This SRS has only the Normal intent, so emitted coverage makes it complete.
    assert row["concrete_tc_complete"] is True
    assert row["test_design_feasible"] is True
    assert row["full_testability_available"] is True
    assert row["intent_complete_status"] == "Complete"
