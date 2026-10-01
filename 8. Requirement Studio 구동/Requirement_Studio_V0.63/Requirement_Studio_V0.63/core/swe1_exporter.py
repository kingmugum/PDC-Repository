from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path
from typing import Any

from core.output_naming import output_filename
from core.swe6_exporter import build_swe6_cases

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


NOT_CONFIRMED = "입력문서에서 확인되지 않음"
NO_DETAIL = "입력문서에서 추가 작동 명세를 확인할 수 없음"


FINAL_FIELDS = [
    "SRS ID",
    "상위 기능",
    "분류",
    "요구사항 내역",
    "동작 조건 / Trigger",
    "작동 명세 정의",
    "사전 조건",
    "예상 결과",
    "검증 기준",
    "출처 / Traceability",
    "기타",
]


def _safe_stem(name: str) -> str:
    stem = Path(name).stem or "document"
    stem = re.sub(r'[<>:"/\\|?*]+', "_", stem)
    stem = re.sub(r"\s+", "_", stem).strip("._ ")
    return stem[:80] or "document"


def _as_list(value: Any) -> list[str]:
    if value is None:
        return []
    if isinstance(value, list):
        result = []
        for item in value:
            if isinstance(item, dict):
                text = "; ".join(f"{k}={v}" for k, v in item.items() if v not in (None, "", [], {}))
            else:
                text = str(item).strip()
            if text:
                result.append(text)
        return result
    if isinstance(value, dict):
        text = "; ".join(f"{k}={v}" for k, v in value.items() if v not in (None, "", [], {}))
        return [text] if text else []
    text = str(value).strip()
    return [text] if text else []


def _text(value: Any, *, empty: str = "") -> str:
    items = _as_list(value)
    return "\n".join(items) if items else empty


def _normalize_classification(value: Any) -> str:
    raw = _text(value).strip()
    low = raw.lower()
    if "비기능" in raw or "non-functional" in low or "nonfunctional" in low:
        return "비기능"
    if raw == "기능" or "functional" in low or ("기능" in raw and "비기능" not in raw):
        return "기능"
    return raw or "검토 필요"


def _format_evidence(items: Any) -> str:
    lines = []
    for idx, item in enumerate(items if isinstance(items, list) else [], start=1):
        if not isinstance(item, dict):
            text = str(item).strip()
            if text:
                lines.append(f"[{idx}] {text}")
            continue
        document = str(item.get("document") or "").strip()
        location = str(item.get("location") or "").strip()
        evidence = str(item.get("text") or "").strip()
        prefix = " / ".join(x for x in (document, location) if x)
        if evidence:
            lines.append(f"[{idx}] {prefix}: {evidence}" if prefix else f"[{idx}] {evidence}")
        elif prefix:
            lines.append(f"[{idx}] {prefix}")
    return "\n".join(lines) if lines else NOT_CONFIRMED


def _behavior_text(req: dict[str, Any]) -> str:
    """Render detailed behavior for general engineers without exposing BF IDs."""
    flows = _as_list(req.get("behavior_flows"))
    if flows:
        rendered = []
        for idx, item in enumerate(flows, start=1):
            clean = re.sub(r"^BF\s*\d+(?:[-.]\d+)?[.:\-]?\s*", "", item, flags=re.IGNORECASE).strip()
            rendered.append(f"{idx}. {clean or item}")
        return "\n".join(rendered)
    processing = _text(req.get("processing_action")).strip()
    if processing:
        return processing
    return NO_DETAIL


def _other_materials(req: dict[str, Any]) -> str:
    """Final-output '기타': only source-backed auxiliary materials such as figures/diagrams/tables."""
    lines: list[str] = []
    for key in (
        "related_artifacts",
        "diagram_refs",
        "figure_refs",
        "table_refs",
        "visual_evidence_refs",
        "attachments",
    ):
        for item in _as_list(req.get(key)):
            if item and item not in lines:
                lines.append(item)
    return "\n".join(lines)


def _source_behavior_units(requirement_data: dict[str, Any]) -> dict[str, list[str]]:
    """Build source-backed atomic behavior lines from the final occurrence matrix.

    This is intentionally source-derived and is used to prevent a grouped SRS from carrying
    Source IDs without the corresponding behavior being visible in the engineer-facing SWE.1.
    """
    out: dict[str, list[str]] = {}
    # V0.58: Canonical atomic units are authoritative for both explicit-ID and semantic-source modes.
    for req in requirement_data.get("requirements") or []:
        if not isinstance(req, dict):
            continue
        keys = [str(req.get("candidate_id") or "").strip(), str(req.get("srs_id") or "").strip()]
        for atom in req.get("source_backed_atomic_behaviors") or []:
            if not isinstance(atom, dict):
                continue
            rid = str(atom.get("source_requirement_id") or atom.get("source_semantic_unit_id") or "").strip()
            loc = str(atom.get("source_location") or "").strip()
            text = re.sub(r"\s+", " ", str(atom.get("behavior_text") or "")).strip()
            if not text:
                continue
            line = f"{rid} · {loc}: {text}" if rid and loc else (f"{rid}: {text}" if rid else text)
            for key in keys:
                if not key:
                    continue
                bucket = out.setdefault(key, [])
                if line not in bucket:
                    bucket.append(line)
    coverage = requirement_data.get("source_coverage") if isinstance(requirement_data.get("source_coverage"), dict) else {}
    for occ in coverage.get("source_requirement_occurrences") or []:
        if not isinstance(occ, dict) or occ.get("occurrence_type") != "declaration":
            continue
        if str(occ.get("coverage_status") or "") not in {"Covered", "Partially Covered"}:
            continue
        sid = str(occ.get("source_req_id") or "").strip()
        loc = str(occ.get("source_location") or occ.get("section_path") or "").strip()
        text = re.sub(r"\s+", " ", str(occ.get("source_excerpt") or "")).strip()
        if not text:
            continue
        line = f"{sid} · {loc}: {text}" if sid and loc else (f"{sid}: {text}" if sid else text)
        for key in list(occ.get("linked_candidate_ids") or []) + list(occ.get("linked_srs_ids") or []):
            key = str(key or "").strip()
            if not key:
                continue
            bucket = out.setdefault(key, [])
            if line not in bucket:
                bucket.append(line)
    return out


def _behavior_text_with_source_units(req: dict[str, Any], units: list[str]) -> str:
    base = _behavior_text(req)
    source_ids = _as_list(req.get("source_requirement_ids"))
    # Only expand the main view when grouping/semantic preservation needs it; otherwise keep the
    # normal engineer-facing output concise.
    needs_units = len(source_ids) > 1 or len(units) > 1 or bool(_as_list(req.get("source_semantic_unit_ids")))
    if not units or not needs_units:
        return base
    rendered = "\n".join(f"- {x}" for x in units)
    return f"{base}\n\n[Source-backed Atomic Behavior]\n{rendered}"


def _trace_sources(req: dict[str, Any]) -> list[tuple[str, str]]:
    """Return explicit source document/location pairs without inventing source IDs."""
    rows: list[tuple[str, str]] = []
    for item in req.get("source_evidence") or []:
        if isinstance(item, dict):
            document = str(item.get("document") or "").strip()
            location = str(item.get("location") or "").strip()
            evidence = str(item.get("text") or "").strip()
            source = document or "입력문서"
            detail = location or evidence
            pair = (source, detail)
        else:
            text = str(item).strip()
            pair = (text or "입력문서", "")
        if pair not in rows:
            rows.append(pair)
    return rows or [("입력문서에서 확인되지 않음", "")]


def build_swe1_records(requirement_data: dict[str, Any]) -> list[dict[str, str]]:
    records: list[dict[str, str]] = []
    requirements = requirement_data.get("requirements") or []
    source_units_by_ref = _source_behavior_units(requirement_data)
    for idx, req in enumerate(requirements, start=1):
        if not isinstance(req, dict):
            continue
        # V0.58: the Main SWE.1 body contains only confirmed SWE.1-eligible software requirements.
        # Review Needed and Not Applicable items are preserved in Allocation Annexes instead.
        if str(req.get("swe1_eligibility") or "Eligible") != "Eligible":
            continue
        records.append({
            "SRS ID": _text(req.get("srs_id"), empty=f"SRS_{idx:03d}") or f"SRS_{idx:03d}",
            "상위 기능": _text(req.get("function_name"), empty="미분류") or "미분류",
            "분류": "검토 필요" if str(req.get("swe1_eligibility") or "") == "Review Needed" else _normalize_classification(req.get("category")),
            "요구사항 내역": _text(req.get("requirement"), empty=NOT_CONFIRMED) or NOT_CONFIRMED,
            "동작 조건 / Trigger": _text(req.get("activation_trigger"), empty=NOT_CONFIRMED) or NOT_CONFIRMED,
            "작동 명세 정의": _behavior_text_with_source_units(
                req,
                source_units_by_ref.get(str(req.get("candidate_id") or ""), [])
                or source_units_by_ref.get(str(req.get("srs_id") or ""), []),
            ),
            "사전 조건": _text(req.get("preconditions"), empty="") or _text(req.get("system_input_preconditions"), empty=NOT_CONFIRMED) or NOT_CONFIRMED,
            "예상 결과": _text(req.get("output"), empty=NOT_CONFIRMED) or NOT_CONFIRMED,
            "검증 기준": _text(req.get("acceptance_criteria"), empty=NOT_CONFIRMED) or NOT_CONFIRMED,
            "출처 / Traceability": _format_evidence(req.get("source_evidence")),
            "기타": _other_materials(req),
            # The following values remain internal only. They are intentionally not exported as final SWE.1 fields.
            "_classification_basis": _text(req.get("classification_basis"), empty=""),
            "_evaluation_method": _text(req.get("evaluation_method"), empty=""),
            "_exception_conditions": _text(req.get("exception_conditions"), empty=""),
            "_clarification_needed": _text(req.get("clarification_needed"), empty=""),
            "_candidate_id": _text(req.get("candidate_id"), empty=""),
            "_source_requirement_ids": _text(req.get("source_requirement_ids"), empty=""),
            "_applicability": req.get("applicability") if isinstance(req.get("applicability"), dict) else {},
            "_external_dependencies": _text(req.get("external_dependencies"), empty=""),
            "_tbd_items": _text(req.get("tbd_items"), empty=""),
            "_conflicts": _text(req.get("conflicts"), empty=""),
            "_open_issue_ids": _text(req.get("open_issue_ids"), empty=""),
            "_requirement_status": _text(req.get("requirement_status"), empty="Draft") or "Draft",
            "_verification_constraints": _text(req.get("verification_constraints"), empty=""),
            "_knowledge_state": _text(req.get("knowledge_state"), empty=""),
            "_source_semantic_unit_ids": _text(req.get("source_semantic_unit_ids"), empty=""),
            "_requirement_level": _text(req.get("requirement_level"), empty=""),
            "_allocation_status": _text(req.get("allocation_status"), empty=""),
            "_swe1_eligibility": _text(req.get("swe1_eligibility"), empty=""),
            "_swe6_eligibility": _text(req.get("swe6_eligibility"), empty=""),
            "_verification_domain": _text(req.get("verification_domain"), empty=""),
            "_source_behavior_units": "\n".join(
                source_units_by_ref.get(str(req.get("candidate_id") or ""), [])
                or source_units_by_ref.get(str(req.get("srs_id") or ""), [])
            ),
        })
    return records


class SWE1Exporter:
    def __init__(self, output_dir: Path):
        self.output_dir = Path(output_dir).resolve()
        self.output_dir.mkdir(parents=True, exist_ok=True)

    def _filename(self, prefix: str, source_document: Path, suffix: str) -> Path:
        return self.output_dir / output_filename(
            prefix, _safe_stem(Path(source_document).name), suffix
        )

    def export_word(self, source_document: Path, requirement_data: dict[str, Any]) -> Path:
        records = build_swe1_records(requirement_data)
        all_requirements = [x for x in (requirement_data.get("requirements") or []) if isinstance(x, dict)]
        if not all_requirements:
            raise ValueError("SWE.1 Word로 내보낼 Requirement/Review Object가 없습니다.")
        out_path = self._filename("SWE.1 요구사항 정리", source_document, "docx")
        doc = Document()
        doc.styles["Normal"].font.name = "Malgun Gothic"
        doc.styles["Normal"].font.size = Pt(9)
        title = doc.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = title.add_run("SWE.1 소프트웨어 요구사항 정리")
        run.bold = True
        run.font.size = Pt(18)
        meta = doc.add_table(rows=4, cols=2)
        meta.style = "Table Grid"
        for row, (key, value) in enumerate([
            ("원본 문서", Path(source_document).name),
            ("생성 일시", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
            ("Main SWE.1 SRS 개수", str(len(records))),
            ("작성 기준", "Requirement Studio 산출물 작성 기준 V0.2 — SWE.1"),
        ]):
            meta.cell(row, 0).text = key
            meta.cell(row, 1).text = value
        note = doc.add_paragraph()
        note_run = note.add_run(
            "※ 입력문서를 factual Source of Truth로 사용하며 확인되지 않은 정보는 임의로 창작하지 않습니다. "
            "분류 근거·평가 방안·예외 조건·Gap/TBD/Conflict와 같은 분석 정보는 내부 Canonical/Review 정보로 유지하고, "
            "최종 SWE.1에는 엔지니어가 직접 활용할 핵심 Requirement 정보만 표시합니다."
        )
        note_run.font.size = Pt(8.5)
        note_run.font.color.rgb = RGBColor(90, 98, 110)

        section = doc.sections[0]
        usable_width = int(section.page_width - section.left_margin - section.right_margin)
        left_width = int(usable_width * 0.24)
        right_width = usable_width - left_width

        current_feature = None
        for record in records:
            feature = record["상위 기능"]
            if feature != current_feature:
                doc.add_heading(feature, level=1)
                current_feature = feature
            table = doc.add_table(rows=len(FINAL_FIELDS), cols=2)
            table.style = "Table Grid"
            table.autofit = False
            table.columns[0].width = left_width
            table.columns[1].width = right_width
            for row_idx, field in enumerate(FINAL_FIELDS):
                left, right = table.cell(row_idx, 0), table.cell(row_idx, 1)
                left.width, right.width = left_width, right_width
                left.text, right.text = field, record[field]
                left.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for paragraph in left.paragraphs:
                    for r in paragraph.runs:
                        r.bold = True
                        r.font.size = Pt(8.5)
                for paragraph in right.paragraphs:
                    for r in paragraph.runs:
                        r.font.size = Pt(8.5)
            doc.add_paragraph()

        # V0.47 Review Annex: keep the engineer-facing main view concise while preventing silent loss.
        annex_items = [r for r in records if any([
            r.get("_exception_conditions"), r.get("_clarification_needed"),
            r.get("_external_dependencies"), r.get("_tbd_items"), r.get("_conflicts"),
            r.get("_open_issue_ids"), r.get("_source_requirement_ids"), r.get("_source_semantic_unit_ids"), r.get("_allocation_status"),
            (r.get("_applicability") or {}).get("vehicle_lines"),
            (r.get("_applicability") or {}).get("baseline_versions"),
            (r.get("_applicability") or {}).get("feature_variants"),
        ])]
        if annex_items:
            doc.add_page_break()
            doc.add_heading("Review Annex — Applicability / Issues / Dependencies", level=1)
            note = doc.add_paragraph(
                "본 Annex는 Final SWE.1 본문에서 간결성을 위해 분리한 Source-backed Review 정보를 보존합니다. "
                "빈 값은 임의 추론하지 않으며 Source 확인이 필요한 항목은 Review Needed/UNKNOWN으로 유지합니다."
            )
            for record in annex_items:
                doc.add_heading(record["SRS ID"], level=2)
                table = doc.add_table(rows=14, cols=2)
                table.style = "Table Grid"
                app = record.get("_applicability") or {}
                rows = [
                    ("Source Requirement ID", record.get("_source_requirement_ids") or ""),
                    ("Semantic Source Unit ID", record.get("_source_semantic_unit_ids") or ""),
                    ("Source-backed Atomic Behavior", record.get("_source_behavior_units") or ""),
                    ("Requirement Level", record.get("_requirement_level") or ""),
                    ("Allocation / Eligibility", " | ".join(x for x in (record.get("_allocation_status") or "", record.get("_swe1_eligibility") or "", record.get("_swe6_eligibility") or "") if x)),
                    ("Verification Domain", record.get("_verification_domain") or ""),
                    ("Vehicle / Variant / Baseline", "; ".join(str(x) for k in ("vehicle_lines","feature_variants","baseline_versions") for x in (app.get(k) or []))),
                    ("Enable / Exclusion", "; ".join(str(x) for k in ("enable_conditions","exclusion_conditions") for x in (app.get(k) or []))),
                    ("Exception Conditions", record.get("_exception_conditions") or ""),
                    ("Clarification Needed", record.get("_clarification_needed") or ""),
                    ("External Dependency", record.get("_external_dependencies") or ""),
                    ("TBD / Conflict", " | ".join(x for x in (record.get("_tbd_items") or "", record.get("_conflicts") or "") if x)),
                    ("Requirement Status", record.get("_requirement_status") or "Draft"),
                    ("Knowledge State", record.get("_knowledge_state") or ""),
                ]
                for ridx, (label, value) in enumerate(rows):
                    table.cell(ridx, 0).text = label
                    table.cell(ridx, 1).text = value
                    for rr in table.cell(ridx, 0).paragraphs[0].runs:
                        rr.bold = True
                doc.add_paragraph()

        pending = [r for r in all_requirements if str(r.get("swe1_eligibility") or "Eligible") == "Review Needed"]
        if pending:
            doc.add_page_break()
            doc.add_heading("Allocation Review Annex — Pending SWE.1 Allocation", level=1)
            doc.add_paragraph(
                "다음 Source-backed 항목은 시스템/다중 요소 책임 또는 불충분한 allocation evidence로 인해 Main SWE.1 본문에 포함하지 않습니다. "
                "Source fact는 보존하며 SW allocation 확인 후 SWE.1 승격 여부를 결정합니다."
            )
            for req in pending:
                doc.add_heading(str(req.get("srs_id") or req.get("candidate_id") or "Review Object"), level=2)
                table = doc.add_table(rows=10, cols=2); table.style = "Table Grid"
                rows = [
                    ("Source-backed Candidate", _text(req.get("requirement"), empty="")),
                    ("Source Traceability", _format_evidence(req.get("source_evidence"))),
                    ("Semantic Source Unit", _text(req.get("source_semantic_unit_ids"), empty="")),
                    ("Source-backed Atomic Behavior / Fact", _text(req.get("source_backed_atomic_behaviors") or req.get("source_backed_facts"), empty="")),
                    ("Requirement Level", _text(req.get("requirement_level"), empty="")),
                    ("Allocation Status", _text(req.get("allocation_status"), empty="")),
                    ("SWE.1 / SWE.6 Eligibility", " | ".join(x for x in (_text(req.get("swe1_eligibility"), empty=""), _text(req.get("swe6_eligibility"), empty="")) if x)),
                    ("Verification Domain", _text(req.get("verification_domain"), empty="")),
                    ("Allocation Rationale", _text(req.get("allocation_rationale"), empty="")),
                    ("Clarification Needed", _text(req.get("clarification_needed"), empty="")),
                ]
                for ridx,(label,value) in enumerate(rows):
                    table.cell(ridx,0).text=label; table.cell(ridx,1).text=value
                    for rr in table.cell(ridx,0).paragraphs[0].runs: rr.bold=True
                doc.add_paragraph()

        excluded = [r for r in all_requirements if str(r.get("swe1_eligibility") or "Eligible") == "Not Applicable"]
        if excluded:
            doc.add_page_break()
            doc.add_heading("Cross-domain Allocation Annex — Source Facts Outside SWE.1", level=1)
            doc.add_paragraph(
                "다음 Source-backed 항목은 삭제된 것이 아니라 Canonical/Review Package에 보존됩니다. "
                "V0.63 Allocation Gate에 따라 Main SWE.1/SWE.6 분모에서 제외되고 해당 전문 verification domain으로 전달됩니다."
            )
            for req in excluded:
                doc.add_heading(str(req.get("srs_id") or req.get("candidate_id") or "Review Object"), level=2)
                table = doc.add_table(rows=9, cols=2); table.style = "Table Grid"
                rows = [
                    ("Source Fact", _text(req.get("requirement"), empty="")),
                    ("Source Traceability", _format_evidence(req.get("source_evidence"))),
                    ("Semantic Source Unit", _text(req.get("source_semantic_unit_ids"), empty="")),
                    ("Source-backed Fact", _text(req.get("source_backed_facts"), empty="")),
                    ("Requirement Level", _text(req.get("requirement_level"), empty="")),
                    ("Allocation Status", _text(req.get("allocation_status"), empty="")),
                    ("Verification Domain", _text(req.get("verification_domain"), empty="")),
                    ("Allocation Rationale", _text(req.get("allocation_rationale"), empty="")),
                    ("Value Relation / Normative Strength", _text(req.get("value_relation_status") or req.get("normative_strength_status"), empty="")),
                ]
                for ridx,(label,value) in enumerate(rows):
                    table.cell(ridx,0).text=label; table.cell(ridx,1).text=value
                    for rr in table.cell(ridx,0).paragraphs[0].runs: rr.bold=True
                doc.add_paragraph()
        doc.save(out_path)
        return out_path

    def export_excel(self, source_document: Path, requirement_data: dict[str, Any]) -> Path:
        records = build_swe1_records(requirement_data)
        all_requirements = [x for x in (requirement_data.get("requirements") or []) if isinstance(x, dict)]
        if not all_requirements:
            raise ValueError("SWE.1 Excel로 내보낼 Requirement/Review Object가 없습니다.")
        out_path = self._filename("SWE.1 요구사항 명세", source_document, "xlsx")
        wb = Workbook()
        ws0 = wb.active
        ws0.title = "00_Overview"
        ws1 = wb.create_sheet("01_SWE1")
        ws2 = wb.create_sheet("02_Traceability")
        ws3 = wb.create_sheet("03_Applicability")
        ws4 = wb.create_sheet("04_Open_Issues_Gaps")
        ws5 = wb.create_sheet("05_Coverage")
        ws6 = wb.create_sheet("06_Requirement_Review_Details")
        ws7 = wb.create_sheet("07_Allocation_Annex")

        blue, light, border_color = "2F75B5", "D9EAF7", "D0D7DE"
        thin = Side(style="thin", color=border_color)
        header_fill = PatternFill("solid", fgColor=blue)
        sub_fill = PatternFill("solid", fgColor=light)
        white_font, bold = Font(color="FFFFFF", bold=True), Font(bold=True)
        wrap = Alignment(vertical="top", wrap_text=True)
        center = Alignment(horizontal="center", vertical="center", wrap_text=True)
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        for row in [
            ["항목", "내용"],
            ["문서", "SWE.1 소프트웨어 요구사항 명세"],
            ["원본 문서", Path(source_document).name],
            ["생성 일시", datetime.now().strftime("%Y-%m-%d %H:%M:%S")],
            ["SRS 개수", len(records)],
            ["작성 기준", "Requirement Studio 산출물 작성 기준 V0.2 — SWE.1"],
            ["핵심 원칙", "01_SWE1은 간결한 Final View로 유지하고 상세 적용성/Issue/Coverage 정보는 별도 Review View에서 추적"],
            ["Source 원칙", "입력문서에 없는 사실은 만들지 않으며 불명확/미정/상충 정보는 Review View 및 Canonical에서 추적"],
        ]:
            ws0.append(row)
        for cell in ws0[1]:
            cell.fill, cell.font, cell.alignment = header_fill, white_font, center
        for row in ws0.iter_rows(min_row=2):
            row[0].font = bold
            for cell in row:
                cell.alignment, cell.border = wrap, border
        ws0.column_dimensions["A"].width, ws0.column_dimensions["B"].width = 18, 90
        ws0.freeze_panes = "A2"

        headers = FINAL_FIELDS
        ws1.append(headers)
        for record in records:
            ws1.append([record[h] for h in headers])
        for cell in ws1[1]:
            cell.fill, cell.font, cell.alignment, cell.border = header_fill, white_font, center, border
        for row in ws1.iter_rows(min_row=2):
            for cell in row:
                cell.alignment, cell.border = wrap, border
        ws1.freeze_panes, ws1.auto_filter.ref = "A2", ws1.dimensions
        widths = [12, 22, 11, 48, 32, 46, 30, 32, 32, 56, 40]
        for idx, width in enumerate(widths, start=1):
            ws1.column_dimensions[get_column_letter(idx)].width = width
        for row in ws1.iter_rows(min_row=2):
            row[0].fill, row[0].font = sub_fill, bold

        # End-to-end traceability view: Source -> SWE.1 SRS -> SWE.6 TC.
        trace_headers = [
            "입력 Source",
            "Source Location / Requirement",
            "SWE.1 SRS ID",
            "SWE.1 Requirement",
            "SWE.6 Test Case ID",
            "상태",
        ]
        ws2.append(trace_headers)
        tc_by_srs: dict[str, list[str]] = {}
        for case in build_swe6_cases(requirement_data):
            tc_by_srs.setdefault(str(case.get("srs_id") or ""), []).append(str(case.get("tc_id") or ""))
        requirements = [x for x in all_requirements if str(x.get("swe1_eligibility") or "Eligible") == "Eligible"]
        for idx, req in enumerate(requirements, start=1):
            srs_id = _text(req.get("srs_id"), empty=f"SRS_{idx:03d}") or f"SRS_{idx:03d}"
            requirement = _text(req.get("requirement"), empty=NOT_CONFIRMED) or NOT_CONFIRMED
            tc_ids = [x for x in tc_by_srs.get(srs_id, []) if x]
            tc_text = ", ".join(tc_ids)
            status = "연결" if tc_ids else "검토 필요"
            for source, location in _trace_sources(req):
                ws2.append([source, location, srs_id, requirement, tc_text, status])
        for cell in ws2[1]:
            cell.fill, cell.font, cell.alignment, cell.border = header_fill, white_font, center, border
        for row in ws2.iter_rows(min_row=2):
            for cell in row:
                cell.alignment, cell.border = wrap, border
        ws2.freeze_panes = "A2"
        for idx, width in enumerate([34, 50, 14, 62, 24, 14], start=1):
            ws2.column_dimensions[get_column_letter(idx)].width = width

        # V0.47 Review View sheets: preserve review-critical Canonical information without bloating 01_SWE1.
        app_headers = ["SRS ID", "Source Requirement ID", "Vehicle Lines", "Baseline Versions", "Feature Variants", "Enable Conditions", "Exclusion Conditions", "Knowledge State", "Source Evidence"]
        ws3.append(app_headers)
        for record in records:
            app = record.get("_applicability") or {}
            ws3.append([
                record["SRS ID"], record.get("_source_requirement_ids") or "",
                _text(app.get("vehicle_lines")), _text(app.get("baseline_versions")), _text(app.get("feature_variants")),
                _text(app.get("enable_conditions")), _text(app.get("exclusion_conditions")),
                _text(app.get("knowledge_state")), _text(app.get("source_evidence")),
            ])

        issue_headers = ["Type", "ID / SRS", "Related SRS", "Status", "Detail", "Source / Link"]
        ws4.append(issue_headers)
        for record in records:
            srs = record["SRS ID"]
            for typ, value in [
                ("Exception", record.get("_exception_conditions")),
                ("Clarification Needed", record.get("_clarification_needed")),
                ("External Dependency", record.get("_external_dependencies")),
                ("TBD", record.get("_tbd_items")),
                ("Conflict", record.get("_conflicts")),
                ("Open Issue", record.get("_open_issue_ids")),
            ]:
                if value:
                    ws4.append([typ, srs, srs, record.get("_requirement_status") or "Draft", value, record.get("출처 / Traceability") or ""])
        for gap in [x for x in (requirement_data.get("gaps") or []) if isinstance(x, dict)]:
            ws4.append([
                "Gap", _text(gap.get("gap_id")), _text(gap.get("related_srs_ids")),
                "Blocked" if gap.get("blocking_for_verification") else "Review Needed",
                _text(gap.get("description") or gap.get("gap") or gap.get("reason")),
                _text(gap.get("related_source_requirement_ids") or gap.get("related_source_occurrence_ids")),
            ])
        for conflict in [x for x in (requirement_data.get("conflict_register") or []) if isinstance(x, dict)]:
            ws4.append([
                conflict.get("classification") or "Conflict", conflict.get("conflict_id") or "",
                _text(conflict.get("affected_srs_ids")), conflict.get("status") or "Open",
                _text(conflict.get("title") or conflict.get("description")),
                _text({"source_a": conflict.get("source_a"), "source_b": conflict.get("source_b")}),
            ])

        coverage = requirement_data.get("source_coverage") if isinstance(requirement_data.get("source_coverage"), dict) else {}
        cov_headers = ["Occurrence ID", "Occurrence Type", "Source REQ ID", "Source Location", "Kind", "Coverage Status", "Candidate IDs", "SRS IDs", "Gap/Issue IDs", "Disposition Reason", "Source Excerpt", "Content Hash"]
        ws5.append(cov_headers)
        for occ in coverage.get("source_requirement_occurrences") or []:
            if not isinstance(occ, dict):
                continue
            ws5.append([
                occ.get("occurrence_id"), occ.get("occurrence_type"), occ.get("source_req_id"), occ.get("source_location"), occ.get("source_kind"),
                occ.get("coverage_status"), _text(occ.get("linked_candidate_ids")), _text(occ.get("linked_srs_ids")),
                _text(occ.get("linked_gap_issue_ids")), occ.get("disposition_reason"), occ.get("source_excerpt"), occ.get("content_hash"),
            ])

        detail_headers = ["SRS ID", "Candidate ID", "Source Requirement ID", "Requirement Status", "Knowledge State", "Exception Conditions", "Clarification Needed", "External Dependencies", "TBD Items", "Conflicts", "Verification Constraints", "Source Evidence", "SWE.6 TC ID"]
        ws6.append(detail_headers)
        for record in records:
            tc_ids = ", ".join(x for x in tc_by_srs.get(record["SRS ID"], []) if x)
            ws6.append([
                record["SRS ID"], record.get("_candidate_id"), record.get("_source_requirement_ids"),
                record.get("_requirement_status"), record.get("_knowledge_state"), record.get("_exception_conditions"),
                record.get("_clarification_needed"), record.get("_external_dependencies"), record.get("_tbd_items"),
                record.get("_conflicts"), record.get("_verification_constraints"), record.get("출처 / Traceability"), tc_ids,
            ])

        alloc_headers = ["SRS ID", "SWE.1 Eligibility", "SWE.6 Eligibility", "Requirement Level", "Allocation Status", "Verification Domain", "Source Fact / Candidate", "Source Traceability", "Semantic Source Unit", "Rationale"]
        ws7.append(alloc_headers)
        for req in all_requirements:
            if str(req.get("swe1_eligibility") or "Eligible") == "Eligible":
                continue
            ws7.append([
                _text(req.get("srs_id")), _text(req.get("swe1_eligibility")), _text(req.get("swe6_eligibility")),
                _text(req.get("requirement_level")), _text(req.get("allocation_status")), _text(req.get("verification_domain")),
                _text(req.get("requirement")), _format_evidence(req.get("source_evidence")), _text(req.get("source_semantic_unit_ids")),
                _text(req.get("allocation_rationale")),
            ])

        for review_ws in (ws3, ws4, ws5, ws6, ws7):
            for cell in review_ws[1]:
                cell.fill, cell.font, cell.alignment, cell.border = header_fill, white_font, center, border
            for row in review_ws.iter_rows(min_row=2):
                for cell in row:
                    cell.alignment, cell.border = wrap, border
            review_ws.freeze_panes = "A2"
            review_ws.auto_filter.ref = review_ws.dimensions
        for idx, width in enumerate([14,24,28,22,28,30,30,18,50], start=1): ws3.column_dimensions[get_column_letter(idx)].width = width
        for idx, width in enumerate([20,18,24,18,70,70], start=1): ws4.column_dimensions[get_column_letter(idx)].width = width
        for idx, width in enumerate([18,18,20,30,16,22,24,24,24,52,70,20], start=1): ws5.column_dimensions[get_column_letter(idx)].width = width
        for idx, width in enumerate([14,18,24,18,16,50,50,50,40,50,45,70,24], start=1): ws6.column_dimensions[get_column_letter(idx)].width = width
        for idx, width in enumerate([14,18,18,24,34,34,70,70,34,70], start=1): ws7.column_dimensions[get_column_letter(idx)].width = width

        for ws in (ws0, ws1, ws2, ws3, ws4, ws5, ws6, ws7):
            ws.sheet_view.showGridLines = False
        wb.save(out_path)
        return out_path
