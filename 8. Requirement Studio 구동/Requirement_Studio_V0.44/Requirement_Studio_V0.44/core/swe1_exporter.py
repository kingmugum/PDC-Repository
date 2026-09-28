from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path
from typing import Any

from core.output_naming import output_filename
from core.swe6_exporter import build_swe6_cases

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


NOT_CONFIRMED = "입력문서에서 확인되지 않음"
NO_DETAIL = "입력문서에서 추가 작동 명세를 확인할 수 없음"


FINAL_FIELDS = [
    "SRS ID",
    "상위 기능",
    "분류",
    "요구사항 내역",
    "동작 조건 / Trigger",
    "작동 명세 정의",
    "사전 조건",
    "예상 결과",
    "검증 기준",
    "출처 / Traceability",
    "기타",
]


def _safe_stem(name: str) -> str:
    stem = Path(name).stem or "document"
    stem = re.sub(r'[<>:"/\\|?*]+', "_", stem)
    stem = re.sub(r"\s+", "_", stem).strip("._ ")
    return stem[:80] or "document"


def _as_list(value: Any) -> list[str]:
    if value is None:
        return []
    if isinstance(value, list):
        result = []
        for item in value:
            if isinstance(item, dict):
                text = "; ".join(f"{k}={v}" for k, v in item.items() if v not in (None, "", [], {}))
            else:
                text = str(item).strip()
            if text:
                result.append(text)
        return result
    if isinstance(value, dict):
        text = "; ".join(f"{k}={v}" for k, v in value.items() if v not in (None, "", [], {}))
        return [text] if text else []
    text = str(value).strip()
    return [text] if text else []


def _text(value: Any, *, empty: str = "") -> str:
    items = _as_list(value)
    return "\n".join(items) if items else empty


def _normalize_classification(value: Any) -> str:
    raw = _text(value).strip()
    low = raw.lower()
    if "비기능" in raw or "non-functional" in low or "nonfunctional" in low:
        return "비기능"
    if raw == "기능" or "functional" in low or ("기능" in raw and "비기능" not in raw):
        return "기능"
    return raw or "검토 필요"


def _format_evidence(items: Any) -> str:
    lines = []
    for idx, item in enumerate(items if isinstance(items, list) else [], start=1):
        if not isinstance(item, dict):
            text = str(item).strip()
            if text:
                lines.append(f"[{idx}] {text}")
            continue
        document = str(item.get("document") or "").strip()
        location = str(item.get("location") or "").strip()
        evidence = str(item.get("text") or "").strip()
        prefix = " / ".join(x for x in (document, location) if x)
        if evidence:
            lines.append(f"[{idx}] {prefix}: {evidence}" if prefix else f"[{idx}] {evidence}")
        elif prefix:
            lines.append(f"[{idx}] {prefix}")
    return "\n".join(lines) if lines else NOT_CONFIRMED


def _behavior_text(req: dict[str, Any]) -> str:
    """Render detailed behavior for general engineers without exposing BF IDs."""
    flows = _as_list(req.get("behavior_flows"))
    if flows:
        rendered = []
        for idx, item in enumerate(flows, start=1):
            clean = re.sub(r"^BF\s*\d+(?:[-.]\d+)?[.:\-]?\s*", "", item, flags=re.IGNORECASE).strip()
            rendered.append(f"{idx}. {clean or item}")
        return "\n".join(rendered)
    processing = _text(req.get("processing_action")).strip()
    if processing:
        return processing
    return NO_DETAIL


def _other_materials(req: dict[str, Any]) -> str:
    """Final-output '기타': only source-backed auxiliary materials such as figures/diagrams/tables."""
    lines: list[str] = []
    for key in (
        "related_artifacts",
        "diagram_refs",
        "figure_refs",
        "table_refs",
        "visual_evidence_refs",
        "attachments",
    ):
        for item in _as_list(req.get(key)):
            if item and item not in lines:
                lines.append(item)
    return "\n".join(lines)


def _trace_sources(req: dict[str, Any]) -> list[tuple[str, str]]:
    """Return explicit source document/location pairs without inventing source IDs."""
    rows: list[tuple[str, str]] = []
    for item in req.get("source_evidence") or []:
        if isinstance(item, dict):
            document = str(item.get("document") or "").strip()
            location = str(item.get("location") or "").strip()
            evidence = str(item.get("text") or "").strip()
            source = document or "입력문서"
            detail = location or evidence
            pair = (source, detail)
        else:
            text = str(item).strip()
            pair = (text or "입력문서", "")
        if pair not in rows:
            rows.append(pair)
    return rows or [("입력문서에서 확인되지 않음", "")]


def build_swe1_records(requirement_data: dict[str, Any]) -> list[dict[str, str]]:
    records: list[dict[str, str]] = []
    requirements = requirement_data.get("requirements") or []
    for idx, req in enumerate(requirements, start=1):
        if not isinstance(req, dict):
            continue
        records.append({
            "SRS ID": _text(req.get("srs_id"), empty=f"SRS_{idx:03d}") or f"SRS_{idx:03d}",
            "상위 기능": _text(req.get("function_name"), empty="미분류") or "미분류",
            "분류": _normalize_classification(req.get("category")),
            "요구사항 내역": _text(req.get("requirement"), empty=NOT_CONFIRMED) or NOT_CONFIRMED,
            "동작 조건 / Trigger": _text(req.get("activation_trigger"), empty=NOT_CONFIRMED) or NOT_CONFIRMED,
            "작동 명세 정의": _behavior_text(req),
            "사전 조건": _text(req.get("preconditions"), empty="") or _text(req.get("system_input_preconditions"), empty=NOT_CONFIRMED) or NOT_CONFIRMED,
            "예상 결과": _text(req.get("output"), empty=NOT_CONFIRMED) or NOT_CONFIRMED,
            "검증 기준": _text(req.get("acceptance_criteria"), empty=NOT_CONFIRMED) or NOT_CONFIRMED,
            "출처 / Traceability": _format_evidence(req.get("source_evidence")),
            "기타": _other_materials(req),
            # The following values remain internal only. They are intentionally not exported as final SWE.1 fields.
            "_classification_basis": _text(req.get("classification_basis"), empty=""),
            "_evaluation_method": _text(req.get("evaluation_method"), empty=""),
            "_exception_conditions": _text(req.get("exception_conditions"), empty=""),
            "_clarification_needed": _text(req.get("clarification_needed"), empty=""),
            "_candidate_id": _text(req.get("candidate_id"), empty=""),
        })
    return records


class SWE1Exporter:
    def __init__(self, output_dir: Path):
        self.output_dir = Path(output_dir).resolve()
        self.output_dir.mkdir(parents=True, exist_ok=True)

    def _filename(self, prefix: str, source_document: Path, suffix: str) -> Path:
        return self.output_dir / output_filename(
            prefix, _safe_stem(Path(source_document).name), suffix
        )

    def export_word(self, source_document: Path, requirement_data: dict[str, Any]) -> Path:
        records = build_swe1_records(requirement_data)
        if not records:
            raise ValueError("SWE.1 Word로 내보낼 Requirement가 없습니다.")
        out_path = self._filename("SWE.1 요구사항 정리", source_document, "docx")
        doc = Document()
        doc.styles["Normal"].font.name = "Malgun Gothic"
        doc.styles["Normal"].font.size = Pt(9)
        title = doc.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = title.add_run("SWE.1 소프트웨어 요구사항 정리")
        run.bold = True
        run.font.size = Pt(18)
        meta = doc.add_table(rows=4, cols=2)
        meta.style = "Table Grid"
        for row, (key, value) in enumerate([
            ("원본 문서", Path(source_document).name),
            ("생성 일시", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
            ("SRS 개수", str(len(records))),
            ("작성 기준", "Requirement Studio 산출물 작성 기준 V0.2 — SWE.1"),
        ]):
            meta.cell(row, 0).text = key
            meta.cell(row, 1).text = value
        note = doc.add_paragraph()
        note_run = note.add_run(
            "※ 입력문서를 factual Source of Truth로 사용하며 확인되지 않은 정보는 임의로 창작하지 않습니다. "
            "분류 근거·평가 방안·예외 조건·Gap/TBD/Conflict와 같은 분석 정보는 내부 Canonical/Review 정보로 유지하고, "
            "최종 SWE.1에는 엔지니어가 직접 활용할 핵심 Requirement 정보만 표시합니다."
        )
        note_run.font.size = Pt(8.5)
        note_run.font.color.rgb = RGBColor(90, 98, 110)

        section = doc.sections[0]
        usable_width = int(section.page_width - section.left_margin - section.right_margin)
        left_width = int(usable_width * 0.24)
        right_width = usable_width - left_width

        current_feature = None
        for record in records:
            feature = record["상위 기능"]
            if feature != current_feature:
                doc.add_heading(feature, level=1)
                current_feature = feature
            table = doc.add_table(rows=len(FINAL_FIELDS), cols=2)
            table.style = "Table Grid"
            table.autofit = False
            table.columns[0].width = left_width
            table.columns[1].width = right_width
            for row_idx, field in enumerate(FINAL_FIELDS):
                left, right = table.cell(row_idx, 0), table.cell(row_idx, 1)
                left.width, right.width = left_width, right_width
                left.text, right.text = field, record[field]
                left.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for paragraph in left.paragraphs:
                    for r in paragraph.runs:
                        r.bold = True
                        r.font.size = Pt(8.5)
                for paragraph in right.paragraphs:
                    for r in paragraph.runs:
                        r.font.size = Pt(8.5)
            doc.add_paragraph()
        doc.save(out_path)
        return out_path

    def export_excel(self, source_document: Path, requirement_data: dict[str, Any]) -> Path:
        records = build_swe1_records(requirement_data)
        if not records:
            raise ValueError("SWE.1 Excel로 내보낼 Requirement가 없습니다.")
        out_path = self._filename("SWE.1 요구사항 명세", source_document, "xlsx")
        wb = Workbook()
        ws0 = wb.active
        ws0.title = "00_Overview"
        ws1 = wb.create_sheet("01_SWE1")
        ws2 = wb.create_sheet("02_Traceability")

        blue, light, border_color = "2F75B5", "D9EAF7", "D0D7DE"
        thin = Side(style="thin", color=border_color)
        header_fill = PatternFill("solid", fgColor=blue)
        sub_fill = PatternFill("solid", fgColor=light)
        white_font, bold = Font(color="FFFFFF", bold=True), Font(bold=True)
        wrap = Alignment(vertical="top", wrap_text=True)
        center = Alignment(horizontal="center", vertical="center", wrap_text=True)
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        for row in [
            ["항목", "내용"],
            ["문서", "SWE.1 소프트웨어 요구사항 명세"],
            ["원본 문서", Path(source_document).name],
            ["생성 일시", datetime.now().strftime("%Y-%m-%d %H:%M:%S")],
            ["SRS 개수", len(records)],
            ["작성 기준", "Requirement Studio 산출물 작성 기준 V0.2 — SWE.1"],
            ["핵심 원칙", "일반 엔지니어가 바로 읽을 수 있는 간결한 SWE.1 View를 제공하고 상세 분석 정보는 Canonical 내부에 유지"],
            ["Source 원칙", "입력문서에 없는 사실은 만들지 않으며 불명확/미정/상충 정보는 내부 Review Flag로 유지"],
        ]:
            ws0.append(row)
        for cell in ws0[1]:
            cell.fill, cell.font, cell.alignment = header_fill, white_font, center
        for row in ws0.iter_rows(min_row=2):
            row[0].font = bold
            for cell in row:
                cell.alignment, cell.border = wrap, border
        ws0.column_dimensions["A"].width, ws0.column_dimensions["B"].width = 18, 90
        ws0.freeze_panes = "A2"

        headers = FINAL_FIELDS
        ws1.append(headers)
        for record in records:
            ws1.append([record[h] for h in headers])
        for cell in ws1[1]:
            cell.fill, cell.font, cell.alignment, cell.border = header_fill, white_font, center, border
        for row in ws1.iter_rows(min_row=2):
            for cell in row:
                cell.alignment, cell.border = wrap, border
        ws1.freeze_panes, ws1.auto_filter.ref = "A2", ws1.dimensions
        widths = [12, 22, 11, 48, 32, 46, 30, 32, 32, 56, 40]
        for idx, width in enumerate(widths, start=1):
            ws1.column_dimensions[get_column_letter(idx)].width = width
        for row in ws1.iter_rows(min_row=2):
            row[0].fill, row[0].font = sub_fill, bold

        # End-to-end traceability view: Source -> SWE.1 SRS -> SWE.6 TC.
        trace_headers = [
            "입력 Source",
            "Source Location / Requirement",
            "SWE.1 SRS ID",
            "SWE.1 Requirement",
            "SWE.6 Test Case ID",
            "상태",
        ]
        ws2.append(trace_headers)
        tc_by_srs: dict[str, list[str]] = {}
        for case in build_swe6_cases(requirement_data):
            tc_by_srs.setdefault(str(case.get("srs_id") or ""), []).append(str(case.get("tc_id") or ""))
        requirements = [x for x in (requirement_data.get("requirements") or []) if isinstance(x, dict)]
        for idx, req in enumerate(requirements, start=1):
            srs_id = _text(req.get("srs_id"), empty=f"SRS_{idx:03d}") or f"SRS_{idx:03d}"
            requirement = _text(req.get("requirement"), empty=NOT_CONFIRMED) or NOT_CONFIRMED
            tc_ids = [x for x in tc_by_srs.get(srs_id, []) if x]
            tc_text = ", ".join(tc_ids)
            status = "연결" if tc_ids else "검토 필요"
            for source, location in _trace_sources(req):
                ws2.append([source, location, srs_id, requirement, tc_text, status])
        for cell in ws2[1]:
            cell.fill, cell.font, cell.alignment, cell.border = header_fill, white_font, center, border
        for row in ws2.iter_rows(min_row=2):
            for cell in row:
                cell.alignment, cell.border = wrap, border
        ws2.freeze_panes = "A2"
        for idx, width in enumerate([34, 50, 14, 62, 24, 14], start=1):
            ws2.column_dimensions[get_column_letter(idx)].width = width

        for ws in (ws0, ws1, ws2):
            ws.sheet_view.showGridLines = False
        wb.save(out_path)
        return out_path
