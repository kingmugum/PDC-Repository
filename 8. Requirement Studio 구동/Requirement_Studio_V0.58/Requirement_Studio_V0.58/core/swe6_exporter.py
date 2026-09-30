from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

from core.output_naming import output_filename

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

NOT_CONFIRMED = "입력문서에서 확인되지 않음"
REVIEW_NEEDED = "검토 필요"


def _safe_stem(name: str) -> str:
    stem = Path(name).stem or "document"
    stem = re.sub(r'[<>:"/\\|?*]+', "_", stem)
    stem = re.sub(r"\s+", "_", stem).strip("._ ")
    return stem[:80] or "document"


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, list):
        return "\n".join(str(x).strip() for x in value if str(x).strip())
    if isinstance(value, dict):
        return "; ".join(f"{k}={v}" for k, v in value.items() if v not in (None, "", [], {}))
    return str(value).strip()


def _first_nonempty(*values: Any) -> str:
    for value in values:
        text = _text(value)
        if text and text not in {NOT_CONFIRMED, "-"}:
            return text
    return ""


def _short_name(text: str, max_len: int = 46) -> str:
    text = re.sub(r"\s+", " ", text or "").strip().rstrip(".")
    for prefix in ("소프트웨어는 ", "Software는 ", "SW는 "):
        if text.startswith(prefix):
            text = text[len(prefix):]
            break
    return text if len(text) <= max_len else text[: max_len - 1].rstrip() + "…"


def _metadata_value(source_text: str, labels: list[str], max_len: int = 80) -> str:
    if not source_text:
        return ""
    for label in labels:
        # Explicitly labelled metadata only. Do not infer from descriptive prose.
        pattern = rf"(?im)^\s*{re.escape(label)}\s*[:：\t|]\s*([^\n|]{{1,{max_len}}})\s*$"
        match = re.search(pattern, source_text)
        if match:
            value = re.sub(r"\s+", " ", match.group(1)).strip()
            if len(value) <= max_len and not re.search(r"[。!?]\s*$", value):
                return value
    return ""


def _to_swe6_artifact_id(value: str) -> str:
    value = (value or "").strip()
    if not value:
        return ""
    # Change only an explicitly embedded SWE.1 process marker. Other IDs are preserved.
    return re.sub(r"(?i)SWE[._ -]?1", "SWE.6", value)


def _section_excerpt(source_text: str, labels: list[str], max_chars: int = 420) -> str:
    lines = [re.sub(r"\s+", " ", x).strip() for x in (source_text or "").splitlines()]
    for idx, line in enumerate(lines):
        if any(label.lower() in line.lower() for label in labels):
            body = []
            for nxt in lines[idx + 1: idx + 8]:
                if not nxt:
                    continue
                if re.match(r"^\d+(?:\.\d+)*\.?\s+", nxt) and body:
                    break
                body.append(nxt)
                if sum(len(x) for x in body) >= max_chars:
                    break
            return " ".join(body)[:max_chars].strip()
    return ""


def extract_metadata(source_text: str) -> dict[str, str]:
    vehicle = _metadata_value(source_text, ["차종", "Vehicle", "Vehicle Model", "Model"], 40)
    oem = _metadata_value(source_text, ["OEM", "고객사", "Customer"], 40)
    project = _metadata_value(source_text, ["프로젝트명", "Project Name", "Project"], 80)
    artifact = _metadata_value(source_text, ["산출물 ID", "문서 ID", "Document ID", "Artifact ID", "Deliverable ID"], 60)
    return {
        "artifact_id": _to_swe6_artifact_id(artifact),
        "vehicle_oem": " / ".join(x for x in (vehicle, oem) if x),
        "project_name": project,
        "hardware": _metadata_value(source_text, ["Hardware", "하드웨어"], 80),
        "software": _metadata_value(source_text, ["Software", "소프트웨어"], 80),
        "mechanical": _metadata_value(source_text, ["Mechanical", "메카니컬", "기구"], 80),
        "environment": _section_excerpt(source_text, ["테스트 환경", "Test Environment"], 420),
    }


def _parse_compare_value(text: str) -> tuple[str, str, str]:
    """Return variable, compare, value only when directly recoverable from text."""
    raw = re.sub(r"\s+", " ", text or "").strip()
    if not raw:
        return "", "", ""
    # Signal-style equality: Foo = 0x1 / Foo(0x1)
    m = re.search(r"([A-Za-z_][A-Za-z0-9_./-]{1,50})\s*(?:=|==)\s*(0x[0-9A-Fa-f]+|[-+]?\d+(?:\.\d+)?(?:\s*[A-Za-z%℃°/]+)?)", raw)
    if m:
        return m.group(1), "=", m.group(2).strip()
    m = re.search(r"([A-Za-z_][A-Za-z0-9_./-]{1,50})\s*\(\s*(0x[0-9A-Fa-f]+)\s*\)", raw)
    if m:
        return m.group(1), "=", m.group(2)

    op_map = {"이상": ">=", "이하": "<=", "초과": ">", "미만": "<"}
    unit = r"(?:ms|msec|s|sec|초|V|mV|A|mA|%|℃|°C|도|Step|step|회|bps|kbps)?"
    m = re.search(rf"([-+]?\d+(?:\.\d+)?)\s*({unit})\s*(이상|이하|초과|미만)", raw)
    if m:
        val = m.group(1) + ((" " + m.group(2)) if m.group(2) else "")
        before = raw[:m.start()].strip(" ,.:;-()")
        variable = _short_name(before.split("에서")[-1].split("경우")[-1], 36)
        return variable, op_map[m.group(3)], val.strip()
    m = re.search(rf"(>=|<=|>|<)\s*([-+]?\d+(?:\.\d+)?)\s*({unit})", raw)
    if m:
        val = m.group(2) + ((" " + m.group(3)) if m.group(3) else "")
        before = raw[:m.start()].strip(" ,.:;-()")
        return _short_name(before, 36), m.group(1), val.strip()
    return "", "", ""


def _technique(req: dict[str, Any]) -> str:
    text = " ".join(_text(req.get(k)) for k in (
        "requirement", "activation_trigger", "preconditions", "processing_action", "output", "acceptance_criteria"
    )).lower()
    methods = ["요구사항 기반 테스트"]
    numeric_boundary = bool(re.search(r"(?:>=|<=|>|<|이상|이하|초과|미만|최대|최소)\s*[-+]?\d|[-+]?\d+(?:\.\d+)?\s*(?:ms|s|초|v|mv|a|ma|%|℃|도)\s*(?:이상|이하|초과|미만)", text))
    if numeric_boundary:
        methods.append("경계값 분석")
    elif any(x in text for x in ("invalid", "reserved", "유효", "무효", "범위", "0x0 ~", "0x0~")):
        methods.append("동등분할")
    elif any(x in text for x in ("state", "mode", "상태전이", "상태 전이", "진입", "복귀", "wake", "sleep")):
        methods.append("상태전이 테스트")
    return " / ".join(methods)


def _testability(req: dict[str, Any]) -> str:
    if str(req.get("swe6_eligibility") or "Eligible") != "Eligible":
        return REVIEW_NEEDED
    category = _text(req.get("category"))
    requirement = _text(req.get("requirement"))
    if not requirement or category == REVIEW_NEEDED:
        return REVIEW_NEEDED
    gap = _text(req.get("clarification_needed"))
    meaningful = any(_first_nonempty(req.get(k)) for k in (
        "activation_trigger", "preconditions", "processing_action", "output", "acceptance_criteria"
    ))
    if not meaningful and gap:
        return REVIEW_NEEDED
    return "대상"


def build_swe6_cases(requirement_data: dict[str, Any]) -> list[dict[str, str]]:
    cases: list[dict[str, str]] = []
    tc_no = 1
    for idx, req in enumerate(requirement_data.get("requirements") or [], start=1):
        if not isinstance(req, dict):
            continue
        srs_id = _text(req.get("srs_id")) or f"SRS_{idx:03d}"
        status = _testability(req)
        if status != "대상":
            continue
        requirement = _text(req.get("requirement"))
        pre = _first_nonempty(req.get("preconditions"), req.get("system_input_preconditions"))
        trigger = _text(req.get("activation_trigger"))
        behavior = _first_nonempty(req.get("behavior_flows"), req.get("processing_action"))
        expected = _first_nonempty(req.get("output"), req.get("acceptance_criteria"))
        prep_var, prep_cmp, prep_val = _parse_compare_value(pre)
        exec_source = trigger or behavior
        exec_var, exec_cmp, exec_val = _parse_compare_value(exec_source)
        exp_var, exp_cmp, exp_val = _parse_compare_value(expected or _text(req.get("acceptance_criteria")))
        case_name = _short_name(_first_nonempty(req.get("function_name"), requirement), 44) or f"{srs_id} 검증"
        desc_parts = [f"{srs_id} 요구사항을 검증한다."]
        if trigger:
            desc_parts.append(f"동작 조건은 '{_short_name(trigger, 80)}'이다.")
        if expected:
            desc_parts.append(f"기대 결과는 '{_short_name(expected, 100)}'이다.")
        cases.append({
            "srs_id": srs_id,
            "tc_id": f"TC_{tc_no:03d}",
            "method": _technique(req),
            "name": case_name,
            "description": " ".join(desc_parts),
            "prep_desc": pre,
            "prep_var": prep_var,
            "prep_compare": prep_cmp,
            "prep_value": prep_val,
            "exec_desc": exec_source,
            "exec_var": exec_var,
            "exec_compare": exec_cmp,
            "exec_value": exec_val,
            "expected_desc": expected,
            "expected_var": exp_var,
            "expected_compare": exp_cmp,
            "expected_value": exp_val,
        })
        tc_no += 1
    return cases


@dataclass
class _Style:
    dark: str = "D9E1F2"
    header: str = "D9D9D9"
    sub: str = "E2F0D9"
    accent: str = "70AD47"
    thin_color: str = "808080"


class SWE6Exporter:
    def __init__(self, output_dir: Path):
        self.output_dir = Path(output_dir).resolve()
        self.output_dir.mkdir(parents=True, exist_ok=True)
        self.style = _Style()

    def _filename(self, source_document: Path) -> Path:
        return self.output_dir / output_filename(
            "SWE.6 적격성 평가", _safe_stem(Path(source_document).name), "xlsx"
        )

    def _border(self):
        side = Side(style="thin", color=self.style.thin_color)
        return Border(left=side, right=side, top=side, bottom=side)

    def _cell_style(self, cell, *, fill=None, bold=False, center=False, font_size=10):
        cell.font = Font(name="Malgun Gothic", size=font_size, bold=bold)
        cell.alignment = Alignment(horizontal="center" if center else "left", vertical="center", wrap_text=True)
        cell.border = self._border()
        if fill:
            cell.fill = PatternFill("solid", fgColor=fill)

    def _merge_label(self, ws, rng: str, text: str, *, fill=None, bold=False, center=False, font_size=10):
        ws.merge_cells(rng)
        c = ws[rng.split(":")[0]]
        c.value = text
        self._cell_style(c, fill=fill, bold=bold, center=center, font_size=font_size)
        # Apply borders to whole merged range.
        for row in ws[rng]:
            for cell in row:
                cell.border = self._border()
                if fill:
                    cell.fill = PatternFill("solid", fgColor=fill)
        return c

    def export_excel(self, source_document: Path, requirement_data: dict[str, Any], *, source_text: str = "") -> Path:
        out = self._filename(source_document)
        metadata = extract_metadata(source_text)
        requirements = [x for x in (requirement_data.get("requirements") or []) if isinstance(x, dict)]
        cases = build_swe6_cases(requirement_data)
        today = datetime.now().strftime("%Y.%m.%d")

        wb = Workbook()
        cover = wb.active
        cover.title = "표지"
        history = wb.create_sheet("0_변경이력")
        summary = wb.create_sheet("1_테스트요약")
        tc = wb.create_sheet("2_테스트 케이스")

        # --- Cover ------------------------------------------------------
        cover.sheet_view.showGridLines = False
        for col, width in {"A":4,"B":12,"C":13,"D":13,"E":13,"F":13,"G":13,"H":16,"I":4}.items():
            cover.column_dimensions[col].width = width
        self._merge_label(cover,"B2:F4","SWE.6 소프트웨어 적격성 테스트",bold=True,center=True,font_size=13)
        for row, (label, value) in enumerate([
            ("산출물 ID", metadata["artifact_id"]),
            ("개정번호", "V.0.0"),
            ("개정일자", today),
        ], start=2):
            cover[f"G{row}"].value=label; self._cell_style(cover[f"G{row}"],fill=self.style.header,bold=True,center=True)
            cover[f"H{row}"].value=value; self._cell_style(cover[f"H{row}"],center=True)
        self._merge_label(cover,"B12:H14","소프트웨어 적격성 테스트",bold=True,center=True,font_size=20)
        self._merge_label(cover,"B17:H19","Software Qualification Test",bold=False,center=True,font_size=18)
        cover["C23"]="차종 / OEM"; self._cell_style(cover["C23"],fill=self.style.header,bold=True,center=True)
        self._merge_label(cover,"D23:F23",metadata["vehicle_oem"],center=False)
        cover["C25"]="프로젝트명"; self._cell_style(cover["C25"],fill=self.style.header,bold=True,center=True)
        self._merge_label(cover,"D25:F25",metadata["project_name"],center=False)
        self._merge_label(cover,"C27:C31","결재 정보",fill=self.style.header,bold=True,center=True)
        for col,label in zip(("D","E","F"),("작성","검토","승인")):
            cover[f"{col}27"]=label; self._cell_style(cover[f"{col}27"],fill=self.style.header,bold=True,center=True)
            self._merge_label(cover,f"{col}28:{col}31","",center=True)
        cover["C34"]="문서 상태:"; cover["C36"]="배포 날짜:"
        cover["D34"]="Draft"; cover["D36"]=""
        status_dv = DataValidation(type="list", formula1='"Draft,Released,Restricted,Expired"', allow_blank=False)
        cover.add_data_validation(status_dv)
        status_dv.add(cover["D34"])
        for c in (cover["C34"],cover["C36"]): c.font=Font(name="Malgun Gothic",size=11,bold=True)
        for c in (cover["D34"],cover["D36"]): self._cell_style(c,center=True)

        # --- History ----------------------------------------------------
        history.sheet_view.showGridLines=False
        history["A1"]="0. 문서 제/개정 이력"; history["A1"].font=Font(name="Malgun Gothic",size=14,bold=True)
        headers=["번호","개정 일자","개정 버전","개정 내용","개정자"]
        history.append([]); history.append(headers)
        for cell in history[3]: self._cell_style(cell,fill=self.style.header,bold=True,center=True)
        for no in range(1,11):
            vals=[no, today if no==1 else "", "V.0.0" if no==1 else "", "초안 작성" if no==1 else "", ""]
            history.append(vals)
            for cell in history[3+no]: self._cell_style(cell,center=(cell.column != 4))
        for col,width in zip("ABCDE",[9,15,15,70,18]): history.column_dimensions[col].width=width
        for row in range(4,14): history.row_dimensions[row].height=26

        # --- Summary ----------------------------------------------------
        summary.sheet_view.showGridLines=False
        summary["A1"]="1. 테스트 요약"; summary["A1"].font=Font(name="Malgun Gothic",size=14,bold=True)
        summary["A3"]="1. 평가 대상"; summary["A3"].font=Font(name="Malgun Gothic",size=11,bold=True)
        summary["A5"]="Project Name"; self._cell_style(summary["A5"],fill=self.style.header,bold=True,center=True)
        self._merge_label(summary,"B5:F5",metadata["project_name"])
        summary["A6"]="Test Level"; self._cell_style(summary["A6"],fill=self.style.header,bold=True,center=True)
        self._merge_label(summary,"B6:F6","소프트웨어 적격성 테스트")
        self._merge_label(summary,"A7:A9","Test Item",fill=self.style.header,bold=True,center=True)
        for rr,label,key in ((7,"Hardware","hardware"),(8,"Software","software"),(9,"Mechanical","mechanical")):
            summary[f"B{rr}"]=label; self._cell_style(summary[f"B{rr}"],fill=self.style.header,bold=True,center=True)
            self._merge_label(summary,f"C{rr}:F{rr}",metadata.get(key,""))
        summary["A10"]="입력물"; self._cell_style(summary["A10"],fill=self.style.header,bold=True,center=True)
        self._merge_label(summary,"B10:F10","소프트웨어 요구사항 명세(SWE.1)")
        self._merge_label(
            summary,"A11:F12",
            "※ 본 문서는 자동 생성된 SWE.6 테스트 케이스 초안입니다. 시험 환경, 외부 DB, 파라미터, 임계값, Variant 및 실행 결과는 시험 담당자 검토 후 확정되어야 합니다. Source에 없는 시험값은 임의 생성하지 않습니다.",
            center=False,font_size=9
        )
        summary["A13"]="2. 테스트 환경"; summary["A13"].font=Font(name="Malgun Gothic",size=11,bold=True)
        environment_text = metadata.get("environment", "")
        if not environment_text:
            environment_text = "테스트 환경 구성도 또는 대표 사진이 필요합니다.\n입력문서에서 Test Environment 관련 시각자료를 확인할 수 없습니다."
        self._merge_label(summary,"A15:F22",environment_text,center=True,font_size=11)
        summary["A24"]="3. 기능 요구사항"; summary["A24"].font=Font(name="Malgun Gothic",size=11,bold=True)

        def add_req_table(start_row:int, wanted_category:str):
            hdr=["번호","분류 1","분류 2","분류 3","식별자","테스트"]
            for col,val in enumerate(hdr,1):
                c=summary.cell(start_row,col,val); self._cell_style(c,fill=self.style.header,bold=True,center=True)
            rr=start_row+1; no=1
            for req in requirements:
                cat=_text(req.get("category"))
                normalized="비기능" if "비기능" in cat else ("기능" if "기능" in cat else REVIEW_NEEDED)
                if normalized != wanted_category: continue
                req_text=_text(req.get("requirement"))
                vals=[no,_text(req.get("function_name")),"",_short_name(req_text,54),_text(req.get("srs_id")),_testability(req)]
                for col,val in enumerate(vals,1):
                    c=summary.cell(rr,col,val); self._cell_style(c,center=(col in (1,5,6)))
                rr+=1; no+=1
            if rr==start_row+1:
                vals=[1,"","","등록된 요구사항 없음","","검토 필요"]
                for col,val in enumerate(vals,1):
                    c=summary.cell(rr,col,val); self._cell_style(c,center=(col in (1,5,6)))
                rr+=1
            return rr

        next_row=add_req_table(26,"기능")
        nr=max(next_row+2,36)
        summary[f"A{nr}"]="4. 비기능 요구사항"; summary[f"A{nr}"].font=Font(name="Malgun Gothic",size=11,bold=True)
        add_req_table(nr+2,"비기능")
        for col,width in zip("ABCDEF",[9,24,24,46,15,14]): summary.column_dimensions[col].width=width
        summary.freeze_panes="A5"

        # --- Test Cases -------------------------------------------------
        tc.sheet_view.showGridLines=False
        tc["A1"]="2. 테스트 케이스"; tc["A1"].font=Font(name="Malgun Gothic",size=14,bold=True)
        # Row 2 groups the detailed headers below into practical test blocks.
        group_headers = [
            ("A2:E2", "Test ID"),
            ("F2:I2", "Test Input"),
            ("J2:M2", "Test Execution"),
            ("N2:Q2", "Test Expect Result"),
            ("R2:S2", "Test Result"),
            ("T2:U2", "Etc"),
        ]
        for rng, label in group_headers:
            self._merge_label(tc, rng, label, fill=self.style.dark, bold=True, center=True, font_size=9)
        tc.row_dimensions[2].height = 23

        headers=[
            "SW 요구사항 ID","Test Case ID","Methods for Testing /\nMethod for deriving TC",
            "Test case name","Test case description",
            "Test preparation\nDescription","Variable","Compare","Value",
            "Test execution\nDescription","Variable","Compare","Value",
            "Expected Result\nDescription","Variable","Compare","Value",
            "Output Value","PASS / FAIL","Comment","Capture\nCANoe (Optional)"
        ]
        for col,val in enumerate(headers,1):
            c=tc.cell(3,col,val); self._cell_style(c,fill=self.style.header,bold=True,center=True,font_size=9)
        tc.row_dimensions[3].height = 38

        def na_triplet(description: str, variable: str, compare: str, value: str) -> tuple[str, str, str]:
            """Use N/A only for unavailable derived V/C/V values, never for intentionally blank result fields."""
            if not _text(description):
                return variable or "", compare or "", value or ""
            return variable or "N/A", compare or "N/A", value or "N/A"

        r=4
        if not cases:
            cases=[{
                "srs_id":"","tc_id":"","method":"","name":"검토 필요","description":"SWE.6 Test Case 자동 생성에 필요한 검증 가능 Requirement가 없습니다.",
                "prep_desc":"","prep_var":"","prep_compare":"","prep_value":"","exec_desc":"","exec_var":"","exec_compare":"","exec_value":"",
                "expected_desc":"","expected_var":"","expected_compare":"","expected_value":""
            }]
        for case in cases:
            prep_var, prep_cmp, prep_val = na_triplet(case["prep_desc"], case["prep_var"], case["prep_compare"], case["prep_value"])
            exec_var, exec_cmp, exec_val = na_triplet(case["exec_desc"], case["exec_var"], case["exec_compare"], case["exec_value"])
            exp_var, exp_cmp, exp_val = na_triplet(case["expected_desc"], case["expected_var"], case["expected_compare"], case["expected_value"])
            vals=[
                case["srs_id"],case["tc_id"],case["method"],case["name"],case["description"],
                case["prep_desc"],prep_var,prep_cmp,prep_val,
                case["exec_desc"],exec_var,exec_cmp,exec_val,
                case["expected_desc"],exp_var,exp_cmp,exp_val,
                "","","","",
            ]
            for col,val in enumerate(vals,1):
                c=tc.cell(r,col,val); self._cell_style(c,center=(col in (1,2,7,8,9,11,12,13,15,16,17,18,19)))
            r+=1

        widths=[16,14,28,28,46,42,20,11,16,42,20,11,16,42,20,11,16,18,14,30,22]
        for idx,width in enumerate(widths,1):
            tc.column_dimensions[chr(64+idx)].width=width

        # If a Variable / Compare / Value column contains only N/A (ignoring deliberate blanks),
        # keep it narrow so empty metadata does not dominate the worksheet.
        for idx in (7,8,9,11,12,13,15,16,17):
            observed=[_text(tc.cell(row_idx, idx).value) for row_idx in range(4, r)]
            nonblank=[value for value in observed if value]
            if nonblank and all(value == "N/A" for value in nonblank):
                tc.column_dimensions[chr(64+idx)].width=8

        tc.freeze_panes="A4"
        tc.auto_filter.ref=f"A3:U{max(4,r-1)}"
        # Editable PASS/FAIL list, left blank on generation.
        dv=DataValidation(type="list", formula1='"PASS,FAIL,N/A"', allow_blank=True)
        tc.add_data_validation(dv); dv.add(f"S4:S{max(100,r+20)}")

        # Common formatting / print setup.
        for ws in (cover,history,summary,tc):
            ws.sheet_properties.pageSetUpPr.fitToPage=True
            ws.page_setup.fitToWidth=1
            ws.page_setup.fitToHeight=0
            ws.page_margins.left=0.25; ws.page_margins.right=0.25; ws.page_margins.top=0.4; ws.page_margins.bottom=0.4
        cover.page_setup.orientation="portrait"
        history.page_setup.orientation="landscape"
        summary.page_setup.orientation="landscape"
        tc.page_setup.orientation="landscape"

        wb.save(out)
        return out
