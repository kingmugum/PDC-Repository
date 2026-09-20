from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


class AnalysisResultExporter:
    def __init__(self, output_dir: Path):
        self.output_dir = Path(output_dir).resolve()
        self.output_dir.mkdir(parents=True, exist_ok=True)

    @staticmethod
    def _safe_stem(name: str) -> str:
        stem = Path(name).stem or "document"
        stem = re.sub(r'[<>:"/\\|?*]+', "_", stem)
        stem = re.sub(r"\s+", "_", stem).strip("._ ")
        return stem[:80] or "document"

    def export_docx(
        self,
        source_document: Path,
        analysis_text: str,
        *,
        model: str,
        api_base: str,
    ) -> Path:
        text = (analysis_text or "").strip()
        if not text:
            raise ValueError("저장할 문서 분석 결과가 없습니다.")

        source_document = Path(source_document)
        source_type = re.sub(r"[^A-Za-z0-9]+", "", source_document.suffix.lstrip(".")) or "doc"
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
        filename = (
            f"RequirementStudio_Analysis_{self._safe_stem(source_document.name)}_{source_type}_{stamp}.docx"
        )
        out_path = self.output_dir / filename

        doc = Document()

        title = doc.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = title.add_run("Requirement Studio 문서 분석 결과")
        run.bold = True
        run.font.size = Pt(18)

        meta = doc.add_table(rows=4, cols=2)
        meta.style = "Table Grid"
        meta.cell(0, 0).text = "원본 문서"
        meta.cell(0, 1).text = source_document.name
        meta.cell(1, 0).text = "생성 일시"
        meta.cell(1, 1).text = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        meta.cell(2, 0).text = "Model"
        meta.cell(2, 1).text = model
        meta.cell(3, 0).text = "API Base"
        meta.cell(3, 1).text = api_base

        doc.add_paragraph()

        for raw_line in text.splitlines():
            line = raw_line.strip()

            if not line:
                doc.add_paragraph()
                continue

            heading = re.match(r"^(?:#{1,4}\s*)?(\d+[\.\)]\s+.+)$", line)
            if heading:
                paragraph = doc.add_paragraph()
                run = paragraph.add_run(heading.group(1))
                run.bold = True
                run.font.size = Pt(12)
                continue

            if line.startswith("#"):
                paragraph = doc.add_paragraph()
                run = paragraph.add_run(line.lstrip("#").strip())
                run.bold = True
                run.font.size = Pt(12)
                continue

            doc.add_paragraph(line)

        doc.save(out_path)
        return out_path
