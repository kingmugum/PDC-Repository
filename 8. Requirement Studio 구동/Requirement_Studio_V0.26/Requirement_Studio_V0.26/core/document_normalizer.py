from __future__ import annotations

import hashlib
import json
import re
import zipfile
from dataclasses import asdict, dataclass
from datetime import datetime
from pathlib import Path
from typing import Iterable
from xml.etree import ElementTree as ET

from docx import Document
from pypdf import PdfReader
from pptx import Presentation


@dataclass
class DocumentChunk:
    chunk_id: str
    kind: str
    location: str
    text: str


@dataclass
class NormalizedDocument:
    source_document: str
    source_path: str
    source_sha256: str
    source_type: str
    chunks: list[DocumentChunk]
    visual_assets: list[dict]
    total_chars: int
    normalized_at: str
    ir_json_path: str
    ir_jsonl_path: str

    def compact_text(self) -> str:
        lines = [
            f"[DOCUMENT] {self.source_document}",
            f"[TYPE] {self.source_type}",
            f"[SOURCE_SHA256] {self.source_sha256}",
            f"[VISUAL_ASSET_COUNT] {len(self.visual_assets)}",
        ]
        if self.visual_assets:
            lines.append(
                "[VISUAL_NOTICE] 문서에 그림/미디어가 존재하지만 v0.16 Text Provider 경로에서는 "
                "이미지 본문을 전송하지 않습니다. 그림을 참조하는 요구사항은 추정하지 말고 Gap/TBD로 남기세요."
            )
        for chunk in self.chunks:
            lines.append(f"\n[SRC {chunk.chunk_id} | {chunk.location} | {chunk.kind}]")
            lines.append(chunk.text)
        return "\n".join(lines).strip()


class DocumentNormalizer:
    """Provider와 무관하게 Source Document를 로컬 Canonical Document IR로 변환한다."""

    def __init__(self, project_root: Path):
        self.project_root = Path(project_root).resolve()
        self.work_root = self.project_root / "work" / "normalized"
        self.work_root.mkdir(parents=True, exist_ok=True)

    @staticmethod
    def _sha256(path: Path) -> str:
        h = hashlib.sha256()
        with path.open("rb") as f:
            for block in iter(lambda: f.read(1024 * 1024), b""):
                h.update(block)
        return h.hexdigest()

    @staticmethod
    def _clean(text: str) -> str:
        text = (text or "").replace("\r\n", "\n").replace("\r", "\n")
        text = re.sub(r"[ \t]+", " ", text)
        text = re.sub(r"\n{3,}", "\n\n", text)
        return text.strip()

    @staticmethod
    def _zip_media(path: Path, prefixes: Iterable[str]) -> list[dict]:
        out: list[dict] = []
        try:
            with zipfile.ZipFile(path) as zf:
                for name in zf.namelist():
                    if any(name.startswith(prefix) for prefix in prefixes) and not name.endswith("/"):
                        out.append({"kind": "embedded_media", "source": name})
        except Exception:
            pass
        return out

    def _pdf(self, path: Path) -> tuple[list[DocumentChunk], list[dict]]:
        reader = PdfReader(str(path))
        chunks: list[DocumentChunk] = []
        visuals: list[dict] = []
        seq = 1
        for page_no, page in enumerate(reader.pages, start=1):
            text = self._clean(page.extract_text() or "")
            if text:
                chunks.append(DocumentChunk(f"DOC-C{seq:04d}", "text", f"Page {page_no}", text))
                seq += 1
            try:
                images = list(page.images)
                for idx, image in enumerate(images, start=1):
                    visuals.append({"kind": "pdf_image", "source": f"Page {page_no} / Image {idx}"})
            except Exception:
                pass
        return chunks, visuals

    def _docx(self, path: Path) -> tuple[list[DocumentChunk], list[dict]]:
        doc = Document(str(path))
        chunks: list[DocumentChunk] = []
        seq = 1
        for idx, paragraph in enumerate(doc.paragraphs, start=1):
            text = self._clean(paragraph.text)
            if text:
                chunks.append(DocumentChunk(f"DOC-C{seq:04d}", "text", f"Paragraph {idx}", text))
                seq += 1
        for t_idx, table in enumerate(doc.tables, start=1):
            for r_idx, row in enumerate(table.rows, start=1):
                cells = [self._clean(cell.text) for cell in row.cells]
                if any(cells):
                    text = " | ".join(cells)
                    chunks.append(DocumentChunk(f"DOC-C{seq:04d}", "table", f"Table {t_idx} / Row {r_idx}", text))
                    seq += 1
        visuals = self._zip_media(path, ("word/media/",))
        return chunks, visuals

    def _pptx(self, path: Path) -> tuple[list[DocumentChunk], list[dict]]:
        prs = Presentation(str(path))
        chunks: list[DocumentChunk] = []
        seq = 1
        for slide_no, slide in enumerate(prs.slides, start=1):
            texts: list[str] = []
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    t = self._clean(getattr(shape, "text", ""))
                    if t:
                        texts.append(t)
                if getattr(shape, "has_table", False):
                    for row in shape.table.rows:
                        cells = [self._clean(cell.text) for cell in row.cells]
                        if any(cells):
                            texts.append(" | ".join(cells))
            merged = self._clean("\n".join(texts))
            if merged:
                chunks.append(DocumentChunk(f"DOC-C{seq:04d}", "slide", f"Slide {slide_no}", merged))
                seq += 1
        visuals = self._zip_media(path, ("ppt/media/",))
        return chunks, visuals

    @staticmethod
    def _xlsx_col_name(cell_ref: str) -> str:
        m = re.match(r"([A-Z]+)", cell_ref or "")
        return m.group(1) if m else ""

    def _xlsx(self, path: Path) -> tuple[list[DocumentChunk], list[dict]]:
        NS = {
            "main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
            "rel": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
            "pkg": "http://schemas.openxmlformats.org/package/2006/relationships",
        }
        chunks: list[DocumentChunk] = []
        seq = 1
        with zipfile.ZipFile(path) as zf:
            shared: list[str] = []
            if "xl/sharedStrings.xml" in zf.namelist():
                root = ET.fromstring(zf.read("xl/sharedStrings.xml"))
                for si in root.findall("main:si", NS):
                    texts = [t.text or "" for t in si.findall(".//main:t", NS)]
                    shared.append(self._clean("".join(texts)))

            wb_root = ET.fromstring(zf.read("xl/workbook.xml"))
            rel_root = ET.fromstring(zf.read("xl/_rels/workbook.xml.rels"))
            rel_map = {
                r.attrib.get("Id"): r.attrib.get("Target", "")
                for r in rel_root.findall("pkg:Relationship", NS)
            }

            for sheet in wb_root.findall("main:sheets/main:sheet", NS):
                sheet_name = sheet.attrib.get("name", "Sheet")
                rel_id = sheet.attrib.get(f"{{{NS['rel']}}}id")
                target = rel_map.get(rel_id, "")
                if not target:
                    continue
                cleaned_target = target.lstrip("/")
                if cleaned_target.startswith("xl/"):
                    sheet_xml = cleaned_target
                else:
                    sheet_xml = "xl/" + cleaned_target
                if sheet_xml not in zf.namelist():
                    continue
                root = ET.fromstring(zf.read(sheet_xml))
                for row in root.findall(".//main:sheetData/main:row", NS):
                    row_no = row.attrib.get("r", "?")
                    values: list[str] = []
                    for cell in row.findall("main:c", NS):
                        ref = cell.attrib.get("r", "")
                        cell_type = cell.attrib.get("t")
                        value = ""
                        if cell_type == "inlineStr":
                            value = "".join(t.text or "" for t in cell.findall(".//main:t", NS))
                        else:
                            v = cell.find("main:v", NS)
                            raw = v.text if v is not None and v.text is not None else ""
                            if cell_type == "s" and raw.isdigit():
                                idx = int(raw)
                                value = shared[idx] if 0 <= idx < len(shared) else raw
                            else:
                                value = raw
                        value = self._clean(value)
                        if value:
                            values.append(f"{ref}={value}")
                    if values:
                        chunks.append(
                            DocumentChunk(
                                f"DOC-C{seq:04d}",
                                "sheet_row",
                                f"Sheet {sheet_name} / Row {row_no}",
                                " | ".join(values),
                            )
                        )
                        seq += 1
        visuals = self._zip_media(path, ("xl/media/",))
        return chunks, visuals

    def normalize(self, source_path: Path) -> NormalizedDocument:
        source = Path(source_path).resolve()
        if not source.is_file():
            raise FileNotFoundError(source)
        suffix = source.suffix.lower()
        if suffix == ".pdf":
            chunks, visuals = self._pdf(source)
        elif suffix == ".docx":
            chunks, visuals = self._docx(source)
        elif suffix == ".pptx":
            chunks, visuals = self._pptx(source)
        elif suffix in {".xlsx", ".xlsm"}:
            chunks, visuals = self._xlsx(source)
        else:
            raise ValueError(f"지원하지 않는 문서 형식: {suffix}")

        if not chunks:
            raise RuntimeError("문서에서 전송 가능한 텍스트/표 내용을 추출하지 못했습니다.")

        sha = self._sha256(source)
        work_dir = self.work_root / f"{source.stem}_{sha[:8]}"
        work_dir.mkdir(parents=True, exist_ok=True)
        ir_json = work_dir / "document_ir.json"
        ir_jsonl = work_dir / "document_ir.jsonl"
        total_chars = sum(len(c.text) for c in chunks)

        result = NormalizedDocument(
            source_document=source.name,
            source_path=str(source),
            source_sha256=sha,
            source_type=suffix.lstrip(".").upper(),
            chunks=chunks,
            visual_assets=visuals,
            total_chars=total_chars,
            normalized_at=datetime.now().isoformat(timespec="seconds"),
            ir_json_path=str(ir_json),
            ir_jsonl_path=str(ir_jsonl),
        )
        payload = asdict(result)
        ir_json.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
        with ir_jsonl.open("w", encoding="utf-8") as f:
            for chunk in chunks:
                f.write(json.dumps(asdict(chunk), ensure_ascii=False) + "\n")
        return result
