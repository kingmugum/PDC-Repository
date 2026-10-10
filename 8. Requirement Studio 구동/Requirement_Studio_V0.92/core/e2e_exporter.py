from __future__ import annotations

import re
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from core.output_naming import output_filename
from core.swe1_exporter import build_engineering_records
from core.verification_single_truth import finalize_verification_single_truth

NOT_CONFIRMED = "입력문서에서 확인되지 않음"


def _list(value: Any) -> list[Any]:
    if value is None:
        return []
    return value if isinstance(value, list) else [value]


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, dict):
        return "\n".join(f"{k}: {_text(v)}" for k, v in value.items() if v not in (None, "", [], {}))
    if isinstance(value, (list, tuple, set)):
        return "\n".join(x for x in (_text(v).strip() for v in value) if x)
    return str(value).strip()


def _safe_stem(name: str) -> str:
    stem = Path(name).stem or "document"
    stem = re.sub(r'[<>:"/\\|?*]+', "_", stem)
    stem = re.sub(r"\s+", "_", stem).strip("._ ")
    return stem[:80] or "document"


def _source_trace(req: dict[str, Any]) -> tuple[str, str, str]:
    sems = [str(x) for x in _list(req.get("source_semantic_unit_ids")) if str(x)]
    fids: list[str] = []
    locs: list[str] = []
    for key in ("source_backed_atomic_behaviors", "fact_level_allocations", "source_fact_fragments"):
        for item in _list(req.get(key)):
            if not isinstance(item, dict):
                continue
            fid = str(item.get("source_fact_fragment_id") or "")
            if fid and fid not in fids:
                fids.append(fid)
            sem = str(item.get("source_semantic_unit_id") or item.get("parent_source_semantic_unit_id") or "")
            if sem and sem not in sems:
                sems.append(sem)
            loc = _text(item.get("source_location"))
            if loc and loc not in locs:
                locs.append(loc)
    for ev in _list(req.get("source_evidence")):
        if isinstance(ev, dict):
            loc = _text(ev.get("location"))
            if loc and loc not in locs:
                locs.append(loc)
    return ", ".join(sems), ", ".join(fids), " | ".join(locs)


def _source_evidence(req: dict[str, Any]) -> str:
    lines=[]
    for ev in _list(req.get("source_evidence")):
        if isinstance(ev, dict):
            loc=_text(ev.get("location")); txt=_text(ev.get("text")); doc=_text(ev.get("document"))
            prefix=" / ".join(x for x in (doc,loc) if x)
            val=f"{prefix}: {txt}" if prefix and txt else (prefix or txt)
        else:
            val=_text(ev)
        if val and val not in lines:
            lines.append(val)
    return "\n".join(lines) or NOT_CONFIRMED


def _scope_label(req: dict[str, Any]) -> str:
    sys1=str(req.get("sys1_eligibility") or "")
    swe1=str(req.get("swe1_eligibility") or "")
    if sys1 in {"Eligible","Review Needed"} and swe1 in {"Eligible","Review Needed"}:
        return "SYS.1 + SWE.1"
    if swe1 in {"Eligible","Review Needed"}:
        return "SWE.1"
    if sys1 in {"Eligible","Review Needed"}:
        return "SYS.1"
    return "Engineering Native / Review"


OBJECT_TYPE_KO = {
    "SYS5_CANDIDATE": "SYS.5 평가 후보",
    "SYS5_PARENT_SCOPE_REVIEW": "SYS.5 상위 범위 검토",
    "SWE6_TC": "SWE.6 테스트 케이스",
    "DEFERRED_INTENT": "보류 검증 의도",
    "ALLOCATION_PENDING_INTENT": "할당 검토 대기",
    "EXTERNAL_DEPENDENCY_INTENT": "외부 의존 검토",
    "ELECTRICAL_VERIFICATION_INTENT": "전기 검증",
    "ENVIRONMENTAL_VERIFICATION_INTENT": "환경 검증",
    "MECHANICAL_VERIFICATION_INTENT": "기구/HW 검증",
    "MANUFACTURING_VERIFICATION_INTENT": "제조 검증",
    "REVIEW_REQUIRED": "사람 검토 필요",
}

LOGIC_KO = {
    "ALL_OF": "모두 만족(AND)",
    "ANY_OF": "하나 이상 만족(OR)",
    "SEQUENCE": "순차 조건",
}

TRANSITION_KO = {
    "power_mode_transition": "전원 모드 전이",
    "capture_mode_transition": "캡처 모드 전이",
    "mirror_fold_transition": "미러 폴딩 전이",
    "ignition_transition": "시동/점화 전이",
    "sleep_wake_transition": "슬립/웨이크 전이",
    "wake_init_transition": "웨이크 초기화 전이",
    "communication_recovery": "통신 복구",
    "environmental_qualification": "환경 적격성",
}

ALLOCATION_KO = {
    "SW_BEHAVIOR_INFERRED_REVIEW_REQUIRED": "SW 동작 추론 - 사람 검토 필요",
    "SYSTEM_REQUIREMENT_PENDING_SW_ALLOCATION": "시스템 요구사항 - SW 할당 검토 대기",
    "SW_IMPLEMENTATION_REQUIREMENT": "SW 구현 요구사항",
    "ELECTRICAL_REQUIREMENT": "전기 요구사항",
    "ENVIRONMENTAL_QUALIFICATION_REQUIREMENT": "환경 적격성 요구사항",
    "MECHANICAL_CONNECTOR_REQUIREMENT": "기구/커넥터 요구사항",
    "MANUFACTURING_PROCESS_REQUIREMENT": "제조 공정 요구사항",
    "EXTERNAL_STANDARD_REFERENCE": "외부 표준 참조",
}

def _display_enum(value: Any, mapping: dict[str, str]) -> str:
    raw = str(value or "")
    return mapping.get(raw, raw)

def _display_eligibility(value: Any) -> str:
    raw = str(value or "")
    return {
        "Eligible": "대상",
        "Review Needed": "검토 필요",
        "Not Applicable": "비대상",
        "Deferred pending SW allocation": "SW 할당 검토 대기",
    }.get(raw, raw)

def _display_readiness(value: Any) -> str:
    raw = str(value or "")
    return {
        "REVIEW_REQUIRED": "검토 필요",
        "PARTIALLY_TESTABLE": "부분 평가 가능",
        "INTENT_DRAFT": "평가 의도 초안",
        "READY": "실행 가능",
    }.get(raw, raw)


REVIEW_STATUS_KO = {
    "HUMAN_DECISION_REQUIRED": "사람의 리뷰 필요",
    "REVIEW_REQUIRED": "사람의 리뷰 필요",
    "HOLD": "사람의 리뷰 필요",
    "SOURCE_BACKED": "리뷰 불필요",
    "APPROVED": "리뷰 완료",
}

REASON_KO = {
    "ALLOCATION_PENDING": "SW 할당 검토 대기",
    "SOURCE_INSUFFICIENT": "입력 문서 정보 부족",
    "EXPORTER_CAPABILITY_PENDING": "자동 구조화 검토 필요",
    "HUMAN_REVIEW_REQUIRED": "사람의 리뷰 필요",
    "STATE_TRANSITION_LOGIC_REVIEW_REQUIRED": "상태 전이 조건 논리 검토 필요",
    "STATE_TRANSITION_TARGET_REVIEW_REQUIRED": "상태 전이 목표 상태 검토 필요",
    "SW_BEHAVIOR_INFERRED_REVIEW_REQUIRED": "SW 동작 추론 결과에 대한 사람 검토 필요",
    "SYSTEM_REQUIREMENT_PENDING_SW_ALLOCATION": "시스템 요구사항의 SW 할당 검토 필요",
}


def _display_review_status(req: dict[str, Any], raw: Any = "") -> str:
    if req.get("human_decision_required"):
        return "사람의 리뷰 필요"
    value = str(raw or req.get("review_status") or "").strip()
    if value in REVIEW_STATUS_KO:
        return REVIEW_STATUS_KO[value]
    upper = value.upper()
    if any(token in upper for token in ("REVIEW", "HOLD", "PENDING", "REQUIRED")):
        return "사람의 리뷰 필요"
    return "리뷰 불필요"


def _format_human_multiline(value: Any) -> str:
    """Presentation-only line breaking. Never changes the Source facts stored in JSON/audit."""
    text = _text(value).replace("\r\n", "\n").replace("\r", "\n").strip()
    if not text:
        return ""
    # Display-only Korean wording; raw Source/JSON values are untouched.
    text = text.replace("Source-defined", "사양에 정의된").replace("Source-backed Atomic Behavior", "출처 기반 세부 동작")
    # Preserve existing lines while breaking common source-table/numbered patterns.
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\s*•\s*", "\n- ", text)
    text = re.sub(r"(?<!\n)\s+(?=(?:\d+(?:\.\d+)*[\)\.]\s+))", "\n", text)
    text = re.sub(r"(?<!\n)\s+(?=0x[0-9A-Fa-f]+\s*[:·-])", "\n", text)
    text = re.sub(r"(?<!\n)\s+(?=[A-Za-z][A-Za-z0-9_]*\s*값\s*송신\s*조건)", "\n", text)
    text = re.sub(r"(?<=다\.)\s+(?=[가-힣A-Za-z0-9])", "\n", text)
    lines=[]
    for line in text.splitlines():
        line=re.sub(r"\s+", " ", line).strip()
        if not line:
            continue
        if line.startswith("•"):
            line="- "+line[1:].strip()
        lines.append(line)
    # Avoid orphan numbering lines such as `3)` followed by the actual item on the next line.
    merged=[]; i=0
    while i < len(lines):
        if re.fullmatch(r"\d+(?:\.\d+)*[\)\.]", lines[i]) and i+1 < len(lines):
            merged.append(lines[i]+" "+lines[i+1]); i += 2
        else:
            merged.append(lines[i]); i += 1
    return "\n".join(merged)


def _display_domain_compact(case: dict[str, Any]) -> str:
    typ=str(case.get("test_object_type") or "")
    domain=_text(case.get("verification_domain")).upper()
    if "SWE6" in typ or "SWE.6" in domain:
        return "SWE.6"
    if "SYS5" in typ or "SYS.5" in domain:
        return "SYS.5"
    return "검토"


def _display_compare(value: Any) -> str:
    raw=str(value or "").strip()
    return {
        "Source-defined": "사양 정의값",
        "source-defined": "사양 정의값",
        "SOURCE_DEFINED": "사양 정의값",
        "transition_to": "상태 전이",
        "순차 입력": "순차",
        "ALL_OF": "모두 만족",
        "ANY_OF": "하나 이상 만족",
    }.get(raw, LOGIC_KO.get(raw, raw))


def _compact_meaning(value: Any) -> str:
    text=re.sub(r"\s+", " ", str(value or "")).strip()
    if not text:
        return ""
    # Remove appended long requirement prose from enum meanings; retain the concise semantic label.
    for sep in (" - ", " – ", " — "):
        if sep in text:
            left,right=text.split(sep,1)
            if len(right)>36 or "HU는" in right or "해야" in right:
                text=left.strip()
                break
    if len(text)>72:
        m=re.match(r"(.{1,72}?)(?:[.!?]|$)", text)
        text=(m.group(1) if m else text[:72]).strip()
    return text


def _compact_input_value(case: dict[str, Any]) -> str:
    vectors=[v for v in _list(case.get("test_vectors")) if isinstance(v,dict)]
    vals=[]
    for i,v in enumerate(vectors,1):
        val=str(v.get("input_value") or "").strip()
        if val:
            vals.append(f"{i}) {val}" if len(vectors)>1 else val)
    if vals:
        return "\n".join(vals)
    raw=_format_human_multiline(case.get("input_value") or case.get("value"))
    cleaned=[]
    for line in raw.splitlines():
        m=re.match(r"(\d+\))\s*([^·]+)", line)
        cleaned.append(f"{m.group(1)} {m.group(2).strip()}" if m else line)
    return "\n".join(cleaned)


def _compact_output_value(case: dict[str, Any]) -> str:
    vectors=[v for v in _list(case.get("test_vectors")) if isinstance(v,dict)]
    vals=[]
    for i,v in enumerate(vectors,1):
        val=str(v.get("expected_output") or "").strip()
        if not val:
            continue
        meaning=_compact_meaning(v.get("output_meaning"))
        item=f"{val} · {meaning}" if meaning else val
        vals.append(f"{i}) {item}" if len(vectors)>1 else item)
    if vals:
        return "\n".join(vals)
    raw=_format_human_multiline(case.get("output_value"))
    out=[]
    for line in raw.splitlines():
        line=re.sub(r"(·\s*[^-\n]{0,60})\s+-\s+HU는.*$", r"\1", line)
        out.append(line)
    return "\n".join(out)


def _display_review_reason(case: dict[str, Any]) -> str:
    parts=[]
    for value in (case.get("deferred_reason_code"), case.get("deferred_reason_detail"), case.get("required_resolution")):
        txt=str(value or "").strip()
        if not txt:
            continue
        txt=REASON_KO.get(txt, txt)
        for k,v in REASON_KO.items():
            txt=txt.replace(k,v)
        if txt not in parts:
            parts.append(txt)
    if not parts and _test_case_review_status(case)=="사람의 리뷰 필요":
        alloc=str(case.get("allocation_status") or "")
        parts.append(REASON_KO.get(alloc, "사람 검토가 필요한 평가 케이스입니다."))
    return _format_human_multiline("\n".join(parts))


def _test_case_review_status(case: dict[str, Any]) -> str:
    if case.get("human_review_required"):
        return "사람의 리뷰 필요"
    readiness=str(case.get("execution_readiness") or "").upper()
    alloc=str(case.get("allocation_status") or "").upper()
    if any(x in readiness for x in ("REVIEW", "DRAFT", "PARTIAL", "HOLD")):
        return "사람의 리뷰 필요"
    if any(x in alloc for x in ("REVIEW", "PENDING", "HOLD")):
        return "사람의 리뷰 필요"
    return "실행 가능"


class E2EExporter:
    """V0.92 reviewer-facing E2E Requirements + structured Basic Functional Evaluation specifications.

    E2E Requirements is the Source-nearest engineering view (SYS.1 + SWE.1 + native/review
    preservation).  E2E Evaluation is the human-facing verification master (SYS.5 + SWE.6 +
    deferred/external/native) while raw objects remain available in the finalized bundle for audit.
    """

    DARK = "1F4E78"
    HEADER = "D9EAF7"
    SUB = "EAF2F8"
    REVIEW = "FFF2CC"
    FAIL = "F4CCCC"
    PASS = "D9EAD3"
    THIN = Side(style="thin", color="B7B7B7")

    def __init__(self, output_dir: Path):
        self.output_dir=Path(output_dir).resolve(); self.output_dir.mkdir(parents=True, exist_ok=True)

    def _filename(self, prefix: str, source: Path, ext: str) -> Path:
        return self.output_dir / output_filename(prefix, _safe_stem(source.name), ext)

    def _style_header(self, cell, fill: str | None = None):
        cell.fill=PatternFill("solid", fgColor=fill or self.DARK)
        cell.font=Font(name="Malgun Gothic", size=9, bold=True, color="FFFFFF" if (fill or self.DARK)==self.DARK else "000000")
        cell.alignment=Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border=Border(left=self.THIN,right=self.THIN,top=self.THIN,bottom=self.THIN)

    def _style_cell(self, cell, *, center=False, fill=None):
        cell.font=Font(name="Malgun Gothic", size=9)
        cell.alignment=Alignment(horizontal="center" if center else "left", vertical="top", wrap_text=True)
        cell.border=Border(left=self.THIN,right=self.THIN,top=self.THIN,bottom=self.THIN)
        if fill: cell.fill=PatternFill("solid", fgColor=fill)

    def _write_table(self, ws, headers, rows, widths, start_row=3):
        for c,h in enumerate(headers,1): self._style_header(ws.cell(start_row,c,h), self.HEADER)
        for r,row in enumerate(rows or [["-"]+[""]*(len(headers)-1)], start_row+1):
            for c,val in enumerate(row,1): self._style_cell(ws.cell(r,c,val), center=c<=3)
        ws.freeze_panes=ws.cell(start_row+1,1)
        ws.auto_filter.ref=f"A{start_row}:{get_column_letter(len(headers))}{start_row+max(1,len(rows))}"
        for i,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width=w

    def _write_group_headers(self, ws):
        groups = [
            (1,6,"종류"),
            (7,8,"사전 준비"),
            (9,12,"입력 / 수행"),
            (13,16,"출력"),
            (17,19,"판단"),
            (20,23,"실행 결과"),
        ]
        for start_col, end_col, label in groups:
            ws.merge_cells(start_row=2,start_column=start_col,end_row=2,end_column=end_col)
            cell=ws.cell(2,start_col,label)
            self._style_header(cell, self.DARK)
            cell.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
        ws.row_dimensions[2].height=23

    def _cover(self, ws, title: str, source: Path, subtitle: str):
        ws.sheet_view.showGridLines=False
        ws.merge_cells("B3:G4"); ws["B3"]=title
        ws["B3"].font=Font(name="Malgun Gothic", size=22, bold=True, color=self.DARK); ws["B3"].alignment=Alignment(horizontal="center",vertical="center")
        ws.merge_cells("B6:G6"); ws["B6"]=subtitle; ws["B6"].font=Font(name="Malgun Gothic", size=11, bold=True, color="595959"); ws["B6"].alignment=Alignment(horizontal="center")
        rows=[("원본 문서",source.name),("생성 일시",datetime.now().strftime("%Y-%m-%d %H:%M:%S")),("Requirement Studio","V0.92"),("Source Fidelity","Source에 없는 값/신호/Timing/Threshold/Test Result를 생성하지 않음")]
        for r,(k,v) in enumerate(rows,9):
            ws.cell(r,2,k); ws.cell(r,3,v); self._style_cell(ws.cell(r,2),center=True,fill=self.SUB); self._style_cell(ws.cell(r,3))
        ws.column_dimensions["B"].width=24; ws.column_dimensions["C"].width=78

    # ------------------------------------------------------------------
    # E2E Requirements
    # ------------------------------------------------------------------
    def export_requirements_word(self, source_document: Path, data: dict[str, Any]) -> Path:
        records=build_engineering_records(data)
        if not records: raise ValueError("E2E 요구사항 명세서로 내보낼 Canonical Engineering Requirement가 없습니다.")
        req_lookup={str(r.get("srs_id") or ""):r for r in _list(data.get("requirements")) if isinstance(r,dict)}
        out=self._filename("E2E_요구사항 명세서", source_document, "docx")
        doc=Document(); doc.styles["Normal"].font.name="Malgun Gothic"; doc.styles["Normal"].font.size=Pt(9)
        p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; run=p.add_run("E2E Engineering Requirements Specification"); run.bold=True; run.font.size=Pt(18)
        meta=doc.add_table(rows=5,cols=2); meta.style="Table Grid"
        meta_rows=[("원본 문서",source_document.name),("생성 일시",datetime.now().strftime("%Y-%m-%d %H:%M:%S")),("Canonical Requirement",str(len(records))),("작성 기준","Requirement Studio V0.92 — Human-facing E2E Requirements"),("표시 원칙","사람이 읽는 본문은 핵심 요구사항만 표시하고 Allocation/Eligibility/Canonical/Fact 메타데이터는 Excel Review/Trace 및 JSON Audit에 보존")]
        for i,(k,v) in enumerate(meta_rows): meta.cell(i,0).text=k; meta.cell(i,1).text=v
        note=doc.add_paragraph(); rr=note.add_run("※ 본 문서는 사람용 E2E 요구사항 Master입니다. 내부 Engineering/Allocation/Eligibility 정보는 삭제하지 않고 Review/Trace/JSON Audit에 보존하되 본문에서는 중복·기술 메타데이터를 최소화합니다."); rr.font.size=Pt(8.5); rr.font.color.rgb=RGBColor(90,98,110)
        section=doc.sections[0]; usable=int(section.page_width-section.left_margin-section.right_margin); left=int(usable*0.24); right=usable-left
        fields=["SRS ID","상위 기능","리뷰 상태","요구사항 내역","동작 조건 / Trigger","작동 명세 정의","사전 조건","예상 결과","검증 기준","출처 / Traceability","기타"]
        multiline={"요구사항 내역","동작 조건 / Trigger","작동 명세 정의","사전 조건","예상 결과","검증 기준","출처 / Traceability","기타"}
        current=None
        for rec in records:
            if rec["상위 기능"]!=current: doc.add_heading(rec["상위 기능"],level=1); current=rec["상위 기능"]
            doc.add_heading(rec["SRS ID"],level=2)
            req=req_lookup.get(rec["SRS ID"],{})
            vals={**rec,"리뷰 상태":_display_review_status(req,rec.get("Review Status"))}
            table=doc.add_table(rows=len(fields),cols=2); table.style="Table Grid"; table.autofit=False; table.columns[0].width=left; table.columns[1].width=right
            for i,f in enumerate(fields):
                a,b=table.cell(i,0),table.cell(i,1); a.width=left; b.width=right; a.text=f
                val=vals.get(f,"")
                b.text=_format_human_multiline(val) if f in multiline else _text(val)
                a.vertical_alignment=b.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for r in a.paragraphs[0].runs: r.bold=True; r.font.size=Pt(8.5)
                for pp in b.paragraphs:
                    for r in pp.runs: r.font.size=Pt(8.5)
            doc.add_paragraph()
        doc.save(out); return out

    def export_requirements_excel(self, source_document: Path, data: dict[str, Any]) -> Path:
        records=build_engineering_records(data); reqs=[x for x in _list(data.get("requirements")) if isinstance(x,dict)]
        if not records: raise ValueError("E2E 요구사항 명세서로 내보낼 Requirement가 없습니다.")
        req_lookup={str(r.get("srs_id") or ""):r for r in reqs}; out=self._filename("E2E_요구사항 명세서",source_document,"xlsx")
        wb=Workbook(); cover=wb.active; cover.title="표지"; hist=wb.create_sheet("0_변경이력"); summary=wb.create_sheet("1_요구사항요약"); main=wb.create_sheet("2_E2E_요구사항"); review=wb.create_sheet("3_검토필요"); trace=wb.create_sheet("4_Source_추적성")
        self._cover(cover,"E2E 요구사항 명세서",source_document,"Human-facing E2E Engineering Requirements")
        self._write_table(hist,["No","Version","Date","Change"],[[1,"V0.92",datetime.now().strftime("%Y-%m-%d"),"V0.91 사람용 Requirement Master 구조 유지. V0.92 변경은 Fact→Basic Functional TC 구조화/추적성 강화에 집중하며 Requirement Main 구조는 변경하지 않음"]],[8,14,14,110])
        sum_rows=[["Canonical Requirement Count",len(records)],["사람의 리뷰 필요",sum(1 for r in reqs if _display_review_status(r)=="사람의 리뷰 필요")],["리뷰 불필요/완료",sum(1 for r in reqs if _display_review_status(r)!="사람의 리뷰 필요")],["내부 메타데이터 보존 위치","3_검토필요 / 4_Source_추적성 / Review Package JSON"]]
        self._write_table(summary,["항목","Count / Status"],sum_rows,[38,52])
        headers=["SRS ID","상위 기능","리뷰 상태","요구사항 내역","동작 조건 / Trigger","작동 명세 정의","사전 조건","예상 결과","검증 기준","출처 / Traceability","기타"]
        rows=[]; review_rows=[]; trace_rows=[]
        multiline_idx={4,5,6,7,8,9,10,11}
        for rec in records:
            req=req_lookup.get(rec["SRS ID"],{}); sem,fids,loc=_source_trace(req)
            vals=[rec["SRS ID"],rec["상위 기능"],_display_review_status(req,rec.get("Review Status")),rec["요구사항 내역"],rec["동작 조건 / Trigger"],rec["작동 명세 정의"],rec["사전 조건"],rec["예상 결과"],rec["검증 기준"],rec["출처 / Traceability"],rec["기타"]]
            rows.append([_format_human_multiline(v) if i in multiline_idx else v for i,v in enumerate(vals,1)])
            if _display_review_status(req,rec.get("Review Status"))=="사람의 리뷰 필요":
                review_rows.append([rec["SRS ID"],_display_review_status(req,rec.get("Review Status")),rec.get("Allocation Status",""),rec.get("SYS.1 Eligibility",""),rec.get("SWE.1 Eligibility",""),rec.get("Canonical State",""),rec.get("Hold Reason",""),_text(req.get("clarification_reference") or req.get("clarification_needed")),"PENDING",""])
            for atom in _list(req.get("source_backed_atomic_behaviors")):
                if isinstance(atom,dict):
                    trace_rows.append([rec["SRS ID"],atom.get("source_semantic_unit_id"),atom.get("source_fact_fragment_id"),atom.get("source_location"),_format_human_multiline(atom.get("behavior_text") or atom.get("source_fact")),rec.get("Engineering Domain",""),rec.get("Requirement Level",""),rec.get("Allocation Status",""),rec.get("SYS.1 Eligibility",""),rec.get("SWE.1 Eligibility",""),rec.get("Canonical State","")])
        self._write_table(main,headers,rows,[16,24,20,56,38,58,38,38,44,54,36])
        self._write_table(review,["SRS ID","리뷰 상태","Allocation Status","SYS.1 Eligibility","SWE.1 Eligibility","Canonical State","Hold Reason","Clarification / Required Resolution","Disposition","Rationale"],review_rows,[16,20,34,18,18,28,50,66,20,56])
        self._write_table(trace,["SRS ID","Semantic Unit ID","Fact Fragment ID","Source Location","Source-backed Fact / Behavior","Engineering Domain","Requirement Level","Allocation Status","SYS.1 Eligibility","SWE.1 Eligibility","Canonical State"],trace_rows,[16,28,30,26,72,30,24,34,18,18,28])
        wb.save(out); return out

    # ------------------------------------------------------------------
    # E2E Evaluation
    # ------------------------------------------------------------------
    def export_evaluation_excel(self, source_document: Path, data: dict[str, Any]) -> Path:
        bundle=finalize_verification_single_truth(data, force=True)
        objective_cases=[x for x in _list(bundle.get("e2e_evaluation_cases")) if isinstance(x,dict)]
        cases=[x for x in _list(bundle.get("e2e_basic_functional_cases")) if isinstance(x,dict)] or objective_cases
        basic_reviews=[x for x in _list(bundle.get("e2e_basic_functional_reviews")) if isinstance(x,dict)]
        basic_audit=bundle.get("basic_functional_design_audit") or {}
        objects=[x for x in _list(bundle.get("integrated_test_objects")) if isinstance(x,dict)]
        reqs={str(r.get("srs_id") or ""):r for r in _list(data.get("requirements")) if isinstance(r,dict)}
        if not cases: raise ValueError("E2E 평가 명세서로 내보낼 평가 Case가 없습니다.")
        out=self._filename("E2E_평가 명세서",source_document,"xlsx")
        wb=Workbook(); cover=wb.active; cover.title="표지"; hist=wb.create_sheet("0_변경이력"); summary=wb.create_sheet("1_평가요약"); main=wb.create_sheet("2_E2E_평가케이스"); review=wb.create_sheet("3_검토필요"); trace=wb.create_sheet("4_요구사항-평가추적성")
        self._cover(cover,"E2E 평가 명세서",source_document,"구조화 기본 기능 검증 — Human-facing Test Case Master")
        self._write_table(hist,["No","Version","Date","Change"],[[1,"V0.92",datetime.now().strftime("%Y-%m-%d"),"Fact 역할 기반 구조화: Test Objective / Precondition / Environment / Input / Procedure / Observable / Expected / PASS Criterion 분리, Split/Keep/Vector/Step 판단, Exact Fact Trace 및 Transition Logic/Target 보존"]],[8,14,14,112])
        structural=bundle.get("single_truth_audit") or {}; semantic=bundle.get("semantic_verification_quality_audit") or {}
        release="PASS" if structural.get("release_gate_status")=="PASS" and semantic.get("status")=="PASS" else "HOLD"
        sum_rows=[["Test Design Mode","구조화 기본 기능 검증"],["원본 통합 평가 객체",len(objects)],["기존 E2E Evaluation Objective",len(objective_cases)],["사람용 E2E Test Case",len(cases)],["추가 사람 검토 항목",len(basic_reviews)],["내부 Trace/Audit","4_요구사항-평가추적성 / Finalized Verification Bundle JSON"],["구조적 단일 진실 Gate",structural.get("release_gate_status")],["의미 검증 품질 Gate",semantic.get("status")],["공식 릴리스",release]]
        self._write_table(summary,["평가 항목","건수 / 상태"],sum_rows,[44,54])

        main_headers=["Test Case ID","요구사항 ID","검증 영역","평가 방법 / 도출 방법","테스트 케이스 이름","테스트 목적","사전 조건","평가 환경 / 준비","입력 변수","비교","입력값","수행 절차","출력 변수 / 관찰 대상","비교","기대 출력값","기대 상태 / 결과","PASS 기준","TC 검토 상태","검토 사유 / 필요 조치","실제 출력값","판정 (PASS/FAIL)","비고","결과 화면 캡처"]
        rows=[]; review_rows=[]; trace_rows=[]
        bf_count: dict[str,int]={}; name_count: dict[str,int]={}
        human_ids=[]
        for global_idx,case in enumerate(cases,1):
            sid=str(case.get("parent_srs_id") or ""); req=reqs.get(sid,{})
            bf_count[sid]=bf_count.get(sid,0)+1
            bf_id=f"{sid}_BF_{bf_count[sid]:02d}" if sid else f"BF_{global_idx:02d}"
            tc_id=f"SRS_TC_{global_idx:03d}"
            human_ids.append(tc_id)
            base_name=str(case.get("evaluation_name") or req.get("function_name") or case.get("test_intent") or case.get("verification_objective") or "기본 기능 확인").strip()
            name_count[base_name]=name_count.get(base_name,0)+1
            tc_name=f"{base_name}_{name_count[base_name]:02d}"
            method_label="기본 기능 검증" if str(case.get("test_design_mode") or "")!="BASIC_FUNCTIONAL_REVIEW" else "검토 필요"
            objective=_format_human_multiline(case.get("test_objective") or case.get("verification_objective") or case.get("evaluation_name") or req.get("function_name"))
            precondition=_format_human_multiline(case.get("precondition") or case.get("test_preparation") or req.get("preconditions") or req.get("activation_trigger"))
            environment=_format_human_multiline(case.get("test_environment"))
            evar=str(case.get("input_variable") or case.get("variable") or "")
            ecomp=_display_compare(case.get("input_compare") or case.get("compare"))
            input_value=_compact_input_value(case)
            exe=_format_human_multiline(case.get("test_execution") or req.get("processing_action"))
            ovar=str(case.get("output_variable") or "")
            ocomp=_display_compare(case.get("output_compare"))
            output_value=_compact_output_value(case)
            expected_state=_format_human_multiline(case.get("expected_state") or case.get("expected_result") or req.get("output"))
            pass_criterion=_format_human_multiline(case.get("pass_criterion") or req.get("acceptance_criteria"))
            review_status=_test_case_review_status(case)
            reason=_display_review_reason(case)
            rows.append([tc_id,bf_id,_display_domain_compact(case),method_label,tc_name,objective,precondition,environment,evar,ecomp,input_value,exe,ovar,ocomp,output_value,expected_state,pass_criterion,review_status,reason,"","","",""])

            typ=str(case.get("test_object_type") or "")
            raw_ids=list(dict.fromkeys([str(case.get("test_object_id") or "")]+[str(x) for x in _list(case.get("related_test_object_ids")) if str(x)]))
            sem=", ".join(str(x) for x in _list(case.get("source_semantic_unit_ids")) if str(x))
            fids=", ".join(str(x) for x in _list(case.get("source_fact_fragment_ids")) if str(x))
            loc=" | ".join(str(x) for x in _list(case.get("source_locations")) if str(x))
            field_trace=case.get("field_fact_trace") if isinstance(case.get("field_fact_trace"),dict) else {}
            field_trace_text=" | ".join(f"{k}={','.join(str(x) for x in _list(v) if str(x))}" for k,v in field_trace.items() if _list(v))
            trace_rows.append([tc_id,bf_id,sid,_display_domain_compact(case),method_label,case.get("evaluation_objective_id") or "",", ".join(x for x in raw_ids if x),typ,sem,fids,loc,str(case.get("split_decision") or ""),str(case.get("condition_logic") or ""),str(case.get("expected_target_state") or ""),str(case.get("exact_field_trace_status") or ""),field_trace_text,str(case.get("allocation_status") or ""),str(case.get("swe6_eligibility") or ""),str(case.get("execution_readiness") or ""),"Y" if case.get("human_review_required") else "N"])
            if review_status=="사람의 리뷰 필요":
                review_rows.append([tc_id,bf_id,sid,"평가 케이스 검토",reason or "사람 검토 필요","관련 내부 메타데이터는 4_요구사항-평가추적성 및 JSON Audit 참조","","PENDING",""])

        widths=[15,20,9,15,32,40,34,30,22,10,18,42,24,10,24,32,42,22,34,18,14,28,20]
        self._write_table(main,main_headers,rows,widths)
        self._write_group_headers(main)
        dv=DataValidation(type="list",formula1='"PASS,FAIL"',allow_blank=True); main.add_data_validation(dv); dv.add(f"U4:U{max(100,main.max_row+20)}")

        review_type_ko={
            "SOURCE_DEFINED_VALUE_WITHOUT_EXPECTED_BEHAVIOR":"Source 기대동작 미정의",
            "EXACT_FIELD_FACT_TRACE_REQUIRED":"Exact Fact Trace 필요",
            "SOURCE_TBD_REVIEW":"Source TBD",
            "SOURCE_CONFLICT_REVIEW":"Source Conflict",
        }
        for n,item in enumerate(basic_reviews,1):
            reason=_format_human_multiline(item.get("reason")); action=_format_human_multiline(item.get("required_resolution"))
            rtype=review_type_ko.get(str(item.get("review_type") or ""),str(item.get("review_type") or "Source 정보 부족"))
            review_rows.append([f"BF-REVIEW-{n:03d}","",item.get("parent_srs_id"),rtype,reason,action,"","PENDING",""])
        self._write_table(review,["검토 ID / Test Case ID","요구사항 ID","상위 SRS ID","검토 유형","사유 / Gap","필요 조치","담당자","처리 상태","판단 근거"],review_rows,[24,20,16,28,66,66,18,22,56])
        dv2=DataValidation(type="list",formula1='"PENDING,APPROVED,REJECTED,NOT_APPLICABLE"',allow_blank=True); review.add_data_validation(dv2); dv2.add(f"H4:H{max(100,review.max_row+20)}")

        self._write_table(trace,["Test Case ID","요구사항 ID","상위 SRS ID","검증 영역","평가 방법","E2E 평가 객체 ID","관련 원본 객체 ID","원본 객체 유형","출처 의미 단위 ID","출처 사실 조각 ID","출처 위치","Split Decision","Condition Logic","Expected Target State","Exact Field Trace Status","Field Fact Trace","Allocation Status","SWE.6 Eligibility","Execution Readiness","Human Review Required"],trace_rows,[15,20,16,10,16,28,40,28,28,34,30,28,20,28,24,70,36,24,24,22])
        wb.save(out); self._verify_evaluation_artifact(out,cases,human_ids); return out

    def _verify_evaluation_artifact(self, path: Path, cases: list[dict[str,Any]], expected_human_ids: list[str] | None = None):
        wb=load_workbook(path,read_only=False,data_only=False)
        try:
            ws=wb["2_E2E_평가케이스"]
            # V0.92 human layout: result fields remain T:W (20~23). They must be physically blank pre-execution.
            for r in range(4,ws.max_row+1):
                if not str(ws.cell(r,1).value or "").strip(): continue
                for c in (20,21,22,23):
                    if ws.cell(r,c).value not in (None,""):
                        ws.cell(r,c).value=None
            wb.save(path)
        finally: wb.close()
        check=load_workbook(path,read_only=True,data_only=False)
        try:
            ws=check["2_E2E_평가케이스"]; violations=[]; ids=[]
            for r in range(4,ws.max_row+1):
                tcid=str(ws.cell(r,1).value or "").strip()
                if not tcid: continue
                ids.append(tcid)
                for c in (20,21,22,23):
                    if ws.cell(r,c).value not in (None,""): violations.append(ws.cell(r,c).coordinate)
            expected=expected_human_ids or [f"SRS_TC_{i:03d}" for i in range(1,len(cases)+1)]
            if ids!=expected:
                raise ValueError(f"E2E human Test Case ID mismatch: expected={expected} actual={ids}")
            if len(set(ids))!=len(ids):
                raise ValueError("E2E human Test Case ID must be unique.")
            if violations: raise ValueError("E2E evaluation pre-execution result fields must be blank: "+", ".join(violations[:20]))
        finally: check.close()

