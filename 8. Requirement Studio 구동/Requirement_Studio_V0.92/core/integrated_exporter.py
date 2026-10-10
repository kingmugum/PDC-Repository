from __future__ import annotations

import re
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from core.output_naming import output_filename
from core.verification_single_truth import finalize_verification_single_truth


NOT_CONFIRMED = "입력문서에서 확인되지 않음"


def _excel_scalar(value: Any) -> Any:
    """Return a value that openpyxl can safely store in one cell.

    V0.85 treats nested list/dict review evidence as structured text instead of passing
    Python containers directly to openpyxl. Numeric/bool values remain native scalars.
    """
    if value is None:
        return ""
    if isinstance(value, (str, int, float, bool, datetime)):
        return value
    if isinstance(value, dict):
        lines: list[str] = []
        preferred = ("document", "location", "text", "description", "reason", "detail")
        used: set[str] = set()
        for key in preferred:
            if key in value and value.get(key) not in (None, "", [], {}):
                used.add(key)
                rendered = _excel_scalar(value.get(key))
                if rendered not in (None, ""):
                    lines.append(f"{key}: {rendered}")
        for key, item in value.items():
            if key in used or item in (None, "", [], {}):
                continue
            rendered = _excel_scalar(item)
            if rendered not in (None, ""):
                lines.append(f"{key}: {rendered}")
        return "\n".join(lines)
    if isinstance(value, (list, tuple, set)):
        lines: list[str] = []
        for item in value:
            rendered = _excel_scalar(item)
            text = str(rendered).strip() if rendered is not None else ""
            if text:
                lines.append(text)
        return "\n---\n".join(lines)
    return str(value).strip()


def _text(value: Any) -> str:
    rendered = _excel_scalar(value)
    return str(rendered).strip() if rendered is not None else ""


def _list(value: Any) -> list[Any]:
    if value is None:
        return []
    return value if isinstance(value, list) else [value]


def _safe_stem(name: str) -> str:
    stem = Path(name).stem or "document"
    stem = re.sub(r'[<>:"/\\|?*]+', "_", stem)
    stem = re.sub(r"\s+", "_", stem).strip("._ ")
    return stem[:80] or "document"


def _join(values: Any, sep: str = " | ") -> str:
    out: list[str] = []
    for value in _list(values):
        if isinstance(value, dict):
            text = _text(value.get("name") or value.get("document") or value.get("text") or value.get("description") or value)
        else:
            text = _text(value)
        if text and text not in out:
            out.append(text)
    return sep.join(out)


def _source_evidence(req: dict[str, Any]) -> str:
    out: list[str] = []
    for item in _list(req.get("source_evidence")):
        if isinstance(item, dict):
            loc = _text(item.get("location"))
            text = _text(item.get("text"))
            value = f"{loc}: {text}" if loc and text else (loc or text)
        else:
            value = _text(item)
        if value and value not in out:
            out.append(value)
    return "\n".join(out)


def _source_trace(req: dict[str, Any]) -> tuple[str, str, str]:
    sem: list[str] = [str(x) for x in _list(req.get("source_semantic_unit_ids")) if str(x)]
    fids: list[str] = []
    locations: list[str] = []
    for key in ("source_fact_fragments", "fact_level_allocations", "source_backed_atomic_behaviors"):
        for item in _list(req.get(key)):
            if not isinstance(item, dict):
                continue
            fid = str(item.get("source_fact_fragment_id") or "")
            if fid and fid not in fids:
                fids.append(fid)
            uid = str(item.get("source_semantic_unit_id") or item.get("parent_source_semantic_unit_id") or "")
            if uid and uid not in sem:
                sem.append(uid)
            loc = _text(item.get("source_location"))
            if loc and loc not in locations:
                locations.append(loc)
    for ev in _list(req.get("source_evidence")):
        if isinstance(ev, dict):
            loc = _text(ev.get("location"))
            if loc and loc not in locations:
                locations.append(loc)
    return ", ".join(sem), ", ".join(fids), " | ".join(locations)


def _atomic_facts(req: dict[str, Any]) -> str:
    out: list[str] = []
    for item in _list(req.get("source_backed_atomic_behaviors")):
        if isinstance(item, dict):
            txt = _text(item.get("behavior_text") or item.get("source_fact"))
        else:
            txt = _text(item)
        if txt and txt not in out:
            out.append(txt)
    return "\n".join(out)


def _req_map(data: dict[str, Any]) -> dict[str, dict[str, Any]]:
    return {str(r.get("srs_id") or ""): r for r in (data.get("requirements") or []) if isinstance(r, dict) and r.get("srs_id")}




def _fact_alloc_by_fragment(req: dict[str, Any]) -> dict[str, dict[str, Any]]:
    return {
        str(x.get("source_fact_fragment_id") or ""): x
        for x in _list(req.get("fact_level_allocations"))
        if isinstance(x, dict) and str(x.get("source_fact_fragment_id") or "")
    }


def _linked_objects_for_fact(
    objects: list[dict[str, Any]],
    *,
    srs_id: str,
    fragment_id: str = "",
    semantic_unit_id: str = "",
) -> list[dict[str, Any]]:
    """Return only verification objects that actually carry the fact/semantic trace.

    V0.85 fixes the V0.82 parent-SRS fan-out bug: a SWE.6 TC must never appear linked to
    a fragment that is absent from that TC's finalized SWE.6 scope.
    """
    out: list[dict[str, Any]] = []
    for obj in objects:
        if str(obj.get("parent_srs_id") or "") != str(srs_id or ""):
            continue
        fids = {str(x) for x in _list(obj.get("source_fact_fragment_ids")) if str(x)}
        sems = {str(x) for x in _list(obj.get("source_semantic_unit_ids")) if str(x)}
        if fragment_id:
            if fragment_id not in fids:
                continue
        elif semantic_unit_id:
            if semantic_unit_id not in sems:
                continue
        out.append(obj)
    return out


def _trace_link_status(parent_objects: list[dict[str, Any]], linked: list[dict[str, Any]]) -> str:
    parent_has_swe6 = any(str(x.get("test_object_type") or "") == "SWE6_TC" for x in parent_objects)
    linked_has_swe6 = any(str(x.get("test_object_type") or "") == "SWE6_TC" for x in linked)
    if parent_has_swe6 and not linked_has_swe6:
        return "EXCLUDED_FROM_SWE6_SCOPE"
    if linked_has_swe6:
        return "IN_SWE6_SCOPE"
    return "NON_SWE6_OR_DEFERRED"

def _test_map(data: dict[str, Any]) -> dict[str, dict[str, Any]]:
    result = data.get("testability_and_decomposition_result")
    if not isinstance(result, dict):
        return {}
    return {str(r.get("srs_id") or ""): r for r in (result.get("by_srs") or []) if isinstance(r, dict) and r.get("srs_id")}


class IntegratedExporter:
    """V0.85 additive exporter for engineer-facing SYS/SWE integrated specifications.

    Existing SWE.1/SWE.6 artifacts remain unchanged. These workbooks are additional views over the
    same finalized Canonical + verification bundle and therefore never change allocation merely to
    increase TC counts.
    """

    HEADER_FILL = "1F4E78"
    SUB_FILL = "D9EAF7"
    REVIEW_FILL = "FFF2CC"
    BLOCK_FILL = "F4CCCC"
    PASS_FILL = "D9EAD3"
    THIN = Side(style="thin", color="B7B7B7")

    def __init__(self, output_dir: Path):
        self.output_dir = Path(output_dir).resolve()
        self.output_dir.mkdir(parents=True, exist_ok=True)

    def _req_filename(self, source_document: Path) -> Path:
        return self.output_dir / output_filename("통합_요구사항 명세서", _safe_stem(source_document.name), "xlsx")

    def _test_filename(self, source_document: Path) -> Path:
        return self.output_dir / output_filename("통합_테스트 명세서", _safe_stem(source_document.name), "xlsx")

    def _style_header(self, cell):
        cell.fill = PatternFill("solid", fgColor=self.HEADER_FILL)
        cell.font = Font(name="Malgun Gothic", color="FFFFFF", bold=True, size=9)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(left=self.THIN, right=self.THIN, top=self.THIN, bottom=self.THIN)

    def _style_cell(self, cell, *, center: bool = False, fill: str | None = None):
        cell.font = Font(name="Malgun Gothic", size=9)
        cell.alignment = Alignment(horizontal="center" if center else "left", vertical="top", wrap_text=True)
        cell.border = Border(left=self.THIN, right=self.THIN, top=self.THIN, bottom=self.THIN)
        if fill:
            cell.fill = PatternFill("solid", fgColor=fill)

    def _write_table(self, ws, headers: list[str], rows: list[list[Any]], widths: list[int] | None = None, *, start_row: int = 3):
        for col, header in enumerate(headers, 1):
            c = ws.cell(start_row, col, header)
            self._style_header(c)
        for r_idx, row in enumerate(rows or [["-"] + [""] * (len(headers) - 1)], start_row + 1):
            for col, value in enumerate(row, 1):
                c = ws.cell(r_idx, col, _excel_scalar(value))
                self._style_cell(c, center=col <= 3)
        last_row = start_row + max(1, len(rows))
        ws.freeze_panes = ws.cell(start_row + 1, 1)
        ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(headers))}{last_row}"
        if widths:
            for idx, width in enumerate(widths, 1):
                ws.column_dimensions[get_column_letter(idx)].width = width
        else:
            for idx in range(1, len(headers) + 1):
                ws.column_dimensions[get_column_letter(idx)].width = 22
        ws.sheet_view.showGridLines = False
        ws.page_setup.orientation = "landscape"
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        return last_row

    def _overview(self, ws, title: str, rows: list[tuple[str, Any]]):
        ws.sheet_view.showGridLines = False
        ws["A1"] = title
        ws["A1"].font = Font(name="Malgun Gothic", size=16, bold=True, color="1F4E78")
        ws.merge_cells("A1:F1")
        ws["A3"] = "항목"; ws["B3"] = "내용"
        self._style_header(ws["A3"]); self._style_header(ws["B3"])
        r = 4
        for label, value in rows:
            ws.cell(r, 1, label); ws.cell(r, 2, _text(value))
            self._style_cell(ws.cell(r, 1), center=True, fill=self.SUB_FILL)
            self._style_cell(ws.cell(r, 2))
            r += 1
        ws.column_dimensions["A"].width = 34
        ws.column_dimensions["B"].width = 110
        ws.freeze_panes = "A4"

    def export_requirements_excel(self, source_document: Path, data: dict[str, Any]) -> Path:
        bundle = finalize_verification_single_truth(data)
        out = self._req_filename(source_document)
        reqs = [r for r in (data.get("requirements") or []) if isinstance(r, dict)]
        tests = _test_map(data)
        objects = [x for x in bundle.get("integrated_test_objects") or [] if isinstance(x, dict)]
        obj_by_srs: dict[str, list[dict[str, Any]]] = {}
        for obj in objects:
            obj_by_srs.setdefault(str(obj.get("parent_srs_id") or ""), []).append(obj)

        wb = Workbook()
        overview = wb.active; overview.title = "00_Overview"
        req_ws = wb.create_sheet("01_Integrated_Eng_Req")  # Excel 31-char limit alias of requested logical name.
        sys1_ws = wb.create_sheet("02_SYS1_View")
        swe1_ws = wb.create_sheet("03_SWE1_View")
        alloc_ws = wb.create_sheet("04_Allocation_Review")
        trace_ws = wb.create_sheet("05_Source_Traceability")
        dep_ws = wb.create_sheet("06_External_Dependencies")
        gap_ws = wb.create_sheet("07_Open_Issues_Gaps")
        sem_ws = wb.create_sheet("08_Semantic_Unit_Coverage")
        fact_ws = wb.create_sheet("09_Fact_Fragment_Allocation")

        self._overview(overview, "통합_요구사항 명세서 — SYS.1 + SWE.1 + Engineering Domain", [
            ("Requirement Studio Version", "V0.92"),
            ("Source", source_document.name),
            ("Canonical Requirement Count", len(reqs)),
            ("핵심 원칙", "Not SWE.1 != Not Requirement. Canonical Engineering Requirement를 보존하고 SYS.1/SWE.1/native domain은 projection으로 표시합니다."),
            ("기존 산출물", "기존 SWE.1 Word/Excel은 유지되며 본 문서는 추가 통합 산출물입니다."),
            ("Excel Sheet Alias", "요청 논리명 01_Integrated_Engineering_Requirements는 Excel 31자 제한으로 01_Integrated_Eng_Req로 저장됩니다."),
            ("Source Fidelity", "Source에 없는 Signal/DB/Timing/Threshold/Variant/Test Value를 생성하지 않습니다."),
            ("Single Truth", bundle.get("schema_version")),
        ])

        headers = [
            "Canonical Requirement ID / SRS ID", "Candidate ID", "상위 기능", "Requirement", "Source Evidence",
            "Source Semantic Unit ID", "Source Fact Fragment ID", "Source-backed Atomic Behavior / Fact",
            "Engineering Domain", "Requirement Level", "Allocation Status", "SYS.1 Eligibility", "SWE.1 Eligibility",
            "SYS.5 Eligibility", "SWE.6 Eligibility", "Verification Domain", "Review Status", "Canonical State",
            "Human Decision Required", "Hold Reason", "Clarification Needed", "Canonical Gap IDs", "External Dependency",
            "Source-backed Test Intent 상태", "연결된 SYS.5 Candidate ID", "연결된 SWE.6 TC ID", "연결된 Deferred Intent ID / Reason",
        ]
        rows: list[list[Any]] = []
        for req in reqs:
            sid = str(req.get("srs_id") or "")
            sem, fids, _ = _source_trace(req)
            test_row = tests.get(sid, {})
            per_obj = obj_by_srs.get(sid, [])
            sys5_ids = [str(x.get("test_object_id") or "") for x in per_obj if x.get("test_object_type") == "SYS5_CANDIDATE"]
            swe6_ids = [str(x.get("test_object_id") or "") for x in per_obj if x.get("test_object_type") == "SWE6_TC"]
            deferred = [x for x in per_obj if x.get("test_object_type") in {"DEFERRED_INTENT", "ALLOCATION_PENDING_INTENT", "EXTERNAL_DEPENDENCY_INTENT", "REVIEW_REQUIRED"} and str(x.get("test_object_id") or "").startswith(("DEF_INTENT_", "AP_INTENT_", "EXT_INTENT_", "REV_INTENT_"))]
            rows.append([
                sid, req.get("candidate_id", ""), req.get("function_name", ""), req.get("requirement", ""), _source_evidence(req),
                sem, fids, _atomic_facts(req), _join(req.get("engineering_domains")), req.get("requirement_level", ""),
                req.get("allocation_status", ""), req.get("sys1_eligibility", ""), req.get("swe1_eligibility", ""),
                req.get("sys5_eligibility", ""), req.get("swe6_eligibility", ""), _join(req.get("verification_domains")) or req.get("verification_domain", ""),
                req.get("review_status", ""), req.get("canonical_state", ""), "Y" if req.get("human_decision_required") else "N",
                req.get("hold_reason", ""), _text(req.get("clarification_reference") or req.get("clarification_needed")), _join(req.get("canonical_gap_ids")), _text(req.get("external_dependencies")),
                test_row.get("intent_complete_status", ""), ", ".join(sys5_ids), ", ".join(swe6_ids),
                "\n".join(f"{x.get('test_object_id')}: {x.get('deferred_reason_code')} - {x.get('deferred_reason_detail')}" for x in deferred),
            ])
        self._write_table(req_ws, headers, rows, [23,18,24,72,72,28,34,72,28,24,34,18,18,18,22,42,26,28,20,54,54,24,48,28,26,24,70])
        req_ws["A1"] = "1. Integrated Engineering Requirements (logical: 01_Integrated_Engineering_Requirements)"; req_ws["A1"].font = Font(name="Malgun Gothic", size=14, bold=True)

        # SYS.1 / SWE.1 projection views are non-destructive filters over the same Canonical rows.
        sys_rows = [row for row, req in zip(rows, reqs) if str(req.get("sys1_eligibility") or "") in {"Eligible", "Review Needed"}]
        swe_rows = [row for row, req in zip(rows, reqs) if str(req.get("swe1_eligibility") or "") in {"Eligible", "Review Needed"}]
        self._write_table(sys1_ws, headers, sys_rows, [23,18,24,72,72] + [28] * (len(headers)-5)); sys1_ws["A1"] = "2. SYS.1 System Requirement View"; sys1_ws["A1"].font = Font(name="Malgun Gothic", size=14, bold=True)
        self._write_table(swe1_ws, headers, swe_rows, [23,18,24,72,72] + [28] * (len(headers)-5)); swe1_ws["A1"] = "3. SWE.1 Software Requirement View"; swe1_ws["A1"].font = Font(name="Malgun Gothic", size=14, bold=True)

        alloc_headers = ["Finding / SRS", "Finding ID", "Type", "Affected SRS", "Allocation Status", "SYS.5", "SWE.6", "Review Disposition", "Review Owner", "Reason", "Required Resolution", "Human Review"]
        alloc_rows: list[list[Any]] = []
        for req in reqs:
            if req.get("human_decision_required") or req.get("cross_domain_bundle_review_required"):
                alloc_rows.append([req.get("srs_id"), "", "REQUIREMENT_ALLOCATION_REVIEW", req.get("srs_id"), req.get("allocation_status"), req.get("sys5_eligibility"), req.get("swe6_eligibility"), "PENDING", "", req.get("hold_reason") or _text(req.get("clarification_reference") or req.get("clarification_needed")), "Approve fact-level System/SW/native-domain ownership; do not auto-promote Source-missing scope.", "Y"])
        for finding in bundle.get("mode_transition_allocation_findings") or []:
            alloc_rows.append([finding.get("finding_type"), finding.get("finding_id"), finding.get("finding_type"), ", ".join(finding.get("affected_srs_ids") or []), _text(finding.get("left_allocation")) + " <> " + _text(finding.get("right_allocation")), "", "", finding.get("review_disposition"), finding.get("review_owner"), finding.get("reason"), finding.get("required_resolution"), "N" if finding.get("review_closed") else "Y"])
        self._write_table(alloc_ws, alloc_headers, alloc_rows, [28,20,34,28,48,18,18,24,20,72,72,18]); alloc_ws["A1"] = "4. Allocation / Human Review"; alloc_ws["A1"].font = Font(name="Malgun Gothic", size=14, bold=True)

        trace_headers = ["SRS ID", "Candidate ID", "Source Semantic Unit ID", "Source Fact Fragment ID", "Source Location", "Source-backed Fact / Behavior", "Requirement", "Allocation Status"]
        trace_rows: list[list[Any]] = []
        fact_rows: list[list[Any]] = []
        fact_headers = ["SRS ID", "Source Semantic Unit ID", "Source Fact Fragment ID", "Source Location", "Source Fact", "Allocation Status", "SYS.5 Eligibility", "SWE.6 Eligibility", "Verification Domain", "Allocation Inheritance", "Override Evidence"]
        for req in reqs:
            sid = str(req.get("srs_id") or "")
            atoms = _list(req.get("source_backed_atomic_behaviors")) or [req.get("requirement")]
            for atom in atoms:
                if isinstance(atom, dict):
                    trace_rows.append([sid, req.get("candidate_id"), atom.get("source_semantic_unit_id"), atom.get("source_fact_fragment_id"), atom.get("source_location"), atom.get("behavior_text") or atom.get("source_fact"), req.get("requirement"), req.get("allocation_status")])
                elif atom:
                    sem, fids, loc = _source_trace(req)
                    trace_rows.append([sid, req.get("candidate_id"), sem, fids, loc, atom, req.get("requirement"), req.get("allocation_status")])
            for alloc in _list(req.get("fact_level_allocations")):
                if not isinstance(alloc, dict): continue
                fact_rows.append([sid, alloc.get("source_semantic_unit_id"), alloc.get("source_fact_fragment_id"), alloc.get("source_location"), alloc.get("source_fact"), alloc.get("allocation_status"), "Eligible" if "sys.5" in str(alloc.get("verification_domain") or "").lower() or "system integration" in str(alloc.get("verification_domain") or "").lower() else "", alloc.get("swe6_eligibility"), alloc.get("verification_domain"), alloc.get("allocation_inheritance"), _join(alloc.get("allocation_override_evidence"))])
        self._write_table(trace_ws, trace_headers, trace_rows, [18,18,30,34,28,72,72,34]); trace_ws["A1"] = "5. Source Traceability"; trace_ws["A1"].font = Font(name="Malgun Gothic", size=14, bold=True)
        self._write_table(fact_ws, fact_headers, fact_rows, [18,30,34,28,72,34,18,22,44,34,54]); fact_ws["A1"] = "9. Fact Fragment Allocation"; fact_ws["A1"].font = Font(name="Malgun Gothic", size=14, bold=True)

        dep_headers = ["SRS ID", "External Dependency", "Verification Domain", "Clarification / Required Resolution", "Source Evidence"]
        dep_rows = []
        for req in reqs:
            deps = _list(req.get("external_dependencies"))
            for dep in deps:
                dep_rows.append([req.get("srs_id"), _text(dep), req.get("verification_domain"), _text(req.get("clarification_reference") or req.get("clarification_needed")), _source_evidence(req)])
        self._write_table(dep_ws, dep_headers, dep_rows, [18,58,42,68,72]); dep_ws["A1"] = "6. External Dependencies"; dep_ws["A1"].font = Font(name="Malgun Gothic", size=14, bold=True)

        gap_headers = ["Type", "ID", "Affected SRS", "Description / Reason", "Source", "Status / Required Resolution"]
        gap_rows = []
        for key, typ in (("gaps", "GAP"), ("conflict_register", "CONFLICT"), ("open_issues", "OPEN_ISSUE")):
            for item in _list(data.get(key)):
                if not isinstance(item, dict): continue
                gap_rows.append([typ, item.get("gap_id") or item.get("conflict_id") or item.get("issue_id") or "", _join(item.get("related_srs_ids") or item.get("affected_srs_ids")), item.get("description") or item.get("reason") or item.get("detail") or "", item.get("source_evidence") or item.get("source_location") or "", item.get("required_resolution") or item.get("status") or ""])
        self._write_table(gap_ws, gap_headers, gap_rows, [18,18,28,76,60,60]); gap_ws["A1"] = "7. Open Issues / Gaps / Conflicts"; gap_ws["A1"].font = Font(name="Malgun Gothic", size=14, bold=True)

        sem_cov = data.get("semantic_source_unit_coverage") if isinstance(data.get("semantic_source_unit_coverage"), dict) else {}
        sem_headers = ["Semantic Unit ID", "Source Location", "Disposition", "Review State", "Terminal", "Linked SRS", "Reason / Required Resolution", "Review Shortcut Suggestion", "Suggested SRS", "Similarity", "Suggested Action"]
        sem_rows = []
        for item in _list(sem_cov.get("open_review_needed_records")):
            if isinstance(item, dict):
                sem_rows.append([
                    item.get("source_semantic_unit_id"), item.get("source_location"), item.get("disposition_status"), item.get("review_state"), item.get("terminal_disposition"),
                    _join(item.get("linked_srs_ids")), item.get("required_resolution") or item.get("reason") or "",
                    item.get("review_shortcut_suggestion"), item.get("suggested_srs_id"), item.get("similarity_score"), item.get("suggested_action"),
                ])
        self._write_table(sem_ws, sem_headers, sem_rows, [30,28,38,20,16,28,72,34,20,14,80]); sem_ws["A1"] = "8. Semantic Unit Coverage / OPEN Review"; sem_ws["A1"].font = Font(name="Malgun Gothic", size=14, bold=True)

        wb.save(out)
        return out

    def export_tests_excel(self, source_document: Path, data: dict[str, Any]) -> Path:
        bundle = finalize_verification_single_truth(data)
        out = self._test_filename(source_document)
        objects = [x for x in bundle.get("integrated_test_objects") or [] if isinstance(x, dict)]
        reqs = _req_map(data)
        audit = bundle.get("single_truth_audit") if isinstance(bundle.get("single_truth_audit"), dict) else {}
        counts = Counter(str(x.get("test_object_type") or "") for x in objects)

        wb = Workbook()
        overview = wb.active; overview.title = "00_Overview"
        all_ws = wb.create_sheet("01_Integrated_Test_Objects")
        sys5_ws = wb.create_sheet("02_SYS5_Candidates")
        swe6_ws = wb.create_sheet("03_SWE6_Test_Cases")
        deferred_ws = wb.create_sheet("04_Deferred_Intents")
        alloc_ws = wb.create_sheet("05_Allocation_Pending_Intents")
        ext_ws = wb.create_sheet("06_External_Dependency_Intents")
        nonsw_ws = wb.create_sheet("07_Non_SW_Verification_Intents")
        trace_ws = wb.create_sheet("08_Source_Intent_Traceability")
        fact_ws = wb.create_sheet("09_Fact_Allocation_Audit")
        result_ws = wb.create_sheet("10_Execution_Result_Entry")
        gate_ws = wb.create_sheet("11_Release_Gate_Summary")

        self._overview(overview, "통합_테스트 명세서 — SYS.5 + SWE.6 + Deferred / Native Verification", [
            ("Requirement Studio Version", "V0.92"),
            ("Source", source_document.name),
            ("Integrated Test Object Count", len(objects)),
            ("SYS5_CANDIDATE count", counts.get("SYS5_CANDIDATE", 0)),
            ("SWE6_TC count", counts.get("SWE6_TC", 0)),
            ("DEFERRED_INTENT count", counts.get("DEFERRED_INTENT", 0)),
            ("ALLOCATION_PENDING_INTENT count", counts.get("ALLOCATION_PENDING_INTENT", 0)),
            ("EXTERNAL_DEPENDENCY_INTENT count", counts.get("EXTERNAL_DEPENDENCY_INTENT", 0)),
            ("Non-SW verification intent count", sum(counts.get(x, 0) for x in ("ELECTRICAL_VERIFICATION_INTENT", "ENVIRONMENTAL_VERIFICATION_INTENT", "MECHANICAL_VERIFICATION_INTENT", "MANUFACTURING_VERIFICATION_INTENT"))),
            ("Single Truth Release Gate", audit.get("release_gate_status")),
            ("핵심 원칙", "Not SWE.6 != Not Verifiable. SWE.6 미생성 Source-backed intent도 SYS.5/Deferred/Allocation Pending/Native Domain object로 보존합니다."),
            ("Execution Results", "Output Value / PASS-FAIL / Comment / Capture는 생성 시 물리적으로 blank입니다."),
            ("기존 산출물", "기존 SWE.6 Excel은 유지되며 본 문서는 추가 통합 산출물입니다."),
        ])

        headers = [
            "Test Object Type", "Test Object ID", "Parent SRS ID", "Candidate ID", "Source Semantic Unit ID",
            "Source Fact Fragment ID", "Source Location", "Source-backed Behavior / Source Fact", "Verification Objective / Test Intent",
            "Engineering Domain", "Verification Domain", "SYS.5 Eligibility", "SWE.6 Eligibility", "Allocation Status",
            "Execution Readiness", "Human Review Required", "Deferred Reason Code", "Deferred Reason Detail", "Required Resolution",
            "Source / External Dependency", "Test Preparation", "Test Execution", "Expected Result", "Variable", "Compare", "Value",
            "Output Value", "PASS / FAIL", "Comment", "Capture",
        ]

        def row_of(obj: dict[str, Any]) -> list[Any]:
            return [
                obj.get("test_object_type"), obj.get("test_object_id"), obj.get("parent_srs_id"), obj.get("candidate_id"),
                ", ".join(str(x) for x in _list(obj.get("source_semantic_unit_ids")) if str(x)),
                ", ".join(str(x) for x in _list(obj.get("source_fact_fragment_ids")) if str(x)),
                " | ".join(str(x) for x in _list(obj.get("source_locations")) if str(x)),
                obj.get("source_backed_behavior"), obj.get("verification_objective") or obj.get("test_intent"),
                obj.get("engineering_domain"), obj.get("verification_domain"), obj.get("sys5_eligibility"), obj.get("swe6_eligibility"),
                obj.get("allocation_status"), obj.get("execution_readiness"), "Y" if obj.get("human_review_required") else "N",
                obj.get("deferred_reason_code"), obj.get("deferred_reason_detail"), obj.get("required_resolution"),
                obj.get("source_or_external_dependency"), obj.get("test_preparation"), obj.get("test_execution"), obj.get("expected_result"),
                obj.get("variable"), obj.get("compare"), obj.get("value"), "", "", "", "",
            ]

        all_rows = [row_of(x) for x in objects]
        widths = [28,22,18,18,30,34,28,72,72,28,44,18,22,34,24,20,28,64,64,54,62,62,62,24,16,22,18,14,32,24]
        self._write_table(all_ws, headers, all_rows, widths); all_ws["A1"] = "1. Integrated Test Objects"; all_ws["A1"].font = Font(name="Malgun Gothic", size=14, bold=True)

        type_sets = {
            sys5_ws: {"SYS5_CANDIDATE"},
            swe6_ws: {"SWE6_TC"},
            deferred_ws: {"DEFERRED_INTENT", "REVIEW_REQUIRED"},
            alloc_ws: {"ALLOCATION_PENDING_INTENT"},
            ext_ws: {"EXTERNAL_DEPENDENCY_INTENT"},
            nonsw_ws: {"ELECTRICAL_VERIFICATION_INTENT", "ENVIRONMENTAL_VERIFICATION_INTENT", "MECHANICAL_VERIFICATION_INTENT", "MANUFACTURING_VERIFICATION_INTENT"},
        }
        titles = {
            sys5_ws: "2. SYS.5 System Qualification Candidates",
            swe6_ws: "3. SWE.6 Software Qualification Test Cases",
            deferred_ws: "4. Deferred / Review Required Intents",
            alloc_ws: "5. Allocation Pending Intents",
            ext_ws: "6. External Dependency Intents",
            nonsw_ws: "7. Non-SW Verification Intents",
        }
        for ws, allowed in type_sets.items():
            subset = [row_of(x) for x in objects if x.get("test_object_type") in allowed]
            self._write_table(ws, headers, subset, widths)
            ws["A1"] = titles[ws]; ws["A1"].font = Font(name="Malgun Gothic", size=14, bold=True)

        # Source Intent Traceability: one row per fact, linked only to objects that actually carry that exact fragment.
        trace_headers = [
            "Parent SRS ID", "Candidate ID", "Source Semantic Unit ID", "Source Fact Fragment ID", "Source Location",
            "Source-backed Behavior / Fact", "Linked Test Object IDs", "Linked Object Types", "Fact Allocation Status",
            "Fact SWE.6 Eligibility", "Fact Verification Domain", "SWE.6 Scope Link Status",
        ]
        trace_rows: list[list[Any]] = []
        objects_by_srs: dict[str, list[dict[str, Any]]] = {}
        for obj in objects:
            objects_by_srs.setdefault(str(obj.get("parent_srs_id") or ""), []).append(obj)
        for sid, req in reqs.items():
            parent_objects = objects_by_srs.get(sid, [])
            alloc_map = _fact_alloc_by_fragment(req)
            atoms = [x for x in _list(req.get("source_backed_atomic_behaviors")) if x not in (None, "", {})]
            if not atoms:
                sem, fids, loc = _source_trace(req)
                linked = _linked_objects_for_fact(objects, srs_id=sid, semantic_unit_id=(sem.split(", ")[0] if sem else ""))
                linked_ids = ", ".join(str(x.get("test_object_id") or "") for x in linked if x.get("test_object_id"))
                linked_types = ", ".join(dict.fromkeys(str(x.get("test_object_type") or "") for x in linked if x.get("test_object_type")))
                trace_rows.append([sid, req.get("candidate_id"), sem, fids, loc, req.get("requirement"), linked_ids, linked_types, req.get("allocation_status"), req.get("swe6_eligibility"), req.get("verification_domain"), _trace_link_status(parent_objects, linked)])
            for atom in atoms:
                if isinstance(atom, dict):
                    fid = str(atom.get("source_fact_fragment_id") or "")
                    sem = str(atom.get("source_semantic_unit_id") or "")
                    alloc = alloc_map.get(fid) or {}
                    linked = _linked_objects_for_fact(objects, srs_id=sid, fragment_id=fid, semantic_unit_id=sem)
                    linked_ids = ", ".join(str(x.get("test_object_id") or "") for x in linked if x.get("test_object_id"))
                    linked_types = ", ".join(dict.fromkeys(str(x.get("test_object_type") or "") for x in linked if x.get("test_object_type")))
                    trace_rows.append([
                        sid, req.get("candidate_id"), sem, fid, atom.get("source_location"), atom.get("behavior_text") or atom.get("source_fact"),
                        linked_ids, linked_types, alloc.get("allocation_status") or req.get("allocation_status"),
                        alloc.get("swe6_eligibility") or req.get("swe6_eligibility"), alloc.get("verification_domain") or req.get("verification_domain"),
                        _trace_link_status(parent_objects, linked),
                    ])
                else:
                    sem, fids, loc = _source_trace(req)
                    linked = _linked_objects_for_fact(objects, srs_id=sid, semantic_unit_id=(sem.split(", ")[0] if sem else ""))
                    linked_ids = ", ".join(str(x.get("test_object_id") or "") for x in linked if x.get("test_object_id"))
                    linked_types = ", ".join(dict.fromkeys(str(x.get("test_object_type") or "") for x in linked if x.get("test_object_type")))
                    trace_rows.append([sid, req.get("candidate_id"), sem, fids, loc, atom, linked_ids, linked_types, req.get("allocation_status"), req.get("swe6_eligibility"), req.get("verification_domain"), _trace_link_status(parent_objects, linked)])
        self._write_table(trace_ws, trace_headers, trace_rows, [18,18,30,34,28,76,54,54,34,22,44,28]); trace_ws["A1"] = "8. Source Intent Traceability"; trace_ws["A1"].font = Font(name="Malgun Gothic", size=14, bold=True)

        fact_headers = ["SRS ID", "Source Semantic Unit ID", "Source Fact Fragment ID", "Source Location", "Source Fact", "Engineering / Allocation Status", "SWE.6 Eligibility", "Verification Domain", "Linked Test Object IDs", "SWE.6 Scope Link Status"]
        fact_rows: list[list[Any]] = []
        for sid, req in reqs.items():
            parent_objects = objects_by_srs.get(sid, [])
            for alloc in _list(req.get("fact_level_allocations")):
                if isinstance(alloc, dict):
                    fid = str(alloc.get("source_fact_fragment_id") or "")
                    sem = str(alloc.get("source_semantic_unit_id") or "")
                    linked = _linked_objects_for_fact(objects, srs_id=sid, fragment_id=fid, semantic_unit_id=sem)
                    linked_ids = ", ".join(str(x.get("test_object_id") or "") for x in linked if x.get("test_object_id"))
                    fact_rows.append([sid, sem, fid, alloc.get("source_location"), alloc.get("source_fact"), alloc.get("allocation_status"), alloc.get("swe6_eligibility"), alloc.get("verification_domain"), linked_ids, _trace_link_status(parent_objects, linked)])
        self._write_table(fact_ws, fact_headers, fact_rows, [18,30,34,28,76,38,22,44,60,28]); fact_ws["A1"] = "9. Fact Allocation Audit"; fact_ws["A1"].font = Font(name="Malgun Gothic", size=14, bold=True)

        # Dedicated execution result entry keeps every result cell physically blank at creation.
        result_headers = ["Test Object Type", "Test Object ID", "Parent SRS ID", "Execution Readiness", "Output Value", "PASS / FAIL", "Comment", "Capture"]
        result_rows = [[x.get("test_object_type"), x.get("test_object_id"), x.get("parent_srs_id"), x.get("execution_readiness"), "", "", "", ""] for x in objects]
        self._write_table(result_ws, result_headers, result_rows, [30,24,18,24,22,16,48,28]); result_ws["A1"] = "10. Execution Result Entry (blank until execution)"; result_ws["A1"].font = Font(name="Malgun Gothic", size=14, bold=True)
        dv = DataValidation(type="list", formula1='"PASS,FAIL"', allow_blank=True)
        result_ws.add_data_validation(dv); dv.add(f"F4:F{max(100, result_ws.max_row + 20)}")

        gate_headers = ["Gate / Finding", "Finding ID", "Status / Severity", "Affected IDs", "Review Disposition", "Review Owner", "Details", "Required Resolution", "Review Rationale / Reviewed At"]
        gate_rows = [["Integrated Single Truth", "", audit.get("release_gate_status"), "", "", "", f"blocking={audit.get('blocking_issue_count', 0)}; objects={audit.get('total_integrated_test_object_count', len(objects))}", "Resolve all RELEASE_BLOCKING single-truth findings before release.", ""]]
        for rec in audit.get("blocking_records") or []:
            gate_rows.append([rec.get("issue"), "", "RELEASE_BLOCKING", rec.get("srs_id") or ", ".join(rec.get("test_object_ids") or []), "", "", _text(rec), rec.get("required_resolution") or "Correct the deterministic cross-output inconsistency.", ""])
        for finding in bundle.get("mode_transition_allocation_findings") or []:
            gate_rows.append([
                finding.get("finding_type"), finding.get("finding_id"), finding.get("severity"), ", ".join(finding.get("affected_srs_ids") or []),
                finding.get("review_disposition"), finding.get("review_owner"), finding.get("reason"), finding.get("required_resolution"),
                " | ".join(x for x in (str(finding.get("review_rationale") or ""), str(finding.get("reviewed_at") or "")) if x),
            ])
        self._write_table(gate_ws, gate_headers, gate_rows, [40,20,24,32,24,20,90,80,60]); gate_ws["A1"] = "11. Release Gate / Human Review Summary"; gate_ws["A1"].font = Font(name="Malgun Gothic", size=14, bold=True)

        wb.save(out)
        self._enforce_and_verify_blank_results(out, objects)
        self._verify_integrated_ids(out, bundle)
        return out

    def _enforce_and_verify_blank_results(self, path: Path, objects: list[dict[str, Any]]) -> None:
        wb = load_workbook(path)
        changed = False
        try:
            # In integrated object sheets result fields are columns 27~30; in result entry they are 5~8.
            for ws_name in ("01_Integrated_Test_Objects", "02_SYS5_Candidates", "03_SWE6_Test_Cases", "04_Deferred_Intents", "05_Allocation_Pending_Intents", "06_External_Dependency_Intents", "07_Non_SW_Verification_Intents"):
                ws = wb[ws_name]
                for row in range(4, ws.max_row + 1):
                    if not str(ws.cell(row, 2).value or "").strip():
                        continue
                    for col in (27, 28, 29, 30):
                        if ws.cell(row, col).value not in (None, ""):
                            ws.cell(row, col).value = None; changed = True
            ws = wb["10_Execution_Result_Entry"]
            for row in range(4, ws.max_row + 1):
                if not str(ws.cell(row, 2).value or "").strip():
                    continue
                for col in (5, 6, 7, 8):
                    if ws.cell(row, col).value not in (None, ""):
                        ws.cell(row, col).value = None; changed = True
            if changed:
                wb.save(path)
        finally:
            wb.close()

        verify = load_workbook(path, read_only=True, data_only=False)
        violations: list[str] = []
        try:
            for ws_name in ("01_Integrated_Test_Objects", "02_SYS5_Candidates", "03_SWE6_Test_Cases", "04_Deferred_Intents", "05_Allocation_Pending_Intents", "06_External_Dependency_Intents", "07_Non_SW_Verification_Intents"):
                ws = verify[ws_name]
                for row in range(4, ws.max_row + 1):
                    if not str(ws.cell(row, 2).value or "").strip(): continue
                    for col in (27, 28, 29, 30):
                        if ws.cell(row, col).value not in (None, ""):
                            violations.append(f"{ws_name}!{ws.cell(row, col).coordinate}")
            ws = verify["10_Execution_Result_Entry"]
            for row in range(4, ws.max_row + 1):
                if not str(ws.cell(row, 2).value or "").strip(): continue
                for col in (5, 6, 7, 8):
                    if ws.cell(row, col).value not in (None, ""):
                        violations.append(f"{ws.title}!{ws.cell(row, col).coordinate}")
        finally:
            verify.close()
        if violations:
            raise ValueError("Integrated test export invariant failure: execution result fields must be blank: " + ", ".join(violations[:20]))

    def _verify_integrated_ids(self, path: Path, bundle: dict[str, Any]) -> None:
        wb = load_workbook(path, read_only=True, data_only=False)
        try:
            ws = wb["01_Integrated_Test_Objects"]
            actual = {str(ws.cell(r, 2).value or "") for r in range(4, ws.max_row + 1) if str(ws.cell(r, 2).value or "")}
            expected = {str(x.get("test_object_id") or "") for x in (bundle.get("integrated_test_objects") or []) if isinstance(x, dict) and x.get("test_object_id")}
            if actual != expected:
                raise ValueError(f"Integrated test export ID mismatch: missing={sorted(expected-actual)} extra={sorted(actual-expected)}")
            sys_ws = wb["02_SYS5_Candidates"]
            actual_sys = {str(sys_ws.cell(r, 2).value or "") for r in range(4, sys_ws.max_row + 1) if str(sys_ws.cell(r, 2).value or "")}
            expected_sys = {str(x.get("sys5_id") or "") for x in (bundle.get("sys5_candidates") or []) if isinstance(x, dict) and x.get("sys5_id")}
            if actual_sys != expected_sys:
                raise ValueError("Integrated SYS.5 Candidate ID mismatch")
            swe_ws = wb["03_SWE6_Test_Cases"]
            actual_tc = {str(swe_ws.cell(r, 2).value or "") for r in range(4, swe_ws.max_row + 1) if str(swe_ws.cell(r, 2).value or "")}
            expected_tc = {str(x.get("tc_id") or "") for x in (bundle.get("swe6_cases") or []) if isinstance(x, dict) and x.get("tc_id")}
            if actual_tc != expected_tc:
                raise ValueError("Integrated SWE.6 TC ID mismatch")

            # V0.85 exact fact-to-object invariant: exported trace links must match the finalized
            # object's own fragment scope.  Parent-level fan-out is release-blocking.
            trace_ws = wb["08_Source_Intent_Traceability"]
            objects = [x for x in (bundle.get("integrated_test_objects") or []) if isinstance(x, dict)]
            expected_by_fact: dict[tuple[str, str], set[str]] = {}
            for obj in objects:
                sid = str(obj.get("parent_srs_id") or "")
                oid = str(obj.get("test_object_id") or "")
                for fid in _list(obj.get("source_fact_fragment_ids")):
                    key = (sid, str(fid or ""))
                    if key[0] and key[1] and oid:
                        expected_by_fact.setdefault(key, set()).add(oid)
            mismatches: list[str] = []
            for r in range(4, trace_ws.max_row + 1):
                sid = str(trace_ws.cell(r, 1).value or "")
                fid = str(trace_ws.cell(r, 4).value or "")
                if not sid or not fid:
                    continue
                actual_links = {x.strip() for x in str(trace_ws.cell(r, 7).value or "").split(",") if x.strip()}
                expected_links = expected_by_fact.get((sid, fid), set())
                if actual_links != expected_links:
                    mismatches.append(f"row={r} {sid}/{fid} expected={sorted(expected_links)} actual={sorted(actual_links)}")
            if mismatches:
                raise ValueError("Integrated Source Intent Traceability scope mismatch: " + " | ".join(mismatches[:10]))
        finally:
            wb.close()
