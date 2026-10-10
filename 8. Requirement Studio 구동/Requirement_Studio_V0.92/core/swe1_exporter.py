from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path
from typing import Any

from core.output_naming import output_filename
from core.swe6_exporter import build_swe6_cases

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


NOT_CONFIRMED = "입력문서에서 확인되지 않음"
NO_DETAIL = "입력문서에서 추가 작동 명세를 확인할 수 없음"


FINAL_FIELDS = [
    "SRS ID",
    "상위 기능",
    "분류",
    "검토 상태",
    "요구사항 내역",
    "동작 조건 / Trigger",
    "작동 명세 정의",
    "사전 조건",
    "예상 결과",
    "검증 기준",
    "출처 / Traceability",
    "기타",
]

ENGINEERING_WORD_FIELDS = [
    "SRS ID", "Engineering Domain", "Requirement Level", "Allocation Status",
    "SYS.1 Eligibility", "SWE.1 Eligibility", "Review Status", "Canonical State",
    "요구사항 내역", "동작 조건 / Trigger", "작동 명세 정의", "사전 조건",
    "예상 결과", "검증 기준", "출처 / Traceability", "Hold Reason", "기타",
]


def _safe_stem(name: str) -> str:
    stem = Path(name).stem or "document"
    stem = re.sub(r'[<>:"/\\|?*]+', "_", stem)
    stem = re.sub(r"\s+", "_", stem).strip("._ ")
    return stem[:80] or "document"


def _as_list(value: Any) -> list[str]:
    if value is None:
        return []
    if isinstance(value, list):
        result = []
        for item in value:
            if isinstance(item, dict):
                text = "; ".join(f"{k}={v}" for k, v in item.items() if v not in (None, "", [], {}))
            else:
                text = str(item).strip()
            if text:
                result.append(text)
        return result
    if isinstance(value, dict):
        text = "; ".join(f"{k}={v}" for k, v in value.items() if v not in (None, "", [], {}))
        return [text] if text else []
    text = str(value).strip()
    return [text] if text else []


def _text(value: Any, *, empty: str = "") -> str:
    items = _as_list(value)
    return "\n".join(items) if items else empty


def _normalize_classification(value: Any) -> str:
    raw = _text(value).strip()
    low = raw.lower()
    if "비기능" in raw or "non-functional" in low or "nonfunctional" in low:
        return "비기능"
    if raw == "기능" or "functional" in low or ("기능" in raw and "비기능" not in raw):
        return "기능"
    return raw or "검토 필요"


def _format_evidence(items: Any) -> str:
    lines = []
    for idx, item in enumerate(items if isinstance(items, list) else [], start=1):
        if not isinstance(item, dict):
            text = str(item).strip()
            if text:
                lines.append(f"[{idx}] {text}")
            continue
        document = str(item.get("document") or "").strip()
        location = str(item.get("location") or "").strip()
        evidence = str(item.get("text") or "").strip()
        prefix = " / ".join(x for x in (document, location) if x)
        if evidence:
            lines.append(f"[{idx}] {prefix}: {evidence}" if prefix else f"[{idx}] {evidence}")
        elif prefix:
            lines.append(f"[{idx}] {prefix}")
    return "\n".join(lines) if lines else NOT_CONFIRMED


def _behavior_text(req: dict[str, Any]) -> str:
    """Render detailed behavior for general engineers without exposing BF IDs."""
    flows = _as_list(req.get("behavior_flows"))
    if flows:
        rendered = []
        for idx, item in enumerate(flows, start=1):
            clean = re.sub(r"^BF\s*\d+(?:[-.]\d+)?[.:\-]?\s*", "", item, flags=re.IGNORECASE).strip()
            rendered.append(f"{idx}. {clean or item}")
        return "\n".join(rendered)
    processing = _text(req.get("processing_action")).strip()
    if processing:
        return processing
    return NO_DETAIL


def _other_materials(req: dict[str, Any]) -> str:
    """Final-output '기타': only source-backed auxiliary materials such as figures/diagrams/tables."""
    lines: list[str] = []
    for key in (
        "related_artifacts",
        "diagram_refs",
        "figure_refs",
        "table_refs",
        "visual_evidence_refs",
        "attachments",
    ):
        for item in _as_list(req.get(key)):
            if item and item not in lines:
                lines.append(item)
    return "\n".join(lines)


def _source_behavior_units(requirement_data: dict[str, Any]) -> dict[str, list[str]]:
    """Build source-backed atomic behavior lines from the final occurrence matrix.

    This is intentionally source-derived and is used to prevent a grouped SRS from carrying
    Source IDs without the corresponding behavior being visible in the engineer-facing SWE.1.
    """
    out: dict[str, list[str]] = {}
    # V0.58: Canonical atomic units are authoritative for both explicit-ID and semantic-source modes.
    for req in requirement_data.get("requirements") or []:
        if not isinstance(req, dict):
            continue
        keys = [str(req.get("candidate_id") or "").strip(), str(req.get("srs_id") or "").strip()]
        for atom in req.get("source_backed_atomic_behaviors") or []:
            if not isinstance(atom, dict):
                continue
            rid = str(atom.get("source_requirement_id") or atom.get("source_semantic_unit_id") or "").strip()
            loc = str(atom.get("source_location") or "").strip()
            text = re.sub(r"\s+", " ", str(atom.get("behavior_text") or "")).strip()
            if not text:
                continue
            line = f"{rid} · {loc}: {text}" if rid and loc else (f"{rid}: {text}" if rid else text)
            for key in keys:
                if not key:
                    continue
                bucket = out.setdefault(key, [])
                if line not in bucket:
                    bucket.append(line)
    coverage = requirement_data.get("source_coverage") if isinstance(requirement_data.get("source_coverage"), dict) else {}
    for occ in coverage.get("source_requirement_occurrences") or []:
        if not isinstance(occ, dict) or occ.get("occurrence_type") != "declaration":
            continue
        if str(occ.get("coverage_status") or "") not in {"Covered", "Partially Covered"}:
            continue
        sid = str(occ.get("source_req_id") or "").strip()
        loc = str(occ.get("source_location") or occ.get("section_path") or "").strip()
        text = re.sub(r"\s+", " ", str(occ.get("source_excerpt") or "")).strip()
        if not text:
            continue
        line = f"{sid} · {loc}: {text}" if sid and loc else (f"{sid}: {text}" if sid else text)
        for key in list(occ.get("linked_candidate_ids") or []) + list(occ.get("linked_srs_ids") or []):
            key = str(key or "").strip()
            if not key:
                continue
            bucket = out.setdefault(key, [])
            if line not in bucket:
                bucket.append(line)
    return out


def _behavior_text_with_source_units(req: dict[str, Any], units: list[str]) -> str:
    base = _behavior_text(req)
    source_ids = _as_list(req.get("source_requirement_ids"))
    # Only expand the main view when grouping/semantic preservation needs it; otherwise keep the
    # normal engineer-facing output concise.
    needs_units = len(source_ids) > 1 or len(units) > 1 or bool(_as_list(req.get("source_semantic_unit_ids")))
    if not units or not needs_units:
        return base
    rendered = "\n".join(f"- {x}" for x in units)
    return f"{base}\n\n[Source-backed Atomic Behavior]\n{rendered}"


def _trace_sources(req: dict[str, Any]) -> list[tuple[str, str]]:
    """Return explicit source document/location pairs without inventing source IDs."""
    rows: list[tuple[str, str]] = []
    for item in req.get("source_evidence") or []:
        if isinstance(item, dict):
            document = str(item.get("document") or "").strip()
            location = str(item.get("location") or "").strip()
            evidence = str(item.get("text") or "").strip()
            source = document or "입력문서"
            detail = location or evidence
            pair = (source, detail)
        else:
            text = str(item).strip()
            pair = (text or "입력문서", "")
        if pair not in rows:
            rows.append(pair)
    return rows or [("입력문서에서 확인되지 않음", "")]


def _clarification_display(req: dict[str, Any]) -> str:
    ref = _text(req.get("clarification_reference"), empty="")
    if ref:
        return ref
    return _text(req.get("clarification_needed"), empty="")


def _main_swe1_fact_scope(req: dict[str, Any]) -> tuple[list[str], list[str]]:
    """Return Source-backed atomic texts partitioned by fact-level SWE.1 eligibility.

    V0.77 prevents a mixed-domain parent requirement from presenting SYS.5/non-SW facts as if
    they were part of Main SWE.1. The Canonical parent stays intact in Review Package; the Main
    view renders only eligible Source-backed atomic facts and carries excluded facts as review
    context.
    """
    alloc_by_fragment = {
        str(x.get("source_fact_fragment_id") or ""): x
        for x in (req.get("fact_level_allocations") or [])
        if isinstance(x, dict) and str(x.get("source_fact_fragment_id") or "")
    }
    if not alloc_by_fragment:
        return [], []
    eligible: list[str] = []
    excluded: list[str] = []
    for atom in (req.get("source_backed_atomic_behaviors") or []):
        if not isinstance(atom, dict):
            continue
        text = re.sub(r"\s+", " ", str(atom.get("behavior_text") or atom.get("source_fact") or "")).strip()
        if not text:
            continue
        fid = str(atom.get("source_fact_fragment_id") or "")
        alloc = alloc_by_fragment.get(fid) if fid else None
        if alloc is None:
            # When exact fragment allocation is unavailable, do not invent a scope split.
            continue
        target = eligible if str(alloc.get("swe1_eligibility") or "") == "Eligible" else excluded
        if text not in target:
            target.append(text)
    return eligible, excluded


def _main_swe1_requirement_text(req: dict[str, Any]) -> tuple[str, str]:
    eligible, excluded = _main_swe1_fact_scope(req)
    if not (eligible and excluded):
        return (_text(req.get("requirement"), empty=NOT_CONFIRMED) or NOT_CONFIRMED, "")
    requirement_text = "[Main SWE.1 Eligible Scope]\n" + "\n".join(f"- {x}" for x in eligible)
    excluded_text = "[Allocation Pending / Non-SWE.1 Source Facts]\n" + "\n".join(f"- {x}" for x in excluded)
    return requirement_text, excluded_text




def _field_token_set(text: str) -> set[str]:
    return {
        x.lower() for x in re.findall(r"[A-Za-z0-9가-힣_.+-]+", str(text or ""))
        if len(x) > 1
    }


def _materially_represents(field_text: str, atom_text: str) -> bool:
    a = _field_token_set(field_text)
    b = _field_token_set(atom_text)
    if not a or not b:
        return False
    return len(a & b) / max(1, min(len(b), 8)) >= 0.25


def _main_swe1_scoped_field(req: dict[str, Any], value: Any, *, empty: str = "") -> str:
    """V0.78 remove pending/non-SWE.1 fact content from every operative Main SWE.1 field.

    V0.77 scoped only the `요구사항 내역` cell.  Trigger/behavior/output/acceptance fields could
    still carry a pending SYS.5 fact from the full Canonical parent.  If an excluded fact is
    materially represented in a field, rebuild that field from Source-backed Eligible fact text
    that is also relevant to the field.  If no field-specific eligible clause exists, fall back to
    the complete Eligible fact set rather than leaking the excluded fact.
    """
    original = _text(value, empty=empty) or empty
    eligible, excluded = _main_swe1_fact_scope(req)
    if not (eligible and excluded) or not original:
        return original
    if not any(_materially_represents(original, x) for x in excluded):
        return original
    relevant = [x for x in eligible if _materially_represents(original, x)]
    safe = relevant or eligible
    if not safe:
        return empty
    return "[Main SWE.1 Eligible Scope]\n" + "\n".join(f"- {x}" for x in safe)

def build_swe1_records(requirement_data: dict[str, Any], *, include_review_required: bool = False) -> list[dict[str, str]]:
    records: list[dict[str, str]] = []
    requirements = requirement_data.get("requirements") or []
    source_units_by_ref = _source_behavior_units(requirement_data)
    for idx, req in enumerate(requirements, start=1):
        if not isinstance(req, dict):
            continue
        eligibility = str(req.get("swe1_eligibility") or "Eligible")
        # Compatibility API: callers that explicitly request the historical SWE.1-only record set
        # still receive Eligible rows only. V0.79 exporter views pass include_review_required=True
        # so REVIEW_REQUIRED/HOLD requirements remain visible in the Main Requirement Specification.
        if eligibility != "Eligible" and not (include_review_required and eligibility == "Review Needed"):
            continue
        main_requirement_text, excluded_scope_text = _main_swe1_requirement_text(req)
        other_text = _other_materials(req)
        if excluded_scope_text:
            other_text = (other_text + "\n\n" if other_text else "") + excluded_scope_text
        records.append({
            "SRS ID": _text(req.get("srs_id"), empty=f"SRS_{idx:03d}") or f"SRS_{idx:03d}",
            "상위 기능": _text(req.get("function_name"), empty="미분류") or "미분류",
            "분류": "검토 필요" if eligibility == "Review Needed" else _normalize_classification(req.get("category")),
            "검토 상태": (
                "사람의 판단 필요 (REVIEW_REQUIRED/HOLD)"
                if eligibility == "Review Needed" or str(req.get("allocation_status") or "") == "SW_BEHAVIOR_INFERRED_REVIEW_REQUIRED"
                else "Source-backed Requirement"
            ),
            "요구사항 내역": main_requirement_text,
            "동작 조건 / Trigger": _main_swe1_scoped_field(req, req.get("activation_trigger"), empty=NOT_CONFIRMED) or NOT_CONFIRMED,
            "작동 명세 정의": _main_swe1_scoped_field(
                req,
                _behavior_text_with_source_units(
                    req,
                    source_units_by_ref.get(str(req.get("candidate_id") or ""), [])
                    or source_units_by_ref.get(str(req.get("srs_id") or ""), []),
                ),
                empty=NO_DETAIL,
            ) or NO_DETAIL,
            "사전 조건": _main_swe1_scoped_field(
                req,
                _text(req.get("preconditions"), empty="") or _text(req.get("system_input_preconditions"), empty=NOT_CONFIRMED),
                empty=NOT_CONFIRMED,
            ) or NOT_CONFIRMED,
            "예상 결과": _main_swe1_scoped_field(req, req.get("output"), empty=NOT_CONFIRMED) or NOT_CONFIRMED,
            "검증 기준": _main_swe1_scoped_field(req, req.get("acceptance_criteria"), empty=NOT_CONFIRMED) or NOT_CONFIRMED,
            "출처 / Traceability": _format_evidence(req.get("source_evidence")),
            "기타": other_text,
            # The following values remain internal only. They are intentionally not exported as final SWE.1 fields.
            "_classification_basis": _text(req.get("classification_basis"), empty=""),
            "_evaluation_method": _text(req.get("evaluation_method"), empty=""),
            "_exception_conditions": _text(req.get("exception_conditions"), empty=""),
            "_clarification_needed": _clarification_display(req),
            "_candidate_id": _text(req.get("candidate_id"), empty=""),
            "_source_requirement_ids": _text(req.get("source_requirement_ids"), empty=""),
            "_applicability": req.get("applicability") if isinstance(req.get("applicability"), dict) else {},
            "_external_dependencies": _text(req.get("external_dependencies"), empty=""),
            "_tbd_items": _text(req.get("tbd_items"), empty=""),
            "_conflicts": _text(req.get("conflicts"), empty=""),
            "_open_issue_ids": _text(req.get("open_issue_ids"), empty=""),
            "_requirement_status": _text(req.get("requirement_status"), empty="Draft") or "Draft",
            "_verification_constraints": _text(req.get("verification_constraints"), empty=""),
            "_knowledge_state": _text(req.get("knowledge_state"), empty=""),
            "_source_semantic_unit_ids": _text(req.get("source_semantic_unit_ids"), empty=""),
            "_requirement_level": _text(req.get("requirement_level"), empty=""),
            "_allocation_status": _text(req.get("allocation_status"), empty=""),
            "_swe1_eligibility": _text(req.get("swe1_eligibility"), empty=""),
            "_swe6_eligibility": _text(req.get("swe6_eligibility"), empty=""),
            "_verification_domain": _text(req.get("verification_domain"), empty=""),
            "_source_behavior_units": "\n".join(
                source_units_by_ref.get(str(req.get("candidate_id") or ""), [])
                or source_units_by_ref.get(str(req.get("srs_id") or ""), [])
            ),
        })
    return records



def build_engineering_records(requirement_data: dict[str, Any]) -> list[dict[str, str]]:
    """V0.80 Canonical Engineering Requirement Main View.

    Every Canonical requirement is visible here regardless of SWE.1 eligibility.  Domain-specific
    SYS.1/SWE.1 views are projections of this single record set.
    """
    records: list[dict[str, str]] = []
    source_units_by_ref = _source_behavior_units(requirement_data)
    for idx, req in enumerate(requirement_data.get("requirements") or [], start=1):
        if not isinstance(req, dict):
            continue
        sid = _text(req.get("srs_id"), empty=f"SRS_{idx:03d}") or f"SRS_{idx:03d}"
        review_state = _text(req.get("review_status"), empty="")
        if not review_state:
            review_state = "HUMAN_DECISION_REQUIRED" if req.get("human_decision_required") else "SOURCE_BACKED"
        records.append({
            "SRS ID": sid,
            "상위 기능": _text(req.get("function_name"), empty="미분류") or "미분류",
            "분류": _normalize_classification(req.get("category")),
            "Engineering Domain": ", ".join(str(x) for x in (req.get("engineering_domains") or [])) or _text(req.get("requirement_level"), empty="Unallocated Engineering"),
            "Requirement Level": _text(req.get("requirement_level")),
            "Allocation Status": _text(req.get("allocation_status")),
            "SYS.1 Eligibility": _text(req.get("sys1_eligibility"), empty="Not Applicable"),
            "SWE.1 Eligibility": _text(req.get("swe1_eligibility"), empty="Not Applicable"),
            "Review Status": review_state,
            "Canonical State": _text(req.get("canonical_state"), empty="CANONICAL_SOURCE_BACKED"),
            "요구사항 내역": _text(req.get("requirement"), empty=NOT_CONFIRMED) or NOT_CONFIRMED,
            "동작 조건 / Trigger": _text(req.get("activation_trigger"), empty=NOT_CONFIRMED) or NOT_CONFIRMED,
            "작동 명세 정의": _behavior_text_with_source_units(
                req,
                source_units_by_ref.get(str(req.get("candidate_id") or ""), [])
                or source_units_by_ref.get(sid, []),
            ),
            "사전 조건": _text(req.get("preconditions"), empty="") or _text(req.get("system_input_preconditions"), empty=NOT_CONFIRMED),
            "예상 결과": _text(req.get("output"), empty=NOT_CONFIRMED) or NOT_CONFIRMED,
            "검증 기준": _text(req.get("acceptance_criteria"), empty=NOT_CONFIRMED) or NOT_CONFIRMED,
            "출처 / Traceability": _format_evidence(req.get("source_evidence")),
            "Hold Reason": _text(req.get("hold_reason")),
            "기타": _other_materials(req),
            "_verification_domains": " | ".join(str(x) for x in (req.get("verification_domains") or [])),
            "_sys5_eligibility": _text(req.get("sys5_eligibility")),
            "_swe6_eligibility": _text(req.get("swe6_eligibility")),
            "_human_decision_required": "Y" if req.get("human_decision_required") else "N",
            "_official_release_eligible": "Y" if req.get("official_release_eligible") else "N",
            # Preserve the same internal review metadata used by the legacy SWE.1 view so
            # Engineering Main can still drive Review/Allocation annexes without hiding
            # non-SWE or REVIEW_REQUIRED Canonical requirements.
            "_exception_conditions": _text(req.get("exception_conditions")),
            "_clarification_needed": _text(req.get("clarification_needed")),
            "_external_dependencies": _text(req.get("external_dependencies")),
            "_tbd_items": _text(req.get("tbd_items")),
            "_conflicts": _text(req.get("conflicts")),
            "_open_issue_ids": _text(req.get("open_issue_ids")),
            "_source_requirement_ids": _text(req.get("source_requirement_ids")),
            "_source_semantic_unit_ids": _text(req.get("source_semantic_unit_ids")),
            "_allocation_status": _text(req.get("allocation_status")),
            "_swe1_eligibility": _text(req.get("swe1_eligibility")),
            "_swe6_eligibility": _text(req.get("swe6_eligibility")),
            "_requirement_level": _text(req.get("requirement_level")),
            "_verification_domain": _text(req.get("verification_domain")),
            "_applicability": req.get("applicability") if isinstance(req.get("applicability"), dict) else {},
            "_source_behavior_units": _text(req.get("source_behavior_units")),
            "_requirement_status": _text(req.get("requirement_status")),
            "_knowledge_state": _text(req.get("knowledge_state")),
        })
    return records


class SWE1Exporter:
    def __init__(self, output_dir: Path):
        self.output_dir = Path(output_dir).resolve()
        self.output_dir.mkdir(parents=True, exist_ok=True)

    def _filename(self, prefix: str, source_document: Path, suffix: str) -> Path:
        return self.output_dir / output_filename(
            prefix, _safe_stem(Path(source_document).name), suffix
        )

    def export_word(self, source_document: Path, requirement_data: dict[str, Any]) -> Path:
        records = build_engineering_records(requirement_data)
        all_requirements = [x for x in (requirement_data.get("requirements") or []) if isinstance(x, dict)]
        if not all_requirements:
            raise ValueError("SWE.1 Word로 내보낼 Requirement/Review Object가 없습니다.")
        out_path = self._filename("SWE.1 요구사항 정리", source_document, "docx")
        doc = Document()
        doc.styles["Normal"].font.name = "Malgun Gothic"
        doc.styles["Normal"].font.size = Pt(9)
        title = doc.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = title.add_run("SWE.1 Software Requirements View — 소프트웨어 요구사항 전문 View")
        run.bold = True
        run.font.size = Pt(18)
        meta = doc.add_table(rows=4, cols=2)
        meta.style = "Table Grid"
        for row, (key, value) in enumerate([
            ("원본 문서", Path(source_document).name),
            ("생성 일시", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
            ("Canonical Engineering Requirement 개수", str(len(records))),
            ("작성 기준", "Requirement Studio V0.92 — SWE.1 Software View + E2E Requirements Master"),
        ]):
            meta.cell(row, 0).text = key
            meta.cell(row, 1).text = value
        note = doc.add_paragraph()
        note_run = note.add_run(
            "※ 본 문서는 SWE.1 관점의 Software View입니다. 최종 사람용 Engineering Master는 E2E 요구사항 명세서(Word/Excel)를 사용합니다. "
            "V0.92에서도 동일 SRS를 다시 반복하던 Word Review Annex는 출력하지 않으며, Review/Allocation 상세는 E2E 요구사항 Excel의 `3_검토필요` 및 Review Package JSON에 보존합니다. "
            "입력 Source에 없는 사실은 임의 생성하지 않고 REVIEW_REQUIRED/HOLD 상태와 Source Trace는 계속 유지합니다."
        )
        note_run.font.size = Pt(8.5)
        note_run.font.color.rgb = RGBColor(90, 98, 110)

        section = doc.sections[0]
        usable_width = int(section.page_width - section.left_margin - section.right_margin)
        left_width = int(usable_width * 0.24)
        right_width = usable_width - left_width

        current_feature = None
        for record in records:
            feature = record["상위 기능"]
            if feature != current_feature:
                doc.add_heading(feature, level=1)
                current_feature = feature
            # Keep the stable SRS identifier visible in normal paragraph flow as well as the table.
            # This supports text-based review/search tools and makes REVIEW_REQUIRED items explicit.
            doc.add_heading(record["SRS ID"], level=2)
            table = doc.add_table(rows=len(ENGINEERING_WORD_FIELDS), cols=2)
            table.style = "Table Grid"
            table.autofit = False
            table.columns[0].width = left_width
            table.columns[1].width = right_width
            for row_idx, field in enumerate(ENGINEERING_WORD_FIELDS):
                left, right = table.cell(row_idx, 0), table.cell(row_idx, 1)
                left.width, right.width = left_width, right_width
                left.text, right.text = field, record[field]
                left.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for paragraph in left.paragraphs:
                    for r in paragraph.runs:
                        r.bold = True
                        r.font.size = Pt(8.5)
                for paragraph in right.paragraphs:
                    for r in paragraph.runs:
                        r.font.size = Pt(8.5)
            doc.add_paragraph()

        # V0.92: human-facing duplicate Review Annex sections were retired.
        # Review/Allocation details remain in the E2E Requirements `3_검토필요` sheet and
        # the Review Package JSON canonical core, so no Source-backed information is lost.
        doc.save(out_path)
        return out_path

    def export_excel(self, source_document: Path, requirement_data: dict[str, Any]) -> Path:
        records = build_swe1_records(requirement_data, include_review_required=True)
        all_requirements = [x for x in (requirement_data.get("requirements") or []) if isinstance(x, dict)]
        if not all_requirements:
            raise ValueError("SWE.1 Excel로 내보낼 Requirement/Review Object가 없습니다.")
        out_path = self._filename("SWE.1 요구사항 명세", source_document, "xlsx")
        wb = Workbook()
        ws0 = wb.active
        ws0.title = "00_Overview"
        ws1 = wb.create_sheet("01_SWE1")
        ws2 = wb.create_sheet("02_Traceability")
        ws3 = wb.create_sheet("03_Applicability")
        ws4 = wb.create_sheet("04_Open_Issues_Gaps")
        ws5 = wb.create_sheet("05_Coverage")
        ws6 = wb.create_sheet("06_Requirement_Review_Details")
        ws7 = wb.create_sheet("07_Allocation_Review_Index")
        ws8 = wb.create_sheet("08_Engineering_Main")
        ws9 = wb.create_sheet("09_SYS1_System_View")
        ws10 = wb.create_sheet("10_SWE1_Software_View")
        ws11 = wb.create_sheet("11_Human_Decision_View")

        blue, light, border_color = "2F75B5", "D9EAF7", "D0D7DE"
        thin = Side(style="thin", color=border_color)
        header_fill = PatternFill("solid", fgColor=blue)
        sub_fill = PatternFill("solid", fgColor=light)
        white_font, bold = Font(color="FFFFFF", bold=True), Font(bold=True)
        wrap = Alignment(vertical="top", wrap_text=True)
        center = Alignment(horizontal="center", vertical="center", wrap_text=True)
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        for row in [
            ["항목", "내용"],
            ["문서", "SWE.1 Software Requirements View (최종 Engineering Master는 E2E 요구사항 명세서)"],
            ["원본 문서", Path(source_document).name],
            ["생성 일시", datetime.now().strftime("%Y-%m-%d %H:%M:%S")],
            ["Main Requirement SRS 개수 (Eligible + Review Required)", len(records)],
            ["작성 기준", "Requirement Studio V0.92 — SWE.1 Software View + E2E Requirements Master"],
            ["핵심 원칙", "Not SWE.1 != Not Requirement. 08_Engineering_Main은 모든 Canonical Engineering Requirement를 보존하며 09_SYS1/10_SWE1은 domain view이다. REVIEW_REQUIRED/HOLD 항목은 Main + Review View에 유지한다."],
            ["Source 원칙", "입력문서에 없는 사실은 만들지 않으며 불명확/미정/상충 정보는 Review View 및 Canonical에서 추적"],
        ]:
            ws0.append(row)
        for cell in ws0[1]:
            cell.fill, cell.font, cell.alignment = header_fill, white_font, center
        for row in ws0.iter_rows(min_row=2):
            row[0].font = bold
            for cell in row:
                cell.alignment, cell.border = wrap, border
        ws0.column_dimensions["A"].width, ws0.column_dimensions["B"].width = 18, 90
        ws0.freeze_panes = "A2"

        headers = FINAL_FIELDS
        ws1.append(headers)
        for record in records:
            ws1.append([record[h] for h in headers])
        for cell in ws1[1]:
            cell.fill, cell.font, cell.alignment, cell.border = header_fill, white_font, center, border
        for row in ws1.iter_rows(min_row=2):
            for cell in row:
                cell.alignment, cell.border = wrap, border
        ws1.freeze_panes, ws1.auto_filter.ref = "A2", ws1.dimensions
        widths = [12, 22, 11, 28, 48, 32, 46, 30, 32, 32, 56, 40]
        for idx, width in enumerate(widths, start=1):
            ws1.column_dimensions[get_column_letter(idx)].width = width
        for row in ws1.iter_rows(min_row=2):
            row[0].fill, row[0].font = sub_fill, bold

        # End-to-end traceability view: Source -> SWE.1 SRS -> SWE.6 TC.
        trace_headers = [
            "입력 Source",
            "Source Location / Requirement",
            "SWE.1 SRS ID",
            "SWE.1 Requirement",
            "SWE.6 Test Case ID",
            "상태",
        ]
        ws2.append(trace_headers)
        tc_by_srs: dict[str, list[str]] = {}
        bundle = requirement_data.get("finalized_verification_bundle") if isinstance(requirement_data.get("finalized_verification_bundle"), dict) else {}
        swe6_cases_for_trace = bundle.get("swe6_cases") if bundle.get("schema_version") == "REQ-STUDIO-VERIFICATION-BUNDLE-0.82" else build_swe6_cases(requirement_data)
        for case in swe6_cases_for_trace or []:
            if isinstance(case, dict):
                tc_by_srs.setdefault(str(case.get("srs_id") or ""), []).append(str(case.get("tc_id") or ""))
        requirements = [x for x in all_requirements if str(x.get("swe1_eligibility") or "Eligible") in {"Eligible", "Review Needed"}]
        for idx, req in enumerate(requirements, start=1):
            srs_id = _text(req.get("srs_id"), empty=f"SRS_{idx:03d}") or f"SRS_{idx:03d}"
            requirement = _text(req.get("requirement"), empty=NOT_CONFIRMED) or NOT_CONFIRMED
            tc_ids = [x for x in tc_by_srs.get(srs_id, []) if x]
            tc_text = ", ".join(tc_ids)
            status = "연결" if tc_ids else "검토 필요"
            for source, location in _trace_sources(req):
                ws2.append([source, location, srs_id, requirement, tc_text, status])
        for cell in ws2[1]:
            cell.fill, cell.font, cell.alignment, cell.border = header_fill, white_font, center, border
        for row in ws2.iter_rows(min_row=2):
            for cell in row:
                cell.alignment, cell.border = wrap, border
        ws2.freeze_panes = "A2"
        for idx, width in enumerate([34, 50, 14, 62, 24, 14], start=1):
            ws2.column_dimensions[get_column_letter(idx)].width = width

        # V0.47 Review View sheets: preserve review-critical Canonical information without bloating 01_SWE1.
        app_headers = ["SRS ID", "Source Requirement ID", "Vehicle Lines", "Baseline Versions", "Feature Variants", "Enable Conditions", "Exclusion Conditions", "Knowledge State", "Source Evidence"]
        ws3.append(app_headers)
        for record in records:
            app = record.get("_applicability") or {}
            ws3.append([
                record["SRS ID"], record.get("_source_requirement_ids") or "",
                _text(app.get("vehicle_lines")), _text(app.get("baseline_versions")), _text(app.get("feature_variants")),
                _text(app.get("enable_conditions")), _text(app.get("exclusion_conditions")),
                _text(app.get("knowledge_state")), _text(app.get("source_evidence")),
            ])

        issue_headers = ["Type", "ID / SRS", "Related SRS", "Status", "Detail", "Source / Link"]
        ws4.append(issue_headers)
        for record in records:
            srs = record["SRS ID"]
            for typ, value in [
                ("Exception", record.get("_exception_conditions")),
                ("Clarification Needed", record.get("_clarification_needed")),
                ("External Dependency", record.get("_external_dependencies")),
                ("TBD", record.get("_tbd_items")),
                ("Conflict", record.get("_conflicts")),
                ("Open Issue", record.get("_open_issue_ids")),
            ]:
                if value:
                    ws4.append([typ, srs, srs, record.get("_requirement_status") or "Draft", value, record.get("출처 / Traceability") or ""])
        for gap in [x for x in (requirement_data.get("gaps") or []) if isinstance(x, dict)]:
            ws4.append([
                "Gap", _text(gap.get("gap_id")), _text(gap.get("related_srs_ids")),
                "Blocked" if gap.get("blocking_for_verification") else "Review Needed",
                _text(gap.get("description") or gap.get("gap") or gap.get("reason")),
                _text(gap.get("related_source_requirement_ids") or gap.get("related_source_occurrence_ids")),
            ])
        for conflict in [x for x in (requirement_data.get("conflict_register") or []) if isinstance(x, dict)]:
            ws4.append([
                conflict.get("classification") or "Conflict", conflict.get("conflict_id") or "",
                _text(conflict.get("affected_srs_ids")), conflict.get("status") or "Open",
                _text(conflict.get("title") or conflict.get("description")),
                _text({"source_a": conflict.get("source_a"), "source_b": conflict.get("source_b")}),
            ])

        coverage = requirement_data.get("source_coverage") if isinstance(requirement_data.get("source_coverage"), dict) else {}
        cov_headers = ["Occurrence ID", "Occurrence Type", "Source REQ ID", "Source Location", "Kind", "Coverage Status", "Candidate IDs", "SRS IDs", "Gap/Issue IDs", "Disposition Reason", "Source Excerpt", "Content Hash"]
        ws5.append(cov_headers)
        for occ in coverage.get("source_requirement_occurrences") or []:
            if not isinstance(occ, dict):
                continue
            ws5.append([
                occ.get("occurrence_id"), occ.get("occurrence_type"), occ.get("source_req_id"), occ.get("source_location"), occ.get("source_kind"),
                occ.get("coverage_status"), _text(occ.get("linked_candidate_ids")), _text(occ.get("linked_srs_ids")),
                _text(occ.get("linked_gap_issue_ids")), occ.get("disposition_reason"), occ.get("source_excerpt"), occ.get("content_hash"),
            ])

        detail_headers = ["SRS ID", "Candidate ID", "Source Requirement ID", "Requirement Status", "Knowledge State", "Exception Conditions", "Clarification Needed", "External Dependencies", "TBD Items", "Conflicts", "Verification Constraints", "Source Evidence", "SWE.6 TC ID"]
        ws6.append(detail_headers)
        for record in records:
            tc_ids = ", ".join(x for x in tc_by_srs.get(record["SRS ID"], []) if x)
            ws6.append([
                record["SRS ID"], record.get("_candidate_id"), record.get("_source_requirement_ids"),
                record.get("_requirement_status"), record.get("_knowledge_state"), record.get("_exception_conditions"),
                record.get("_clarification_needed"), record.get("_external_dependencies"), record.get("_tbd_items"),
                record.get("_conflicts"), record.get("_verification_constraints"), record.get("출처 / Traceability"), tc_ids,
            ])

        # V0.92: keep only a lightweight allocation-review index in the SWE.1 workbook.
        # Full review detail is intentionally not repeated here; it is preserved in the E2E
        # Requirements `3_검토필요` sheet and Review Package JSON.
        alloc_headers = ["SRS ID", "SWE.1 적격성", "SWE.6 적격성", "할당 상태", "검증 영역", "상세 검토 위치"]
        ws7.append(alloc_headers)
        for req in all_requirements:
            if str(req.get("swe1_eligibility") or "Eligible") == "Eligible":
                continue
            ws7.append([
                _text(req.get("srs_id")), _text(req.get("swe1_eligibility")), _text(req.get("swe6_eligibility")),
                _text(req.get("allocation_status")), _text(req.get("verification_domain")),
                "E2E 요구사항 명세서 > 3_검토필요 / Review Package JSON",
            ])

        # V0.80 Canonical Engineering Main + SYS/SWE domain views.
        engineering_records = build_engineering_records(requirement_data)
        eng_headers = [
            "SRS ID", "Engineering Domain", "Requirement Level", "Allocation Status",
            "SYS.1 Eligibility", "SWE.1 Eligibility", "SYS.5 Eligibility", "SWE.6 Eligibility",
            "Verification Domains", "Review Status", "Canonical State", "Human Decision Required",
            "Official Release Eligible", "Requirement", "Source / Traceability", "Hold Reason",
        ]
        ws8.append(eng_headers)
        for rec in engineering_records:
            ws8.append([
                rec["SRS ID"], rec["Engineering Domain"], rec["Requirement Level"], rec["Allocation Status"],
                rec["SYS.1 Eligibility"], rec["SWE.1 Eligibility"], rec.get("_sys5_eligibility", ""), rec.get("_swe6_eligibility", ""),
                rec.get("_verification_domains", ""), rec["Review Status"], rec["Canonical State"], rec.get("_human_decision_required", "N"),
                rec.get("_official_release_eligible", "N"), rec["요구사항 내역"], rec["출처 / Traceability"], rec["Hold Reason"],
            ])

        sys_headers = ["SRS ID", "Engineering Domain", "SYS.1 Eligibility", "Review Status", "Requirement", "Verification Domains", "Source / Traceability", "Hold Reason"]
        ws9.append(sys_headers)
        for rec in engineering_records:
            if rec["SYS.1 Eligibility"] not in {"Eligible", "Review Needed"}:
                continue
            ws9.append([rec["SRS ID"], rec["Engineering Domain"], rec["SYS.1 Eligibility"], rec["Review Status"], rec["요구사항 내역"], rec.get("_verification_domains", ""), rec["출처 / Traceability"], rec["Hold Reason"]])

        sw_headers = ["SRS ID", "Engineering Domain", "SWE.1 Eligibility", "Review Status", "Requirement", "Verification Domains", "Source / Traceability", "Hold Reason"]
        ws10.append(sw_headers)
        for rec in engineering_records:
            if rec["SWE.1 Eligibility"] not in {"Eligible", "Review Needed"}:
                continue
            ws10.append([rec["SRS ID"], rec["Engineering Domain"], rec["SWE.1 Eligibility"], rec["Review Status"], rec["요구사항 내역"], rec.get("_verification_domains", ""), rec["출처 / Traceability"], rec["Hold Reason"]])

        decision_headers = ["SRS ID", "Review Status", "Canonical State", "Human Decision Required", "Allocation Status", "SYS.1", "SWE.1", "SYS.5", "SWE.6", "Requirement", "Hold Reason", "Source / Traceability"]
        ws11.append(decision_headers)
        for rec in engineering_records:
            if rec.get("_human_decision_required") != "Y" and rec["Review Status"] not in {"HUMAN_DECISION_REQUIRED", "REVIEW_REQUIRED"}:
                continue
            ws11.append([
                rec["SRS ID"], rec["Review Status"], rec["Canonical State"], rec.get("_human_decision_required", "N"), rec["Allocation Status"],
                rec["SYS.1 Eligibility"], rec["SWE.1 Eligibility"], rec.get("_sys5_eligibility", ""), rec.get("_swe6_eligibility", ""),
                rec["요구사항 내역"], rec["Hold Reason"], rec["출처 / Traceability"],
            ])

        for review_ws in (ws3, ws4, ws5, ws6, ws7):
            for cell in review_ws[1]:
                cell.fill, cell.font, cell.alignment, cell.border = header_fill, white_font, center, border
            for row in review_ws.iter_rows(min_row=2):
                for cell in row:
                    cell.alignment, cell.border = wrap, border
            review_ws.freeze_panes = "A2"
            review_ws.auto_filter.ref = review_ws.dimensions
        for idx, width in enumerate([14,24,28,22,28,30,30,18,50], start=1): ws3.column_dimensions[get_column_letter(idx)].width = width
        for idx, width in enumerate([20,18,24,18,70,70], start=1): ws4.column_dimensions[get_column_letter(idx)].width = width
        for idx, width in enumerate([18,18,20,30,16,22,24,24,24,52,70,20], start=1): ws5.column_dimensions[get_column_letter(idx)].width = width
        for idx, width in enumerate([14,18,24,18,16,50,50,50,40,50,45,70,24], start=1): ws6.column_dimensions[get_column_letter(idx)].width = width
        for idx, width in enumerate([14,18,18,34,34,64], start=1): ws7.column_dimensions[get_column_letter(idx)].width = width

        for review_ws in (ws8, ws9, ws10, ws11):
            for cell in review_ws[1]:
                cell.fill, cell.font, cell.alignment, cell.border = header_fill, white_font, center, border
            for row in review_ws.iter_rows(min_row=2):
                for cell in row:
                    cell.alignment, cell.border = wrap, border
            review_ws.freeze_panes = "A2"
            review_ws.auto_filter.ref = review_ws.dimensions
        for idx, width in enumerate([14,28,26,34,18,18,18,18,54,28,30,22,22,72,72,60], start=1): ws8.column_dimensions[get_column_letter(idx)].width = width
        for idx, width in enumerate([14,28,18,26,72,54,72,60], start=1): ws9.column_dimensions[get_column_letter(idx)].width = width
        for idx, width in enumerate([14,28,18,26,72,54,72,60], start=1): ws10.column_dimensions[get_column_letter(idx)].width = width
        for idx, width in enumerate([14,26,30,22,34,16,16,16,16,72,60,72], start=1): ws11.column_dimensions[get_column_letter(idx)].width = width

        for ws in (ws0, ws1, ws2, ws3, ws4, ws5, ws6, ws7, ws8, ws9, ws10, ws11):
            ws.sheet_view.showGridLines = False
        wb.save(out_path)
        return out_path
