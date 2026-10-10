from pathlib import Path
from docx import Document
from openpyxl import load_workbook

from core.e2e_exporter import E2EExporter
from core.swe1_exporter import SWE1Exporter
from tests.test_v082_integrated_specs_single_truth import _mode_fixture


def test_v089_swe1_word_removes_duplicate_review_annex_but_keeps_main_requirements(tmp_path: Path):
    data = _mode_fixture()
    src = tmp_path / "source.docx"
    Document().save(src)
    out = SWE1Exporter(tmp_path).export_word(src, data)
    doc = Document(out)
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "SWE.1 Software Requirements View" in text
    assert "Review Annex — Applicability / Issues / Dependencies" not in text
    assert "Allocation Review Annex — Pending SWE.1 Allocation" not in text
    assert "Cross-domain Allocation Annex — Source Facts Outside SWE.1" not in text
    assert "SRS_004" in text and "SRS_005" in text


def test_v089_swe1_excel_allocation_review_is_pointer_index_not_duplicate_requirement_text(tmp_path: Path):
    data = _mode_fixture()
    src = tmp_path / "source.docx"
    src.write_text("x", encoding="utf-8")
    out = SWE1Exporter(tmp_path).export_excel(src, data)
    wb = load_workbook(out, read_only=True, data_only=False)
    try:
        assert "07_Allocation_Review_Index" in wb.sheetnames
        assert "07_Allocation_Annex" not in wb.sheetnames
        ws = wb["07_Allocation_Review_Index"]
        headers = [str(ws.cell(1, c).value or "") for c in range(1, ws.max_column + 1)]
        assert headers == ["SRS ID", "SWE.1 적격성", "SWE.6 적격성", "할당 상태", "검증 영역", "상세 검토 위치"]
        for r in range(2, ws.max_row + 1):
            assert "E2E 요구사항 명세서 > 3_검토필요 / Review Package JSON" == str(ws.cell(r, 6).value or "")
    finally:
        wb.close()


def test_v089_e2e_evaluation_row2_groups_and_row3_korean_headers(tmp_path: Path):
    data = _mode_fixture()
    src = tmp_path / "source.docx"
    src.write_text("x", encoding="utf-8")
    out = E2EExporter(tmp_path).export_evaluation_excel(src, data)
    wb = load_workbook(out, read_only=False, data_only=False)
    try:
        ws = wb["2_E2E_평가케이스"]
        merged = {str(rng) for rng in ws.merged_cells.ranges}
        assert {"A2:F2", "G2:H2", "I2:L2", "M2:P2", "Q2:S2", "T2:W2"}.issubset(merged)
        assert [ws["A2"].value, ws["G2"].value, ws["I2"].value, ws["M2"].value, ws["Q2"].value, ws["T2"].value] == [
            "종류", "사전 준비", "입력 / 수행", "출력", "판단", "실행 결과"
        ]
        headers = [str(ws.cell(3, c).value or "") for c in range(1, ws.max_column + 1)]
        for label in ("Test Case ID", "요구사항 ID", "검증 영역", "테스트 목적", "사전 조건", "수행 절차", "출력 변수 / 관찰 대상", "PASS 기준", "TC 검토 상태", "실제 출력값", "판정 (PASS/FAIL)"):
            assert label in headers
        for removed in ("평가 객체 유형","평가 객체 ID","관련 원본 객체 ID","상위 요구사항 ID","할당 상태","SWE.6 적격성","실행 준비도","사람 검토 필요"):
            assert removed not in headers
        # Internal EVAL/Raw IDs remain stable in the Trace/Audit sheet.
        tr=wb["4_요구사항-평가추적성"]
        th=[str(tr.cell(3,c).value or "") for c in range(1,tr.max_column+1)]
        eid_col=th.index("E2E 평가 객체 ID")+1
        values=[str(tr.cell(r,eid_col).value or "") for r in range(4,tr.max_row+1)]
        assert all(v.startswith("EVAL_") for v in values if v)
    finally:
        wb.close()
