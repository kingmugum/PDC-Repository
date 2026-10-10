from pathlib import Path
from zipfile import ZipFile

from docx import Document
from openpyxl import load_workbook

from core.e2e_exporter import E2EExporter


def _req_fixture():
    req={
        "srs_id":"SRS_001",
        "candidate_id":"REQ_001",
        "function_name":"실시간 감시모드 옵션 및 전원 조건 판단",
        "requirement":"1) SVM의 실시간 감시 모드 적용 여부: SVM_Typ2의 bit 1 값 2) ICMU의 실시간 감시 모드 적용 및 On/Off 여부: ICMU_RealTimeMonSet 신호 3) ICMU_RealTimeMonSet 값 송신 조건 0x0:Default 0x1:Off 0x2:On 0x3:Invalid",
        "activation_trigger":"ACC OFF 조건에서 기능 요청이 수신되면 동작해야 한다. IGN3 ON 조건도 Source에 정의된 경우 적용한다.",
        "processing_action":"요청 조건을 판단하고 Source-defined 상태를 처리해야 한다.",
        "preconditions":["통신 정상","EOL coding 완료"],
        "output":"Source-defined 상태를 출력해야 한다.",
        "acceptance_criteria":"요청 조건과 출력 상태가 Source 정의와 일치해야 한다.",
        "source_evidence":[{"location":"Page 5, 2.1","text":"SVM_Typ2 bit 1 / ICMU_RealTimeMonSet 0x0 Default 0x1 Off 0x2 On 0x3 Invalid"}],
        "source_backed_atomic_behaviors":[
            {"source_fact_fragment_id":"F1","source_semantic_unit_id":"U1","source_location":"Page 5","behavior_text":"SVM_Typ2 bit 1 값으로 적용 여부를 판단한다."},
            {"source_fact_fragment_id":"F2","source_semantic_unit_id":"U1","source_location":"Page 5","behavior_text":"ICMU_RealTimeMonSet 0x2는 On 상태이다."},
        ],
        "engineering_domains":["System","Software"],
        "requirement_level":"System/Software",
        "allocation_status":"SW_BEHAVIOR_INFERRED_REVIEW_REQUIRED",
        "sys1_eligibility":"Eligible",
        "swe1_eligibility":"Eligible",
        "sys5_eligibility":"Review Needed",
        "swe6_eligibility":"Eligible",
        "verification_domain":"SWE.6 Software Qualification",
        "review_status":"HUMAN_DECISION_REQUIRED",
        "canonical_state":"CANONICAL_REVIEW_REQUIRED",
        "human_decision_required":True,
        "hold_reason":"Software allocation pending human review",
        "category":"기능",
        "scenario_candidate_id":"SCN_001",
    }
    return req


def _capture_req():
    source='''2.9 실외 이미지 캡처\nHU는 ADAS_PRK에 신호(SVM_CaptureModeViewRequest)를 송신하여 각 뷰 모드에 대해 Front->Rear->Left->Right 순서로 영상을 요청해야 한다.\nHU는 ADAS_PRK로부터 신호(SVM_CaptureModeViewState)를 수신 받아 Front->Rear->Left->Right 순서의 영상 송출 여부를 확인한다.\nSVM_CaptureModeViewRequest값 송신 조건\n0x1:Front View Printing Request 전방 요청\n0x2:Rear View Printing Request 후방 요청\n0x3:Left View Printing Request 좌측 요청\n0x4:Right View Printing Request 우측 요청\n0x5:Top View Printing Request 탑뷰 요청\n0x7:Invalid\nSVM_CaptureModeViewState값 송신 조건\n0x1:Front 전방 응답\n0x2:Rear 후방 응답\n0x3:Left 좌측 응답\n0x4:Right 우측 응답\n0x5:Top 탑뷰 응답\n0x6:Camera Input Error - HU는 이 값에 대한 동작이 이 절에 정의되지 않는다.\n0x7:Invalid - HU는 이 값에 대한 동작이 이 절에 정의되지 않는다.\n'''
    req={
        "srs_id":"SRS_009","candidate_id":"REQ_009","scenario_candidate_id":"SCN_009",
        "function_name":"실외 이미지 순차 캡처","requirement":"실외 이미지 순차 캡처","category":"기능",
        "processing_action":"실외 이미지 캡처 수행","output":"각 뷰의 영상 송출 확인","acceptance_criteria":"각 뷰의 영상 송출 확인",
        "allocation_status":"SW_BEHAVIOR_INFERRED_REVIEW_REQUIRED","swe6_eligibility":"Eligible","sys5_eligibility":"Review Needed",
        "source_evidence":[{"location":"Page 12, 2.9","text":source}],
        "source_backed_atomic_behaviors":[
            {"source_fact_fragment_id":"F1","source_semantic_unit_id":"U1","source_location":"Page 12","behavior_text":"HU는 SVM_CaptureModeViewRequest로 Front->Rear->Left->Right 순서로 요청해야 한다."},
            {"source_fact_fragment_id":"F2","source_semantic_unit_id":"U1","source_location":"Page 12","behavior_text":"HU는 SVM_CaptureModeViewState로 Front->Rear->Left->Right 응답을 확인해야 한다."},
            {"source_fact_fragment_id":"F3","source_semantic_unit_id":"U1","source_location":"Page 12","behavior_text":"SVM_CaptureModeViewRequest 0x1 Front 0x2 Rear 0x3 Left 0x4 Right 0x5 Top 0x7 Invalid"},
            {"source_fact_fragment_id":"F4","source_semantic_unit_id":"U1","source_location":"Page 12","behavior_text":"SVM_CaptureModeViewState 0x1 Front 0x2 Rear 0x3 Left 0x4 Right 0x5 Top 0x6 Camera Input Error 0x7 Invalid"},
        ],
        "fact_level_allocations":[
            {"source_fact_fragment_id":f"F{i}","source_semantic_unit_id":"U1","source_location":"Page 12","source_fact":"capture fact","allocation_status":"SW_BEHAVIOR_INFERRED_REVIEW_REQUIRED","swe6_eligibility":"Eligible","verification_domain":"SWE.6 Software Qualification"}
            for i in range(1,5)
        ],
        "semantic_provenance_status":"COMPLETE","source_semantic_unit_ids":["U1"],
        "human_decision_required":True,
    }
    return source,req


def test_v091_requirements_word_and_excel_use_slim_human_view_and_keep_audit_elsewhere(tmp_path: Path):
    req=_req_fixture(); data={"requirements":[req]}
    src=tmp_path/"source.docx"; src.write_text("dummy",encoding="utf-8")
    ex=E2EExporter(tmp_path)
    word=ex.export_requirements_word(src,data)
    excel=ex.export_requirements_excel(src,data)

    doc=Document(word)
    all_text="\n".join(p.text for p in doc.paragraphs)+"\n"+"\n".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    for forbidden in ["E2E Requirement Scope","Engineering Domain","Requirement Level","Allocation Status","SYS.1 Eligibility","SWE.1 Eligibility","Canonical State","Human Review / Hold Reason"]:
        assert forbidden not in all_text
    assert "리뷰 상태" in all_text
    assert "사람의 리뷰 필요" in all_text

    wb=load_workbook(excel,read_only=True,data_only=False)
    try:
        ws=wb["2_E2E_요구사항"]
        headers=[str(ws.cell(3,c).value or "") for c in range(1,ws.max_column+1)]
        assert headers==["SRS ID","상위 기능","리뷰 상태","요구사항 내역","동작 조건 / Trigger","작동 명세 정의","사전 조건","예상 결과","검증 기준","출처 / Traceability","기타"]
        assert "\n" in str(ws.cell(4,4).value or "")
        trace=wb["4_Source_추적성"]
        th=[str(trace.cell(3,c).value or "") for c in range(1,trace.max_column+1)]
        assert "Allocation Status" in th and "Canonical State" in th
    finally:
        wb.close()


def test_v091_evaluation_main_is_compact_and_human_readable(tmp_path: Path):
    source,req=_capture_req()
    data={"requirements":[req],"testability_and_decomposition_result":{"by_srs":[{"srs_id":"SRS_009","swe6_eligibility":"Eligible","verification_domain":"SWE.6 Software Qualification"}]}}
    src=tmp_path/"capture.docx"; src.write_text(source,encoding="utf-8")
    out=E2EExporter(tmp_path).export_evaluation_excel(src,data)
    wb=load_workbook(out,read_only=False,data_only=False)
    try:
        ws=wb["2_E2E_평가케이스"]
        headers=[str(ws.cell(3,c).value or "") for c in range(1,ws.max_column+1)]
        assert headers==["Test Case ID","요구사항 ID","검증 영역","평가 방법 / 도출 방법","테스트 케이스 이름","테스트 목적","사전 조건","평가 환경 / 준비","입력 변수","비교","입력값","수행 절차","출력 변수 / 관찰 대상","비교","기대 출력값","기대 상태 / 결과","PASS 기준","TC 검토 상태","검토 사유 / 필요 조치","실제 출력값","판정 (PASS/FAIL)","비고","결과 화면 캡처"]
        assert ws["A2"].value=="종류" and ws["G2"].value=="사전 준비" and ws["I2"].value=="입력 / 수행" and ws["M2"].value=="출력" and ws["Q2"].value=="판단" and ws["T2"].value=="실행 결과"
        assert str(ws["A4"].value)=="SRS_TC_001"
        assert str(ws["B4"].value).startswith("SRS_009_BF_")
        assert ws["C4"].value in {"SWE.6","SYS.5","검토"}
        assert ws["D4"].value=="기본 기능 검증"
        assert str(ws["E4"].value).endswith("_01")
        # Input is compact value-only; output may retain concise enum meaning but not long HU prose.
        input_text="\n".join(str(ws.cell(r,11).value or "") for r in range(4,ws.max_row+1))
        output_text="\n".join(str(ws.cell(r,15).value or "") for r in range(4,ws.max_row+1))
        assert "Front View Printing Request" not in input_text
        assert "0x1" in input_text
        assert "HU는" not in output_text
        assert "Front" in output_text or "0x1" in output_text
        # Result columns are blank pre-execution.
        for r in range(4,ws.max_row+1):
            for c in range(20,24):
                assert ws.cell(r,c).value in (None,"")
        # Narrow columns requested by the human-facing design.
        assert ws.column_dimensions["J"].width <= 10
        assert ws.column_dimensions["K"].width <= 20
        assert ws.column_dimensions["N"].width <= 10
        assert ws.column_dimensions["S"].width <= 36
        trace=wb["4_요구사항-평가추적성"]
        trace_headers=[str(trace.cell(3,c).value or "") for c in range(1,trace.max_column+1)]
        assert "관련 원본 객체 ID" in trace_headers
        assert "출처 사실 조각 ID" in trace_headers
        review=wb["3_검토필요"]
        assert any(str(review.cell(r,4).value or "") in {"Source 기대동작 미정의","Exact Fact Trace 필요","Source TBD","Source Conflict"} for r in range(4,review.max_row+1))
    finally:
        wb.close()
