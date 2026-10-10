from pathlib import Path

from openpyxl import load_workbook

from core.cross_document_semantics import apply_allocation_gate, build_semantic_source_units
from core.quality_audit import (
    build_swe6_export_preservation_audit,
    build_testability_result,
    finalize_test_intent_coverage,
    normalize_requirement_extensions,
)
from core.swe6_exporter import SWE6Exporter, build_swe6_cases


def _req(srs="SRS_001", text="소프트웨어는 값을 저장해야 한다."):
    req = {
        "candidate_id": "C1", "srs_id": srs, "scenario_candidate_id": "SCN-CAND-001",
        "category": "기능", "function_name": "Feature", "requirement": text,
        "user_input": "", "system_input_preconditions": "", "processing_action": text,
        "output": text, "acceptance_criteria": text, "failure_situations": [], "user_intervention_points": [],
        "derivation_type": "explicit", "derivation_reason": "", "clarification_needed": [],
        "source_evidence": [{"document": "system.docx", "location": "Paragraph 10", "text": text}],
        "confidence": 0.95, "classification_basis": "입력문서 명시", "activation_trigger": "",
        "preconditions": "", "behavior_flows": [], "evaluation_method": "", "exception_conditions": [],
        "related_artifacts": [],
    }
    normalize_requirement_extensions(req)
    return req


def _data(reqs):
    if not isinstance(reqs, list):
        reqs = [reqs]
    return {
        "schema_version": "REQ-STUDIO-CANONICAL-REQ-1.9",
        "source_document": "system.docx",
        "scenario_candidates": [{
            "scenario_candidate_id": "SCN-CAND-001", "scenario_name": "x", "user_goal_context": "",
            "scenario_flow": [], "expected_outcome": "",
            "source_evidence": [{"document": "system.docx", "location": "Paragraph 1", "text": "x"}],
        }],
        "requirements": reqs, "gaps": [],
    }


def test_reference_row_with_table_ordinal_and_numbered_heading_are_context():
    compact = """[DOCUMENT] system.docx
[SRC DOC-C1 | Table 14 / Row 31 | table]
31 | MS201-02 | 유해물질 금지 및 신고 - 부품 및 재료 |
[SRC DOC-C2 | Paragraph 100 | text]
5.4.2.1 커넥터 타입
[SRC DOC-C3 | Paragraph 101 | text]
커넥터(DIP) 타입을 적용하고 LEAD-WIRE는 금지한다.
"""
    units = build_semantic_source_units(compact)
    by_loc = {x["source_location"]: x for x in units}
    assert by_loc["Table 14 / Row 31"]["source_unit_type"] == "external_reference_only"
    assert by_loc["Table 14 / Row 31"]["coverage_eligibility"] == "review_context"
    assert by_loc["Paragraph 100"]["source_unit_type"] == "heading_or_document_context"
    assert by_loc["Paragraph 100"]["coverage_eligibility"] == "review_context"
    assert by_loc["Paragraph 101"]["coverage_eligibility"] == "semantic_unit"


def test_allocation_pending_has_nonempty_deferred_intent_and_reason():
    req = _req(text="제어기는 Low Power Mode로 진입해야 한다.")
    req.update({
        "requirement_level": "System", "allocation_status": "SYSTEM_REQUIREMENT_PENDING_SW_ALLOCATION",
        "swe1_eligibility": "Review Needed", "swe6_eligibility": "Deferred pending SW allocation",
        "verification_domain": "System Integration / SYS.5",
        "source_requirement_ids": [], "source_semantic_unit_ids": ["SRC-SEM-X"],
        "source_backed_atomic_behaviors": [{"source_semantic_unit_id": "SRC-SEM-X", "behavior_text": req["requirement"], "knowledge_state": "KNOWN"}],
        "semantic_provenance_status": "COMPLETE",
    })
    data = _data(req)
    data["testability_and_decomposition_result"] = build_testability_result(data)
    finalize_test_intent_coverage(data, [])
    row = data["testability_and_decomposition_result"]["by_srs"][0]
    assert row["required_test_intents"]
    assert row["not_generated_test_intents"]
    assert all(x["reason_code"] == "ALLOCATION_PENDING" for x in row["not_generated_test_intents"])
    assert data["test_intent_coverage"]["allocation_pending_intent_count"] >= 1
    assert data["test_intent_coverage"]["allocation_pending_reason_completeness_percent"] == 100.0
    audit = data["swe6_export_preservation_audit"]
    assert audit["passed"] is True
    assert audit["rows"][0]["allocation_pending_deferred_governance_ok"] is True


def test_cross_domain_source_bundle_keeps_parent_and_fact_level_allocations():
    req = _req(text="최대 소비전류 MAX 100mA, 사용 온도 -40℃ ~ +85℃, 보존 온도 -40℃ ~ +95℃")
    req["source_requirement_ids"] = []
    units = [
        {"source_semantic_unit_id":"SRC-SEM-E", "source_chunk_id":"C1", "source_location":"Table 1 / Row 1", "source_excerpt":"최대 소비전류 | MAX 100mA"},
        {"source_semantic_unit_id":"SRC-SEM-T1", "source_chunk_id":"C2", "source_location":"Table 1 / Row 2", "source_excerpt":"사용 온도 | -40℃ ~ +85℃"},
        {"source_semantic_unit_id":"SRC-SEM-T2", "source_chunk_id":"C3", "source_location":"Table 1 / Row 3", "source_excerpt":"보존 온도 | -40℃ ~ +95℃"},
    ]
    apply_allocation_gate(req, units)
    assert req["cross_domain_bundle_review_required"] is True
    assert req["allocation_status"] == "CROSS_DOMAIN_FACT_BUNDLE_REVIEW_REQUIRED"
    assert req["swe1_eligibility"] == "Not Applicable"
    assert req["swe6_eligibility"] == "Not Applicable"
    domains = set(req["cross_domain_verification_domains"])
    assert "Electrical Verification" in domains
    assert "Environmental Qualification" in domains
    assert len(req["fact_level_allocations"]) == 3


def test_child_intent_has_coverage_disposition_not_automatic_complete():
    req = _req(text="입력을 수신하고 NVM에 저장한다.")
    req.update({
        "requirement_level": "Software", "allocation_status": "SW_IMPLEMENTATION_REQUIREMENT",
        "swe1_eligibility": "Eligible", "swe6_eligibility": "Eligible",
        "verification_domain": "SWE.6 Software Qualification", "source_requirement_ids": [],
        "source_semantic_unit_ids": ["SRC-SEM-1", "SRC-SEM-2"],
        "source_backed_atomic_behaviors": [
            {"source_semantic_unit_id":"SRC-SEM-1", "behavior_text":"입력을 수신한다.", "knowledge_state":"KNOWN"},
            {"source_semantic_unit_id":"SRC-SEM-2", "behavior_text":"수신한 설정을 NVM에 저장한다.", "knowledge_state":"KNOWN"},
        ],
        "semantic_provenance_status": "COMPLETE",
        "activation_trigger": "입력 수신 시", "processing_action": "입력을 수신한다.", "output": "입력을 수신한다.",
    })
    data = _data(req)
    data["testability_and_decomposition_result"] = build_testability_result(data)
    cases = build_swe6_cases(data)
    finalize_test_intent_coverage(data, cases)
    row = data["swe6_export_preservation_audit"]["rows"][0]
    assert row["source_backed_child_intents"]
    statuses = {x["intent_status"] for x in row["source_backed_child_intents"]}
    assert statuses <= {"COVERED_BY_GENERIC_TC", "REPRESENTED_IN_GENERIC_TC", "INDEPENDENTLY_COVERED", "DEFERRED", "REVIEW_REQUIRED"}
    assert all("related_tc_ids" in x for x in row["source_backed_child_intents"])
    # Presence of one generic TC must not silently set every child to a synthetic COMPLETE state.
    assert "COMPLETE" not in statuses


def test_swe6_audit_flags_canonical_numeric_fact_missing_from_provenance():
    req = _req(text="사용 온도는 -40℃ ~ +85℃이고 보존 온도는 -40℃ ~ +95℃이다.")
    req.update({
        "requirement_level": "Environmental Qualification", "allocation_status": "ENVIRONMENTAL_QUALIFICATION_REQUIREMENT",
        "swe1_eligibility": "Not Applicable", "swe6_eligibility": "Not Applicable",
        "verification_domain": "Environmental Qualification", "source_requirement_ids": [],
        "source_semantic_unit_ids": ["SRC-SEM-T1"],
        "source_backed_facts": [{"source_semantic_unit_id":"SRC-SEM-T1", "source_fact":"사용 온도 | -40℃ ~ +85℃", "knowledge_state":"KNOWN"}],
        "semantic_provenance_status": "COMPLETE",
    })
    data = _data(req)
    data["testability_and_decomposition_result"] = build_testability_result(data)
    finalize_test_intent_coverage(data, [])
    audit = data["swe6_export_preservation_audit"]
    assert audit["rows"][0]["source_fact_linkage_ok"] is False
    assert audit["rows"][0]["missing_source_fact_tokens"]
    assert any(x["issue"] == "CANONICAL_SOURCE_FACT_PROVENANCE_GAP" for x in audit["review_issue_records"])


def test_derived_verification_method_is_labelled_and_intent_audit_sheet_exists(tmp_path):
    src = "Watchdog Timer는 lowest task 이후 동작하며 ISR 내 존재할 수 없다."
    req = _req(text=src)
    req.update({
        "requirement_level": "Software", "allocation_status": "SW_IMPLEMENTATION_REQUIREMENT",
        "swe1_eligibility": "Eligible", "swe6_eligibility": "Eligible",
        "verification_domain": "SWE.6 Software Qualification", "source_requirement_ids": ["REQ-1"],
        "source_backed_atomic_behaviors": [{"behavior_text":src, "knowledge_state":"KNOWN"}],
        "output": "Watchdog Timer 동작", "acceptance_criteria": "lowest task 이후 동작하며 ISR 내 존재하지 않는다.",
    })
    data = _data(req)
    data["testability_and_decomposition_result"] = build_testability_result(data)
    cases = build_swe6_cases(data)
    assert "DERIVED 후보" in cases[0]["method"]
    assert cases[0]["verification_method_candidates"][0]["knowledge_state"] == "DERIVED"
    source = tmp_path / "system.docx"
    source.write_bytes(b"dummy")
    out = SWE6Exporter(tmp_path).export_excel(source, data, source_text="")
    wb = load_workbook(out, read_only=True)
    try:
        assert "3_Deferred Intent" in wb.sheetnames
        assert "4_Source_Intent_Audit" in wb.sheetnames
    finally:
        wb.close()
