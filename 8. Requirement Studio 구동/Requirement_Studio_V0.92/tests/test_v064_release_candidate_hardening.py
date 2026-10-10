from pathlib import Path

from openpyxl import load_workbook

from core.cross_document_semantics import (
    _structured_source_facts,
    attach_semantic_traceability,
)
from core.quality_audit import _source_backed_child_intents, normalize_requirement_extensions
from core.regression_engine import build_regression_report
from core.review_exchange import ReviewExchangeBuilder
from core.swe6_exporter import SWE6Exporter


def _req(text: str, source_text: str | None = None, srs: str = "SRS_001") -> dict:
    row = {
        "candidate_id": "C1",
        "srs_id": srs,
        "scenario_candidate_id": "SCN-CAND-001",
        "category": "기능",
        "function_name": "Feature",
        "requirement": text,
        "user_input": "",
        "system_input_preconditions": "",
        "processing_action": text,
        "output": text,
        "acceptance_criteria": text,
        "failure_situations": [],
        "user_intervention_points": [],
        "derivation_type": "explicit",
        "derivation_reason": "",
        "clarification_needed": [],
        "source_evidence": [{"document": "x.docx", "location": "Paragraph 1", "text": source_text or text}],
        "confidence": 0.95,
        "classification_basis": "source",
        "activation_trigger": "",
        "preconditions": "",
        "behavior_flows": [],
        "evaluation_method": "",
        "exception_conditions": [],
        "related_artifacts": [],
        "source_requirement_ids": [],
    }
    normalize_requirement_extensions(row)
    return row


def _data(req: dict) -> dict:
    return {
        "schema_version": "REQ-STUDIO-CANONICAL-REQ-2.0",
        "source_document": "x.docx",
        "scenario_candidates": [],
        "requirements": [req],
        "gaps": [],
    }


def test_actual_xlsx_verifier_detects_prepopulated_result_cells_and_export_boundary_can_clean_them(tmp_path):
    req = _req("입력을 확인한다.")
    req.update({
        "requirement_level": "Software",
        "allocation_status": "SW_IMPLEMENTATION_REQUIREMENT",
        "swe1_eligibility": "Eligible",
        "swe6_eligibility": "Eligible",
        "verification_domain": "SWE.6 Software Qualification",
        "source_semantic_unit_ids": ["S1"],
        "source_backed_atomic_behaviors": [{"source_semantic_unit_id": "S1", "behavior_text": "입력을 확인한다.", "knowledge_state": "KNOWN"}],
        "semantic_provenance_status": "COMPLETE",
    })
    exporter = SWE6Exporter(tmp_path)
    out = exporter.export_excel(tmp_path / "x.docx", _data(req), source_text="입력을 확인한다.")

    wb = load_workbook(out)
    try:
        ws = wb["2_테스트 케이스"]
        ws.cell(4, 18, "N/A")
        ws.cell(4, 19, "N/A")
        ws.cell(4, 20, "N/A")
        ws.cell(4, 21, "N/A")
        wb.save(out)
    finally:
        wb.close()

    cases = [{"srs_id": "SRS_001", "tc_id": "TC_001"}]
    failed = ReviewExchangeBuilder._verify_swe6_excel_artifact(out, cases)
    assert failed["passed"] is False
    assert failed["prepopulated_result_tc_ids"] == ["TC_001"]
    assert {x["field"] for x in failed["prepopulated_result_cells"]} == {
        "Output Value", "PASS / FAIL", "Comment", "Capture CANoe"
    }
    assert failed["artifact_sha256"]

    exporter._enforce_blank_execution_results(out)
    passed = ReviewExchangeBuilder._verify_swe6_excel_artifact(out, cases)
    assert passed["passed"] is True
    assert passed["prepopulated_result_cells"] == []


def test_numbered_source_list_backfills_missing_fourth_condition_and_child_fragment_trace():
    compact = """[DOCUMENT] x.docx
[SRC C1 | Paragraph 1 | text]
High Power Mode 조건: 1) ON 명령 수신 시 2) B-CAN WAKE UP 상태 3) TRANSCEIVER의 ERR PIN이 HIGH에서 LOW 상태 변환 시 4) 최초 B+ 인가 시
"""
    req = _req(
        "ON 명령 수신, B-CAN WAKE UP 상태, TRANSCEIVER ERR PIN HIGH에서 LOW 상태 변환, 최초 B+ 인가 조건 중 하나라도 참인 경우 High Power Mode로 진입한다.",
        source_text="High Power Mode 조건",
        srs="SRS_004",
    )
    # Simulate the V0.63 failure shape: generic OR + first three Source conditions already exist,
    # while the fourth Source condition is absent from atomic/child-intent trace.
    req["source_backed_atomic_behaviors"] = [
        {"source_semantic_unit_id": "", "behavior_text": "High Power Mode OR condition", "knowledge_state": "KNOWN"},
        {"source_semantic_unit_id": "", "behavior_text": "ON 명령 수신 시", "knowledge_state": "KNOWN"},
        {"source_semantic_unit_id": "", "behavior_text": "B-CAN WAKE UP 상태", "knowledge_state": "KNOWN"},
        {"source_semantic_unit_id": "", "behavior_text": "TRANSCEIVER의 ERR PIN이 HIGH에서 LOW 상태 변환 시", "knowledge_state": "KNOWN"},
    ]
    data = _data(req)
    attach_semantic_traceability(data, compact)
    actual = data["requirements"][0]

    fragments = actual.get("source_fact_fragments") or []
    condition_fragments = [f for f in fragments if f.get("source_list_marker")]
    assert len(condition_fragments) == 4
    assert [f.get("source_list_marker") for f in condition_fragments] == ["1)", "2)", "3)", "4)"]
    assert any("최초 B+ 인가" in str(f.get("source_excerpt") or "") for f in condition_fragments)
    assert not any(str(f.get("source_excerpt") or "").endswith("조건:") for f in fragments)

    atoms = actual.get("source_backed_atomic_behaviors") or []
    bplus_atoms = [a for a in atoms if "최초 B+ 인가" in str(a.get("behavior_text") or "")]
    assert len(bplus_atoms) == 1
    assert bplus_atoms[0].get("source_fact_fragment_id")

    children = _source_backed_child_intents(actual)
    bplus_children = [c for c in children if "최초 B+ 인가" in str(c.get("source_backed_behavior") or "")]
    assert len(bplus_children) == 1
    assert bplus_children[0].get("source_fact_fragment_id") == bplus_atoms[0].get("source_fact_fragment_id")


def test_structured_fact_type_separates_bitrate_and_external_spec_identifier():
    bitrate = _structured_source_facts("CAN 500 Kbit/s, LIN 19.2 Kbit/s", unit_id="U1", location="Table 1 / Row 1")[0]
    assert bitrate["fact_type"] == "COMMUNICATION_BITRATE"
    assert bitrate["timing_values"] == []

    external = _structured_source_facts("ES995480-00 기준을 따른다.", unit_id="U2", location="Table 1 / Row 2")[0]
    assert external["fact_type"] == "EXTERNAL_SPEC_IDENTIFIER"
    assert "ES995480-00" in external["external_spec_identifiers"]
    assert external["numeric_values"] == []
    assert external["range_values"] == []


def test_no_baseline_is_explicit_registration_required_not_false_pass():
    report = build_regression_report({}, None, "No same-source approved baseline")
    assert report["regression_gate"]["status"] == "NOT_EVALUATED"
    assert report["regression_gate"]["passed"] is False
    assert report["baseline_registration_required"] is True
    assert "never auto-promote" in report["baseline_registration_policy"]
