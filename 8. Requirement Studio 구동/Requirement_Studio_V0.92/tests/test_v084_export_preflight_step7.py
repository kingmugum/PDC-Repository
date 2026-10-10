from __future__ import annotations

import json
from pathlib import Path
from types import SimpleNamespace

from openpyxl import load_workbook

from core.evaluation_orchestrator import MultiModelEvaluationOrchestrator
from core.integrated_exporter import IntegratedExporter
from core.review_exchange import ReviewExchangeBuilder
from core.swe1_exporter import SWE1Exporter
from core.swe6_exporter import SWE6Exporter
from core.verification_single_truth import finalize_verification_single_truth
from tests.test_v082_integrated_specs_single_truth import _mode_fixture


def _meta():
    return SimpleNamespace(display_name="H-Chat / GPT", provider_id="hchat_gpt", model="gpt-5.6-terra", endpoint_family="hchat", auth_mode="key", project_header_enabled=False, response_parsing_mode="text")


def _project(tmp_path: Path) -> Path:
    root = tmp_path / "project"
    (root / "history" / "change_decisions").mkdir(parents=True)
    (root / "history" / "change_decisions" / "04_CHANGE_DECISION_V0.83_to_V0.84.txt").write_text("decision", encoding="utf-8")
    (root / "config").mkdir(parents=True)
    (root / "config" / "evaluation_api_config.json").write_text(json.dumps({"enabled": True, "require_all_providers": False}), encoding="utf-8")
    return root


def _artifacts(tmp_path: Path, data: dict):
    out = tmp_path / "out"
    source = tmp_path / "MLM.docx"
    source.write_text("placeholder", encoding="utf-8")
    swe1, swe6, integrated = SWE1Exporter(out), SWE6Exporter(out), IntegratedExporter(out)
    return source, {
        "swe1_word": swe1.export_word(source, data),
        "swe1_excel": swe1.export_excel(source, data),
        "swe6_excel": swe6.export_excel(source, data, source_text="Low Power Mode High Power Mode"),
        "integrated_requirements_excel": integrated.export_requirements_excel(source, data),
        "integrated_tests_excel": integrated.export_tests_excel(source, data),
    }


def test_v084_integrated_requirements_serializes_nested_source_evidence(tmp_path: Path):
    data = _mode_fixture()
    data["gaps"] = [{
        "gap_id": "GAP_STRUCTURED",
        "description": "interface decision pending",
        "source_evidence": [{
            "document": "sample.pdf",
            "location": "Page 10, Page 14, Page 23",
            "text": "SVM_CaptureModeCMD / ICMU_CaptureModeCMD proposal",
        }],
        "required_resolution": "approve authoritative signal",
    }]
    source = tmp_path / "sample.pdf"
    source.write_text("placeholder", encoding="utf-8")
    path = IntegratedExporter(tmp_path).export_requirements_excel(source, data)
    wb = load_workbook(path, read_only=True)
    try:
        ws = wb["07_Open_Issues_Gaps"]
        rows = [r for r in range(4, ws.max_row + 1) if str(ws.cell(r, 2).value or "") == "GAP_STRUCTURED"]
        assert rows
        value = ws.cell(rows[0], 5).value
        assert isinstance(value, str)
        assert "sample.pdf" in value and "Page 10, Page 14, Page 23" in value and "ICMU_CaptureModeCMD" in value
    finally:
        wb.close()


def test_v084_integrated_table_scalar_guard_handles_arbitrary_nested_values(tmp_path: Path):
    exporter = IntegratedExporter(tmp_path)
    # Exercise the common cell-write guard directly so future integrated sheets inherit protection.
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active
    exporter._write_table(ws, ["A", "B"], [[[1, {"x": 2}], {"document": "a.pdf", "location": "P1"}]])
    assert isinstance(ws.cell(4, 1).value, str)
    assert isinstance(ws.cell(4, 2).value, str)


def test_v084_unified_preflight_skips_external_reviewers_but_emits_deterministic_fail(tmp_path: Path):
    data = _mode_fixture(); finalize_verification_single_truth(data)
    source, artifacts = _artifacts(tmp_path, data)
    artifacts.pop("integrated_requirements_excel")
    root = _project(tmp_path)
    result = ReviewExchangeBuilder(root, app_version="v0.84").build(
        source_document=source, provider_metadata=_meta(), analysis_text="", requirement_data=data, evaluation={}, normalized_summary={},
        quality_review_config={"enabled": True, "evaluation_artifact": "unified"}, artifact_paths=artifacts, evaluation_mode="initial",
    )
    assert result["unified_artifact_contract"]["status"] == "FAIL"
    orchestrator = MultiModelEvaluationOrchestrator(root, app_version="v0.84")
    def forbidden(*args, **kwargs):
        raise AssertionError("external reviewer must not be called after deterministic Unified preflight failure")
    orchestrator._call_one = forbidden
    out = orchestrator.run(Path(result["run_dir"]))
    assert out["status"] == "PRECHECK_FAILED"
    assert out["providers_completed"] == []
    assert out["tool_quality_gate"] == "FAIL"
    assert out["official_release"] == "HOLD"
    txt = (Path(out["output_dir"]) / "01_GPT_EVALUATION.txt").read_text(encoding="utf-8")
    assert "preflight failed" in txt.lower()


def test_v084_step7_log_contract_separates_internal_save_and_artifact_export():
    root = Path(__file__).resolve().parents[1]
    runner = (root / "core" / "ai_job_runner.py").read_text(encoding="utf-8")
    main = (root / "main.py").read_text(encoding="utf-8")
    assert "내부 분석 결과 저장 완료 · 선택 산출물 Export는 후처리에서 계속" in runner
    assert "STEP 7B 시작 · 선택 산출물 및 평가용 Artifact Export" in main
    assert "STEP 7B 완료 · 선택/평가 산출물 Export 및 Review Package 준비 완료" in main
