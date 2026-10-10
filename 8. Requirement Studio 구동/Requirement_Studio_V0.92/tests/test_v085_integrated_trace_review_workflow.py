from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from core.integrated_exporter import IntegratedExporter
from core.verification_single_truth import BUNDLE_SCHEMA_VERSION, detect_mode_transition_allocation_findings
from core.review_decision_registry import record_mode_transition_decision
from core.cross_document_semantics import _semantic_review_shortcut
from core.quality_audit import _attach_canonical_gap_references


def _mixed_data() -> dict:
    req = {
        "srs_id": "SRS_MIX",
        "candidate_id": "REQ-CAND-MIX",
        "function_name": "Mixed behavior",
        "requirement": "Controller performs SW action while system power state is verified separately.",
        "source_semantic_unit_ids": ["U1"],
        "source_evidence": [{"location": "Paragraph 10", "text": "mixed"}],
        "source_backed_atomic_behaviors": [
            {"source_semantic_unit_id": "U1", "source_fact_fragment_id": "F_SW", "source_location": "Paragraph 10", "behavior_text": "SW action"},
            {"source_semantic_unit_id": "U1", "source_fact_fragment_id": "F_SYS", "source_location": "Paragraph 11", "behavior_text": "system power state"},
        ],
        "fact_level_allocations": [
            {"source_semantic_unit_id": "U1", "source_fact_fragment_id": "F_SW", "source_location": "Paragraph 10", "source_fact": "SW action", "allocation_status": "SW_BEHAVIOR_INFERRED_REVIEW_REQUIRED", "swe6_eligibility": "Eligible", "verification_domain": "SWE.6 Software Qualification"},
            {"source_semantic_unit_id": "U1", "source_fact_fragment_id": "F_SYS", "source_location": "Paragraph 11", "source_fact": "system power state", "allocation_status": "SYSTEM_REQUIREMENT_PENDING_SW_ALLOCATION", "swe6_eligibility": "Deferred pending SW allocation", "verification_domain": "System Integration / SYS.5"},
        ],
        "allocation_status": "SW_BEHAVIOR_INFERRED_REVIEW_REQUIRED",
        "swe6_eligibility": "Eligible",
        "sys5_eligibility": "Eligible",
        "verification_domain": "Mixed",
        "engineering_domains": ["System", "Software"],
        "human_decision_required": True,
    }
    objects = [
        {"test_object_type": "SWE6_TC", "test_object_id": "TC_001", "parent_srs_id": "SRS_MIX", "candidate_id": "REQ-CAND-MIX", "source_semantic_unit_ids": ["U1"], "source_fact_fragment_ids": ["F_SW"], "source_locations": ["Paragraph 10"], "source_backed_behavior": "SW action", "verification_objective": "verify SW action", "engineering_domain": "Software", "verification_domain": "SWE.6 Software Qualification", "sys5_eligibility": "Eligible", "swe6_eligibility": "Eligible", "allocation_status": "SW_BEHAVIOR_INFERRED_REVIEW_REQUIRED", "execution_readiness": "REVIEW_REQUIRED", "human_review_required": True},
        {"test_object_type": "SYS5_CANDIDATE", "test_object_id": "SYS5_TC_001", "parent_srs_id": "SRS_MIX", "candidate_id": "REQ-CAND-MIX", "source_semantic_unit_ids": ["U1"], "source_fact_fragment_ids": ["F_SYS"], "source_locations": ["Paragraph 11"], "source_backed_behavior": "system power state", "verification_objective": "verify system power state", "engineering_domain": "System", "verification_domain": "System Integration / SYS.5", "sys5_eligibility": "Eligible", "swe6_eligibility": "Deferred pending SW allocation", "allocation_status": "SYSTEM_REQUIREMENT_PENDING_SW_ALLOCATION", "execution_readiness": "REVIEW_REQUIRED", "human_review_required": True},
        {"test_object_type": "ALLOCATION_PENDING_INTENT", "test_object_id": "AP_INTENT_001", "parent_srs_id": "SRS_MIX", "candidate_id": "REQ-CAND-MIX", "source_semantic_unit_ids": ["U1"], "source_fact_fragment_ids": ["F_SYS"], "source_locations": ["Paragraph 11"], "source_backed_behavior": "system power state", "verification_objective": "verify system power state", "engineering_domain": "System", "verification_domain": "System Integration / SYS.5", "sys5_eligibility": "Eligible", "swe6_eligibility": "Deferred pending SW allocation", "allocation_status": "SYSTEM_REQUIREMENT_PENDING_SW_ALLOCATION", "execution_readiness": "REVIEW_REQUIRED", "human_review_required": True, "deferred_reason_code": "ALLOCATION_PENDING", "deferred_reason_detail": "pending", "required_resolution": "review"},
    ]
    bundle = {
        "schema_version": BUNDLE_SCHEMA_VERSION,
        "swe6_cases": [{"srs_id": "SRS_MIX", "tc_id": "TC_001", "swe6_scope_fragment_ids": ["F_SW"]}],
        "sys5_candidates": [{"srs_id": "SRS_MIX", "sys5_id": "SYS5_TC_001", "source_fact_fragment_ids": ["F_SYS"]}],
        "integrated_test_objects": objects,
        "mode_transition_allocation_findings": [],
        "single_truth_audit": {"release_gate_status": "PASS", "blocking_issue_count": 0, "blocking_records": [], "total_integrated_test_object_count": len(objects)},
        "summary": {},
    }
    return {"requirements": [req], "finalized_verification_bundle": bundle, "testability_and_decomposition_result": {"by_srs": []}}


def test_v085_integrated_trace_links_only_exact_fact_scope(tmp_path: Path):
    data = _mixed_data()
    src = tmp_path / "source.docx"; src.write_text("x", encoding="utf-8")
    path = IntegratedExporter(tmp_path).export_tests_excel(src, data)
    wb = load_workbook(path, read_only=True, data_only=False)
    try:
        ws = wb["08_Source_Intent_Traceability"]
        rows = {str(ws.cell(r, 4).value or ""): r for r in range(4, ws.max_row + 1)}
        sw = rows["F_SW"]; sys = rows["F_SYS"]
        assert "TC_001" in str(ws.cell(sw, 7).value or "")
        sys_links = {x.strip() for x in str(ws.cell(sys, 7).value or "").split(",") if x.strip()}
        assert "TC_001" not in sys_links
        assert "SYS5_TC_001" in sys_links
        assert str(ws.cell(sys, 12).value or "") == "EXCLUDED_FROM_SWE6_SCOPE"
        assert str(ws.cell(sys, 9).value or "") == "SYSTEM_REQUIREMENT_PENDING_SW_ALLOCATION"
    finally:
        wb.close()


def _mode_req(sid: str, text: str, paragraph: int, allocation: str, swe6: str, sys5: str = "Eligible") -> dict:
    return {
        "srs_id": sid,
        "function_name": text,
        "requirement": f"HU는 {text} 해야 한다.",
        "allocation_status": allocation,
        "swe6_eligibility": swe6,
        "sys5_eligibility": sys5,
        "source_evidence": [{"location": f"Paragraph {paragraph}", "text": text}],
    }


def test_v085_mode_transition_findings_are_family_scoped_not_actor_cross_product(monkeypatch, tmp_path: Path):
    monkeypatch.setenv("REQUIREMENT_STUDIO_REVIEW_REGISTRY_DIR", str(tmp_path / "registry"))
    data = {"requirements": [
        _mode_req("SRS_LOW", "Low Power Mode 전환", 125, "SYSTEM_REQUIREMENT_PENDING_SW_ALLOCATION", "Deferred pending SW allocation"),
        _mode_req("SRS_HIGH", "High Power Mode 전환", 132, "SW_BEHAVIOR_INFERRED_REVIEW_REQUIRED", "Eligible"),
        _mode_req("SRS_CAP", "Capture Mode 전환", 140, "SW_BEHAVIOR_INFERRED_REVIEW_REQUIRED", "Eligible"),
        _mode_req("SRS_MIR", "Mirror Unfold 전환", 142, "SYSTEM_REQUIREMENT_PENDING_SW_ALLOCATION", "Deferred pending SW allocation"),
        _mode_req("SRS_IGN", "IGN3 Off 전환", 144, "SW_BEHAVIOR_INFERRED_REVIEW_REQUIRED", "Eligible"),
    ]}
    findings = detect_mode_transition_allocation_findings(data)
    assert len(findings) == 1
    assert set(findings[0]["affected_srs_ids"]) == {"SRS_LOW", "SRS_HIGH"}
    assert findings[0]["mode_transition_family"] == "power_mode_transition"


def test_v085_mode_transition_human_disposition_persists(monkeypatch, tmp_path: Path):
    monkeypatch.setenv("REQUIREMENT_STUDIO_REVIEW_REGISTRY_DIR", str(tmp_path / "registry"))
    data = {"requirements": [
        _mode_req("SRS_LOW", "Low Power Mode 전환", 125, "SYSTEM_REQUIREMENT_PENDING_SW_ALLOCATION", "Deferred pending SW allocation"),
        _mode_req("SRS_HIGH", "High Power Mode 전환", 132, "SW_BEHAVIOR_INFERRED_REVIEW_REQUIRED", "Eligible"),
    ]}
    finding = detect_mode_transition_allocation_findings(data)[0]
    record_mode_transition_decision(finding["finding_id"], disposition="CONFIRMED_INTENTIONAL", owner="Reviewer", rationale="Approved domain split")
    rerun = detect_mode_transition_allocation_findings(data)[0]
    assert rerun["review_closed"] is True
    assert rerun["review_disposition"] == "CONFIRMED_INTENTIONAL"
    assert rerun["review_owner"] == "Reviewer"


def test_v085_open_semantic_unit_gets_advisory_existing_srs_shortcut_only():
    unit = {"source_excerpt": "마스터 제어기는 CAN 통신 신호를 수신하여 무드램프 동작을 판단하고 LIN 통신으로 슬레이브 제어기를 제어한다."}
    reqs = [{"srs_id": "SRS_001", "candidate_id": "REQ_001", "requirement": "마스터 제어기는 CAN 통신 신호를 수신하여 무드램프 동작을 판단하고 LIN 통신으로 슬레이브 제어기를 제어해야 한다.", "source_backed_atomic_behaviors": []}]
    suggestion = _semantic_review_shortcut(unit, reqs)
    assert suggestion["review_shortcut_suggestion"] == "POSSIBLE_EXISTING_SRS_COVERAGE"
    assert suggestion["suggested_srs_id"] == "SRS_001"
    assert suggestion["similarity_score"] >= 0.55
    assert "human" in suggestion["suggested_action"].lower()


def test_v085_canonical_gap_ids_are_attached_without_deleting_original_clarification():
    data = {
        "requirements": [{"srs_id": "SRS_007", "clarification_needed": ["ICMU command approval needed"]}, {"srs_id": "SRS_013", "clarification_needed": ["same issue restated"]}],
        "gaps": [{"gap_id": "GAP_001", "related_srs_ids": ["SRS_007", "SRS_013"], "description": "ICMU command baseline unresolved"}],
    }
    _attach_canonical_gap_references(data)
    for req in data["requirements"]:
        assert req["canonical_gap_ids"] == ["GAP_001"]
        assert "GAP_001" in req["clarification_reference"]
        assert req["clarification_needed_detail"]


def test_v085_mode_review_excel_import_updates_external_registry(monkeypatch, tmp_path: Path):
    from openpyxl import Workbook
    from tools.import_mode_transition_reviews import import_reviews
    from core.review_decision_registry import load_mode_transition_decisions

    monkeypatch.setenv("REQUIREMENT_STUDIO_REVIEW_REGISTRY_DIR", str(tmp_path / "registry"))
    xlsx = tmp_path / "integrated.xlsx"
    wb = Workbook(); ws = wb.active; ws.title = "11_Release_Gate_Summary"
    headers = ["Gate / Finding", "Finding ID", "Status / Severity", "Affected IDs", "Review Disposition", "Review Owner", "Details", "Required Resolution", "Review Rationale / Reviewed At"]
    for c, h in enumerate(headers, 1): ws.cell(3, c, h)
    ws.append(["MODE_TRANSITION_ALLOCATION_INCONSISTENCY", "MT-ABC123", "RELEASE_REVIEW", "SRS_A,SRS_B", "HARMONIZED", "Owner", "detail", "resolution", "approved after review"])
    wb.save(xlsx); wb.close()

    assert import_reviews(xlsx) == 1
    decisions = load_mode_transition_decisions()
    assert decisions["MT-ABC123"]["disposition"] == "HARMONIZED"
    assert decisions["MT-ABC123"]["owner"] == "Owner"
