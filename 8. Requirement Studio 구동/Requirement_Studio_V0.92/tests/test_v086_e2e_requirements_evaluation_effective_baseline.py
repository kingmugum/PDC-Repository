from pathlib import Path
from openpyxl import load_workbook
from docx import Document

from core.e2e_exporter import E2EExporter
from core.swe6_exporter import build_sys5_candidates
from core.verification_single_truth import finalize_verification_single_truth
from tests.test_v082_integrated_specs_single_truth import _mode_fixture


def test_v086_e2e_requirements_word_excel_preserve_sys1_and_swe1(tmp_path: Path):
    data = _mode_fixture()
    src = tmp_path / "source.docx"; src.write_text("x", encoding="utf-8")
    ex = E2EExporter(tmp_path)
    w = ex.export_requirements_word(src, data)
    x = ex.export_requirements_excel(src, data)
    doc = Document(w)
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "E2E Engineering Requirements Specification" in text
    assert "SRS_004" in text and "SRS_005" in text
    assert len(doc.tables) >= 3
    wb = load_workbook(x, read_only=True, data_only=False)
    try:
        assert wb.sheetnames == ["표지","0_변경이력","1_요구사항요약","2_E2E_요구사항","3_검토필요","4_Source_추적성"]
        ws = wb["2_E2E_요구사항"]
        headers = [str(ws.cell(3,c).value or "") for c in range(1, ws.max_column+1)]
        # V0.91 removes eligibility metadata from the human Main but preserves it in Source Trace.
        assert "SYS.1 Eligibility" not in headers and "SWE.1 Eligibility" not in headers
        ids = {str(ws.cell(r,1).value or "") for r in range(4,ws.max_row+1)}
        assert {"SRS_004","SRS_005"}.issubset(ids)
        trace = wb["4_Source_추적성"]
        th = [str(trace.cell(3,c).value or "") for c in range(1, trace.max_column+1)]
        assert "SYS.1 Eligibility" in th and "SWE.1 Eligibility" in th
    finally:
        wb.close()


def test_v086_e2e_evaluation_is_compact_swe6_style_and_preserves_low_high_logic(tmp_path: Path):
    data = _mode_fixture()
    # Make Low and High both system-pending to mirror the intended human-reviewed example.
    high = data["requirements"][1]
    high["allocation_status"] = "SYSTEM_REQUIREMENT_PENDING_SW_ALLOCATION"
    high["swe1_eligibility"] = "Review Needed"
    high["swe6_eligibility"] = "Deferred pending SW allocation"
    high["sys5_eligibility"] = "Eligible"
    for a in high["fact_level_allocations"]:
        a["allocation_status"] = "SYSTEM_REQUIREMENT_PENDING_SW_ALLOCATION"
        a["swe6_eligibility"] = "Deferred pending SW allocation"
        a["verification_domain"] = "System Integration / SYS.5"
    data["testability_and_decomposition_result"]["by_srs"][1]["swe6_eligibility"] = "Deferred pending SW allocation"
    data["testability_and_decomposition_result"]["by_srs"][1]["verification_domain"] = "System Integration / SYS.5"
    data.pop("finalized_verification_bundle", None)
    src=tmp_path/"source.docx"; src.write_text("x",encoding="utf-8")
    path=E2EExporter(tmp_path).export_evaluation_excel(src,data)
    wb=load_workbook(path,read_only=True,data_only=False)
    try:
        assert wb.sheetnames == ["표지","0_변경이력","1_평가요약","2_E2E_평가케이스","3_검토필요","4_요구사항-평가추적성"]
        ws=wb["2_E2E_평가케이스"]
        header={str(ws.cell(3,c).value or ""):c for c in range(1,ws.max_column+1)}
        # Human Main keeps compact IDs/domain and blank result fields.
        by_srs={}
        for r in range(4,ws.max_row+1):
            human_req=str(ws.cell(r,header["요구사항 ID"]).value or "")
            sid=human_req.split("_BF_")[0] if "_BF_" in human_req else ""
            if sid: by_srs.setdefault(sid,[]).append(r)
        assert len(by_srs["SRS_004"]) == 1
        assert len(by_srs["SRS_005"]) == 1
        for r in (by_srs["SRS_004"][0], by_srs["SRS_005"][0]):
            assert ws.cell(r,header["검증 영역"]).value == "SYS.5"
            for name in ("실제 출력값","판정 (PASS/FAIL)","비고","결과 화면 캡처"):
                assert ws.cell(r,header[name]).value in (None,"")
        # Transition logic/target are intentionally off the human Main but remain in Trace/Audit.
        tr=wb["4_요구사항-평가추적성"]
        th={str(tr.cell(3,c).value or ""):c for c in range(1,tr.max_column+1)}
        trace_by_srs={}
        for r in range(4,tr.max_row+1):
            sid=str(tr.cell(r,th["상위 SRS ID"]).value or "")
            if sid: trace_by_srs.setdefault(sid,[]).append(r)
        low=trace_by_srs["SRS_004"][0]; highr=trace_by_srs["SRS_005"][0]
        assert tr.cell(low,th["원본 객체 유형"]).value == "SYS5_CANDIDATE"
        assert tr.cell(highr,th["원본 객체 유형"]).value == "SYS5_CANDIDATE"
        assert tr.cell(low,th["Condition Logic"]).value == "ALL_OF"
        assert tr.cell(highr,th["Condition Logic"]).value == "ANY_OF"
        assert "Low Power Mode" in str(tr.cell(low,th["Expected Target State"]).value or "")
        assert "High Power Mode" in str(tr.cell(highr,th["Expected Target State"]).value or "")
    finally:
        wb.close()


def test_v086_sys5_parent_fallback_is_review_object_not_executable_candidate():
    data = {"requirements":[{
        "srs_id":"SRS_ONLY_SW", "requirement":"ECU는 값을 계산해야 한다.",
        "sys5_eligibility":"Review Needed", "swe6_eligibility":"Eligible",
        "fact_level_allocations":[{
            "source_fact_fragment_id":"F_SW", "source_semantic_unit_id":"U1", "source_location":"Page 1",
            "source_fact":"값을 계산한다", "swe6_eligibility":"Eligible", "verification_domain":"SWE.6 Software Qualification",
        }],
    }], "testability_and_decomposition_result":{"by_srs":[]}}
    bundle=finalize_verification_single_truth(data, force=True)
    e2e=[x for x in bundle["e2e_evaluation_cases"] if x["parent_srs_id"]=="SRS_ONLY_SW"]
    assert any(x["test_object_type"] == "SYS5_PARENT_SCOPE_REVIEW" for x in e2e)
    # Legacy/raw SYS.5 builder remains backward compatible for prior artifacts.
    raw=build_sys5_candidates(data)
    assert raw and raw[0]["sys5_id"].startswith("SYS5_TC_")


def test_v086_structural_and_semantic_gates_are_separate():
    data=_mode_fixture()
    bundle=finalize_verification_single_truth(data, force=True)
    assert bundle["single_truth_audit"]["release_gate_status"] == "PASS"
    assert bundle["semantic_verification_quality_audit"]["status"] in {"PASS","REVIEW_REQUIRED"}
    assert "e2e_evaluation_cases" in bundle



def test_v086_full_e2e_workbook_snapshot_has_all_sheets_and_rows(tmp_path: Path):
    data = _mode_fixture()
    src=tmp_path/"source.docx"; src.write_text("x",encoding="utf-8")
    ex=E2EExporter(tmp_path)
    req_x=ex.export_requirements_excel(src,data)
    eval_x=ex.export_evaluation_excel(src,data)
    from core.evaluation_orchestrator import MultiModelEvaluationOrchestrator
    req_snap=MultiModelEvaluationOrchestrator._xlsx_snapshot(req_x, profile="integrated_requirements")
    eval_snap=MultiModelEvaluationOrchestrator._xlsx_snapshot(eval_x, profile="integrated_tests")
    assert set(req_snap["sheets"]) == {"표지","0_변경이력","1_요구사항요약","2_E2E_요구사항","3_검토필요","4_Source_추적성"}
    assert set(eval_snap["sheets"]) == {"표지","0_변경이력","1_평가요약","2_E2E_평가케이스","3_검토필요","4_요구사항-평가추적성"}
    assert req_snap["evaluation_scope_complete"] is True and req_snap["evaluation_coverage_percent"] == 100.0
    assert eval_snap["evaluation_scope_complete"] is True and eval_snap["evaluation_coverage_percent"] == 100.0
    assert not req_snap["truncated_sheets"] and not eval_snap["truncated_sheets"]


def test_v086_requirements_rev70_has_effective_projection_and_resolved_direct_conflicts():
    from openpyxl import load_workbook
    root=Path(__file__).resolve().parents[1]
    path=root/"history"/"requirements"/"Requirement_Studio_Requirements_Management_rev70.xlsx"
    wb=load_workbook(path,read_only=True,data_only=False)
    try:
        assert "AI_11_현재유효요구사항" in wb.sheetnames
        hist=wb["AI_02_요구사항목록"]
        status={str(row[0].value or ""): str(row[5].value or "") for row in hist.iter_rows(min_row=4) if row and row[0].value}
        for fr in ("FR-011","FR-012","FR-045","FR-050","FR-051","FR-054","FR-060","FR-091","FR-101","FR-103","FR-120"):
            assert status.get(fr)=="대체됨"
        eff=wb["AI_11_현재유효요구사항"]
        active={str(row[0].value or "") for row in eff.iter_rows(min_row=5) if row and row[0].value}
        assert "FR-011" not in active and "FR-050" not in active and "FR-091" not in active
        assert {"FR-128","FR-098","FR-115","FR-170","FR-188","FR-458","FR-468"}.issubset(active)
        lock=wb["AI_05_변경금지조건"]
        assert "Requirement_Studio_V<Major>.<Minor>.zip" in str(lock["B9"].value)
    finally:
        wb.close()



def test_v086_unified_contract_contains_legacy_plus_e2e_artifacts(tmp_path: Path):
    import json
    from types import SimpleNamespace
    from core.review_exchange import ReviewExchangeBuilder
    from core.evaluation_orchestrator import MultiModelEvaluationOrchestrator
    from core.swe1_exporter import SWE1Exporter
    from core.swe6_exporter import SWE6Exporter

    data=_mode_fixture(); finalize_verification_single_truth(data, force=True)
    source=tmp_path/"source.docx"; source.write_text("source placeholder",encoding="utf-8")
    out=tmp_path/"out"; ex=E2EExporter(out); s1=SWE1Exporter(out); s6=SWE6Exporter(out)
    artifacts={
        "swe1_word":s1.export_word(source,data),
        "swe1_excel":s1.export_excel(source,data),
        "swe6_excel":s6.export_excel(source,data,source_text="Low Power Mode High Power Mode"),
        "e2e_requirements_word":ex.export_requirements_word(source,data),
        "e2e_requirements_excel":ex.export_requirements_excel(source,data),
        "e2e_evaluation_excel":ex.export_evaluation_excel(source,data),
    }
    root=tmp_path/"project"; (root/"history"/"change_decisions").mkdir(parents=True); (root/"config").mkdir(parents=True)
    (root/"config"/"evaluation_api_config.json").write_text("{}",encoding="utf-8")
    meta=SimpleNamespace(display_name="H-Chat / GPT",provider_id="hchat_gpt",model="gpt-5.6-terra",endpoint_family="hchat",auth_mode="key",project_header_enabled=False,response_parsing_mode="text")
    result=ReviewExchangeBuilder(root,app_version="v0.86").build(
        source_document=source,provider_metadata=meta,analysis_text="",requirement_data=data,evaluation={},normalized_summary={},
        quality_review_config={"enabled":True,"evaluation_artifact":"unified","selected_outputs":{k:True for k in artifacts}},
        artifact_paths=artifacts,evaluation_mode="initial")
    run_dir=Path(result["run_dir"])
    payload=json.loads((run_dir/"02_REQUIREMENT_STUDIO_REVIEW_PACKAGE.json").read_text(encoding="utf-8"))
    contract=payload["review_handoff"]["unified_artifact_contract"]
    assert contract["status"]=="PASS"
    assert set(contract["required_artifacts"])==set(artifacts)
    evidence,_=MultiModelEvaluationOrchestrator(root,app_version="v0.86").build_evidence_package(run_dir)
    assert set(evidence["actual_artifacts"])==set(artifacts)
    assert evidence["actual_artifacts"]["e2e_requirements_excel"]["evaluation_scope_complete"] is True
    assert evidence["actual_artifacts"]["e2e_evaluation_excel"]["evaluation_scope_complete"] is True
    assert evidence["actual_artifacts"]["e2e_evaluation_excel"]["evaluation_coverage_percent"]==100.0
