from pathlib import Path

from core.cross_document_semantics import apply_allocation_gate, build_semantic_source_units, attach_semantic_traceability
from core.quality_audit import apply_quality_audits, build_testability_result, finalize_test_intent_coverage, normalize_requirement_extensions
from core.swe6_exporter import build_swe6_cases, SWE6Exporter
from openpyxl import load_workbook


def _req(text, source_text=None, srs='SRS_001'):
    r={
        'candidate_id':'C1','srs_id':srs,'scenario_candidate_id':'SCN-CAND-001','category':'기능','function_name':'Feature',
        'requirement':text,'user_input':'','system_input_preconditions':'','processing_action':text,'output':text,
        'acceptance_criteria':text,'failure_situations':[],'user_intervention_points':[],'derivation_type':'explicit',
        'derivation_reason':'','clarification_needed':[],'source_evidence':[{'document':'x.docx','location':'Paragraph 1','text':source_text or text}],
        'confidence':0.95,'classification_basis':'source','activation_trigger':'','preconditions':'','behavior_flows':[],
        'evaluation_method':'','exception_conditions':[],'related_artifacts':[],'source_requirement_ids':[],
    }
    normalize_requirement_extensions(r)
    return r


def _data(reqs):
    if not isinstance(reqs,list): reqs=[reqs]
    return {'schema_version':'REQ-STUDIO-CANONICAL-REQ-1.9','source_document':'x.docx','scenario_candidates':[{
        'scenario_candidate_id':'SCN-CAND-001','scenario_name':'x','user_goal_context':'','scenario_flow':[],
        'expected_outcome':'','source_evidence':[]}], 'requirements':reqs,'gaps':[]}


def test_parent_pending_child_inherits_without_explicit_override():
    req=_req('시스템은 Low Power Mode로 진입한다.','시스템은 Low Power Mode로 진입한다.')
    unit={'source_semantic_unit_id':'SRC-SEM-X','source_chunk_id':'C1','source_location':'Paragraph 1','source_excerpt':'밝기 설정값을 저장한다.'}
    apply_allocation_gate(req,[unit])
    assert req['swe6_eligibility'].startswith('Deferred')
    child=req['fact_level_allocations'][0]
    assert child['allocation_status']=='INHERIT_PARENT_ALLOCATION_PENDING'
    assert child['swe6_eligibility'].startswith('Deferred')
    assert child['allocation_override'] is False


def test_cross_domain_primary_domains_and_external_dependency_are_separate_axes():
    req=_req('Source facts review','Source facts review')
    units=[
        {'source_semantic_unit_id':'E','source_chunk_id':'C1','source_location':'Table 1 / Row 1','source_excerpt':'최대 소비 전류 MAX 100mA'},
        {'source_semantic_unit_id':'T','source_chunk_id':'C2','source_location':'Table 1 / Row 2','source_excerpt':'사용 온도 -40℃ ~ +85℃'},
        {'source_semantic_unit_id':'X','source_chunk_id':'C3','source_location':'Table 1 / Row 3','source_excerpt':'ES95400-30 규격을 참조한다'},
    ]
    apply_allocation_gate(req,units)
    assert req['cross_domain_bundle_review_required'] is True
    assert 'Electrical Verification' in req['cross_domain_verification_domains']
    assert 'Environmental Qualification' in req['cross_domain_verification_domains']
    assert 'External Standard Evidence Required' in req['external_dependency_domains']
    assert req['mixed_external_dependency_review_required'] is True


def test_table_governance_and_variant_headers_are_review_context():
    compact='''[DOCUMENT] x.docx
[SRC C1 | Paragraph 411 | text]
실제 기능 구현 시 차종에 맞게 parameter 종류 및 값을 설정해야 한다.
[SRC C2 | Table 4 / Row 1 | table]
세부 기능 | 차종 Variant | 적용 차종
[SRC C3 | Table 5 / Row 2 | table]
Par_NotPConfirmationTime | 600 ms
'''
    units=build_semantic_source_units(compact)
    by={u['source_location']:u for u in units}
    assert by['Paragraph 411']['coverage_eligibility']=='review_context'
    assert by['Paragraph 411']['source_unit_type']=='parameter_governance_context'
    assert by['Table 4 / Row 1']['coverage_eligibility']=='review_context'
    assert by['Table 5 / Row 2']['coverage_eligibility']=='semantic_unit'


def test_source_table_fact_join_can_feed_existing_source_value_to_tc(tmp_path):
    compact='''[DOCUMENT] x.docx
[SRC C1 | Table 5 / Row 2 | table]
Par_NotPConfirmationTime = 600 ms
'''
    req=_req('Par_NotPConfirmationTime 조건을 확인한다.','Par_NotPConfirmationTime = 600 ms')
    data=_data(req)
    attach_semantic_traceability(data,compact)
    r=data['requirements'][0]
    # force SW eligibility to test exporter mechanics after provenance join
    r.update({'requirement_level':'Software','allocation_status':'SW_IMPLEMENTATION_REQUIREMENT','swe1_eligibility':'Eligible','swe6_eligibility':'Eligible','verification_domain':'SWE.6 Software Qualification','semantic_provenance_status':'COMPLETE'})
    if not r['source_backed_atomic_behaviors']:
        r['source_backed_atomic_behaviors']=[{'behavior_text':'Par_NotPConfirmationTime = 600 ms','knowledge_state':'KNOWN','source_semantic_unit_id':r['source_semantic_unit_ids'][0]}]
    data['testability_and_decomposition_result']=build_testability_result(data)
    cases=build_swe6_cases(data)
    assert cases
    c=cases[0]
    triples=[(c['prep_var'],c['prep_value']),(c['exec_var'],c['exec_value']),(c['expected_var'],c['expected_value'])]
    assert any(v=='Par_NotPConfirmationTime' and '600' in val for v,val in triples)
    out=SWE6Exporter(tmp_path).export_excel(tmp_path/'x.docx',data,source_text='')
    wb=load_workbook(out,read_only=True)
    try:
        assert '5_Source_Fact_Audit' in wb.sheetnames
    finally:
        wb.close()


def test_repeated_missing_explicit_requirement_is_recovered_only_with_runtime_source_identity():
    compact='''[DOCUMENT] x.docx
[SRC C1 | Paragraph 10 | text]
[REQ03_09] 풋램프 제어기는 적용 상태에 따라 RGB Type 또는 White LED Type으로 동작한다.
[SRC C2 | Paragraph 30 | text]
[REQ03_09] 풋램프 제어기는 적용 상태에 따라 RGB Type 또는 White LED Type으로 동작한다.
'''
    data={'source_document':'x.docx','requirements':[],'gaps':[],'scenario_candidates':[]}
    out=apply_quality_audits(data,compact)
    assert out['explicit_missing_recovery']['recovered_candidate_count']==1
    assert any('REQ03_09' in [str(x).upper() for x in r.get('source_requirement_ids',[])] for r in out['requirements'])
    assert out['source_coverage']['actual_missing_behavior_count']==0


def test_swe6_audit_uses_three_level_status_and_child_semantics():
    req=_req('입력을 수신하고 NVM에 저장한다.')
    req.update({'requirement_level':'Software','allocation_status':'SW_IMPLEMENTATION_REQUIREMENT','swe1_eligibility':'Eligible','swe6_eligibility':'Eligible','verification_domain':'SWE.6 Software Qualification','source_semantic_unit_ids':['S1','S2'],'source_backed_atomic_behaviors':[{'source_semantic_unit_id':'S1','behavior_text':'입력을 수신한다.','knowledge_state':'KNOWN'},{'source_semantic_unit_id':'S2','behavior_text':'NVM에 저장한다.','knowledge_state':'KNOWN'}],'semantic_provenance_status':'COMPLETE'})
    data=_data(req)
    data['testability_and_decomposition_result']=build_testability_result(data)
    cases=build_swe6_cases(data)
    finalize_test_intent_coverage(data,cases)
    audit=data['swe6_export_preservation_audit']
    assert audit['audit_status'] in {'PASS','PASS_WITH_REVIEW_ITEMS','FAIL'}
    row=audit['rows'][0]
    assert all(x['intent_status'] in {'REPRESENTED_IN_GENERIC_TC','INDEPENDENTLY_COVERED','DEFERRED','REVIEW_REQUIRED'} for x in row['source_backed_child_intents'])
