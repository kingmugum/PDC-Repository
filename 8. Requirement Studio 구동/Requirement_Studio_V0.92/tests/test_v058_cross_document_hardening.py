from pathlib import Path

from docx import Document

from core.cross_document_semantics import apply_allocation_gate
from core.quality_audit import apply_quality_audits
from core.requirement_engine import RequirementEngine
from core.swe1_exporter import SWE1Exporter, build_swe1_records


def _req(cid, srs, loc, source_text, requirement=None):
    requirement = requirement or source_text
    return {
        "candidate_id": cid, "srs_id": srs, "scenario_candidate_id": "SCN-CAND-001",
        "category": "비기능", "function_name": "Cross Domain", "requirement": requirement,
        "user_input": "", "system_input_preconditions": "", "processing_action": requirement,
        "output": "", "acceptance_criteria": requirement, "failure_situations": [], "user_intervention_points": [],
        "derivation_type": "explicit", "derivation_reason": "", "clarification_needed": [],
        "source_evidence": [{"document": "system.docx", "location": loc, "text": source_text}],
        "confidence": 0.95, "classification_basis": "입력문서 명시", "activation_trigger": "",
        "preconditions": "", "behavior_flows": [], "evaluation_method": "", "exception_conditions": [],
        "related_artifacts": [],
    }


def _data(reqs, gaps=None):
    return {
        "schema_version": "REQ-STUDIO-CANONICAL-REQ-1.9",
        "source_document": "system.docx",
        "scenario_candidates": [{
            "scenario_candidate_id": "SCN-CAND-001", "scenario_name": "x", "user_goal_context": "",
            "scenario_flow": [], "expected_outcome": "",
            "source_evidence": [{"document": "system.docx", "location": "Paragraph 1", "text": "x"}],
        }],
        "requirements": reqs, "gaps": gaps or [],
    }


def test_high_confidence_non_sw_domain_assignment_ignores_noun_false_positives():
    cases = [
        ("제어기 커넥터는 DIP 타입을 적용하고 LEAD-WIRE 적용을 금지한다.", "MECHANICAL_CONNECTOR_REQUIREMENT"),
        ("제어기 자동납땜을 적용한다.", "MANUFACTURING_PROCESS_REQUIREMENT"),
        ("제어기 동작 전압의 정격 전압은 DC 13.5V이다.", "ELECTRICAL_REQUIREMENT"),
        ("제어기 사용 온도 범위는 -40℃ ~ +85℃이다.", "ENVIRONMENTAL_QUALIFICATION_REQUIREMENT"),
    ]
    for i, (text, expected) in enumerate(cases, 1):
        req = _req(f"C{i}", f"SRS_{i:03d}", f"Paragraph {i}", text)
        apply_allocation_gate(req)
        assert req["allocation_status"] == expected
        assert req["swe1_eligibility"] == "Not Applicable"
        assert req["swe6_eligibility"] == "Not Applicable"


def test_semantic_disposition_matrix_and_gap_scope_filter():
    compact = """[DOCUMENT] system.docx
[SRC DOC-C0001 | Paragraph 10 | text]
B+ 인가 후 Slave 초기화 시간 60ms를 대기한다.
[SRC DOC-C0002 | Paragraph 20 | text]
커넥터는 DIP 타입을 적용한다.
[SRC DOC-C0003 | Paragraph 30 | text]
CAN DB의 상세 Signal 정의는 차종별 자료를 따른다.
"""
    reqs = [
        _req("C1", "SRS_001", "Paragraph 10", "B+ 인가 후 Slave 초기화 시간 60ms를 대기한다."),
        _req("C2", "SRS_002", "Paragraph 20", "커넥터는 DIP 타입을 적용한다."),
    ]
    # Artificially broad input link; V0.58 must reduce it after enrichment according to section scope.
    for i in range(3, 11):
        reqs.append(_req(f"C{i}", f"SRS_{i:03d}", f"Paragraph {100+i}", f"독립 기능 {i}를 수행한다."))
    gaps = [{
        "gap_id": "G1", "description": "CAN DB 상세 Signal 정의가 필요하다.", "gap_scope": "section",
        "related_candidate_ids": [r["candidate_id"] for r in reqs],
        "related_srs_ids": [r["srs_id"] for r in reqs],
        "source_evidence": [{"document": "system.docx", "location": "Paragraph 30", "text": "CAN DB의 상세 Signal 정의는 차종별 자료를 따른다."}],
    }]
    data = _data(reqs, gaps)
    apply_quality_audits(data, compact)
    sem = data["semantic_source_unit_coverage"]
    assert sem["eligible_unit_count"] >= 2
    assert len(sem["disposition_records"]) == sem["eligible_unit_count"]
    assert sem["disposition_coverage_percent"] == 100.0
    assert all(x.get("disposition_status") for x in sem["disposition_records"])
    gap = data["gaps"][0]
    assert len(gap.get("related_srs_ids") or []) < len(reqs)
    assert len(gap.get("related_candidate_ids") or []) < len(reqs)


def test_numeric_relation_strengthening_is_removed():
    compact = """[DOCUMENT] system.docx
[SRC DOC-C0001 | Table 10 / Row 2 | table]
최대 소비 전류 | MAX 100mA
[SRC DOC-C0002 | Table 10 / Row 3 | table]
암전류 규제 : 0.1mA
"""
    reqs = [
        _req("C1", "SRS_001", "Table 10 / Row 2", "최대 소비 전류 | MAX 100mA", "최대 소비 전류는 100mA 이하이어야 한다."),
        _req("C2", "SRS_002", "Table 10 / Row 3", "암전류 규제 : 0.1mA", "암전류는 0.1mA이어야 한다."),
    ]
    data = _data(reqs)
    apply_quality_audits(data, compact)
    for req in reqs:
        assert req["normative_strength_status"] == "SOURCE_RELATION_AMBIGUOUS_GENERATED_STRENGTHENING_REMOVED"
        assert req["value_relation_status"]
        assert "Source 표기 보존" in req["requirement"]
        assert "확인" in req["acceptance_criteria"]
        assert req["clarification_needed"]


def test_main_requirement_export_preserves_review_required_and_annexes_preserve_others(tmp_path: Path):
    eligible = _req("C1", "SRS_001", "Paragraph 1", "소프트웨어는 EEPROM Check-sum을 점검해야 한다.")
    eligible.update({"requirement_level":"Software","allocation_status":"SW_IMPLEMENTATION_REQUIREMENT","swe1_eligibility":"Eligible","swe6_eligibility":"Eligible","verification_domain":"SWE.6 Software Qualification"})
    pending = _req("C2", "SRS_002", "Paragraph 2", "Master는 CAN 오류 시 Low Power Mode로 진입한다.")
    pending.update({"requirement_level":"System","allocation_status":"SYSTEM_REQUIREMENT_PENDING_SW_ALLOCATION","swe1_eligibility":"Review Needed","swe6_eligibility":"Deferred pending SW allocation","verification_domain":"System Integration / SYS.5"})
    notsw = _req("C3", "SRS_003", "Paragraph 3", "커넥터는 DIP 타입을 적용한다.")
    notsw.update({"requirement_level":"Hardware/Mechanical","allocation_status":"MECHANICAL_CONNECTOR_REQUIREMENT","swe1_eligibility":"Not Applicable","swe6_eligibility":"Not Applicable","verification_domain":"Mechanical / Connector Verification"})
    data = _data([eligible, pending, notsw])
    assert {x["SRS ID"] for x in build_swe1_records(data)} == {"SRS_001"}
    src = tmp_path / "system.docx"
    Document().save(src)
    out = SWE1Exporter(tmp_path).export_word(src, data)
    doc = Document(out)
    all_text = "\n".join(p.text for p in doc.paragraphs)
    # V0.89: the duplicate human-facing Review Annex sections are retired.
    assert "Review Annex — Applicability / Issues / Dependencies" not in all_text
    assert "Allocation Review Annex — Pending SWE.1 Allocation" not in all_text
    assert "Cross-domain Allocation Annex — Source Facts Outside SWE.1" not in all_text
    # REVIEW_REQUIRED/native requirements remain visible in the main engineering view;
    # detailed review context is preserved in E2E Requirements/Review Package JSON.
    assert "SRS_002" in all_text
    assert "SRS_003" in all_text


def test_structure_uses_statement_candidate_name_and_semantic_invariant(tmp_path: Path):
    req = _req("C1", "SRS_001", "Paragraph 1", "소프트웨어는 값을 저장해야 한다.")
    # Complete required extension fields but deliberately omit semantic provenance for a no-ID Eligible item.
    from core.quality_audit import normalize_requirement_extensions
    normalize_requirement_extensions(req)
    req.update({
        "requirement_level":"Software","allocation_status":"SW_IMPLEMENTATION_REQUIREMENT","swe1_eligibility":"Eligible",
        "swe6_eligibility":"Eligible","verification_domain":"SWE.6 Software Qualification","allocation_rationale":"x",
        "semantic_provenance_status":"INCOMPLETE_REVIEW_REQUIRED", "value_relation_status":[],
        "normative_strength_status":"SOURCE_RELATION_EXPLICIT_OR_NOT_NUMERIC", "normative_strength_findings":[],
    })
    data = _data([req])
    # supply top-level audit fields required by 1.6
    for k, v in {
        "source_coverage":{},"conflict_register":[],"naming_issue_register":[],"reference_integrity_result":{},
        "export_preservation_audit":{},"unsupported_generation_report":{},"testability_and_decomposition_result":{},
        "review_findings":[],"reference_completeness_result":{},"test_intent_coverage":{},"regression_report":{},
        "semantic_source_units":[],"semantic_source_unit_coverage":{},"semantic_provenance_audit":{},
        "semantic_source_unit_classifier_audit":{},
    }.items(): data[k]=v
    engine = RequirementEngine(tmp_path)
    engine.schema_path = Path(__file__).resolve().parents[1] / "contracts" / "canonical_requirement_schema.json"
    result = engine.evaluate_structure(data)
    assert "explicit_statement_candidate_count" in result["summary"]
    assert "explicit_count" not in result["summary"]
    assert any("Semantic provenance" in x for x in result["errors"])
