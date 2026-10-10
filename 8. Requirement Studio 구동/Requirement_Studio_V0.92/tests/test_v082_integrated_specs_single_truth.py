from pathlib import Path

from openpyxl import load_workbook

from core.cross_document_semantics import apply_allocation_gate
from core.integrated_exporter import IntegratedExporter
from core.verification_single_truth import (
    detect_mode_transition_allocation_findings,
    finalize_verification_single_truth,
)


def _fact(uid: str, fid: str, text: str, location: str, *, swe6: str, domain: str, status: str):
    return {
        "source_semantic_unit_id": uid,
        "source_fact_fragment_id": fid,
        "source_location": location,
        "source_fact": text,
        "swe6_eligibility": swe6,
        "verification_domain": domain,
        "allocation_status": status,
    }


def _atom(uid: str, fid: str, text: str, location: str):
    return {
        "source_semantic_unit_id": uid,
        "source_fact_fragment_id": fid,
        "source_location": location,
        "behavior_text": text,
    }


def _mode_fixture():
    low_atoms = [
        _atom("U-LOW", "F-LOW-1", "OFF 명령 수신 시", "Paragraph 125"),
        _atom("U-LOW", "F-LOW-2", "무드램프 OFF 상태", "Paragraph 126"),
        _atom("U-LOW", "F-LOW-3", "슬레이브 제어기 SLEEP(OFF) 상태", "Paragraph 127"),
        _atom("U-LOW", "F-LOW-4", "B-CAN SLEEP MODE 상태", "Paragraph 128"),
        _atom("U-LOW", "F-LOW-5", "모든 조건(&)이 참이면 Low Power Mode(Sleep Mode)로 전환한다.", "Paragraph 129"),
    ]
    low_facts = [
        _fact("U-LOW", a["source_fact_fragment_id"], a["behavior_text"], a["source_location"],
              swe6="Deferred pending SW allocation", domain="System Integration / SYS.5",
              status="SYSTEM_REQUIREMENT_PENDING_SW_ALLOCATION")
        for a in low_atoms
    ]
    low = {
        "srs_id": "SRS_004",
        "candidate_id": "REQ-CAND-004",
        "function_name": "Low Power Mode",
        "category": "기능",
        "requirement": "무드램프 마스터 제어기는 OFF 명령 수신, 무드램프 OFF, Slave Sleep(OFF), B-CAN Sleep 조건이 모두 참이면 Low Power Mode(Sleep Mode)로 전환해야 한다.",
        "processing_action": "모든 조건(&)이 참이면 Low Power Mode(Sleep Mode)로 전환한다.",
        "acceptance_criteria": "모든 조건(&)이 참이면 Low Power Mode(Sleep Mode)로 전환한다.",
        "source_evidence": [{"location": "Paragraph 125-129", "text": "Low Power Mode 전환 조건"}],
        "source_semantic_unit_ids": ["U-LOW"],
        "source_fact_fragments": [{
            "source_fact_fragment_id": a["source_fact_fragment_id"],
            "parent_source_semantic_unit_id": "U-LOW",
            "source_location": a["source_location"],
            "source_excerpt": a["behavior_text"],
        } for a in low_atoms],
        "source_backed_atomic_behaviors": low_atoms,
        "fact_level_allocations": low_facts,
        "semantic_provenance_status": "COMPLETE",
        "source_requirement_ids": [],
        "engineering_domains": ["System"],
        "requirement_level": "System",
        "allocation_status": "SYSTEM_REQUIREMENT_PENDING_SW_ALLOCATION",
        "sys1_eligibility": "Eligible",
        "swe1_eligibility": "Review Needed",
        "sys5_eligibility": "Eligible",
        "swe6_eligibility": "Deferred pending SW allocation",
        "verification_domain": "System Integration / SYS.5",
        "verification_domains": ["SYS.5 System Qualification"],
        "human_decision_required": True,
        "review_status": "HUMAN_DECISION_REQUIRED",
        "canonical_state": "CANONICAL_REVIEW_REQUIRED",
        "hold_reason": "Software allocation is not explicitly approved.",
        "clarification_needed": [],
        "external_dependencies": [],
    }

    high_atoms = [
        _atom("U-HIGH", "F-HIGH-1", "ON 명령 수신 시", "Paragraph 132"),
        _atom("U-HIGH", "F-HIGH-2", "B-CAN WAKE UP 상태", "Paragraph 133"),
        _atom("U-HIGH", "F-HIGH-3", "TRANSCEIVER의 ERR PIN이 HIGH에서 LOW 상태 변환 시", "Paragraph 134"),
        _atom("U-HIGH", "F-HIGH-4", "최초 B+ 인가 시", "Paragraph 135"),
        _atom("U-HIGH", "F-HIGH-5", "조건 중 하나라도(or) 참이면 High Power Mode로 전환한다.", "Paragraph 136"),
    ]
    high_facts = [
        _fact("U-HIGH", a["source_fact_fragment_id"], a["behavior_text"], a["source_location"],
              swe6="Eligible", domain="SWE.6 Software Qualification (Allocation Review Required)",
              status="SW_BEHAVIOR_INFERRED_REVIEW_REQUIRED")
        for a in high_atoms
    ]
    high = {
        "srs_id": "SRS_005",
        "candidate_id": "REQ-CAND-005",
        "function_name": "High Power Mode",
        "category": "기능",
        "requirement": "무드램프 마스터 제어기는 ON 명령, B-CAN Wake, ERR PIN 변화 또는 최초 B+ 인가 중 하나라도 참이면 High Power Mode로 전환해야 한다.",
        "processing_action": "조건 중 하나라도(or) 참이면 High Power Mode로 전환한다.",
        "acceptance_criteria": "조건 중 하나라도(or) 참이면 High Power Mode로 전환한다.",
        "source_evidence": [{"location": "Paragraph 132-136", "text": "High Power Mode 전환 조건"}],
        "source_semantic_unit_ids": ["U-HIGH"],
        "source_fact_fragments": [{
            "source_fact_fragment_id": a["source_fact_fragment_id"],
            "parent_source_semantic_unit_id": "U-HIGH",
            "source_location": a["source_location"],
            "source_excerpt": a["behavior_text"],
        } for a in high_atoms],
        "source_backed_atomic_behaviors": high_atoms,
        "fact_level_allocations": high_facts,
        "semantic_provenance_status": "COMPLETE",
        "source_requirement_ids": [],
        "engineering_domains": ["Software", "System"],
        "requirement_level": "Software Candidate",
        "allocation_status": "SW_BEHAVIOR_INFERRED_REVIEW_REQUIRED",
        "sys1_eligibility": "Eligible",
        "swe1_eligibility": "Eligible",
        "sys5_eligibility": "Review Needed",
        "swe6_eligibility": "Eligible",
        "verification_domain": "SWE.6 Software Qualification (Allocation Review Required)",
        "verification_domains": ["SWE.6 Software Qualification", "SYS.5 System Qualification"],
        "human_decision_required": True,
        "review_status": "HUMAN_DECISION_REQUIRED",
        "canonical_state": "CANONICAL_REVIEW_REQUIRED",
        "hold_reason": "Inferred software allocation requires human approval.",
        "clarification_needed": [],
        "external_dependencies": [],
    }

    data = {
        "schema_version": "REQ-STUDIO-CANONICAL-REQ-2.0",
        "source_document": "MLM.docx",
        "requirements": [low, high],
        "gaps": [],
        "semantic_source_unit_coverage": {"open_review_needed_records": []},
        "testability_and_decomposition_result": {
            "by_srs": [
                {
                    "srs_id": "SRS_004",
                    "swe6_eligibility": "Deferred pending SW allocation",
                    "verification_domain": "System Integration / SYS.5",
                    "required_test_intents": ["Normal / Positive", "State Transition"],
                },
                {
                    "srs_id": "SRS_005",
                    "swe6_eligibility": "Eligible",
                    "verification_domain": "SWE.6 Software Qualification (Allocation Review Required)",
                    "required_test_intents": ["Normal / Positive", "State Transition"],
                    "normal_test_available": True,
                    "test_design_feasible": True,
                },
            ],
            "summary": {},
        },
    }
    return data


def test_v082_low_power_is_visible_as_sys5_and_allocation_pending_even_without_swe6_tc():
    data = _mode_fixture()
    bundle = finalize_verification_single_truth(data)
    low_objects = [x for x in bundle["integrated_test_objects"] if x["parent_srs_id"] == "SRS_004"]
    assert any(x["test_object_type"] == "SYS5_CANDIDATE" for x in low_objects)
    assert any(x["test_object_type"] == "ALLOCATION_PENDING_INTENT" for x in low_objects)
    assert not any(x["test_object_type"] == "SWE6_TC" for x in low_objects)
    assert all(x.get("source_semantic_unit_ids") for x in low_objects)
    assert all(x.get("source_fact_fragment_ids") for x in low_objects)
    assert bundle["single_truth_audit"]["release_gate_status"] == "PASS"


def test_v082_low_power_source_facts_remain_traceable():
    data = _mode_fixture()
    bundle = finalize_verification_single_truth(data)
    low_objects = [x for x in bundle["integrated_test_objects"] if x["parent_srs_id"] == "SRS_004"]
    combined = "\n".join(str(x.get("source_backed_behavior") or "") + "\n" + str(x.get("test_intent") or "") for x in low_objects)
    for literal in ("OFF 명령", "무드램프 OFF", "SLEEP(OFF)", "B-CAN SLEEP", "Low Power Mode"):
        assert literal in combined


def test_v082_low_high_power_allocation_divergence_creates_review_finding_not_reallocation():
    data = _mode_fixture()
    findings = detect_mode_transition_allocation_findings(data)
    assert findings
    assert any(x["finding_type"] == "MODE_TRANSITION_ALLOCATION_INCONSISTENCY" for x in findings)
    assert any(set(x["affected_srs_ids"]) == {"SRS_004", "SRS_005"} for x in findings)
    assert data["requirements"][0]["allocation_status"] == "SYSTEM_REQUIREMENT_PENDING_SW_ALLOCATION"
    assert data["requirements"][1]["allocation_status"] == "SW_BEHAVIOR_INFERRED_REVIEW_REQUIRED"


def test_v082_trigger_only_is_not_positive_software_allocation_evidence():
    req = {
        "srs_id": "SRS_TRIGGER",
        "requirement": "TRANSCEIVER의 ERR PIN이 HIGH에서 LOW 상태 변환 시",
        "function_name": "ERR PIN Trigger",
        "processing_action": "",
        "acceptance_criteria": "",
        "source_evidence": [{"location": "Paragraph 134", "text": "TRANSCEIVER의 ERR PIN이 HIGH에서 LOW 상태 변환 시"}],
        "source_requirement_ids": [],
        "source_semantic_unit_ids": ["U-TRIG"],
        "source_backed_atomic_behaviors": [_atom("U-TRIG", "F-TRIG", "TRANSCEIVER의 ERR PIN이 HIGH에서 LOW 상태 변환 시", "Paragraph 134")],
        "semantic_provenance_status": "COMPLETE",
    }
    apply_allocation_gate(req, [{
        "source_semantic_unit_id": "U-TRIG",
        "source_fact_fragment_id": "F-TRIG",
        "source_location": "Paragraph 134",
        "source_excerpt": req["requirement"],
    }])
    assert req.get("allocation_status") != "SW_BEHAVIOR_INFERRED_REVIEW_REQUIRED"
    assert req.get("allocation_status") != "SW_IMPLEMENTATION_REQUIREMENT"


def test_v082_integrated_workbooks_are_additive_and_have_requested_views(tmp_path: Path):
    data = _mode_fixture()
    source = tmp_path / "MLM.docx"
    source.write_text("source placeholder", encoding="utf-8")
    exporter = IntegratedExporter(tmp_path)
    req_path = exporter.export_requirements_excel(source, data)
    test_path = exporter.export_tests_excel(source, data)

    req_wb = load_workbook(req_path, read_only=True)
    try:
        assert {
            "00_Overview", "01_Integrated_Eng_Req", "02_SYS1_View", "03_SWE1_View", "04_Allocation_Review",
            "05_Source_Traceability", "06_External_Dependencies", "07_Open_Issues_Gaps", "08_Semantic_Unit_Coverage",
            "09_Fact_Fragment_Allocation",
        }.issubset(set(req_wb.sheetnames))
        ws = req_wb["01_Integrated_Eng_Req"]
        srs_ids = {str(ws.cell(r, 1).value or "") for r in range(4, ws.max_row + 1)}
        assert {"SRS_004", "SRS_005"}.issubset(srs_ids)
        alloc_ws = req_wb["04_Allocation_Review"]
        assert any("MODE_TRANSITION_ALLOCATION_INCONSISTENCY" in {str(alloc_ws.cell(r, c).value or "") for c in range(1, min(6, alloc_ws.max_column) + 1)} for r in range(4, alloc_ws.max_row + 1))
    finally:
        req_wb.close()

    test_wb = load_workbook(test_path, read_only=True)
    try:
        assert {
            "00_Overview", "01_Integrated_Test_Objects", "02_SYS5_Candidates", "03_SWE6_Test_Cases", "04_Deferred_Intents",
            "05_Allocation_Pending_Intents", "06_External_Dependency_Intents", "07_Non_SW_Verification_Intents",
            "08_Source_Intent_Traceability", "09_Fact_Allocation_Audit", "10_Execution_Result_Entry", "11_Release_Gate_Summary",
        }.issubset(set(test_wb.sheetnames))
        alloc_ws = test_wb["05_Allocation_Pending_Intents"]
        low_rows = [r for r in range(4, alloc_ws.max_row + 1) if str(alloc_ws.cell(r, 3).value or "") == "SRS_004"]
        assert low_rows
        assert all(str(alloc_ws.cell(r, 17).value or "") == "ALLOCATION_PENDING" for r in low_rows)
        sys5_ws = test_wb["02_SYS5_Candidates"]
        assert any(str(sys5_ws.cell(r, 3).value or "") == "SRS_004" for r in range(4, sys5_ws.max_row + 1))
        swe6_ws = test_wb["03_SWE6_Test_Cases"]
        assert any(str(swe6_ws.cell(r, 3).value or "") == "SRS_005" for r in range(4, swe6_ws.max_row + 1))
    finally:
        test_wb.close()


def test_v082_integrated_test_execution_result_cells_are_physically_blank(tmp_path: Path):
    data = _mode_fixture()
    source = tmp_path / "MLM.docx"
    source.write_text("source placeholder", encoding="utf-8")
    path = IntegratedExporter(tmp_path).export_tests_excel(source, data)
    wb = load_workbook(path, read_only=True, data_only=False)
    try:
        for ws_name in ("01_Integrated_Test_Objects", "02_SYS5_Candidates", "03_SWE6_Test_Cases", "04_Deferred_Intents", "05_Allocation_Pending_Intents"):
            ws = wb[ws_name]
            for r in range(4, ws.max_row + 1):
                if not str(ws.cell(r, 2).value or "").strip():
                    continue
                assert all(ws.cell(r, c).value in (None, "") for c in (27, 28, 29, 30))
        result_ws = wb["10_Execution_Result_Entry"]
        for r in range(4, result_ws.max_row + 1):
            if not str(result_ws.cell(r, 2).value or "").strip():
                continue
            assert all(result_ws.cell(r, c).value in (None, "") for c in (5, 6, 7, 8))
    finally:
        wb.close()


def test_v082_integrated_test_ids_and_counts_share_one_finalized_bundle(tmp_path: Path):
    data = _mode_fixture()
    bundle = finalize_verification_single_truth(data)
    assert bundle["summary"]["swe6_tc_count"] == 1
    assert bundle["summary"]["sys5_candidate_count"] >= 1
    assert bundle["single_truth_audit"]["expected_deferred_intent_count"] == bundle["single_truth_audit"]["integrated_deferred_intent_count"]
    assert bundle["single_truth_audit"]["blocking_issue_count"] == 0
    # Calling again must reuse the exact finalized in-memory object, preventing count/ID drift between exporters.
    assert finalize_verification_single_truth(data) is bundle


def test_v082_unified_review_exchange_carries_integrated_artifacts_and_single_truth(tmp_path: Path):
    import json
    from types import SimpleNamespace
    from core.review_exchange import ReviewExchangeBuilder

    data = _mode_fixture()
    finalize_verification_single_truth(data)
    source = tmp_path / "MLM.docx"
    source.write_text("source placeholder", encoding="utf-8")
    out_dir = tmp_path / "out"
    exporter = IntegratedExporter(out_dir)
    req_xlsx = exporter.export_requirements_excel(source, data)
    test_xlsx = exporter.export_tests_excel(source, data)

    root = tmp_path / "project"
    (root / "history" / "change_decisions").mkdir(parents=True)
    (root / "history" / "change_decisions" / "04_CHANGE_DECISION_V0.81_to_V0.82.txt").write_text("decision", encoding="utf-8")
    meta = SimpleNamespace(
        display_name="H-Chat / GPT", provider_id="hchat_gpt", model="gpt-5.6-terra",
        endpoint_family="hchat", auth_mode="key", project_header_enabled=False,
        response_parsing_mode="text",
    )
    result = ReviewExchangeBuilder(root, app_version="v0.82").build(
        source_document=source,
        provider_metadata=meta,
        analysis_text="",
        requirement_data=data,
        evaluation={},
        normalized_summary={},
        quality_review_config={"enabled": True, "evaluation_artifact": "unified", "selected_outputs": {
            "integrated_requirements_excel": True, "integrated_tests_excel": True,
        }},
        artifact_paths={"integrated_requirements_excel": req_xlsx, "integrated_tests_excel": test_xlsx},
        evaluation_mode="initial",
    )
    run_dir = Path(result["run_dir"])
    payload = json.loads((run_dir / "02_REQUIREMENT_STUDIO_REVIEW_PACKAGE.json").read_text(encoding="utf-8"))
    assert set(payload["review_handoff"]["unified_artifacts"]) == {"integrated_requirements_excel", "integrated_tests_excel"}
    assert list(run_dir.glob("03_INTEGRATED_REQ_*.xlsx"))
    assert list(run_dir.glob("03_INTEGRATED_TEST_*.xlsx"))
    calc = payload["calculated_review_evidence"]
    assert calc["integrated_single_truth_audit"]["release_gate_status"] == "PASS"
    assert calc["mode_transition_allocation_findings"]
    assert str(calc["finalized_verification_bundle"]["schema_version"]).startswith("REQ-STUDIO-VERIFICATION-BUNDLE-")
