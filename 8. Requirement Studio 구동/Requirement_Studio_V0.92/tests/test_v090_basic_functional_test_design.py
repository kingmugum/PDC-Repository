from pathlib import Path

from openpyxl import load_workbook

from core.basic_functional_test_design import BASIC_FUNCTIONAL_MODE, build_basic_functional_cases
from core.e2e_exporter import E2EExporter


def _base(sid="SRS_001"):
    return {
        "parent_srs_id": sid,
        "test_object_type": "SWE6_TC",
        "test_object_id": "TC_001",
        "source_fact_fragment_ids": ["F1", "F2"],
        "verification_objective": "Source-backed basic behavior",
        "verification_domain": "SWE.6 Software Qualification",
        "swe6_eligibility": "Eligible",
        "allocation_status": "SW_BEHAVIOR_INFERRED_REVIEW_REQUIRED",
        "execution_readiness": "REVIEW_REQUIRED",
        "human_review_required": True,
    }


def test_v090_one_srs_can_expand_to_multiple_basic_functional_cases_without_srs_split():
    req = {
        "srs_id": "SRS_001",
        "function_name": "전원 출력 제어",
        "requirement": "입력 전원 상태에 따라 시스템 출력 전원을 제어해야 한다.",
        "source_backed_atomic_behaviors": [
            {"source_fact_fragment_id": "F1", "source_semantic_unit_id": "U1", "source_location": "P1", "behavior_text": "입력 전원이 12V가 되면 시스템 출력 전원이 5V가 되어야 한다"},
            {"source_fact_fragment_id": "F2", "source_semantic_unit_id": "U1", "source_location": "P1", "behavior_text": "입력전원이 OFF 일 경우 출력 전원이 0V가 되어야 한다"},
        ],
    }
    data={"requirements":[req]}
    cases,reviews,audit=build_basic_functional_cases(data,[_base()])
    assert len(data["requirements"]) == 1
    assert len(cases) == 2
    assert reviews == []
    assert audit["mode"] == BASIC_FUNCTIONAL_MODE
    assert [(c["input_value"],c["output_value"]) for c in cases] == [("12V","5V"),("OFF","0V")]
    assert all(c["test_design_mode"] == BASIC_FUNCTIONAL_MODE for c in cases)


def _capture_fixture():
    source='''2.9 실외 이미지 캡처
HU는 ADAS_PRK에 신호(SVM_CaptureModeViewRequest)를 송신하여 각 뷰 모드에 대해 Front->Rear->Left->Right 순서로 영상을 요청해야 한다.
HU는 ADAS_PRK로부터 신호(SVM_CaptureModeViewState)를 수신 받아 Front->Rear->Left->Right 순서의 영상 송출 여부를 확인한다.
SVM_CaptureModeViewRequest값 송신 조건
0x0:All Black Screen Printing Request
0x1:Front View Printing Request 전방 요청
0x2:Rear View Printing Request 후방 요청
0x3:Left View Printing Request 좌측 요청
0x4:Right View Printing Request 우측 요청
0x5:Top View Printing Request 탑뷰 요청
0x6:SVM Video Request
0x7:Invalid
SVM_CaptureModeViewState값 송신 조건
0x0:Default
0x1:Front 전방 응답
0x2:Rear 후방 응답
0x3:Left 좌측 응답
0x4:Right 우측 응답
0x5:Top 탑뷰 응답
0x6:Camera Input Error 카메라 고장 응답
0x7:Invalid
'''
    req={
        "srs_id":"SRS_009", "function_name":"실외 이미지 순차 캡처", "requirement":"실외 이미지 순차 캡처",
        "allocation_status":"SW_BEHAVIOR_INFERRED_REVIEW_REQUIRED", "swe6_eligibility":"Eligible", "sys5_eligibility":"Review Needed",
        "source_evidence":[{"location":"Page 12, 2.9","text":source}],
        "source_backed_atomic_behaviors":[
            {"source_fact_fragment_id":"F1","source_semantic_unit_id":"U1","source_location":"Page 12","behavior_text":"HU는 ADAS_PRK에 신호(SVM_CaptureModeViewRequest)를 송신하여 각 뷰 모드에 대해 Front->Rear->Left->Right 순서로 영상을 요청해야 한다."},
            {"source_fact_fragment_id":"F2","source_semantic_unit_id":"U1","source_location":"Page 12","behavior_text":"HU는 ADAS_PRK로부터 신호(SVM_CaptureModeViewState)를 수신 받아 Front->Rear->Left->Right 순서의 영상 송출 여부를 확인한다."},
            {"source_fact_fragment_id":"F3","source_semantic_unit_id":"U1","source_location":"Page 12","behavior_text":"SVM_CaptureModeViewRequest값 0x1 Front 0x2 Rear 0x3 Left 0x4 Right 0x5 Top 0x7 Invalid"},
            {"source_fact_fragment_id":"F4","source_semantic_unit_id":"U1","source_location":"Page 12","behavior_text":"SVM_CaptureModeViewState값 0x1 Front 0x2 Rear 0x3 Left 0x4 Right 0x5 Top 0x6 Camera Input Error 0x7 Invalid"},
        ],
        "fact_level_allocations":[
            {"source_fact_fragment_id":f"F{i}","source_semantic_unit_id":"U1","source_location":"Page 12","source_fact":"capture fact","allocation_status":"SW_BEHAVIOR_INFERRED_REVIEW_REQUIRED","swe6_eligibility":"Eligible","verification_domain":"SWE.6 Software Qualification"}
            for i in range(1,5)
        ],
        "semantic_provenance_status":"COMPLETE",
        "source_semantic_unit_ids":["U1"],
    }
    base=_base("SRS_009"); base["source_fact_fragment_ids"]=["F1","F2","F3","F4"]; base["test_object_id"]="TC_008"
    return source,req,base


def test_v090_capture_sequence_is_one_tc_with_vectors_top_is_separate_and_invalid_is_review():
    _,req,base=_capture_fixture()
    cases,reviews,audit=build_basic_functional_cases({"requirements":[req]},[base])
    assert len(cases)==2
    seq=next(c for c in cases if c["basic_case_kind"]=="SEQUENCE_REQUEST_RESPONSE")
    top=next(c for c in cases if c["basic_case_kind"]=="ENUM_REQUEST_RESPONSE")
    assert seq["input_variable"]=="SVM_CaptureModeViewRequest"
    assert seq["output_variable"]=="SVM_CaptureModeViewState"
    assert [x["input_value"] for x in seq["test_vectors"]]==["0x1","0x2","0x3","0x4"]
    assert [x["expected_output"] for x in seq["test_vectors"]]==["0x1","0x2","0x3","0x4"]
    assert top["input_value"]=="0x5" and top["output_value"]=="0x5"
    assert any(r["value"]=="0x7" for r in reviews)
    assert any(r["value"]=="0x6" and "Error" in r["meaning"] for r in reviews)
    assert all("0x7" not in str(c.get("input_value")) for c in cases)
    assert any("No EP/BVA/fault-injection expansion" in x for x in audit["guardrails"])


def test_v090_e2e_workbook_puts_source_backed_input_and_output_in_separate_columns(tmp_path: Path):
    source,req,_=_capture_fixture()
    # Finalizer needs enough raw fields to generate a SWE.6 case; this fixture mirrors the current V0.90 basic flow.
    req.update({"category":"기능","candidate_id":"REQ_009","scenario_candidate_id":"SCN_001","processing_action":req["requirement"],"output":"각 뷰의 영상 송출 확인","acceptance_criteria":"각 뷰의 영상 송출 확인"})
    data={"requirements":[req],"testability_and_decomposition_result":{"by_srs":[{"srs_id":"SRS_009","swe6_eligibility":"Eligible","verification_domain":"SWE.6 Software Qualification"}]}}
    src=tmp_path/"source.docx"; src.write_text(source,encoding="utf-8")
    out=E2EExporter(tmp_path).export_evaluation_excel(src,data)
    wb=load_workbook(out,read_only=True,data_only=False)
    try:
        ws=wb["2_E2E_평가케이스"]
        h={str(ws.cell(3,c).value or ""):c for c in range(1,ws.max_column+1)}
        rows=[r for r in range(4,ws.max_row+1) if str(ws.cell(r,h["요구사항 ID"]).value or "").startswith("SRS_009_BF_")]
        assert len(rows)>=2
        seq=next(r for r in rows if "순차 요청" in str(ws.cell(r,h["테스트 케이스 이름"]).value or ""))
        assert ws.cell(seq,h["입력 변수"]).value=="SVM_CaptureModeViewRequest"
        assert "0x1" in str(ws.cell(seq,h["입력값"]).value or "") and "0x4" in str(ws.cell(seq,h["입력값"]).value or "")
        assert ws.cell(seq,h["출력 변수 / 관찰 대상"]).value=="SVM_CaptureModeViewState"
        assert "0x1" in str(ws.cell(seq,h["기대 출력값"]).value or "") and "0x4" in str(ws.cell(seq,h["기대 출력값"]).value or "")
        review=wb["3_검토필요"]
        assert any(str(review.cell(r,4).value or "") in {"Source 기대동작 미정의","Exact Fact Trace 필요","Source TBD","Source Conflict"} for r in range(4,review.max_row+1))
    finally:
        wb.close()
