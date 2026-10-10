from pathlib import Path

from openpyxl import load_workbook

from core.basic_functional_test_design import (
    BASIC_FUNCTIONAL_MODE,
    build_basic_functional_cases,
    classify_source_backed_facts,
)
from core.e2e_exporter import E2EExporter


def _base(sid: str, fids: list[str]):
    return {
        "parent_srs_id": sid,
        "test_object_type": "SWE6_TC",
        "test_object_id": f"TC_{sid}",
        "source_fact_fragment_ids": fids,
        "verification_objective": "Source-backed basic behavior",
        "verification_domain": "SWE.6 Software Qualification",
        "swe6_eligibility": "Eligible",
        "allocation_status": "SW_BEHAVIOR_INFERRED_REVIEW_REQUIRED",
        "execution_readiness": "REVIEW_REQUIRED",
        "human_review_required": True,
    }


def _req(sid: str, name: str, facts: list[tuple[str, str]], *, evidence: str = "", requirement: str = ""):
    req = {
        "srs_id": sid,
        "candidate_id": f"REQ_{sid}",
        "scenario_candidate_id": f"SCN_{sid}",
        "function_name": name,
        "requirement": requirement or name,
        "category": "기능",
        "processing_action": requirement or name,
        "output": "Source-defined 결과",
        "acceptance_criteria": "Source 정의와 일치해야 한다.",
        "allocation_status": "SW_BEHAVIOR_INFERRED_REVIEW_REQUIRED",
        "swe6_eligibility": "Eligible",
        "sys5_eligibility": "Review Needed",
        "source_backed_atomic_behaviors": [
            {
                "source_fact_fragment_id": fid,
                "source_semantic_unit_id": f"U_{sid}",
                "source_location": "Page 1",
                "behavior_text": text,
            }
            for fid, text in facts
        ],
        "fact_level_allocations": [
            {
                "source_fact_fragment_id": fid,
                "source_semantic_unit_id": f"U_{sid}",
                "source_location": "Page 1",
                "source_fact": text,
                "allocation_status": "SW_BEHAVIOR_INFERRED_REVIEW_REQUIRED",
                "swe6_eligibility": "Eligible",
                "verification_domain": "SWE.6 Software Qualification",
            }
            for fid, text in facts
        ],
        "source_semantic_unit_ids": [f"U_{sid}"],
        "semantic_provenance_status": "COMPLETE",
    }
    if evidence:
        req["source_evidence"] = [{"location": "Page 1", "text": evidence}]
    else:
        req["source_evidence"] = [{"location": "Page 1", "text": "\n".join(t for _, t in facts)}]
    return req


def _build(req):
    fids = [x["source_fact_fragment_id"] for x in req.get("source_backed_atomic_behaviors", [])]
    return build_basic_functional_cases({"requirements": [req]}, [_base(req["srs_id"], fids)])


def _exec_cases(cases):
    return [c for c in cases if c.get("test_design_mode") == BASIC_FUNCTIONAL_MODE]


def test_v092_fixture_01_conditional_io_splits_independent_conditions_and_has_pass_criteria():
    req = _req("SRS_001", "전원 출력 제어", [
        ("F001", "입력 전원이 12V가 되면 시스템 출력 전원이 5V가 되어야 한다"),
        ("F002", "입력전원이 OFF 일 경우 출력 전원이 0V가 되어야 한다"),
    ])
    cases, reviews, audit = _build(req)
    cases = _exec_cases(cases)
    assert len(cases) == 2
    assert reviews == []
    assert {(c["input_value"], c["output_value"]) for c in cases} == {("12V", "5V"), ("OFF", "0V")}
    assert all(c["pass_criterion"].endswith("PASS") for c in cases)
    assert all(c["field_fact_trace"]["input_value"] and c["field_fact_trace"]["output_value"] for c in cases)
    assert audit["schema_version"].endswith("0.92")


def test_v092_fixture_02_wakeup_and_variant_are_separate_primary_objectives():
    req = _req("SRS_002", "System 및 B-CAN wake-up", [
        ("F001", "HU는 CCS로부터 내 차 주변 영상 전송 요청을 수신하면 System wake-up 해야 한다."),
        ("F002", "HU는 System wake-up 완료 후 신호(C_PreCrankReq)를 0x1(Req on)으로 송출하여 B-CAN wake-up 해야 한다."),
        ("F003", "HU는 HEV/PHEV/EV 차종인 경우 IGN3 On 해야 한다."),
    ])
    cases, _, _ = _build(req)
    cases = _exec_cases(cases)
    assert len(cases) == 2
    precrank = next(c for c in cases if c.get("output_variable") == "C_PreCrankReq")
    variant = next(c for c in cases if c.get("output_variable") == "IGN3")
    assert precrank["output_value"].startswith("0x1")
    assert "System wake-up" in precrank["precondition"]
    assert "F002" in precrank["field_fact_trace"]["output_value"]
    assert variant["input_variable"] == "차종"
    assert variant["input_value"] == "HEV/PHEV/EV"
    assert variant["output_value"] == "On"
    assert variant["split_decision"] == "SPLIT_BY_VARIANT_OBJECTIVE"


def test_v092_fixture_03_mirror_control_and_state_decision_are_two_objectives_but_table_is_vectorized():
    evidence = """미러 상태 판단 표
Mirror_UnFldCtrlSta | Mirror_FldCtrlSta | 미러 상태
Any | 0x1,0x2 | 접힘
0x1 | Any | 접힘
0x0,0x2,0x3 | 0x0,0x3 | 펼침"""
    req = _req("SRS_003", "아웃사이드 미러 제어 및 상태 판단", [
        ("F001", "HU는 실시간 감시모드 동작 Command 수신 시 OMUOSMirCtrl을 0x2(Unfold)로 송신하고 종료 Command 수신 시 OMUOSMirCtrl을 0x1(Fold)로 송신한다."),
        ("F002", "Mirror_UnFldCtrlSta 및 Mirror_FldCtrlSta를 수신하여 아웃사이드 미러 상태를 판단한다."),
        ("F003", "Mirror_UnFldCtrlSta와 Mirror_FldCtrlSta 조합에 따라 접힘 또는 펼침 상태를 판단한다."),
    ], evidence=evidence)
    cases, _, _ = _build(req)
    cases = _exec_cases(cases)
    assert len(cases) == 2
    control = next(c for c in cases if c.get("output_variable") == "OMUOSMirCtrl")
    decision = next(c for c in cases if c.get("basic_case_kind") == "DECISION_TABLE")
    assert len(control["test_vectors"]) == 2
    assert {v["expected_output"] for v in control["test_vectors"]} == {"0x1", "0x2"}
    assert decision["split_decision"] == "KEEP_AS_VECTOR"
    assert len(decision["test_vectors"]) == 3
    assert decision["output_variable"] == "미러 상태"


def test_v092_fixture_04_normal_and_timeout_exception_are_separate():
    req = _req("SRS_004", "부저 제어 및 부저 미응답 처리", [
        ("F001", "아웃사이드 미러 펼침 제어 이후 서버 Command의 부저 출력 주기에 따라 HU_Buzzer_OnOffReq를 0x1(Active)로 송신하여 부저를 출력한다."),
        ("F002", "부저 출력 요청 후 피드백 신호가 5초 동안 미수신 시 부저 미응답으로 판단하고 촬영을 중단한다. 부저 미응답에는 별도 Retry 없음."),
    ])
    cases, _, _ = _build(req)
    cases = _exec_cases(cases)
    assert len(cases) == 2
    normal = next(c for c in cases if c.get("output_variable") == "HU_Buzzer_OnOffReq")
    timeout = next(c for c in cases if c.get("basic_case_kind") == "TIMEOUT_EXCEPTION")
    assert normal["output_value"].startswith("0x1")
    assert timeout["input_value"] == "5초"
    assert "촬영 중단" in timeout["expected_state"]
    assert "Retry 없음" in timeout["expected_state"]
    assert timeout["split_decision"] == "SPLIT_EXCEPTION_PATH"


def _capture_req():
    source = """2.9 실외 이미지 View 요청/응답
HU는 SVM_CaptureModeViewRequest를 송신하여 Front->Rear->Left->Right 순서로 영상을 요청해야 한다.
HU는 SVM_CaptureModeViewState를 수신하여 동일 순서의 영상 송출 여부를 확인해야 한다.
SVM_CaptureModeViewRequest값
0x1:Front 전방 요청
0x2:Rear 후방 요청
0x3:Left 좌측 요청
0x4:Right 우측 요청
0x5:Top 탑뷰 요청
0x7:Invalid
SVM_CaptureModeViewState값
0x1:Front 전방 응답
0x2:Rear 후방 응답
0x3:Left 좌측 응답
0x4:Right 우측 응답
0x5:Top 탑뷰 응답
0x6:Camera Input Error
0x7:Invalid
"""
    return _req("SRS_005", "실외 이미지 View 요청/응답", [
        ("F001", "SVM_CaptureModeViewRequest는 Front->Rear->Left->Right 순서로 요청해야 한다."),
        ("F002", "SVM_CaptureModeViewState는 Front->Rear->Left->Right 순서의 응답을 확인해야 한다."),
        ("F003", "SVM_CaptureModeViewRequest값 0x1 Front 0x2 Rear 0x3 Left 0x4 Right 0x5 Top 0x7 Invalid"),
        ("F004", "SVM_CaptureModeViewState값 0x1 Front 0x2 Rear 0x3 Left 0x4 Right 0x5 Top 0x6 Camera Input Error 0x7 Invalid"),
    ], evidence=source, requirement="실외 이미지 View 요청/응답")


def test_v092_fixture_05_sequence_is_one_tc_top_is_separate_and_undefined_values_are_review():
    req = _capture_req()
    cases, reviews, _ = _build(req)
    cases = _exec_cases(cases)
    assert len(cases) == 2
    seq = next(c for c in cases if c["basic_case_kind"] == "SEQUENCE_REQUEST_RESPONSE")
    top = next(c for c in cases if c["basic_case_kind"] == "ENUM_REQUEST_RESPONSE")
    assert seq["split_decision"] == "KEEP_AS_SEQUENCE"
    assert [v["input_value"] for v in seq["test_vectors"]] == ["0x1", "0x2", "0x3", "0x4"]
    assert top["input_value"] == "0x5" and top["output_value"] == "0x5"
    assert any(r["review_type"] == "SOURCE_DEFINED_VALUE_WITHOUT_EXPECTED_BEHAVIOR" and r["value"] == "0x7" for r in reviews)
    assert any(r["value"] == "0x6" for r in reviews)


def test_v092_fixture_06_resize_properties_and_size_each_keep_exact_fact_trace():
    req = _req("SRS_006", "이미지 처리 및 저장", [
        ("F001", "Top View 입력 영상은 500×800이다."),
        ("F002", "HU는 Top View를 450×720으로 resizing하여 사용해야 한다."),
        ("F003", "HU는 처리된 영상을 JPEG 형식으로 저장해야 한다."),
        ("F004", "저장 이미지 해상도는 1920×3240이어야 한다."),
        ("F005", "저장 이미지 크기는 0.25 MB 이하이어야 한다."),
    ])
    cases, reviews, audit = _build(req)
    cases = _exec_cases(cases)
    assert len(cases) == 3
    assert reviews == []
    resize = next(c for c in cases if c["basic_case_kind"] == "DIMENSION_TRANSFORM")
    props = next(c for c in cases if c["basic_case_kind"] == "OUTPUT_ARTIFACT_PROPERTIES")
    size = next(c for c in cases if c["basic_case_kind"] == "OUTPUT_ARTIFACT_SIZE_CONSTRAINT")
    assert resize["field_fact_trace"]["input_value"] == ["F001"]
    assert resize["field_fact_trace"]["output_value"] == ["F002"]
    assert set(props["field_fact_trace"]["output_value"]) == {"F003", "F004"}
    assert size["field_fact_trace"]["output_value"] == ["F005"]
    assert all(c["exact_field_trace_status"] == "PASS" for c in cases)
    assert audit["exact_field_trace_review_count"] == 0


def test_v092_fixture_07_conflict_and_tbd_do_not_become_executable_cases():
    req = _req("SRS_007", "Capture Mode Interface / Top View TBD", [
        ("F001", "기존 사양은 SVM_CaptureModeCMD를 사용하고 변경 검토안은 ICMU_CaptureModeCMD 추가를 제안하며 최종 승인 Production Signal은 미확정이다."),
        ("F002", "SVM_TopViewInfo는 Signal(TBD)로 표시되어 최종 Signal 정의가 확정되지 않았다."),
    ])
    cases, reviews, _ = _build(req)
    assert _exec_cases(cases) == []
    assert any(c.get("test_design_mode") == "BASIC_FUNCTIONAL_REVIEW" for c in cases)
    assert len(reviews) >= 2
    assert all("재생성" in r["required_resolution"] or "disposition" in r["required_resolution"] for r in reviews)


def test_v092_fixture_08_state_transition_logic_target_and_inhibit_are_structured():
    req = _req("SRS_008", "전원 모드 상태전이", [
        ("F001", "ACC=OFF이고 IGN1=OFF인 경우 시스템은 Low Power 상태로 전이해야 한다."),
        ("F002", "ACC=ON 또는 IGN1=ON인 경우 시스템은 High Power 상태로 전이해야 한다."),
        ("F003", "Monitoring 상태에서는 Low Power 전이를 수행하지 않는다."),
    ])
    cases, _, _ = _build(req)
    cases = _exec_cases(cases)
    assert len(cases) == 3
    by_logic = {c["condition_logic"]: c for c in cases}
    assert by_logic["ALL_OF"]["expected_target_state"] == "Low Power"
    assert by_logic["ANY_OF"]["expected_target_state"] == "High Power"
    assert by_logic["INHIBIT"]["expected_target_state"] == "NOT Low Power"
    assert all(c["pass_criterion"] for c in cases)


def test_v092_fact_role_classification_is_multilabel_and_source_preserving():
    req = _req("SRS_009", "Role Classification", [
        ("F001", "부저 출력 요청 후 피드백 신호가 5초 동안 미수신되면 촬영을 중단하고 Retry를 수행하지 않는다."),
    ])
    roles = classify_source_backed_facts(req)[0]["roles"]
    assert "PRECONDITION" in roles or "TRIGGER" in roles
    assert "EXCEPTION" in roles
    assert "NEGATIVE_REQUIREMENT" in roles
    assert "CONSTRAINT" in roles


def test_v092_e2e_human_layout_uses_structured_fields_and_keeps_internal_trace(tmp_path: Path):
    req = _capture_req()
    data = {
        "requirements": [req],
        "testability_and_decomposition_result": {
            "by_srs": [{"srs_id": "SRS_005", "swe6_eligibility": "Eligible", "verification_domain": "SWE.6 Software Qualification"}]
        },
    }
    src = tmp_path / "source.docx"
    src.write_text(req["source_evidence"][0]["text"], encoding="utf-8")
    out = E2EExporter(tmp_path).export_evaluation_excel(src, data)
    wb = load_workbook(out, read_only=False, data_only=False)
    try:
        ws = wb["2_E2E_평가케이스"]
        headers = [str(ws.cell(3, c).value or "") for c in range(1, ws.max_column + 1)]
        assert headers == [
            "Test Case ID", "요구사항 ID", "검증 영역", "평가 방법 / 도출 방법", "테스트 케이스 이름", "테스트 목적",
            "사전 조건", "평가 환경 / 준비", "입력 변수", "비교", "입력값", "수행 절차",
            "출력 변수 / 관찰 대상", "비교", "기대 출력값", "기대 상태 / 결과", "PASS 기준",
            "TC 검토 상태", "검토 사유 / 필요 조치", "실제 출력값", "판정 (PASS/FAIL)", "비고", "결과 화면 캡처",
        ]
        assert {str(x) for x in ws.merged_cells.ranges}.issuperset({"A2:F2", "G2:H2", "I2:L2", "M2:P2", "Q2:S2", "T2:W2"})
        h = {str(ws.cell(3, c).value or ""): c for c in range(1, ws.max_column + 1)}
        data_rows = [r for r in range(4, ws.max_row + 1) if ws.cell(r, 1).value]
        assert any(str(ws.cell(r, h["테스트 목적"]).value or "").strip() for r in data_rows)
        assert any(str(ws.cell(r, h["PASS 기준"]).value or "").strip() for r in data_rows)
        for r in data_rows:
            for c in range(20, 24):
                assert ws.cell(r, c).value in (None, "")
        tr = wb["4_요구사항-평가추적성"]
        th = [str(tr.cell(3, c).value or "") for c in range(1, tr.max_column + 1)]
        assert "Split Decision" in th
        assert "Exact Field Trace Status" in th
        assert "Field Fact Trace" in th
    finally:
        wb.close()
