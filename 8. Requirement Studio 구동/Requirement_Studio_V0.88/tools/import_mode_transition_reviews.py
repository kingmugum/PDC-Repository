from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import load_workbook

from core.review_decision_registry import load_mode_transition_decisions, save_mode_transition_decisions


def import_reviews(path: Path) -> int:
    wb = load_workbook(path, read_only=True, data_only=False)
    try:
        ws = wb["11_Release_Gate_Summary"]
        decisions = load_mode_transition_decisions()
        imported = 0
        # V0.85 columns: A type, B finding id, C severity, D affected, E disposition,
        # F owner, G details, H resolution, I rationale/reviewed-at note.
        for row in range(4, ws.max_row + 1):
            finding_type = str(ws.cell(row, 1).value or "").strip()
            finding_id = str(ws.cell(row, 2).value or "").strip()
            disposition = str(ws.cell(row, 5).value or "").strip().upper()
            if finding_type != "MODE_TRANSITION_ALLOCATION_INCONSISTENCY" or not finding_id:
                continue
            if not disposition or disposition == "PENDING":
                continue
            decisions[finding_id] = {
                "disposition": disposition,
                "owner": str(ws.cell(row, 6).value or "").strip(),
                "rationale": str(ws.cell(row, 9).value or "").strip(),
                "reviewed_at": "",
            }
            imported += 1
        if imported:
            save_mode_transition_decisions(decisions)
        return imported
    finally:
        wb.close()


def main() -> int:
    parser = argparse.ArgumentParser(description="Import mode-transition human review decisions from Integrated Test XLSX.")
    parser.add_argument("xlsx", type=Path)
    args = parser.parse_args()
    count = import_reviews(args.xlsx)
    print(f"Imported mode-transition decisions: {count}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
