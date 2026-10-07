from __future__ import annotations

import re
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from core.output_naming import output_filename
from core.swe1_exporter import build_engineering_records
from core.verification_single_truth import finalize_verification_single_truth

NOT_CONFIRMED = "입력문서에서 확인되지 않음"


def _list(value: Any) -> list[Any]:
    if value is None:
        return []
    return value if isinstance(value, list) else [value]


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, dict):
        return "\n".join(f"{k}: {_text(v)}" for k, v in value.items() if v not in (None, "", [], {}))
    if isinstance(value, (list, tuple, set)):
        return "\n".join(x for x in (_text(v).strip() for v in value) if x)
    return str(value).strip()


def _safe_stem(name: str) -> str:
    stem = Path(name).stem or "document"
    stem = re.sub(r'[<>:"/\\|?*]+', "_", stem)
    stem = re.sub(r"\s+", "_", stem).strip("._ ")
    return stem[:80] or "document"


def _source_trace(req: dict[str, Any]) -> tuple[str, str, str]:
    sems = [str(x) for x in _list(req.get("source_semantic_unit_ids")) if str(x)]
    fids: list[str] = []
    locs: list[str] = []
    for key in ("source_backed_atomic_behaviors", "fact_level_allocations", "source_fact_fragments"):
        for item in _list(req.get(key)):
            if not isinstance(item, dict):
                continue
            fid = str(item.get("source_fact_fragment_id") or "")
            if fid and fid not in fids:
                fids.append(fid)
            sem = str(item.get("source_semantic_unit_id") or item.get("parent_source_semantic_unit_id") or "")
            if sem and sem not in sems:
                sems.append(sem)
            loc = _text(item.get("source_location"))
            if loc and loc not in locs:
                locs.append(loc)
    for ev in _list(req.get("source_evidence")):
        if isinstance(ev, dict):
            loc = _text(ev.get("location"))
            if loc and loc not in locs:
                locs.append(loc)
    return ", ".join(sems), ", ".join(fids), " | ".join(locs)


def _source_evidence(req: dict[str, Any]) -> str:
    lines=[]
    for ev in _list(req.get("source_evidence")):
        if isinstance(ev, dict):
            loc=_text(ev.get("location")); txt=_text(ev.get("text")); doc=_text(ev.get("document"))
            prefix=" / ".join(x for x in (doc,loc) if x)
            val=f"{prefix}: {txt}" if prefix and txt else (prefix or txt)
        else:
            val=_text(ev)
        if val and val not in lines:
            lines.append(val)
    return "\n".join(lines) or NOT_CONFIRMED


def _scope_label(req: dict[str, Any]) -> str:
    sys1=str(req.get("sys1_eligibility") or "")
    swe1=str(req.get("swe1_eligibility") or "")
    if sys1 in {"Eligible","Review Needed"} and swe1 in {"Eligible","Review Needed"}:
        return "SYS.1 + SWE.1"
    if swe1 in {"Eligible","Review Needed"}:
        return "SWE.1"
    if sys1 in {"Eligible","Review Needed"}:
        return "SYS.1"
    return "Engineering Native / Review"


class E2EExporter:
    """V0.88 reviewer-facing E2E Requirements + Evaluation specifications.

    E2E Requirements is the Source-nearest engineering view (SYS.1 + SWE.1 + native/review
    preservation).  E2E Evaluation is the human-facing verification master (SYS.5 + SWE.6 +
    deferred/external/native) while raw objects remain available in the finalized bundle for audit.
    """

    DARK = "1F4E78"
    HEADER = "D9EAF7"
    SUB = "EAF2F8"
    REVIEW = "FFF2CC"
    FAIL = "F4CCCC"
    PASS = "D9EAD3"
    THIN = Side(style="thin", color="B7B7B7")

    def __init__(self, output_dir: Path):
        self.output_dir=Path(output_dir).resolve(); self.output_dir.mkdir(parents=True, exist_ok=True)

    def _filename(self, prefix: str, source: Path, ext: str) -> Path:
        return self.output_dir / output_filename(prefix, _safe_stem(source.name), ext)

    def _style_header(self, cell, fill: str | None = None):
        cell.fill=PatternFill("solid", fgColor=fill or self.DARK)
        cell.font=Font(name="Malgun Gothic", size=9, bold=True, color="FFFFFF" if (fill or self.DARK)==self.DARK else "000000")
        cell.alignment=Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border=Border(left=self.THIN,right=self.THIN,top=self.THIN,bottom=self.THIN)

    def _style_cell(self, cell, *, center=False, fill=None):
        cell.font=Font(name="Malgun Gothic", size=9)
        cell.alignment=Alignment(horizontal="center" if center else "left", vertical="top", wrap_text=True)
        cell.border=Border(left=self.THIN,right=self.THIN,top=self.THIN,bottom=self.THIN)
        if fill: cell.fill=PatternFill("solid", fgColor=fill)

    def _write_table(self, ws, headers, rows, widths, start_row=3):
        for c,h in enumerate(headers,1): self._style_header(ws.cell(start_row,c,h), self.HEADER)
        for r,row in enumerate(rows or [["-"]+[""]*(len(headers)-1)], start_row+1):
            for c,val in enumerate(row,1): self._style_cell(ws.cell(r,c,val), center=c<=3)
        ws.freeze_panes=ws.cell(start_row+1,1)
        ws.auto_filter.ref=f"A{start_row}:{get_column_letter(len(headers))}{start_row+max(1,len(rows))}"
        for i,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width=w

    def _cover(self, ws, title: str, source: Path, subtitle: str):
        ws.sheet_view.showGridLines=False
        ws.merge_cells("B3:G4"); ws["B3"]=title
        ws["B3"].font=Font(name="Malgun Gothic", size=22, bold=True, color=self.DARK); ws["B3"].alignment=Alignment(horizontal="center",vertical="center")
        ws.merge_cells("B6:G6"); ws["B6"]=subtitle; ws["B6"].font=Font(name="Malgun Gothic", size=11, bold=True, color="595959"); ws["B6"].alignment=Alignment(horizontal="center")
        rows=[("원본 문서",source.name),("생성 일시",datetime.now().strftime("%Y-%m-%d %H:%M:%S")),("Requirement Studio","V0.88"),("Source Fidelity","Source에 없는 값/신호/Timing/Threshold/Test Result를 생성하지 않음")]
        for r,(k,v) in enumerate(rows,9):
            ws.cell(r,2,k); ws.cell(r,3,v); self._style_cell(ws.cell(r,2),center=True,fill=self.SUB); self._style_cell(ws.cell(r,3))
        ws.column_dimensions["B"].width=24; ws.column_dimensions["C"].width=78

    # ------------------------------------------------------------------
    # E2E Requirements
    # ------------------------------------------------------------------
    def export_requirements_word(self, source_document: Path, data: dict[str, Any]) -> Path:
        records=build_engineering_records(data)
        if not records: raise ValueError("E2E 요구사항 명세서로 내보낼 Canonical Engineering Requirement가 없습니다.")
        req_lookup={str(r.get("srs_id") or ""):r for r in _list(data.get("requirements")) if isinstance(r,dict)}
        out=self._filename("E2E_요구사항 명세서", source_document, "docx")
        doc=Document(); doc.styles["Normal"].font.name="Malgun Gothic"; doc.styles["Normal"].font.size=Pt(9)
        p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; run=p.add_run("E2E Engineering Requirements Specification — SYS.1 + SWE.1"); run.bold=True; run.font.size=Pt(18)
        meta=doc.add_table(rows=5,cols=2); meta.style="Table Grid"
        meta_rows=[("원본 문서",source_document.name),("생성 일시",datetime.now().strftime("%Y-%m-%d %H:%M:%S")),("Canonical Requirement",str(len(records))),("작성 기준","Requirement Studio V0.88 — E2E Requirements"),("범위","SYS.1 + SWE.1을 기본으로 하되 Source-backed native/review requirement도 누락 방지를 위해 보존")]
        for i,(k,v) in enumerate(meta_rows): meta.cell(i,0).text=k; meta.cell(i,1).text=v
        note=doc.add_paragraph(); rr=note.add_run("※ 본 문서는 입력 Source에 가장 가까운 E2E 요구사항 View입니다. Not SWE.1 != Not Requirement 원칙에 따라 SYS.1/SWE.1/native/review 범위를 삭제하지 않으며, 사람이 판단해야 하는 항목은 Review Required로 보존합니다."); rr.font.size=Pt(8.5); rr.font.color.rgb=RGBColor(90,98,110)
        section=doc.sections[0]; usable=int(section.page_width-section.left_margin-section.right_margin); left=int(usable*0.24); right=usable-left
        fields=["SRS ID","상위 기능","E2E Requirement Scope","Engineering Domain","Requirement Level","Allocation Status","SYS.1 Eligibility","SWE.1 Eligibility","Review Status","Canonical State","요구사항 내역","동작 조건 / Trigger","작동 명세 정의","사전 조건","예상 결과","검증 기준","출처 / Traceability","Human Review / Hold Reason","기타"]
        current=None
        for rec in records:
            if rec["상위 기능"]!=current: doc.add_heading(rec["상위 기능"],level=1); current=rec["상위 기능"]
            doc.add_heading(rec["SRS ID"],level=2)
            req=req_lookup.get(rec["SRS ID"],{})
            vals={**rec,"E2E Requirement Scope":_scope_label(req),"Human Review / Hold Reason":" | ".join(x for x in (rec.get("Review Status",""),rec.get("Hold Reason",""),_text(req.get("clarification_reference") or req.get("clarification_needed"))) if x)}
            table=doc.add_table(rows=len(fields),cols=2); table.style="Table Grid"; table.autofit=False; table.columns[0].width=left; table.columns[1].width=right
            for i,f in enumerate(fields):
                a,b=table.cell(i,0),table.cell(i,1); a.width=left; b.width=right; a.text=f; b.text=_text(vals.get(f,"")); a.vertical_alignment=b.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for r in a.paragraphs[0].runs: r.bold=True; r.font.size=Pt(8.5)
                for pp in b.paragraphs:
                    for r in pp.runs: r.font.size=Pt(8.5)
            doc.add_paragraph()
        doc.save(out); return out

    def export_requirements_excel(self, source_document: Path, data: dict[str, Any]) -> Path:
        records=build_engineering_records(data); reqs=[x for x in _list(data.get("requirements")) if isinstance(x,dict)]
        if not records: raise ValueError("E2E 요구사항 명세서로 내보낼 Requirement가 없습니다.")
        req_lookup={str(r.get("srs_id") or ""):r for r in reqs}; out=self._filename("E2E_요구사항 명세서",source_document,"xlsx")
        wb=Workbook(); cover=wb.active; cover.title="표지"; hist=wb.create_sheet("0_변경이력"); summary=wb.create_sheet("1_요구사항요약"); main=wb.create_sheet("2_E2E_요구사항"); review=wb.create_sheet("3_검토필요"); trace=wb.create_sheet("4_Source_추적성")
        self._cover(cover,"E2E 요구사항 명세서",source_document,"SYS.1 + SWE.1 + Source-backed Engineering Requirement")
        self._write_table(hist,["No","Version","Date","Change"],[[1,"V0.88",datetime.now().strftime("%Y-%m-%d"),"E2E 요구사항 Word/Excel 신규: SYS.1+SWE.1+native/review 보존"]],[8,14,14,90])
        counts=Counter(_scope_label(r) for r in reqs)
        sum_rows=[["Canonical Requirement Count",len(records)],["SYS.1 + SWE.1",counts.get("SYS.1 + SWE.1",0)],["SYS.1",counts.get("SYS.1",0)],["SWE.1",counts.get("SWE.1",0)],["Engineering Native / Review",counts.get("Engineering Native / Review",0)],["Human Decision Required",sum(1 for r in reqs if r.get("human_decision_required"))]]
        self._write_table(summary,["항목","Count / Status"],sum_rows,[38,24])
        headers=["SRS ID","Candidate ID","상위 기능","E2E Requirement Scope","Engineering Domain","Requirement Level","Allocation Status","SYS.1 Eligibility","SWE.1 Eligibility","Review Status","Canonical State","요구사항 내역","동작 조건 / Trigger","작동 명세 정의","사전 조건","예상 결과","검증 기준","Source Semantic Unit ID","Source Fact Fragment ID(s)","출처 / Traceability","Human Review Required","Hold / Clarification","기타"]
        rows=[]; review_rows=[]; trace_rows=[]
        for rec in records:
            req=req_lookup.get(rec["SRS ID"],{}); sem,fids,loc=_source_trace(req)
            rows.append([rec["SRS ID"],req.get("candidate_id",""),rec["상위 기능"],_scope_label(req),rec["Engineering Domain"],rec["Requirement Level"],rec["Allocation Status"],rec["SYS.1 Eligibility"],rec["SWE.1 Eligibility"],rec["Review Status"],rec["Canonical State"],rec["요구사항 내역"],rec["동작 조건 / Trigger"],rec["작동 명세 정의"],rec["사전 조건"],rec["예상 결과"],rec["검증 기준"],sem,fids,rec["출처 / Traceability"],"Y" if req.get("human_decision_required") else "N"," | ".join(x for x in (rec.get("Hold Reason",""),_text(req.get("clarification_reference") or req.get("clarification_needed"))) if x),rec["기타"]])
            if req.get("human_decision_required") or rec["Review Status"] not in {"SOURCE_BACKED","APPROVED"}:
                review_rows.append([rec["SRS ID"],_scope_label(req),rec["Allocation Status"],rec["SYS.1 Eligibility"],rec["SWE.1 Eligibility"],rec["Review Status"],rec.get("Hold Reason",""),_text(req.get("clarification_reference") or req.get("clarification_needed")),"PENDING",""])
            for atom in _list(req.get("source_backed_atomic_behaviors")):
                if isinstance(atom,dict): trace_rows.append([rec["SRS ID"],atom.get("source_semantic_unit_id"),atom.get("source_fact_fragment_id"),atom.get("source_location"),atom.get("behavior_text") or atom.get("source_fact"),_scope_label(req),rec["Allocation Status"]])
        self._write_table(main,headers,rows,[16,18,24,24,28,22,34,18,18,24,26,68,46,70,44,44,50,26,30,64,18,58,36])
        self._write_table(review,["SRS ID","E2E Scope","Allocation Status","SYS.1","SWE.1","Review Status","Hold Reason","Clarification / Required Resolution","Disposition","Rationale"],review_rows,[16,24,34,16,16,24,54,72,20,60])
        self._write_table(trace,["SRS ID","Semantic Unit ID","Fact Fragment ID","Source Location","Source-backed Fact / Behavior","E2E Scope","Allocation Status"],trace_rows,[16,28,30,26,78,24,38])
        wb.save(out); return out

    # ------------------------------------------------------------------
    # E2E Evaluation
    # ------------------------------------------------------------------
    def export_evaluation_excel(self, source_document: Path, data: dict[str, Any]) -> Path:
        bundle=finalize_verification_single_truth(data, force=True)
        cases=[x for x in _list(bundle.get("e2e_evaluation_cases")) if isinstance(x,dict)]
        objects=[x for x in _list(bundle.get("integrated_test_objects")) if isinstance(x,dict)]
        reqs={str(r.get("srs_id") or ""):r for r in _list(data.get("requirements")) if isinstance(r,dict)}
        if not cases: raise ValueError("E2E 평가 명세서로 내보낼 평가 Case가 없습니다.")
        out=self._filename("E2E_평가 명세서",source_document,"xlsx")
        wb=Workbook(); cover=wb.active; cover.title="표지"; hist=wb.create_sheet("0_변경이력"); summary=wb.create_sheet("1_평가요약"); main=wb.create_sheet("2_E2E_평가케이스"); review=wb.create_sheet("3_검토필요"); trace=wb.create_sheet("4_요구사항-평가추적성")
        self._cover(cover,"E2E 평가 명세서",source_document,"SYS.5 + SWE.6 + Deferred / External / Native Verification")
        self._write_table(hist,["No","Version","Date","Change"],[[1,"V0.88",datetime.now().strftime("%Y-%m-%d"),"기존 11시트 통합 테스트를 SWE.6 스타일 E2E 평가 Master + 요약/검토/추적성으로 단순화"]],[8,14,14,100])
        raw_counts=Counter(str(x.get("test_object_type") or "") for x in objects); case_counts=Counter(str(x.get("test_object_type") or "") for x in cases)
        structural=bundle.get("single_truth_audit") or {}; semantic=bundle.get("semantic_verification_quality_audit") or {}
        release="PASS" if structural.get("release_gate_status")=="PASS" and semantic.get("status")=="PASS" else "HOLD"
        sum_rows=[["Raw Integrated Test Objects",len(objects)],["E2E Evaluation Cases",len(cases)],["SYS5_CANDIDATE",case_counts.get("SYS5_CANDIDATE",0)],["SYS5_PARENT_SCOPE_REVIEW",case_counts.get("SYS5_PARENT_SCOPE_REVIEW",0)],["SWE6_TC",case_counts.get("SWE6_TC",0)],["Deferred/Allocation/External Review Cases",sum(case_counts.get(x,0) for x in ("DEFERRED_INTENT","ALLOCATION_PENDING_INTENT","EXTERNAL_DEPENDENCY_INTENT","REVIEW_REQUIRED"))],["Native Verification Cases",sum(case_counts.get(x,0) for x in ("ELECTRICAL_VERIFICATION_INTENT","ENVIRONMENTAL_VERIFICATION_INTENT","MECHANICAL_VERIFICATION_INTENT","MANUFACTURING_VERIFICATION_INTENT"))],["Structural Single Truth Gate",structural.get("release_gate_status")],["Semantic Verification Quality Gate",semantic.get("status")],["Official Release",release]]
        self._write_table(summary,["평가 항목","Count / Status"],sum_rows,[44,28])
        main_headers=["Parent Requirement ID","Test Object Type","Evaluation Object ID","Related Raw Object IDs","Verification Domain","Methods for Testing / Method for deriving","Evaluation name","Evaluation description","Source Location","Source Semantic Unit ID","Source Fact Fragment ID(s)","Transition Family","Condition Logic","Expected Target State","Test preparation Description","Variable","Compare / Value","Test execution Description","Variable","Compare","Value","Expected Result Description","Variable","Compare","Value","Allocation Status","SWE.6 Eligibility","Execution Readiness","Human Review Required","Deferred / Review Reason","Output Value","PASS / FAIL","Comment","Capture"]
        rows=[]; review_rows=[]
        raw_by_id={str(x.get("test_object_id") or ""):x for x in objects}
        for case in cases:
            sid=str(case.get("parent_srs_id") or ""); req=reqs.get(sid,{})
            typ=str(case.get("test_object_type") or "")
            method={"SWE6_TC":"Software Qualification / SWE.6","SYS5_CANDIDATE":"System Qualification / SYS.5","SYS5_PARENT_SCOPE_REVIEW":"Parent SYS.5 Scope Review","ALLOCATION_PENDING_INTENT":"Allocation Review","DEFERRED_INTENT":"Deferred Verification Intent","EXTERNAL_DEPENDENCY_INTENT":"External Dependency Review","ELECTRICAL_VERIFICATION_INTENT":"Electrical Verification","ENVIRONMENTAL_VERIFICATION_INTENT":"Environmental Verification","MECHANICAL_VERIFICATION_INTENT":"Mechanical/HW Verification","MANUFACTURING_VERIFICATION_INTENT":"Manufacturing Verification"}.get(typ,"Source-backed Verification Review")
            related=list(dict.fromkeys([str(case.get("test_object_id") or "")]+[str(x) for x in _list(case.get("related_test_object_ids")) if str(x)]))
            sem=", ".join(str(x) for x in _list(case.get("source_semantic_unit_ids")) if str(x)); fids=", ".join(str(x) for x in _list(case.get("source_fact_fragment_ids")) if str(x)); loc=" | ".join(str(x) for x in _list(case.get("source_locations")) if str(x))
            logic=str(case.get("condition_logic") or ""); target=str(case.get("expected_target_state") or "")
            prep=str(case.get("test_preparation") or ""); pvar=str(case.get("variable") or ""); pcomp=str(case.get("compare") or ""); pval=str(case.get("value") or "")
            exe=str(case.get("test_execution") or ""); evar=str(case.get("variable") or ""); ecomp=str(case.get("compare") or ""); eval_=str(case.get("value") or "")
            expected=str(case.get("expected_result") or ""); xvar=str(case.get("variable") or ""); xcomp=str(case.get("compare") or ""); xval=str(case.get("value") or "")
            if case.get("transition_family"):
                condition_text=" / ".join(str(x) for x in _list(case.get("condition_texts")) if str(x))
                prep=prep or (f"Source-backed transition 조건 구성: {condition_text}" if condition_text else "")
                pvar=pvar or "Condition Set"; pcomp=pcomp or logic; pval=pval or ("Source-backed conditions" if logic else "")
                exe=exe or "Source-backed 조건을 성립시키고 Target State 전환 여부를 관찰한다."
                evar=evar or "Target State"; ecomp=ecomp or "transition_to"; eval_=eval_ or target
                expected=expected or (f"Source에 정의된 {target} 전환이 성립해야 한다." if target else "")
                xvar=xvar or "Target State"; xcomp=xcomp or "="; xval=xval or target
            reason=" | ".join(x for x in (str(case.get("deferred_reason_code") or ""),str(case.get("deferred_reason_detail") or ""),str(case.get("required_resolution") or "")) if x)
            rows.append([sid,typ,case.get("evaluation_objective_id"),", ".join(related),case.get("verification_domain"),method,req.get("function_name") or case.get("test_intent") or case.get("verification_objective"),case.get("verification_objective") or case.get("test_intent"),loc,sem,fids,case.get("transition_family"),logic,target,prep,pvar,pcomp,exe,evar,ecomp,eval_,expected,xvar,xcomp,xval,case.get("allocation_status"),case.get("swe6_eligibility"),case.get("execution_readiness"),"Y" if case.get("human_review_required") else "N",reason,"","","",""])
            if case.get("human_review_required") or typ in {"SYS5_PARENT_SCOPE_REVIEW","DEFERRED_INTENT","ALLOCATION_PENDING_INTENT","EXTERNAL_DEPENDENCY_INTENT","REVIEW_REQUIRED"}:
                review_rows.append([case.get("evaluation_objective_id"),sid,typ,", ".join(related),reason,case.get("required_resolution") or "Human engineering review required.","","PENDING",""])
        widths=[16,30,22,34,30,34,28,56,26,26,32,24,16,26,52,20,22,52,20,16,24,52,20,16,24,36,24,24,18,64,18,14,36,24]
        self._write_table(main,main_headers,rows,widths)
        # Result fields are editable but must start physically blank.
        dv=DataValidation(type="list",formula1='"PASS,FAIL"',allow_blank=True); main.add_data_validation(dv); dv.add(f"AF4:AF{max(100,main.max_row+20)}")
        self._write_table(review,["Review ID / Evaluation ID","Parent SRS ID","Review Type","Related Raw Object IDs","Reason / Gap","Required Resolution","Owner","Disposition","Rationale"],review_rows,[24,16,32,40,72,72,18,22,60])
        dv2=DataValidation(type="list",formula1='"PENDING,APPROVED,REJECTED,NOT_APPLICABLE"',allow_blank=True); review.add_data_validation(dv2); dv2.add(f"H4:H{max(100,review.max_row+20)}")
        by_srs={}
        for case in cases: by_srs.setdefault(str(case.get("parent_srs_id") or ""),[]).append(case)
        trace_rows=[]
        for sid,req in reqs.items():
            sem,fids,loc=_source_trace(req); per=by_srs.get(sid,[])
            eval_ids=", ".join(str(x.get("evaluation_objective_id") or "") for x in per if x.get("evaluation_objective_id")); paths=", ".join(dict.fromkeys(str(x.get("test_object_type") or "") for x in per if x.get("test_object_type")))
            status="COVERED" if per else "REVIEW_REQUIRED"
            trace_rows.append([sid,req.get("requirement"),paths,eval_ids,sem,fids,loc,"Y" if req.get("human_decision_required") else "N",status])
        self._write_table(trace,["Requirement ID","Requirement Summary","Verification Path(s)","E2E Evaluation Object IDs","Source Semantic Unit","Fact Fragment Scope","Source Location","Human Review","Coverage Status"],trace_rows,[16,72,46,52,28,34,30,18,24])
        wb.save(out); self._verify_evaluation_artifact(out,cases); return out

    def _verify_evaluation_artifact(self, path: Path, cases: list[dict[str,Any]]):
        wb=load_workbook(path,read_only=False,data_only=False)
        try:
            ws=wb["2_E2E_평가케이스"]
            # Result columns AE? Main header is 34 columns; Output/Pass/Comment/Capture = 31~34.
            for r in range(4,ws.max_row+1):
                if not str(ws.cell(r,3).value or "").strip(): continue
                for c in (31,32,33,34):
                    if ws.cell(r,c).value not in (None,""):
                        ws.cell(r,c).value=None
            wb.save(path)
        finally: wb.close()
        check=load_workbook(path,read_only=True,data_only=False)
        try:
            ws=check["2_E2E_평가케이스"]; violations=[]
            ids=set()
            for r in range(4,ws.max_row+1):
                eid=str(ws.cell(r,3).value or "").strip()
                if not eid: continue
                ids.add(eid)
                for c in (31,32,33,34):
                    if ws.cell(r,c).value not in (None,""): violations.append(ws.cell(r,c).coordinate)
            expected={str(x.get("evaluation_objective_id") or "") for x in cases if x.get("evaluation_objective_id")}
            if ids!=expected: raise ValueError(f"E2E evaluation ID mismatch: missing={sorted(expected-ids)} extra={sorted(ids-expected)}")
            if violations: raise ValueError("E2E evaluation pre-execution result fields must be blank: "+", ".join(violations[:20]))
        finally: check.close()
