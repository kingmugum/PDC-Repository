from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

from core.output_naming import output_filename
from core.quality_audit import finalize_test_intent_coverage

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

NOT_CONFIRMED = "입력문서에서 확인되지 않음"
REVIEW_NEEDED = "검토 필요"


def _safe_stem(name: str) -> str:
    stem = Path(name).stem or "document"
    stem = re.sub(r'[<>:"/\\|?*]+', "_", stem)
    stem = re.sub(r"\s+", "_", stem).strip("._ ")
    return stem[:80] or "document"


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, list):
        return "\n".join(str(x).strip() for x in value if str(x).strip())
    if isinstance(value, dict):
        return "; ".join(f"{k}={v}" for k, v in value.items() if v not in (None, "", [], {}))
    return str(value).strip()


def _first_nonempty(*values: Any) -> str:
    for value in values:
        text = _text(value)
        if text and text not in {NOT_CONFIRMED, "-"}:
            return text
    return ""


def _short_name(text: str, max_len: int = 46) -> str:
    text = re.sub(r"\s+", " ", text or "").strip().rstrip(".")
    for prefix in ("소프트웨어는 ", "Software는 ", "SW는 "):
        if text.startswith(prefix):
            text = text[len(prefix):]
            break
    return text if len(text) <= max_len else text[: max_len - 1].rstrip() + "…"


def _metadata_value(source_text: str, labels: list[str], max_len: int = 80) -> str:
    if not source_text:
        return ""
    for label in labels:
        # Explicitly labelled metadata only. Do not infer from descriptive prose.
        pattern = rf"(?im)^\s*{re.escape(label)}\s*[:：\t|]\s*([^\n|]{{1,{max_len}}})\s*$"
        match = re.search(pattern, source_text)
        if match:
            value = re.sub(r"\s+", " ", match.group(1)).strip()
            if len(value) <= max_len and not re.search(r"[。!?]\s*$", value):
                return value
    return ""


def _to_swe6_artifact_id(value: str) -> str:
    value = (value or "").strip()
    if not value:
        return ""
    # Change only an explicitly embedded SWE.1 process marker. Other IDs are preserved.
    return re.sub(r"(?i)SWE[._ -]?1", "SWE.6", value)


def _section_excerpt(source_text: str, labels: list[str], max_chars: int = 420) -> str:
    lines = [re.sub(r"\s+", " ", x).strip() for x in (source_text or "").splitlines()]
    for idx, line in enumerate(lines):
        if any(label.lower() in line.lower() for label in labels):
            body = []
            for nxt in lines[idx + 1: idx + 8]:
                if not nxt:
                    continue
                if re.match(r"^\d+(?:\.\d+)*\.?\s+", nxt) and body:
                    break
                body.append(nxt)
                if sum(len(x) for x in body) >= max_chars:
                    break
            return " ".join(body)[:max_chars].strip()
    return ""


def extract_metadata(source_text: str) -> dict[str, str]:
    vehicle = _metadata_value(source_text, ["차종", "Vehicle", "Vehicle Model", "Model"], 40)
    oem = _metadata_value(source_text, ["OEM", "고객사", "Customer"], 40)
    project = _metadata_value(source_text, ["프로젝트명", "Project Name", "Project"], 80)
    artifact = _metadata_value(source_text, ["산출물 ID", "문서 ID", "Document ID", "Artifact ID", "Deliverable ID"], 60)
    return {
        "artifact_id": _to_swe6_artifact_id(artifact),
        "vehicle_oem": " / ".join(x for x in (vehicle, oem) if x),
        "project_name": project,
        "hardware": _metadata_value(source_text, ["Hardware", "하드웨어"], 80),
        "software": _metadata_value(source_text, ["Software", "소프트웨어"], 80),
        "mechanical": _metadata_value(source_text, ["Mechanical", "메카니컬", "기구"], 80),
        "environment": _section_excerpt(source_text, ["테스트 환경", "Test Environment"], 420),
    }


def _parse_compare_value(text: str) -> tuple[str, str, str]:
    """Return variable, compare, value only when directly recoverable from text."""
    raw = re.sub(r"\s+", " ", text or "").strip()
    if not raw:
        return "", "", ""
    # Signal-style equality: Foo = 0x1 / Foo(0x1)
    m = re.search(r"([A-Za-z_][A-Za-z0-9_./-]{1,50})\s*(?:=|==)\s*(0x[0-9A-Fa-f]+|[-+]?\d+(?:\.\d+)?(?:\s*[A-Za-z%℃°/]+)?)", raw)
    if m:
        return m.group(1), "=", m.group(2).strip()
    m = re.search(r"([A-Za-z_][A-Za-z0-9_./-]{1,50})\s*\(\s*(0x[0-9A-Fa-f]+)\s*\)", raw)
    if m:
        return m.group(1), "=", m.group(2)

    op_map = {"이상": ">=", "이하": "<=", "초과": ">", "미만": "<"}
    unit = r"(?:ms|msec|s|sec|초|V|mV|A|mA|%|℃|°C|도|Step|step|회|bps|kbps)?"
    m = re.search(rf"([-+]?\d+(?:\.\d+)?)\s*({unit})\s*(이상|이하|초과|미만)", raw)
    if m:
        val = m.group(1) + ((" " + m.group(2)) if m.group(2) else "")
        before = raw[:m.start()].strip(" ,.:;-()")
        variable = _short_name(before.split("에서")[-1].split("경우")[-1], 36)
        return variable, op_map[m.group(3)], val.strip()
    m = re.search(rf"(>=|<=|>|<)\s*([-+]?\d+(?:\.\d+)?)\s*({unit})", raw)
    if m:
        val = m.group(2) + ((" " + m.group(3)) if m.group(3) else "")
        before = raw[:m.start()].strip(" ,.:;-()")
        return _short_name(before, 36), m.group(1), val.strip()
    return "", "", ""


def _source_backed_atomic_texts(req: dict[str, Any]) -> list[str]:
    out: list[str] = []
    seen: set[str] = set()
    atoms = req.get("source_backed_atomic_behaviors") if isinstance(req.get("source_backed_atomic_behaviors"), list) else []
    for atom in atoms:
        if isinstance(atom, dict):
            text = _text(atom.get("behavior_text") or atom.get("source_fact"))
        else:
            text = _text(atom)
        text = re.sub(r"\s+", " ", text).strip()
        if text and text.lower() not in seen:
            seen.add(text.lower())
            out.append(text)
    return out


def _swe6_eligible_atomic_texts(req: dict[str, Any]) -> tuple[list[str], list[str], list[str]]:
    """Return (eligible, deferred_or_other, eligible_fragment_ids) for mixed-domain bundles.

    V0.66 narrows a generated SWE.6 TC to Source facts that are themselves SWE.6 Eligible.
    System/SYS.5 or external-dependency facts remain in the parent audit/deferred records but are
    not copied into the software qualification execution/expected scope.
    """
    alloc_by_fragment = {
        str(x.get("source_fact_fragment_id") or ""): x
        for x in (req.get("fact_level_allocations") or [])
        if isinstance(x, dict) and str(x.get("source_fact_fragment_id") or "")
    }
    eligible: list[str] = []
    other: list[str] = []
    eligible_ids: list[str] = []
    atoms = req.get("source_backed_atomic_behaviors") if isinstance(req.get("source_backed_atomic_behaviors"), list) else []
    for atom in atoms:
        if isinstance(atom, dict):
            text = _text(atom.get("behavior_text") or atom.get("source_fact"))
            fid = str(atom.get("source_fact_fragment_id") or "")
        else:
            text = _text(atom)
            fid = ""
        text = re.sub(r"\s+", " ", text).strip()
        if not text:
            continue
        alloc = alloc_by_fragment.get(fid) if fid else None
        if alloc is not None:
            if str(alloc.get("swe6_eligibility") or "") == "Eligible":
                if text not in eligible:
                    eligible.append(text)
                if fid and fid not in eligible_ids:
                    eligible_ids.append(fid)
            else:
                if text not in other:
                    other.append(text)
        else:
            # Explicit-ID/legacy atoms without fact-level allocation preserve prior behavior.
            if not alloc_by_fragment and text not in eligible:
                eligible.append(text)
            elif text not in other:
                other.append(text)
    # V0.80 strict scope reconciliation: a fact-level allocation can be deferred even when an
    # atom/fragment link is missing or contextualized. Preserve every non-SWE.6 Source fact in the
    # excluded set so parent prose can never leak it back into a Software Qualification TC.
    for alloc in (req.get("fact_level_allocations") or []):
        if not isinstance(alloc, dict):
            continue
        text = re.sub(r"\s+", " ", _text(alloc.get("source_fact"))).strip()
        if not text:
            continue
        if str(alloc.get("swe6_eligibility") or "") == "Eligible":
            fid = str(alloc.get("source_fact_fragment_id") or "")
            if text not in eligible:
                eligible.append(text)
            if fid and fid not in eligible_ids:
                eligible_ids.append(fid)
        elif text not in other:
            other.append(text)
    return eligible, other, eligible_ids


def _important_source_literals(req: dict[str, Any]) -> list[str]:
    """Return Source-backed literals worth preserving in SWE.6 prose.

    This is not a requirement generator.  It only exposes facts already present in source-backed
    atomic behaviors / value relation records so generic Expected Result wording does not erase
    the very value or prohibition being qualified.
    """
    texts = _source_backed_atomic_texts(req)
    for item in req.get("value_relation_status") or []:
        if isinstance(item, dict):
            literal = _text(item.get("source_value_text") or item.get("source_literal") or item.get("source_clause"))
            if literal and literal not in texts:
                texts.append(literal)
    return texts


def _compose_expected_description(req: dict[str, Any], base_expected: str) -> str:
    base = re.sub(r"\s+", " ", base_expected or "").strip()
    literals = _important_source_literals(req)
    if not literals:
        return base
    combined = " ".join([base] + literals).lower()
    # Keep source-backed numeric/interface/prohibition facts visible when the generated base is
    # generic.  The literal clause remains reviewable rather than being normalized into an
    # unsupported comparator/value.
    selected: list[str] = []
    for text in literals:
        tokens = [x.lower() for x in re.findall(r"[A-Za-z0-9가-힣_.+-]+", text) if len(x) > 1]
        if not base or any(re.search(r"\d|금지|않|off|low power|watchdog|checksum|kbit/s|kbps|nm\b|isr\b", tok, re.I) for tok in tokens):
            selected.append(text)
        else:
            overlap = sum(1 for tok in set(tokens) if tok in base.lower())
            if tokens and overlap / max(1, len(set(tokens))) < 0.45:
                selected.append(text)
    if not selected:
        return base
    source_line = " / ".join(selected[:6])
    if base and source_line.lower() in base.lower():
        return base
    return (base + ("\n" if base else "") + "Source-backed 확인 항목: " + source_line).strip()


def _compose_execution_description(req: dict[str, Any], trigger: str, behavior: str) -> str:
    base = trigger or behavior
    if base:
        return base
    # A source-backed requirement clause is safer than inventing a signal/value/action.  It is
    # explicitly labelled as a verification intent so the user knows further execution design may
    # still be required.
    atoms = _source_backed_atomic_texts(req)
    if atoms:
        return "Source-backed 검증 intent: " + atoms[0]
    return ""



def _material_overlap(a: str, b: str) -> bool:
    ta = {x.lower() for x in re.findall(r"[A-Za-z0-9가-힣_.+-]+", a or "") if len(x) > 1}
    tb = {x.lower() for x in re.findall(r"[A-Za-z0-9가-힣_.+-]+", b or "") if len(x) > 1}
    if not ta or not tb:
        return False
    return len(ta & tb) / max(1, min(len(ta), len(tb), 10)) >= 0.35


def _swe6_eligible_semantic_units(req: dict[str, Any]) -> set[str]:
    return {
        str(x.get("source_semantic_unit_id") or "")
        for x in (req.get("fact_level_allocations") or [])
        if isinstance(x, dict)
        and str(x.get("swe6_eligibility") or "") == "Eligible"
        and str(x.get("source_semantic_unit_id") or "")
    }


def _source_table_fact_assignment(req: dict[str, Any], phase: str, context_text: str) -> tuple[str, str, str, list[dict[str, Any]]]:
    """Return a conservative Variable/Compare/Value candidate from HIGH-confidence Source table facts.

    Exact identifier / linked-unit evidence is required upstream.  Ambiguous multi-value rows are
    kept as audit evidence instead of being forced into a concrete TC field.
    """
    matches = [x for x in (req.get("source_table_fact_matches") or []) if isinstance(x, dict) and str(x.get("join_confidence") or "") == "HIGH"]
    fact_allocs = [x for x in (req.get("fact_level_allocations") or []) if isinstance(x, dict)]
    if fact_allocs and str(req.get("allocation_status") or "") != "SW_IMPLEMENTATION_REQUIREMENT":
        eligible_units = _swe6_eligible_semantic_units(req)
        matches = [x for x in matches if str(x.get("source_semantic_unit_id") or "") in eligible_units]
    if not matches:
        return "", "", "", []
    ctx = (context_text or "").lower()
    candidates = []
    for fact in matches:
        ids = [str(x) for x in fact.get("identifiers") or [] if str(x)]
        nums = [str(x).strip() for x in fact.get("numeric_values") or [] if str(x).strip()]
        enums = [x for x in fact.get("enum_mappings") or [] if isinstance(x, dict)]
        ranges = [str(x).strip() for x in fact.get("range_values") or [] if str(x).strip()]
        matched_ids = [i for i in ids if i.lower() in ctx]
        if phase == "expected":
            preferred = [i for i in ids if i.lower().startswith(("output_", "o_"))]
        else:
            preferred = [i for i in ids if i.lower().startswith(("input_", "par_", "param_", "parameter_", "i_", "p_"))]
        var = (matched_ids or preferred or ids[:1])
        if not var:
            continue
        variable = var[0]
        compare = ""
        value = ""
        literal = str(fact.get("source_literal") or "")
        value_role = str(fact.get("value_role") or "SOURCE_LITERAL_ONLY")
        # V0.65: enum/range/encoding/invalid-marker rows are evidence, not a single test value.
        # Only a source-defined single required value or an explicit timing criterion may be
        # conservatively auto-injected into Variable/Compare/Value.
        if value_role not in {"SINGLE_REQUIRED_VALUE", "TIMING_CRITERION"}:
            candidates.append(("", "", "", fact))
            continue
        if len(nums) == 1:
            compare = "=" if re.search(rf"{re.escape(variable)}\s*(?:=|:|：)", literal, re.I) else "Source literal"
            value = nums[0]
        else:
            continue
        candidates.append((variable, compare, value, fact))
    # Require a unique structured candidate. Multiple candidates remain visible in the audit sheet
    # but are not collapsed into an arbitrary test value.
    usable = [x for x in candidates if x[0] and x[2]]
    if len(usable) == 1:
        v, c, val, fact = usable[0]
        # When the requirement has multiple joined facts, never auto-use a value from an
        # unrelated identifier merely because it is the only auto-usable role.
        if len(candidates) == 1 or v.lower() in ctx:
            return v, c, val, [fact]
    if usable:
        context_hits = [x for x in usable if x[0].lower() in ctx]
        if len(context_hits) == 1:
            v, c, val, fact = context_hits[0]
            return v, c, val, [fact]
    return "", "", "", [x[3] for x in candidates]


def _technique(req: dict[str, Any]) -> str:
    text = " ".join(_text(req.get(k)) for k in (
        "requirement", "activation_trigger", "preconditions", "processing_action", "output", "acceptance_criteria"
    )).lower()
    methods = ["요구사항 기반 테스트"]
    structural_constraint = bool(re.search(r"watchdog|\btask\b|\bisr\b|interrupt|배치|위치|존재할\s*수\s*없", text, re.I))
    if structural_constraint:
        methods.append("정적 분석 / 코드·설정 검토 (DERIVED 후보)")
    numeric_boundary = bool(re.search(r"(?:>=|<=|>|<|이상|이하|초과|미만|최대|최소)\s*[-+]?\d|[-+]?\d+(?:\.\d+)?\s*(?:ms|s|초|v|mv|a|ma|%|℃|도)\s*(?:이상|이하|초과|미만)", text))
    if numeric_boundary:
        methods.append("경계값 분석")
    elif any(x in text for x in ("invalid", "reserved", "유효", "무효", "범위", "0x0 ~", "0x0~")):
        methods.append("동등분할")
    elif any(x in text for x in ("state", "mode", "상태전이", "상태 전이", "진입", "복귀", "wake", "sleep")):
        methods.append("상태전이 테스트")
    return " / ".join(methods)


def _testability(req: dict[str, Any]) -> str:
    if str(req.get("swe6_eligibility") or "Eligible") != "Eligible":
        return REVIEW_NEEDED
    fact_allocs = [x for x in (req.get("fact_level_allocations") or []) if isinstance(x, dict)]
    if (
        fact_allocs
        and not any(str(x.get("swe6_eligibility") or "") == "Eligible" for x in fact_allocs)
        and str(req.get("allocation_status") or "") != "SW_IMPLEMENTATION_REQUIREMENT"
    ):
        return REVIEW_NEEDED
    # V0.65 retains the V0.59 rule: for no-ID sources, do not generate a concrete SWE.6 TC until the
    # Source Semantic Unit -> Atomic Behavior -> SWE.1 provenance chain is complete.
    # Explicit-ID Gold flows retain their existing behavior.
    source_ids = req.get("source_requirement_ids") if isinstance(req.get("source_requirement_ids"), list) else []
    if not source_ids:
        sem_ids = req.get("source_semantic_unit_ids") if isinstance(req.get("source_semantic_unit_ids"), list) else []
        atoms = req.get("source_backed_atomic_behaviors") if isinstance(req.get("source_backed_atomic_behaviors"), list) else []
        if str(req.get("semantic_provenance_status") or "") != "COMPLETE" or not sem_ids or not atoms:
            return REVIEW_NEEDED
    category = _text(req.get("category"))
    requirement = _text(req.get("requirement"))
    if not requirement or category == REVIEW_NEEDED:
        return REVIEW_NEEDED
    gap = _text(req.get("clarification_needed"))
    meaningful = any(_first_nonempty(req.get(k)) for k in (
        "activation_trigger", "preconditions", "processing_action", "output", "acceptance_criteria"
    ))
    if not meaningful and gap:
        return REVIEW_NEEDED
    return "대상"


def build_swe6_cases(requirement_data: dict[str, Any]) -> list[dict[str, str]]:
    cases: list[dict[str, str]] = []
    tc_no = 1
    for idx, req in enumerate(requirement_data.get("requirements") or [], start=1):
        if not isinstance(req, dict):
            continue
        srs_id = _text(req.get("srs_id")) or f"SRS_{idx:03d}"
        status = _testability(req)
        if status != "대상":
            continue
        requirement = _text(req.get("requirement"))
        pre = _first_nonempty(req.get("preconditions"), req.get("system_input_preconditions"))
        trigger = _text(req.get("activation_trigger"))
        behavior = _first_nonempty(req.get("behavior_flows"), req.get("processing_action"))
        expected_base = _first_nonempty(req.get("output"), req.get("acceptance_criteria"))
        eligible_atoms, excluded_atoms, eligible_fragment_ids = _swe6_eligible_atomic_texts(req)
        mixed_fact_scope = bool(excluded_atoms)
        # V0.66: when a parent SRS bundles SWE.6 and non-SWE.6 Source facts, the TC is scoped only
        # to the eligible Source facts. The parent requirement remains traceable, but a SYS.5 or
        # external-dependency criterion is never silently asserted as a SWE.6 expected result.
        if mixed_fact_scope:
            expected = ("Source-backed SWE.6 scope: " + " | ".join(eligible_atoms[:6])) if eligible_atoms else ""
            # Parent-level precondition/output text may contain a fact explicitly deferred to SYS.5.
            # Keep only a source-backed eligible scope in the executable fields.
            if any(_material_overlap(pre, x) for x in excluded_atoms):
                pre = ""
        else:
            expected = _compose_expected_description(req, expected_base)
        prep_var, prep_cmp, prep_val = _parse_compare_value(pre)
        if mixed_fact_scope:
            exec_source = ("Source-backed SWE.6 검증 intent: " + eligible_atoms[0]) if eligible_atoms else ""
        else:
            exec_source = _compose_execution_description(req, trigger, behavior)
        exec_var, exec_cmp, exec_val = _parse_compare_value(exec_source)
        exp_var, exp_cmp, exp_val = _parse_compare_value(expected if mixed_fact_scope else (expected_base or _text(req.get("acceptance_criteria"))))
        table_audit = []
        if not prep_var or not prep_val:
            tv, tc, tval, ta = _source_table_fact_assignment(req, "preparation", " ".join([pre, trigger]))
            table_audit.extend(ta)
            if tv and tval:
                prep_var, prep_cmp, prep_val = prep_var or tv, prep_cmp or tc, prep_val or tval
        if not exec_var or not exec_val:
            tv, tc, tval, ta = _source_table_fact_assignment(req, "execution", " ".join([trigger, behavior]))
            table_audit.extend(ta)
            if tv and tval:
                exec_var, exec_cmp, exec_val = exec_var or tv, exec_cmp or tc, exec_val or tval
        if not exp_var or not exp_val:
            tv, tc, tval, ta = _source_table_fact_assignment(req, "expected", " ".join([expected_base, _text(req.get("acceptance_criteria"))]))
            table_audit.extend(ta)
            if tv and tval:
                exp_var, exp_cmp, exp_val = exp_var or tv, exp_cmp or tc, exp_val or tval
        case_name = _short_name(_first_nonempty(req.get("function_name"), requirement), 44) or f"{srs_id} 검증"
        desc_parts = [f"{srs_id} 요구사항의 Source-backed Normal / Positive 검증 초안이다."]
        if trigger:
            desc_parts.append(f"동작 조건은 '{_short_name(trigger, 80)}'이다.")
        atoms = eligible_atoms if mixed_fact_scope else _source_backed_atomic_texts(req)
        if atoms:
            label = "SWE.6-eligible Source-backed atomic intent" if mixed_fact_scope else "Source-backed atomic intent"
            desc_parts.append(label + ": " + " | ".join(f"{i+1}) {_short_name(t, 110)}" for i, t in enumerate(atoms[:6])))
        if excluded_atoms:
            desc_parts.append(f"Excluded from SWE.6 TC scope: {len(excluded_atoms)} non-SWE.6 Source fact(s); see Fact Allocation/Deferred audit.")
        # V0.69: when the parent SRS contains mixed verification domains, do not echo the
        # parent-level expected/result prose into the SWE.6 TC description. That text can include
        # a SYS.5/external-dependency acceptance criterion even though the executable SWE.6 scope
        # is correctly limited above. Keep cross-domain context only in the explicit excluded audit.
        if expected_base and not mixed_fact_scope:
            desc_parts.append(f"기대 결과 기본 표현은 '{_short_name(expected_base, 100)}'이다.")
        cases.append({
            "srs_id": srs_id,
            "tc_id": f"TC_{tc_no:03d}",
            "method": _technique(req),
            "name": case_name,
            "description": " ".join(desc_parts),
            "prep_desc": pre,
            "prep_var": prep_var,
            "prep_compare": prep_cmp,
            "prep_value": prep_val,
            "exec_desc": exec_source,
            "exec_var": exec_var,
            "exec_compare": exec_cmp,
            "exec_value": exec_val,
            "expected_desc": expected,
            "expected_var": exp_var,
            "expected_compare": exp_cmp,
            "expected_value": exp_val,
            "verification_domain": _text(req.get("verification_domain")),
            "verification_method_candidates": ([{"method": "정적 분석 / 코드·설정 검토", "knowledge_state": "DERIVED", "basis": "Requirement characteristic indicates a structural/placement constraint."}] if "DERIVED 후보" in _technique(req) else []),
            "source_semantic_unit_ids": list(req.get("source_semantic_unit_ids") or []),
            "source_backed_atomic_behaviors": atoms,
            "swe6_scope_fragment_ids": eligible_fragment_ids,
            "excluded_non_swe6_source_facts": excluded_atoms,
            "mixed_fact_scope_limited": mixed_fact_scope,
            "atomicity_review_needed": len(atoms) > 1,
            "source_table_fact_matches": list(req.get("source_table_fact_matches") or []),
            "source_table_fact_auto_used": list({str(x.get("fact_id") or ""): x for x in table_audit if isinstance(x, dict) and x.get("fact_id")}.values()),
        })
        tc_no += 1
    return cases



def build_sys5_candidates(requirement_data: dict[str, Any]) -> list[dict[str, Any]]:
    """Build Source-backed System Qualification candidate intents without inventing test data.

    V0.80 intentionally does not claim these are executable SYS.5 procedures.  They are the
    System-level counterpart to SWE.6 intent drafts and preserve facts that software-only export
    previously deferred or hid.
    """
    out: list[dict[str, Any]] = []
    no = 1
    for idx, req in enumerate(requirement_data.get("requirements") or [], start=1):
        if not isinstance(req, dict):
            continue
        sys5 = str(req.get("sys5_eligibility") or "")
        if sys5 not in {"Eligible", "Review Needed"}:
            continue
        sid = _text(req.get("srs_id")) or f"SRS_{idx:03d}"
        system_facts: list[str] = []
        source_ids: list[str] = []
        for alloc in (req.get("fact_level_allocations") or []):
            if not isinstance(alloc, dict):
                continue
            domain = str(alloc.get("verification_domain") or "").lower()
            is_system = "sys.5" in domain or "system integration" in domain or "interface allocation" in domain
            if not is_system and str(alloc.get("swe6_eligibility") or "") == "Eligible":
                continue
            text = re.sub(r"\s+", " ", _text(alloc.get("source_fact"))).strip()
            if text and text not in system_facts:
                system_facts.append(text)
            fid = str(alloc.get("source_fact_fragment_id") or "")
            if fid and fid not in source_ids:
                source_ids.append(fid)
        parent_scope_review = not bool(system_facts)
        if not system_facts:
            requirement = re.sub(r"\s+", " ", _text(req.get("requirement"))).strip()
            if requirement:
                system_facts.append(requirement)
        if not system_facts:
            continue
        out.append({
            "srs_id": sid,
            "sys5_id": f"SYS5_TC_{no:03d}",
            "verification_domain": "SYS.5 System Qualification Candidate",
            "objective": " | ".join(system_facts[:6]),
            "source_fact_fragment_ids": source_ids,
            "source_scope_basis": "PARENT_FALLBACK" if parent_scope_review else "FACT_LEVEL_SYS5",
            "status": "CANDIDATE_REVIEW_REQUIRED",
            "execution_readiness": "REVIEW_REQUIRED",
            "human_decision_required": bool(req.get("human_decision_required")) or sys5 == "Review Needed",
            "missing_data_or_dependency": _first_nonempty(req.get("clarification_needed"), req.get("external_dependencies"), req.get("tbd_items")),
            "execution_result": "",
            "pass_fail": "",
        })
        no += 1
    return out


@dataclass
class _Style:
    dark: str = "D9E1F2"
    header: str = "D9D9D9"
    sub: str = "E2F0D9"
    accent: str = "70AD47"
    thin_color: str = "808080"


class SWE6Exporter:
    def __init__(self, output_dir: Path):
        self.output_dir = Path(output_dir).resolve()
        self.output_dir.mkdir(parents=True, exist_ok=True)
        self.style = _Style()

    def _filename(self, source_document: Path) -> Path:
        return self.output_dir / output_filename(
            "SWE.6 적격성 평가", _safe_stem(Path(source_document).name), "xlsx"
        )

    def _border(self):
        side = Side(style="thin", color=self.style.thin_color)
        return Border(left=side, right=side, top=side, bottom=side)

    def _cell_style(self, cell, *, fill=None, bold=False, center=False, font_size=10):
        cell.font = Font(name="Malgun Gothic", size=font_size, bold=bold)
        cell.alignment = Alignment(horizontal="center" if center else "left", vertical="center", wrap_text=True)
        cell.border = self._border()
        if fill:
            cell.fill = PatternFill("solid", fgColor=fill)

    @staticmethod
    def _enforce_blank_execution_results(path: Path) -> None:
        """Enforce the pre-execution result invariant on the *saved* workbook artifact.

        V0.65 verifies the file after serialization, not only the in-memory row values. If any
        Output Value / PASS-FAIL / Comment / Capture cell was populated by an upstream/default
        path, it is cleared and the workbook is saved once more. A second read must prove all
        result cells are physically blank; otherwise export fails instead of producing a false
        release PASS.
        """
        path = Path(path)
        wb = load_workbook(path, data_only=False)
        changed = False
        try:
            if "2_테스트 케이스" not in wb.sheetnames:
                raise ValueError("SWE.6 export invariant failure: 2_테스트 케이스 sheet missing")
            ws = wb["2_테스트 케이스"]
            for row in range(4, ws.max_row + 1):
                tc_id = str(ws.cell(row, 2).value or "").strip()
                if not tc_id:
                    continue
                for col in (18, 19, 20, 21):
                    value = ws.cell(row, col).value
                    if value not in (None, ""):
                        ws.cell(row, col).value = None
                        changed = True
            if changed:
                wb.save(path)
        finally:
            wb.close()

        verify = load_workbook(path, data_only=False, read_only=True)
        try:
            ws = verify["2_테스트 케이스"]
            violations = []
            for row in range(4, ws.max_row + 1):
                tc_id = str(ws.cell(row, 2).value or "").strip()
                if not tc_id:
                    continue
                for col in (18, 19, 20, 21):
                    value = ws.cell(row, col).value
                    if value not in (None, ""):
                        violations.append(f"{ws.cell(row, col).coordinate}={value!r}")
            if violations:
                raise ValueError("SWE.6 export invariant failure: prepopulated execution result cells: " + ", ".join(violations[:20]))
        finally:
            verify.close()

    def _merge_label(self, ws, rng: str, text: str, *, fill=None, bold=False, center=False, font_size=10):
        ws.merge_cells(rng)
        c = ws[rng.split(":")[0]]
        c.value = text
        self._cell_style(c, fill=fill, bold=bold, center=center, font_size=font_size)
        # Apply borders to whole merged range.
        for row in ws[rng]:
            for cell in row:
                cell.border = self._border()
                if fill:
                    cell.fill = PatternFill("solid", fgColor=fill)
        return c

    def export_excel(self, source_document: Path, requirement_data: dict[str, Any], *, source_text: str = "") -> Path:
        out = self._filename(source_document)
        metadata = extract_metadata(source_text)
        requirements = [x for x in (requirement_data.get("requirements") or []) if isinstance(x, dict)]
        bundle = requirement_data.get("finalized_verification_bundle") if isinstance(requirement_data.get("finalized_verification_bundle"), dict) else {}
        if str(bundle.get("schema_version") or "").startswith("REQ-STUDIO-VERIFICATION-BUNDLE-"):
            cases = [x for x in (bundle.get("swe6_cases") or []) if isinstance(x, dict)]
            sys5_cases = [x for x in (bundle.get("sys5_candidates") or []) if isinstance(x, dict)]
        else:
            cases = build_swe6_cases(requirement_data)
            sys5_cases = build_sys5_candidates(requirement_data)
            finalize_test_intent_coverage(requirement_data, cases)
        today = datetime.now().strftime("%Y.%m.%d")

        wb = Workbook()
        cover = wb.active
        cover.title = "표지"
        history = wb.create_sheet("0_변경이력")
        summary = wb.create_sheet("1_테스트요약")
        tc = wb.create_sheet("2_테스트 케이스")
        deferred_ws = wb.create_sheet("3_Deferred Intent")
        intent_audit_ws = wb.create_sheet("4_Source_Intent_Audit")
        fact_audit_ws = wb.create_sheet("5_Source_Fact_Audit")
        sys5_ws = wb.create_sheet("6_SYS5_Candidates")
        eng_trace_ws = wb.create_sheet("7_Engineering_Verification")

        # --- Cover ------------------------------------------------------
        cover.sheet_view.showGridLines = False
        for col, width in {"A":4,"B":12,"C":13,"D":13,"E":13,"F":13,"G":13,"H":16,"I":4}.items():
            cover.column_dimensions[col].width = width
        self._merge_label(cover,"B2:F4","Engineering Verification Specification — SYS.5 / SWE.6 통합 검토본",bold=True,center=True,font_size=13)
        for row, (label, value) in enumerate([
            ("산출물 ID", metadata["artifact_id"]),
            ("개정번호", "V.0.0"),
            ("개정일자", today),
        ], start=2):
            cover[f"G{row}"].value=label; self._cell_style(cover[f"G{row}"],fill=self.style.header,bold=True,center=True)
            cover[f"H{row}"].value=value; self._cell_style(cover[f"H{row}"],center=True)
        self._merge_label(cover,"B12:H14","시스템 / 소프트웨어 적격성 검증",bold=True,center=True,font_size=20)
        self._merge_label(cover,"B17:H19","SYS.5 / SWE.6 Integrated Verification",bold=False,center=True,font_size=18)
        cover["C23"]="차종 / OEM"; self._cell_style(cover["C23"],fill=self.style.header,bold=True,center=True)
        self._merge_label(cover,"D23:F23",metadata["vehicle_oem"],center=False)
        cover["C25"]="프로젝트명"; self._cell_style(cover["C25"],fill=self.style.header,bold=True,center=True)
        self._merge_label(cover,"D25:F25",metadata["project_name"],center=False)
        self._merge_label(cover,"C27:C31","결재 정보",fill=self.style.header,bold=True,center=True)
        for col,label in zip(("D","E","F"),("작성","검토","승인")):
            cover[f"{col}27"]=label; self._cell_style(cover[f"{col}27"],fill=self.style.header,bold=True,center=True)
            self._merge_label(cover,f"{col}28:{col}31","",center=True)
        cover["C34"]="문서 상태:"; cover["C36"]="배포 날짜:"
        cover["D34"]="Draft"; cover["D36"]=""
        status_dv = DataValidation(type="list", formula1='"Draft,Released,Restricted,Expired"', allow_blank=False)
        cover.add_data_validation(status_dv)
        status_dv.add(cover["D34"])
        for c in (cover["C34"],cover["C36"]): c.font=Font(name="Malgun Gothic",size=11,bold=True)
        for c in (cover["D34"],cover["D36"]): self._cell_style(c,center=True)

        # --- History ----------------------------------------------------
        history.sheet_view.showGridLines=False
        history["A1"]="0. 문서 제/개정 이력"; history["A1"].font=Font(name="Malgun Gothic",size=14,bold=True)
        headers=["번호","개정 일자","개정 버전","개정 내용","개정자"]
        history.append([]); history.append(headers)
        for cell in history[3]: self._cell_style(cell,fill=self.style.header,bold=True,center=True)
        for no in range(1,11):
            vals=[no, today if no==1 else "", "V.0.0" if no==1 else "", "초안 작성" if no==1 else "", ""]
            history.append(vals)
            for cell in history[3+no]: self._cell_style(cell,center=(cell.column != 4))
        for col,width in zip("ABCDE",[9,15,15,70,18]): history.column_dimensions[col].width=width
        for row in range(4,14): history.row_dimensions[row].height=26

        # --- Summary ----------------------------------------------------
        summary.sheet_view.showGridLines=False
        summary["A1"]="1. 테스트 요약"; summary["A1"].font=Font(name="Malgun Gothic",size=14,bold=True)
        summary["A3"]="1. 평가 대상"; summary["A3"].font=Font(name="Malgun Gothic",size=11,bold=True)
        summary["A5"]="Project Name"; self._cell_style(summary["A5"],fill=self.style.header,bold=True,center=True)
        self._merge_label(summary,"B5:F5",metadata["project_name"])
        summary["A6"]="Test Level"; self._cell_style(summary["A6"],fill=self.style.header,bold=True,center=True)
        self._merge_label(summary,"B6:F6","SYS.5 시스템 검증 Candidate + SWE.6 소프트웨어 적격성 테스트 Candidate")
        self._merge_label(summary,"A7:A9","Test Item",fill=self.style.header,bold=True,center=True)
        for rr,label,key in ((7,"Hardware","hardware"),(8,"Software","software"),(9,"Mechanical","mechanical")):
            summary[f"B{rr}"]=label; self._cell_style(summary[f"B{rr}"],fill=self.style.header,bold=True,center=True)
            self._merge_label(summary,f"C{rr}:F{rr}",metadata.get(key,""))
        summary["A10"]="입력물"; self._cell_style(summary["A10"],fill=self.style.header,bold=True,center=True)
        self._merge_label(summary,"B10:F10","Engineering Requirements Specification (SYS.1 / SWE.1 Integrated Review)")
        self._merge_label(
            summary,"A11:F12",
            "※ 본 문서는 V0.88 Engineering E2E 검증 초안입니다. SYS.5/SWE.6 Candidate를 함께 보존하며, REVIEW_REQUIRED/HOLD 항목은 삭제하지 않습니다. 시험 환경, 외부 DB, 파라미터, 임계값, Variant 및 실행 결과는 담당자 검토 후 확정되어야 하고 Source에 없는 시험값은 임의 생성하지 않습니다.",
            center=False,font_size=9
        )
        summary["A13"]="2. 테스트 환경"; summary["A13"].font=Font(name="Malgun Gothic",size=11,bold=True)
        environment_text = metadata.get("environment", "")
        if not environment_text:
            environment_text = "테스트 환경 구성도 또는 대표 사진이 필요합니다.\n입력문서에서 Test Environment 관련 시각자료를 확인할 수 없습니다."
        self._merge_label(summary,"A15:F22",environment_text,center=True,font_size=11)
        summary["A24"]="3. 기능 요구사항"; summary["A24"].font=Font(name="Malgun Gothic",size=11,bold=True)

        def add_req_table(start_row:int, wanted_category:str):
            hdr=["번호","분류 1","분류 2","분류 3","식별자","테스트"]
            for col,val in enumerate(hdr,1):
                c=summary.cell(start_row,col,val); self._cell_style(c,fill=self.style.header,bold=True,center=True)
            rr=start_row+1; no=1
            for req in requirements:
                cat=_text(req.get("category"))
                normalized="비기능" if "비기능" in cat else ("기능" if "기능" in cat else REVIEW_NEEDED)
                if normalized != wanted_category: continue
                req_text=_text(req.get("requirement"))
                vals=[no,_text(req.get("function_name")),"",_short_name(req_text,54),_text(req.get("srs_id")),_testability(req)]
                for col,val in enumerate(vals,1):
                    c=summary.cell(rr,col,val); self._cell_style(c,center=(col in (1,5,6)))
                rr+=1; no+=1
            if rr==start_row+1:
                vals=[1,"","","등록된 요구사항 없음","","검토 필요"]
                for col,val in enumerate(vals,1):
                    c=summary.cell(rr,col,val); self._cell_style(c,center=(col in (1,5,6)))
                rr+=1
            return rr

        next_row=add_req_table(26,"기능")
        nr=max(next_row+2,36)
        summary[f"A{nr}"]="4. 비기능 요구사항"; summary[f"A{nr}"].font=Font(name="Malgun Gothic",size=11,bold=True)
        add_req_table(nr+2,"비기능")
        for col,width in zip("ABCDEF",[9,24,24,46,15,14]): summary.column_dimensions[col].width=width
        summary.freeze_panes="A5"

        # --- Test Cases -------------------------------------------------
        tc.sheet_view.showGridLines=False
        tc["A1"]="2. 테스트 케이스"; tc["A1"].font=Font(name="Malgun Gothic",size=14,bold=True)
        # Row 2 groups the detailed headers below into practical test blocks.
        group_headers = [
            ("A2:E2", "Test ID"),
            ("F2:I2", "Test Input"),
            ("J2:M2", "Test Execution"),
            ("N2:Q2", "Test Expect Result"),
            ("R2:S2", "Test Result"),
            ("T2:U2", "Etc"),
            ("V2:X2", "TC Readiness"),
        ]
        for rng, label in group_headers:
            self._merge_label(tc, rng, label, fill=self.style.dark, bold=True, center=True, font_size=9)
        tc.row_dimensions[2].height = 23

        headers=[
            "SW 요구사항 ID","Test Case ID","Methods for Testing /\nMethod for deriving TC",
            "Test case name","Test case description",
            "Test preparation\nDescription","Variable","Compare","Value",
            "Test execution\nDescription","Variable","Compare","Value",
            "Expected Result\nDescription","Variable","Compare","Value",
            "Output Value","PASS / FAIL","Comment","Capture\nCANoe (Optional)",
            "TC Maturity","Independent Coverage","Execution Readiness"
        ]
        for col,val in enumerate(headers,1):
            c=tc.cell(3,col,val); self._cell_style(c,fill=self.style.header,bold=True,center=True,font_size=9)
        tc.row_dimensions[3].height = 38

        def na_triplet(description: str, variable: str, compare: str, value: str) -> tuple[str, str, str]:
            """Use N/A only for unavailable derived V/C/V values, never for intentionally blank result fields."""
            if not _text(description):
                return variable or "", compare or "", value or ""
            return variable or "N/A", compare or "N/A", value or "N/A"

        r=4
        if not cases:
            cases=[{
                "srs_id":"","tc_id":"","method":"","name":"검토 필요","description":"SWE.6 Test Case 자동 생성에 필요한 검증 가능 Requirement가 없습니다.",
                "prep_desc":"","prep_var":"","prep_compare":"","prep_value":"","exec_desc":"","exec_var":"","exec_compare":"","exec_value":"",
                "expected_desc":"","expected_var":"","expected_compare":"","expected_value":""
            }]
        for case in cases:
            prep_var, prep_cmp, prep_val = na_triplet(case["prep_desc"], case["prep_var"], case["prep_compare"], case["prep_value"])
            exec_var, exec_cmp, exec_val = na_triplet(case["exec_desc"], case["exec_var"], case["exec_compare"], case["exec_value"])
            exp_var, exp_cmp, exp_val = na_triplet(case["expected_desc"], case["expected_var"], case["expected_compare"], case["expected_value"])
            vals=[
                case["srs_id"],case["tc_id"],case["method"],case["name"],case["description"],
                case["prep_desc"],prep_var,prep_cmp,prep_val,
                case["exec_desc"],exec_var,exec_cmp,exec_val,
                case["expected_desc"],exp_var,exp_cmp,exp_val,
                "","","","",
                case.get("tc_maturity", "INTENT_DRAFT"),
                case.get("independent_coverage", "NOT_ESTABLISHED"),
                case.get("execution_readiness", "REVIEW_REQUIRED"),
            ]
            for col,val in enumerate(vals,1):
                c=tc.cell(r,col,val); self._cell_style(c,center=(col in (1,2,7,8,9,11,12,13,15,16,17,18,19,22,23,24)))
            r+=1

        widths=[16,14,28,28,46,42,20,11,16,42,20,11,16,42,20,11,16,18,14,30,22,20,22,20]
        for idx,width in enumerate(widths,1):
            tc.column_dimensions[chr(64+idx)].width=width

        # If a Variable / Compare / Value column contains only N/A (ignoring deliberate blanks),
        # keep it narrow so empty metadata does not dominate the worksheet.
        for idx in (7,8,9,11,12,13,15,16,17):
            observed=[_text(tc.cell(row_idx, idx).value) for row_idx in range(4, r)]
            nonblank=[value for value in observed if value]
            if nonblank and all(value == "N/A" for value in nonblank):
                tc.column_dimensions[chr(64+idx)].width=8

        tc.freeze_panes="A4"
        tc.auto_filter.ref=f"A3:X{max(4,r-1)}"
        # Editable PASS/FAIL list, left blank on generation.
        dv=DataValidation(type="list", formula1='"PASS,FAIL"', allow_blank=True)
        tc.add_data_validation(dv); dv.add(f"S4:S{max(100,r+20)}")

        # --- Deferred Test Intent -----------------------------------------
        deferred_ws.sheet_view.showGridLines = False
        deferred_ws["A1"] = "3. Deferred Test Intent"
        deferred_ws["A1"].font = Font(name="Malgun Gothic", size=14, bold=True)
        deferred_headers = [
            "SW 요구사항 ID", "Test Intent", "Reason Code", "Reason Detail",
            "Source / External Dependency", "Human Review Required",
        ]
        for col, val in enumerate(deferred_headers, 1):
            c = deferred_ws.cell(3, col, val)
            self._cell_style(c, fill=self.style.header, bold=True, center=True, font_size=9)
        dr = 4
        testability = requirement_data.get("testability_and_decomposition_result") if isinstance(requirement_data.get("testability_and_decomposition_result"), dict) else {}
        for row in testability.get("by_srs") or []:
            if not isinstance(row, dict):
                continue
            sid = _text(row.get("srs_id"))
            for item in row.get("not_generated_test_intents") or []:
                if not isinstance(item, dict):
                    continue
                vals = [
                    sid,
                    _text(item.get("intent")),
                    _text(item.get("reason_code")),
                    _text(item.get("reason")),
                    _text(item.get("source_or_dependency")),
                    "Y" if item.get("human_review_required") else "N",
                ]
                for col, val in enumerate(vals, 1):
                    c = deferred_ws.cell(dr, col, val)
                    self._cell_style(c, center=(col in (1, 3, 6)), font_size=9)
                dr += 1
        if dr == 4:
            deferred_ws.cell(4, 1, "-")
            deferred_ws.cell(4, 2, "Deferred intent 없음")
            for col in range(1, 7):
                self._cell_style(deferred_ws.cell(4, col), center=(col in (1, 3, 6)), font_size=9)
            dr = 5
        for idx, width in enumerate([18, 28, 28, 64, 52, 22], 1):
            deferred_ws.column_dimensions[chr(64 + idx)].width = width
        deferred_ws.freeze_panes = "A4"
        deferred_ws.auto_filter.ref = f"A3:F{max(4, dr-1)}"

        # --- Source / Child Intent Audit -----------------------------------
        intent_audit_ws.sheet_view.showGridLines = False
        intent_audit_ws["A1"] = "4. Source / Child Intent Audit"
        intent_audit_ws["A1"].font = Font(name="Malgun Gothic", size=14, bold=True)
        ia_headers = [
            "SW 요구사항 ID", "Child Intent ID", "Intent Status", "Related TC IDs",
            "Source Semantic Unit", "Source Location", "Source-backed Behavior",
            "Source Text Represented", "Generic TC Related", "Independently Observable", "Independent TC Coverage",
            "Deferred Reason Code", "Deferred / Review Detail",
        ]
        for col, val in enumerate(ia_headers, 1):
            c = intent_audit_ws.cell(3, col, val)
            self._cell_style(c, fill=self.style.header, bold=True, center=True, font_size=9)
        ir = 4
        for row in testability.get("by_srs") or []:
            if not isinstance(row, dict):
                continue
            sid = _text(row.get("srs_id"))
            for child in row.get("source_backed_child_intents") or []:
                if not isinstance(child, dict):
                    continue
                vals = [
                    sid, _text(child.get("child_intent_id")), _text(child.get("intent_status")),
                    ", ".join(str(x) for x in (child.get("related_tc_ids") or [])),
                    _text(child.get("source_semantic_unit_id")), _text(child.get("source_location")),
                    _text(child.get("source_backed_behavior")),
                    "Y" if child.get("source_text_represented") else "N",
                    "Y" if child.get("generic_tc_related") else "N",
                    _text(child.get("independently_observable")), _text(child.get("independent_tc_coverage")),
                    _text(child.get("deferred_reason_code")), _text(child.get("deferred_reason_detail")),
                ]
                for col, val in enumerate(vals, 1):
                    c = intent_audit_ws.cell(ir, col, val)
                    self._cell_style(c, center=(col in (1, 2, 3, 4, 5, 8, 9, 10, 11, 12)), font_size=9)
                ir += 1
        if ir == 4:
            intent_audit_ws.cell(4, 1, "-")
            intent_audit_ws.cell(4, 7, "Source-backed child intent 없음")
            for col in range(1, 14):
                self._cell_style(intent_audit_ws.cell(4, col), center=(col in (1, 2, 3, 4, 5, 8, 9, 10, 11, 12)), font_size=9)
            ir = 5
        for idx, width in enumerate([18, 22, 28, 24, 28, 24, 72, 18, 18, 24, 24, 28, 60], 1):
            intent_audit_ws.column_dimensions[chr(64 + idx)].width = width
        intent_audit_ws.freeze_panes = "A4"
        intent_audit_ws.auto_filter.ref = f"A3:M{max(4, ir-1)}"

        # --- Source Fact / Allocation Audit --------------------------------
        fact_audit_ws.sheet_view.showGridLines = False
        fact_audit_ws["A1"] = "5. Source Fact / Allocation Audit"
        fact_audit_ws["A1"].font = Font(name="Malgun Gothic", size=14, bold=True)
        fa_headers = [
            "SW 요구사항 ID", "Source Semantic Unit", "Source Location", "Source Fact",
            "Fact Type", "Value Role", "Auto TC Value Allowed", "Identifiers", "Numeric / Timing / Range", "Allocation Status",
            "SWE.6 Eligibility", "Verification Domain", "Allocation Inheritance",
            "Override Evidence", "Table Join Confidence", "Table Auto-use in TC",
        ]
        for col, val in enumerate(fa_headers, 1):
            c = fact_audit_ws.cell(3, col, val)
            self._cell_style(c, fill=self.style.header, bold=True, center=True, font_size=9)
        fr = 4
        case_by_srs = {}
        for case in cases:
            case_by_srs.setdefault(_text(case.get("srs_id")), []).append(case)
        for req in requirement_data.get("requirements") or []:
            if not isinstance(req, dict):
                continue
            sid = _text(req.get("srs_id"))
            joined = {str(x.get("fact_id") or ""): x for x in (req.get("source_table_fact_matches") or []) if isinstance(x, dict)}
            auto_used = {str(x.get("fact_id") or "") for c in case_by_srs.get(sid, []) for x in (c.get("source_table_fact_auto_used") or []) if isinstance(x, dict)}
            for alloc in req.get("fact_level_allocations") or []:
                if not isinstance(alloc, dict):
                    continue
                structured = alloc.get("structured_source_facts") or [{}]
                if not structured:
                    structured = [{}]
                for fact in structured:
                    if not isinstance(fact, dict):
                        fact = {}
                    fid = str(fact.get("fact_id") or "")
                    nums = list(fact.get("numeric_values") or []) + list(fact.get("timing_values") or []) + list(fact.get("range_values") or [])
                    vals = [
                        sid, _text(alloc.get("source_semantic_unit_id")), _text(alloc.get("source_location")), _text(alloc.get("source_fact")),
                        _text(fact.get("fact_type")), _text(fact.get("value_role")), "Y" if fact.get("auto_tc_value_allowed") else "N",
                        ", ".join(str(x) for x in (fact.get("identifiers") or [])),
                        ", ".join(dict.fromkeys(str(x) for x in nums if str(x))), _text(alloc.get("allocation_status")),
                        _text(alloc.get("swe6_eligibility")), _text(alloc.get("verification_domain")), _text(alloc.get("allocation_inheritance")),
                        "; ".join(str(x) for x in (alloc.get("allocation_override_evidence") or [])),
                        _text((joined.get(fid) or {}).get("join_confidence")), "Y" if fid in auto_used else "N",
                    ]
                    for col, val in enumerate(vals, 1):
                        c = fact_audit_ws.cell(fr, col, val)
                        self._cell_style(c, center=(col in (1, 2, 5, 6, 7, 10, 11, 12, 13, 15, 16)), font_size=9)
                    fr += 1
        if fr == 4:
            fact_audit_ws.cell(4, 1, "-")
            fact_audit_ws.cell(4, 4, "Source fact allocation 없음")
            for col in range(1, 17):
                self._cell_style(fact_audit_ws.cell(4, col), center=(col in (1, 2, 5, 6, 7, 10, 11, 12, 13, 15, 16)), font_size=9)
            fr = 5
        for idx, width in enumerate([18, 28, 24, 70, 28, 24, 18, 42, 42, 34, 30, 38, 34, 48, 22, 22], 1):
            fact_audit_ws.column_dimensions[get_column_letter(idx)].width = width
        fact_audit_ws.freeze_panes = "A4"
        fact_audit_ws.auto_filter.ref = f"A3:P{max(4, fr-1)}"

        # --- SYS.5 System Qualification Candidates --------------------------
        sys5_ws.sheet_view.showGridLines = False
        sys5_ws["A1"] = "6. SYS.5 System Qualification Candidate"
        sys5_ws["A1"].font = Font(name="Malgun Gothic", size=14, bold=True)
        sys5_headers = [
            "Requirement ID", "SYS.5 Candidate ID", "Verification Domain", "Source-backed Objective",
            "Source Fact Fragment IDs", "Status", "Execution Readiness", "Human Decision Required",
            "Missing Data / Dependency", "Execution Result", "PASS / FAIL",
        ]
        for col, val in enumerate(sys5_headers, 1):
            c = sys5_ws.cell(3, col, val)
            self._cell_style(c, fill=self.style.header, bold=True, center=True, font_size=9)
        sr = 4
        if not sys5_cases:
            sys5_cases = [{
                "srs_id": "", "sys5_id": "", "verification_domain": "",
                "objective": "SYS.5 Candidate 대상 없음", "source_fact_fragment_ids": [],
                "status": "", "execution_readiness": "", "human_decision_required": False,
                "missing_data_or_dependency": "", "execution_result": "", "pass_fail": "",
            }]
        for case in sys5_cases:
            vals = [
                case.get("srs_id", ""), case.get("sys5_id", ""), case.get("verification_domain", ""),
                case.get("objective", ""), ", ".join(case.get("source_fact_fragment_ids") or []),
                case.get("status", ""), case.get("execution_readiness", ""),
                "Y" if case.get("human_decision_required") else "N",
                case.get("missing_data_or_dependency", ""), "", "",
            ]
            for col, val in enumerate(vals, 1):
                c = sys5_ws.cell(sr, col, val)
                self._cell_style(c, center=(col in (1,2,3,6,7,8,10,11)), font_size=9)
            sr += 1
        for idx, width in enumerate([18,18,34,72,44,28,24,22,60,18,14], 1):
            sys5_ws.column_dimensions[get_column_letter(idx)].width = width
        sys5_ws.freeze_panes = "A4"
        sys5_ws.auto_filter.ref = f"A3:K{max(4, sr-1)}"
        sys5_dv = DataValidation(type="list", formula1='"PASS,FAIL"', allow_blank=True)
        sys5_ws.add_data_validation(sys5_dv)
        sys5_dv.add(f"K4:K{max(100, sr+20)}")

        # --- Integrated Engineering Verification Trace ---------------------
        eng_trace_ws.sheet_view.showGridLines = False
        eng_trace_ws["A1"] = "7. Engineering Verification Trace — SYS.5 / SWE.6"
        eng_trace_ws["A1"].font = Font(name="Malgun Gothic", size=14, bold=True)
        ev_headers = [
            "Requirement ID", "Engineering Domains", "SYS.1 Eligibility", "SWE.1 Eligibility",
            "SYS.5 Eligibility", "SWE.6 Eligibility", "Verification Domains", "Review Status",
            "Canonical State", "SYS.5 Candidate IDs", "SWE.6 TC IDs", "Official Release Eligible", "Hold Reason",
        ]
        for col, val in enumerate(ev_headers, 1):
            c = eng_trace_ws.cell(3, col, val)
            self._cell_style(c, fill=self.style.header, bold=True, center=True, font_size=9)
        swe6_ids_by_srs: dict[str, list[str]] = {}
        for case in cases:
            if case.get("srs_id") and case.get("tc_id"):
                swe6_ids_by_srs.setdefault(str(case.get("srs_id")), []).append(str(case.get("tc_id")))
        sys5_ids_by_srs: dict[str, list[str]] = {}
        for case in sys5_cases:
            if case.get("srs_id") and case.get("sys5_id"):
                sys5_ids_by_srs.setdefault(str(case.get("srs_id")), []).append(str(case.get("sys5_id")))
        er = 4
        for req in requirements:
            sid = _text(req.get("srs_id"))
            vals = [
                sid, ", ".join(str(x) for x in (req.get("engineering_domains") or [])),
                _text(req.get("sys1_eligibility")), _text(req.get("swe1_eligibility")),
                _text(req.get("sys5_eligibility")), _text(req.get("swe6_eligibility")),
                " | ".join(str(x) for x in (req.get("verification_domains") or [])),
                _text(req.get("review_status")), _text(req.get("canonical_state")),
                ", ".join(sys5_ids_by_srs.get(sid, [])), ", ".join(swe6_ids_by_srs.get(sid, [])),
                "Y" if req.get("official_release_eligible") else "N", _text(req.get("hold_reason")),
            ]
            for col, val in enumerate(vals, 1):
                c = eng_trace_ws.cell(er, col, val)
                self._cell_style(c, center=(col in (1,3,4,5,6,8,9,10,11,12)), font_size=9)
            er += 1
        for idx, width in enumerate([18,34,18,18,18,18,54,26,28,28,24,22,60], 1):
            eng_trace_ws.column_dimensions[get_column_letter(idx)].width = width
        eng_trace_ws.freeze_panes = "A4"
        eng_trace_ws.auto_filter.ref = f"A3:M{max(4, er-1)}"

        # Common formatting / print setup.
        for ws in (cover,history,summary,tc,deferred_ws,intent_audit_ws,fact_audit_ws,sys5_ws,eng_trace_ws):
            ws.sheet_properties.pageSetUpPr.fitToPage=True
            ws.page_setup.fitToWidth=1
            ws.page_setup.fitToHeight=0
            ws.page_margins.left=0.25; ws.page_margins.right=0.25; ws.page_margins.top=0.4; ws.page_margins.bottom=0.4
        cover.page_setup.orientation="portrait"
        history.page_setup.orientation="landscape"
        summary.page_setup.orientation="landscape"
        tc.page_setup.orientation="landscape"
        deferred_ws.page_setup.orientation="landscape"
        intent_audit_ws.page_setup.orientation="landscape"
        fact_audit_ws.page_setup.orientation="landscape"
        sys5_ws.page_setup.orientation="landscape"
        eng_trace_ws.page_setup.orientation="landscape"

        wb.save(out)
        self._enforce_blank_execution_results(out)
        return out
