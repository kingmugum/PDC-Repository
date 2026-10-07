from __future__ import annotations

import copy
import hashlib
import json
import os
import re
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime
from pathlib import Path
from typing import Any

import requests
from openpyxl import load_workbook

from core.document_normalizer import DocumentNormalizer
from core.human_policy import HumanPolicyRegistry
from providers.config_store import ProviderConfigStore
from providers.hchat_provider import HChatProvider


EVALUATION_SCHEMA_VERSION = "1.0"

EVALUATION_SCHEMA: dict[str, Any] = {
    "type": "object",
    "additionalProperties": False,
    "properties": {
        "schema_version": {"type": "string"},
        "provider": {"type": "string"},
        "model": {"type": "string"},
        "verdict": {"type": "string", "enum": [
            "PASS", "PASS_WITH_REVIEW_ITEMS", "PASS_WITH_MAJOR_REVIEW_ITEMS", "FAIL", "NOT_EVALUATED"
        ]},
        "tool_quality_gate": {"type": "string", "enum": ["PASS", "PASS_WITH_REVIEW_ITEMS", "FAIL", "NOT_EVALUATED"]},
        "artifact_readiness_gate": {"type": "string", "enum": ["PASS", "REVIEW_REQUIRED", "FAIL", "NOT_EVALUATED"]},
        "official_release": {"type": "string", "enum": ["PASS", "HOLD", "NOT_EVALUATED"]},
        "executive_summary": {"type": "array", "items": {"type": "string"}},
        "positive_checks": {"type": "array", "items": {"type": "string"}},
        "findings": {
            "type": "array",
            "items": {
                "type": "object",
                "additionalProperties": False,
                "properties": {
                    "finding_key": {"type": "string"},
                    "severity": {"type": "string", "enum": ["RELEASE_BLOCKING", "RELEASE_REVIEW", "ADVISORY"]},
                    "gate_scope": {"type": "string", "enum": ["TOOL_QUALITY", "ARTIFACT_READINESS", "RELEASE_GOVERNANCE", "ADVISORY"]},
                    "category": {"type": "string"},
                    "title": {"type": "string"},
                    "affected_ids": {"type": "array", "items": {"type": "string"}},
                    "source_references": {"type": "array", "items": {"type": "string"}},
                    "evidence": {"type": "array", "items": {"type": "string"}},
                    "problem": {"type": "string"},
                    "recommendation": {"type": "string"},
                    "confidence": {"type": "string", "enum": ["HIGH", "MEDIUM", "LOW"]},
                },
                "required": [
                    "finding_key", "severity", "gate_scope", "category", "title", "affected_ids",
                    "source_references", "evidence", "problem", "recommendation", "confidence"
                ],
            },
        },
        "release_gates": {
            "type": "array",
            "items": {
                "type": "object",
                "additionalProperties": False,
                "properties": {
                    "gate": {"type": "string"},
                    "result": {"type": "string"},
                    "evidence": {"type": "string"},
                },
                "required": ["gate", "result", "evidence"],
            },
        },
        "next_version_recommendations": {
            "type": "array",
            "items": {
                "type": "object",
                "additionalProperties": False,
                "properties": {
                    "priority": {"type": "string", "enum": ["P0", "P1", "P2"]},
                    "title": {"type": "string"},
                    "detail": {"type": "string"},
                },
                "required": ["priority", "title", "detail"],
            },
        },
        "limitations": {"type": "array", "items": {"type": "string"}},
    },
    "required": [
        "schema_version", "provider", "model", "verdict", "tool_quality_gate", "artifact_readiness_gate",
        "official_release", "executive_summary", "positive_checks", "findings", "release_gates",
        "next_version_recommendations", "limitations"
    ],
}


DEFAULT_CONFIG: dict[str, Any] = {
    "enabled": True,
    "transport": "hchat",
    "run_after_review_package": True,
    "require_all_providers": False,
    "final_result_chunk_lines": 110,
    "situation_check_chunk_lines": 110,
    "max_source_chars": 180000,
    "max_review_json_chars": 260000,
    "timeout_seconds": 240,
    "providers": {
        "gpt": {
            "enabled": True,
            "direct_model": "gpt-6.1-sol",
            "direct_base_url": "https://api.openai.com/v1",
            "api_key_env": "OPENAI_API_KEY",
        },
        "gemini": {
            "enabled": True,
            "direct_model": "gemini-3.8-flash",
            "direct_base_url": "https://generativelanguage.googleapis.com/v1beta",
            "api_key_env": "GEMINI_API_KEY",
        },
        "claude": {
            "enabled": True,
            "direct_model": "claude-sonnet-5-5",
            "direct_base_url": "https://api.anthropic.com",
            "api_key_env": "ANTHROPIC_API_KEY",
            "anthropic_version": "2023-06-01",
        },
    },
}


def _deep_merge(base: dict[str, Any], overlay: dict[str, Any]) -> dict[str, Any]:
    out = copy.deepcopy(base)
    for key, value in (overlay or {}).items():
        if isinstance(value, dict) and isinstance(out.get(key), dict):
            out[key] = _deep_merge(out[key], value)
        else:
            out[key] = value
    return out


def _normalize_space(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def _safe_json_extract(text: str) -> dict[str, Any]:
    raw = (text or "").strip()
    if raw.startswith("```"):
        raw = re.sub(r"^```(?:json)?\s*", "", raw, flags=re.I)
        raw = re.sub(r"\s*```$", "", raw)
    try:
        data = json.loads(raw)
    except Exception:
        a, b = raw.find("{"), raw.rfind("}")
        if a < 0 or b <= a:
            raise ValueError("평가 응답에서 JSON object를 찾을 수 없습니다.")
        data = json.loads(raw[a:b + 1])
    if not isinstance(data, dict):
        raise ValueError("평가 응답의 최상위 값은 JSON object여야 합니다.")
    return data


def _schema_validate_minimum(data: dict[str, Any]) -> None:
    required = EVALUATION_SCHEMA["required"]
    missing = [x for x in required if x not in data]
    if missing:
        raise ValueError("평가 JSON 필수 필드 누락: " + ", ".join(missing))
    if not isinstance(data.get("findings"), list):
        raise ValueError("findings는 list여야 합니다.")


def _hash_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def _atomic_write_text(path: Path, text: str) -> None:
    """Durably publish one evaluator artifact.

    V0.68 writes to a sibling temporary file, flushes/fsyncs it, and only then
    replaces the visible target.  This prevents the UI from reporting a
    completed evaluation while the user-visible result file is missing or
    partially written.
    """
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name(path.name + ".tmp")
    with tmp.open("w", encoding="utf-8", newline="") as f:
        f.write(text)
        f.flush()
        os.fsync(f.fileno())
    os.replace(tmp, path)


class MultiModelEvaluationOrchestrator:
    """Independent GPT/Gemini/Claude evaluator for one Review Exchange run.

    The generator model and reviewer models remain independent.  Each reviewer receives the exact
    same evidence package and does not see the other reviewers' outputs.  Provider JSON is persisted
    first, then rendered to human-readable TXT.  Consensus is calculated only after all independent
    reviews finish.
    """

    def __init__(self, project_root: Path, *, app_version: str = "v0.88", provider_config: dict[str, Any] | None = None):
        self.project_root = Path(project_root).resolve()
        self.app_version = app_version
        self.config_path = self.project_root / "config" / "evaluation_api_config.json"
        self.config = self._load_config()
        self.provider_config = provider_config or ProviderConfigStore(self.project_root).load()

    def _load_config(self) -> dict[str, Any]:
        try:
            user = json.loads(self.config_path.read_text(encoding="utf-8")) if self.config_path.is_file() else {}
        except Exception:
            user = {}
        return _deep_merge(DEFAULT_CONFIG, user)

    def is_enabled(self) -> bool:
        return bool(self.config.get("enabled", True))

    def _source_text(self, source_path: Path) -> str:
        try:
            normalized = DocumentNormalizer(self.project_root).normalize(source_path)
            return normalized.compact_text()[: int(self.config.get("max_source_chars") or 180000)]
        except Exception as exc:
            return f"[SOURCE EXTRACTION FAILED] {exc}"

    @staticmethod
    def _xlsx_snapshot(path: Path, *, profile: str = "generic") -> dict[str, Any]:
        wb = load_workbook(path, data_only=False, read_only=True)
        try:
            swe6_keep = {"2_테스트 케이스", "3_Deferred Intent", "4_Source_Intent_Audit", "5_Source_Fact_Audit", "6_SYS5_Candidates", "7_Engineering_Verification"}
            integrated_req_keep = {"표지", "0_변경이력", "1_요구사항요약", "2_E2E_요구사항", "3_검토필요", "4_Source_추적성"}
            integrated_test_keep = {"표지", "0_변경이력", "1_평가요약", "2_E2E_평가케이스", "3_검토필요", "4_요구사항-평가추적성"}
            sheets: dict[str, Any] = {}
            sheet_row_counts: dict[str, int] = {}
            truncated_sheets: list[str] = []
            for ws in wb.worksheets:
                if profile == "swe6" and ws.title not in swe6_keep:
                    continue
                if profile == "integrated_requirements" and ws.title not in integrated_req_keep:
                    continue
                if profile == "integrated_tests" and ws.title not in integrated_test_keep:
                    continue
                rows = []
                # V0.88 E2E workbooks are evaluated in full. Legacy snapshots retain bounded rows.
                row_limit = None if profile in {"integrated_requirements", "integrated_tests"} else (600 if profile == "swe6" else 260)
                for row in ws.iter_rows(values_only=True):
                    values = [None if v is None else str(v) for v in row]
                    if any(v not in (None, "") for v in values):
                        rows.append(values)
                    if row_limit is not None and len(rows) >= row_limit:
                        rows.append([f"[TRUNCATED AFTER {row_limit} NONEMPTY ROWS]"])
                        truncated_sheets.append(ws.title)
                        break
                if rows:
                    sheets[ws.title] = rows
                    sheet_row_counts[ws.title] = len([r for r in rows if not (isinstance(r, list) and r and str(r[0] or "").startswith("[TRUNCATED"))])
            return {
                "artifact": path.name, "sha256": _hash_file(path), "profile": profile, "sheets": sheets,
                "sheet_row_counts": sheet_row_counts, "truncated_sheets": truncated_sheets,
                "evaluation_scope_complete": not truncated_sheets,
                "evaluation_coverage_percent": 100.0 if not truncated_sheets else None,
            }
        finally:
            wb.close()

    @staticmethod
    def _docx_snapshot(path: Path) -> dict[str, Any]:
        from docx import Document
        doc = Document(path)
        lines = [_normalize_space(p.text) for p in doc.paragraphs if _normalize_space(p.text)]
        tables = []
        for tidx, table in enumerate(doc.tables, 1):
            rows = []
            for row in table.rows:
                vals = [_normalize_space(cell.text) for cell in row.cells]
                if any(vals):
                    rows.append(vals)
            if rows:
                tables.append({"table_index": tidx, "rows": rows[:240]})
            if len(tables) >= 120:
                break
        return {
            "artifact": path.name,
            "sha256": _hash_file(path),
            "paragraphs": lines[:1600],
            "tables": tables,
            "table_count": len(doc.tables),
        }

    @staticmethod
    def _trace_graph(review: dict[str, Any]) -> dict[str, Any]:
        canonical = review.get("canonical_requirement") or {}
        reqs = []
        for req in (canonical.get("requirements") or []):
            if not isinstance(req, dict):
                continue
            reqs.append({
                "srs_id": req.get("srs_id"),
                "candidate_id": req.get("candidate_id"),
                "requirement": req.get("requirement"),
                "allocation_status": req.get("allocation_status"),
                "swe1_eligibility": req.get("swe1_eligibility"),
                "swe6_eligibility": req.get("swe6_eligibility"),
                "verification_domain": req.get("verification_domain"),
                "source_semantic_unit_ids": req.get("source_semantic_unit_ids"),
                "source_fact_fragments": [
                    {
                        "id": f.get("source_fact_fragment_id"),
                        "location": f.get("source_location"),
                        "excerpt": f.get("source_excerpt_raw") or f.get("source_excerpt"),
                        "ownership_match_score": f.get("ownership_match_score"),
                    }
                    for f in (req.get("source_fact_fragments") or []) if isinstance(f, dict)
                ],
                "child_intents": [
                    {
                        "id": a.get("atomic_behavior_id") or a.get("child_intent_id"),
                        "fragment_id": a.get("source_fact_fragment_id"),
                        "semantic_unit_id": a.get("source_semantic_unit_id"),
                        "behavior": a.get("source_backed_behavior") or a.get("behavior"),
                    }
                    for a in (req.get("source_backed_atomic_behaviors") or []) if isinstance(a, dict)
                ],
            })
        return {
            "requirements": reqs,
            "calculated_review_evidence": review.get("calculated_review_evidence") or {},
        }

    def build_evidence_package(self, run_dir: Path) -> tuple[dict[str, Any], Path]:
        run_dir = Path(run_dir).resolve()
        review_path = run_dir / "02_REQUIREMENT_STUDIO_REVIEW_PACKAGE.json"
        if not review_path.is_file():
            raise FileNotFoundError(review_path)
        review = json.loads(review_path.read_text(encoding="utf-8"))
        handoff = review.get("review_handoff") or {}
        evaluation_request = review.get("evaluation_request") or {}
        focus = str(handoff.get("evaluation_artifact_focus") or "unified").lower()
        source_candidates = sorted(run_dir.glob("01_SOURCE_*"))
        source_path = source_candidates[0] if source_candidates else None
        decision_candidates = sorted(run_dir.glob("04_CHANGE_DECISION_*.txt"))
        summary_path = run_dir / "05_REQUIREMENT_STUDIO_REVIEW_SUMMARY.txt"

        actual_artifacts: dict[str, Any] = {}
        is_v086 = str(self.app_version).lower().lstrip("v") >= "0.86"
        artifact_map = ({
            "swe1_word": handoff.get("file_3_swe1_word"),
            "swe1_excel": handoff.get("file_3_swe1_excel"),
            "swe6_excel": handoff.get("file_3_swe6_excel"),
            "e2e_requirements_word": handoff.get("file_3_e2e_requirements_word"),
            "e2e_requirements_excel": handoff.get("file_3_e2e_requirements_excel") or handoff.get("file_3_integrated_requirements_excel"),
            "e2e_evaluation_excel": handoff.get("file_3_e2e_evaluation_excel") or handoff.get("file_3_integrated_tests_excel"),
        } if is_v086 else {
            "swe1_word": handoff.get("file_3_swe1_word"),
            "swe1_excel": handoff.get("file_3_swe1_excel"),
            "swe6_excel": handoff.get("file_3_swe6_excel"),
            "integrated_requirements_excel": handoff.get("file_3_integrated_requirements_excel"),
            "integrated_tests_excel": handoff.get("file_3_integrated_tests_excel"),
        })
        # Backward compatibility with pre-unified Review Exchange folders.
        if not any(artifact_map.values()):
            candidates = sorted(run_dir.glob("03_*"))
            for candidate in candidates:
                lower = candidate.name.lower()
                if candidate.suffix.lower() == ".docx": artifact_map["swe1_word"] = candidate.name
                elif candidate.suffix.lower() in {".xlsx", ".xlsm"}: artifact_map["swe6_excel" if "swe.6" in lower or "swe6" in lower else "swe1_excel"] = candidate.name
        for key, filename in artifact_map.items():
            if not filename:
                continue
            path = run_dir / str(filename)
            if not path.is_file():
                continue
            if key in {"swe1_word", "e2e_requirements_word"}:
                actual_artifacts[key] = self._docx_snapshot(path)
            elif key == "swe6_excel":
                actual_artifacts[key] = self._xlsx_snapshot(path, profile="swe6")
            elif key in {"e2e_requirements_excel", "integrated_requirements_excel"}:
                actual_artifacts[key] = self._xlsx_snapshot(path, profile="integrated_requirements")
            elif key in {"e2e_evaluation_excel", "integrated_tests_excel"}:
                actual_artifacts[key] = self._xlsx_snapshot(path, profile="integrated_tests")
            else:
                actual_artifacts[key] = self._xlsx_snapshot(path, profile="swe1")

        source_text = self._source_text(source_path) if source_path else "[SOURCE FILE MISSING]"
        review_text = json.dumps(review, ensure_ascii=False, indent=2)
        max_review = int(self.config.get("max_review_json_chars") or 260000)
        if len(review_text) > max_review:
            compact_review = {
                "run": review.get("run"),
                "review_handoff": handoff,
                "evaluation_request": evaluation_request,
                "local_qa": review.get("local_qa"),
                "quality_review_config": review.get("quality_review_config"),
                "canonical_requirement": review.get("canonical_requirement"),
                "calculated_review_evidence": review.get("calculated_review_evidence"),
                "reviewer_findings_policy": review.get("reviewer_findings_policy"),
            }
            review_text = json.dumps(compact_review, ensure_ascii=False, indent=2)
            if len(review_text) > max_review:
                review_text = review_text[:max_review] + "\n[REVIEW JSON TRUNCATED BY EVALUATION EVIDENCE LIMIT]"

        policy_registry = HumanPolicyRegistry(self.project_root).load()
        active_human_policies = [
            {
                "policy_id": str(p.get("policy_id") or ""),
                "title": str(p.get("title") or ""),
                "decision": str(p.get("decision") or ""),
                "reuse_scope": str(p.get("reuse_scope") or ""),
                "notes": str(p.get("notes") or ""),
            }
            for p in (policy_registry.get("policies") or [])
            if isinstance(p, dict) and str(p.get("status") or "ACTIVE").upper() == "ACTIVE"
        ]
        evidence = {
            "evidence_schema_version": "2.0",
            "requirement_studio_version": self.app_version,
            "evaluation_mode": focus,
            "run_dir": run_dir.name,
            "source": {
                "file": source_path.name if source_path else "",
                "sha256": _hash_file(source_path) if source_path else "",
                "normalized_text": source_text,
            },
            "trace_graph": self._trace_graph(review),
            "review_package_json": review_text,
            "actual_artifacts": actual_artifacts,
            "artifact_presence": {k: bool(v) for k, v in artifact_map.items()},
            "selected_outputs": (review.get("quality_review_config") or {}).get("selected_outputs") or {},
            "unified_artifact_contract": handoff.get("unified_artifact_contract") or ((review.get("calculated_review_evidence") or {}).get("unified_evaluation_artifact_contract") if isinstance(review.get("calculated_review_evidence"), dict) else {}) or {},
            "change_decision": decision_candidates[0].read_text(encoding="utf-8", errors="replace") if decision_candidates else "",
            "human_review_summary": summary_path.read_text(encoding="utf-8", errors="replace") if summary_path.is_file() else "",
            "human_policy_registry": {
                "schema_version": policy_registry.get("schema_version"),
                "active_policies": active_human_policies,
                "review_preservation_policy_active": any(p.get("policy_id") == "HP-PRESERVE-001" and p.get("decision") == "YES" for p in active_human_policies),
                "policy_reuse_is_not_per_srs_official_allocation": True,
            },
            "review_rules": {
                "independent_reviewers": True,
                "do_not_invent_missing_source_facts": True,
                "do_not_penalize_correct_deferral": True,
                "distinguish_tool_quality_from_artifact_readiness": True,
                "gold_regression_requires_approved_baseline": True,
                "unselected_or_not_generated_artifact_is_not_a_defect": True,
                "unified_mode_requires_cross_output_consistency_review": focus == "unified",
                "unified_mode_requires_legacy_and_integrated_actual_artifacts": focus == "unified",
                "exact_fragment_ownership_is_union_of_all_linked_fragments": True,
                "do_not_reflag_ownership_when_owned_fragment_fact_token_sources_proves_coverage": True,
                "open_review_needed_is_not_tool_defect_when_explicitly_nonterminal": True,
                "inferred_controller_sw_behavior_is_review_required_modeling_not_explicit_allocation": True,
                "respect_deterministic_parent_containment_disambiguation": True,
                "review_required_hold_remains_visible_in_main_and_review_views": True,
                "human_policy_reuse_does_not_equal_per_srs_official_allocation": True,
                "e2e_requirements_preserve_sys1_swe1_and_native_domains": True,
                "e2e_evaluation_preserves_sys5_swe6_deferred_and_native_intents": True,
                "e2e_evaluation_all_sheets_all_rows_are_evaluation_scope": True,
                "structural_gate_is_separate_from_semantic_verification_quality_gate": True,
                "mode_transition_allocation_difference_requires_review_not_auto_reallocation": True,
                "integrated_test_execution_result_fields_must_be_blank_pre_execution": True,
            },
        }
        canonical = json.dumps(evidence, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
        evidence["evidence_package_sha256"] = hashlib.sha256(canonical.encode("utf-8")).hexdigest()
        out_path = run_dir / "automatic_evaluation" / "06_EVALUATION_EVIDENCE.json"
        _atomic_write_text(out_path, json.dumps(evidence, ensure_ascii=False, indent=2))
        return evidence, out_path

    @staticmethod
    def _count_deferred_snapshot_rows(rows: list[Any]) -> int:
        """Count serialized Deferred Intent data rows without counting title/header/sentinel rows."""
        if not isinstance(rows, list) or not rows:
            return 0
        header_index = None
        for i, row in enumerate(rows):
            if not isinstance(row, list):
                continue
            first = str(row[0] or "").strip() if row else ""
            second = str(row[1] or "").strip() if len(row) > 1 else ""
            if first in {"SW 요구사항 ID", "요구사항 ID", "Requirement ID"} and "Intent" in second:
                header_index = i
                break
        data = rows[(header_index + 1) if header_index is not None else 0:]
        count = 0
        for row in data:
            if not isinstance(row, list) or not row:
                continue
            first = str(row[0] or "").strip()
            if not first or first == "-" or first.startswith("[TRUNCATED AFTER"):
                continue
            # Deferred rows are keyed by a requirement/SRS identifier. Avoid counting workbook titles.
            if re.match(r"^(?:SRS|REQ|SYS|SWE)[_-]?[A-Za-z0-9]+", first, re.I):
                count += 1
        return count

    @staticmethod
    def _build_situation_check(review: dict[str, Any], evidence: dict[str, Any]) -> dict[str, Any]:
        """Build a deterministic policy-impact diagnosis for unexpectedly low SWE.1/SWE.6 output.

        V0.88 retains Engineering E2E/single truth and makes E2E Requirements Word/Excel plus the complete E2E Evaluation workbook mandatory evidence in recommended Unified AI evaluation while preserving Source-fidelity, ownership, and exact-trace safeguards.
        """
        canonical = review.get("canonical_requirement") if isinstance(review.get("canonical_requirement"), dict) else {}
        calc = review.get("calculated_review_evidence") if isinstance(review.get("calculated_review_evidence"), dict) else {}
        reqs = [x for x in (canonical.get("requirements") or []) if isinstance(x, dict)]
        sem_cov = canonical.get("semantic_source_unit_coverage") if isinstance(canonical.get("semantic_source_unit_coverage"), dict) else {}
        testability = calc.get("testability_and_decomposition_result") if isinstance(calc.get("testability_and_decomposition_result"), dict) else {}
        test_summary = testability.get("summary") if isinstance(testability.get("summary"), dict) else {}
        swe6_audit = calc.get("swe6_export_preservation_audit") if isinstance(calc.get("swe6_export_preservation_audit"), dict) else {}
        verification_bundle = calc.get("finalized_verification_bundle") if isinstance(calc.get("finalized_verification_bundle"), dict) else {}
        integrated_summary = verification_bundle.get("summary") if isinstance(verification_bundle.get("summary"), dict) else {}
        integrated_counts = integrated_summary.get("test_object_type_counts") if isinstance(integrated_summary.get("test_object_type_counts"), dict) else {}
        integrated_audit = calc.get("integrated_single_truth_audit") if isinstance(calc.get("integrated_single_truth_audit"), dict) else {}
        conflicts = calc.get("conflict_register") if isinstance(calc.get("conflict_register"), list) else []

        def _eligible(value: Any) -> str:
            return str(value or "").strip()

        main_swe1 = [r for r in reqs if _eligible(r.get("swe1_eligibility")) == "Eligible"]
        pending_swe1 = [r for r in reqs if _eligible(r.get("swe1_eligibility")) == "Review Needed"]
        not_applicable = [r for r in reqs if _eligible(r.get("swe1_eligibility")) == "Not Applicable"]
        swe6_eligible = [r for r in reqs if _eligible(r.get("swe6_eligibility")) == "Eligible"]
        swe6_deferred = [r for r in reqs if _eligible(r.get("swe6_eligibility")).startswith("Deferred")]
        sys1_eligible = [r for r in reqs if _eligible(r.get("sys1_eligibility")) == "Eligible"]
        sys1_review = [r for r in reqs if _eligible(r.get("sys1_eligibility")) == "Review Needed"]
        sys5_candidates = [r for r in reqs if _eligible(r.get("sys5_eligibility")) in {"Eligible", "Review Needed"}]
        engineering_main = [r for r in reqs if bool(r.get("main_spec_visibility", True))]
        human_decision = [r for r in reqs if bool(r.get("human_decision_required"))]
        positive_sw = [r for r in reqs if (r.get("positive_sw_allocation_evidence") or []) or _eligible(r.get("allocation_status")) == "SW_IMPLEMENTATION_REQUIREMENT"]
        inferred_sw = [r for r in reqs if _eligible(r.get("allocation_status")) == "SW_BEHAVIOR_INFERRED_REVIEW_REQUIRED"]
        pending_due_no_positive_sw = [
            r for r in pending_swe1
            if "PENDING_SW_ALLOCATION" in _eligible(r.get("allocation_status"))
            and not (r.get("positive_sw_allocation_evidence") or [])
        ]

        allocation_counts: dict[str, int] = {}
        for r in reqs:
            key = _eligible(r.get("allocation_status")) or "UNCLASSIFIED"
            allocation_counts[key] = allocation_counts.get(key, 0) + 1

        disposition_counts = sem_cov.get("disposition_counts") if isinstance(sem_cov.get("disposition_counts"), dict) else {}
        potential_extraction_loss = int(disposition_counts.get("SEMANTIC_MISSING_POTENTIAL_EXTRACTION_LOSS") or 0)
        open_semantic = int(sem_cov.get("open_review_needed_count") or 0)
        generated_tc = int(test_summary.get("total_generated_tc_count") or swe6_audit.get("generated_tc_count") or 0)
        deferred_intents_summary = int(test_summary.get("deferred_intent_count") or 0)

        actual_artifacts = evidence.get("actual_artifacts") if isinstance(evidence.get("actual_artifacts"), dict) else {}
        swe6_actual = actual_artifacts.get("swe6_excel") if isinstance(actual_artifacts.get("swe6_excel"), dict) else {}
        swe6_sheets = swe6_actual.get("sheets") if isinstance(swe6_actual.get("sheets"), dict) else {}
        deferred_rows = swe6_sheets.get("3_Deferred Intent") if isinstance(swe6_sheets.get("3_Deferred Intent"), list) else []
        actual_deferred_intents = MultiModelEvaluationOrchestrator._count_deferred_snapshot_rows(deferred_rows)
        deferred_intents = actual_deferred_intents if actual_deferred_intents > 0 else deferred_intents_summary

        total = len(reqs)
        policy_held_count = len(pending_due_no_positive_sw)
        historical_gate_affected_count = policy_held_count + len(inferred_sw)
        policy_ratio = round((historical_gate_affected_count / total * 100.0), 2) if total else 0.0
        main_ratio = round((len(main_swe1) / total * 100.0), 2) if total else 0.0

        # A low Main/SWE.6 count is primarily policy-driven when the Canonical SRS still exist and
        # at least half are explicitly held in the allocation-pending path rather than lost.
        if total and historical_gate_affected_count >= max(1, (total + 1) // 2) and potential_extraction_loss == 0:
            hypothesis = "SUPPORTED"
            hypothesis_reason = (
                "Canonical SRS는 유지되고 있으며 다수가 기존 positive-SW-only gate의 영향권에 있습니다. "
                "V0.88에서도 controller software-like behavior는 review-required inferred candidate로 승격될 수 있고, "
                "나머지 항목만 Pending으로 유지됩니다. 요구사항 후보 소실보다 allocation/promotion policy 영향이 우세합니다."
            )
        elif historical_gate_affected_count:
            hypothesis = "PARTIALLY_SUPPORTED"
            hypothesis_reason = (
                "Allocation policy가 Main SWE.1/SWE.6 감소에 유의미하게 기여하지만, OPEN Semantic Unit/추출 손실/비-SW 분류 등 "
                "다른 원인도 함께 확인해야 합니다."
            )
        else:
            hypothesis = "NOT_SUPPORTED"
            hypothesis_reason = (
                "현재 Run에서는 positive-SW allocation gate가 낮은 산출물 개수의 주원인이라는 근거가 충분하지 않습니다. "
                "Source extraction, semantic linkage, testability 또는 다른 gate를 우선 확인해야 합니다."
            )

        pending_details = []
        for r in pending_swe1:
            evs = [e for e in (r.get("source_evidence") or []) if isinstance(e, dict)]
            locations = [str(e.get("location") or "") for e in evs if str(e.get("location") or "").strip()]
            pending_details.append({
                "srs_id": r.get("srs_id"),
                "requirement": _normalize_space(r.get("requirement")),
                "allocation_status": r.get("allocation_status"),
                "swe6_eligibility": r.get("swe6_eligibility"),
                "verification_domain": r.get("verification_domain"),
                "positive_sw_evidence": list(r.get("positive_sw_allocation_evidence") or []),
                "allocation_rationale": _normalize_space(r.get("allocation_rationale")),
                "source_locations": locations[:4],
            })

        open_semantic_records = []
        for x in (sem_cov.get("open_review_needed_records") or []):
            if not isinstance(x, dict):
                continue
            open_semantic_records.append({
                "source_semantic_unit_id": x.get("source_semantic_unit_id"),
                "source_location": x.get("source_location"),
                "disposition_status": x.get("disposition_status"),
                "disposition_reason": _normalize_space(x.get("disposition_reason")),
                "required_resolution": _normalize_space(x.get("required_resolution")),
            })

        source_meta = evidence.get("source") if isinstance(evidence.get("source"), dict) else {}
        return {
            "schema_version": "1.0",
            "generated_at": datetime.now().isoformat(timespec="seconds"),
            "handoff_identity": {
                "requirement_studio_version": evidence.get("requirement_studio_version") or "",
                "evaluation_mode": evidence.get("evaluation_mode") or "",
                "source_file": source_meta.get("file") or "",
                "source_sha256": source_meta.get("sha256") or "",
                "run_id": evidence.get("run_dir") or "",
                "evidence_package_sha256": evidence.get("evidence_package_sha256") or "",
            },
            "observed_counts": {
                "canonical_srs_count": total,
                "engineering_main_requirement_count": len(engineering_main),
                "sys1_eligible_count": len(sys1_eligible),
                "sys1_review_needed_count": len(sys1_review),
                "sys5_candidate_requirement_count": len(sys5_candidates),
                "human_decision_required_count": len(human_decision),
                "main_swe1_eligible_count": len(main_swe1),
                "pending_swe1_annex_count": len(pending_swe1),
                "not_applicable_swe1_count": len(not_applicable),
                "swe6_eligible_srs_count": len(swe6_eligible),
                "swe6_deferred_srs_count": len(swe6_deferred),
                "generated_tc_count": generated_tc,
                "integrated_test_object_count": int(integrated_summary.get("integrated_test_object_count") or 0),
                "integrated_sys5_candidate_count": int(integrated_counts.get("SYS5_CANDIDATE") or 0),
                "integrated_swe6_tc_count": int(integrated_counts.get("SWE6_TC") or 0),
                "integrated_allocation_pending_intent_count": int(integrated_counts.get("ALLOCATION_PENDING_INTENT") or 0),
                "integrated_external_dependency_intent_count": int(integrated_counts.get("EXTERNAL_DEPENDENCY_INTENT") or 0),
                "integrated_single_truth_gate": str(integrated_audit.get("release_gate_status") or integrated_summary.get("release_gate_status") or ""),
                "mode_transition_allocation_finding_count": int(integrated_summary.get("mode_transition_allocation_finding_count") or 0),
                "deferred_intent_count": deferred_intents,
                "deferred_intent_count_review_summary": deferred_intents_summary,
                "deferred_intent_count_actual_xlsx": actual_deferred_intents,
                "deferred_intent_count_consistent": (actual_deferred_intents == 0 or deferred_intents_summary == 0 or actual_deferred_intents == deferred_intents_summary),
                "positive_sw_evidence_srs_count": len(positive_sw),
                "inferred_sw_behavior_srs_count": len(inferred_sw),
                "policy_held_no_positive_sw_count": policy_held_count,
                "historical_positive_sw_gate_affected_count": historical_gate_affected_count,
                "main_swe1_ratio_percent": main_ratio,
                "policy_held_ratio_percent": policy_ratio,
            },
            "allocation_status_counts": allocation_counts,
            "semantic_coverage": {
                "eligible_unit_count": sem_cov.get("eligible_unit_count"),
                "covered_unit_count": sem_cov.get("covered_unit_count"),
                "open_review_needed_count": open_semantic,
                "potential_extraction_loss_count": potential_extraction_loss,
                "terminal_disposition_complete": sem_cov.get("terminal_disposition_complete"),
            },
            "source_conflict_count": len(conflicts),
            "hypothesis_verdict": hypothesis,
            "hypothesis_reason": hypothesis_reason,
            "policy_interpretation": {
                "current_rule": "V0.88: Canonical Engineering Requirements are preserved before domain projection and exposed in additive integrated requirements/tests. System requirements remain available to SYS.1/SYS.5, software requirements to SWE.1/SWE.6, and native Engineering domains remain visible without inventing missing facts.",
                "primary_gate": "V0.88 Engineering E2E + Structural Single Truth + Semantic Verification Quality + Unified Artifact Contract: Canonical Engineering Requirement → SYS.1/SWE.1/native → SYS.5/SWE.6/Deferred/native verification with protected Source-fidelity safeguards",
                "hypothetical_main_swe1_upper_bound_if_only_allocation_gate_relaxed": len(main_swe1) + policy_held_count,
                "upper_bound_warning": "The inferred promotion is a review candidate, not approved software allocation. Concrete SWE.6 cases remain subject to source-backed testability, conflict, provenance, and execution-readiness gates.",
            },
            "pending_requirements": pending_details,
            "open_semantic_units": open_semantic_records,
            "strategy_guardrail": (
                "V0.88 preserves the Engineering E2E model and E2E requirement/evaluation specifications, and requires those integrated workbooks in recommended Unified AI evaluation while retaining Human Policy Registry reuse. Not SWE.1 does not remove a Requirement and Not SWE.6 does not remove a verifiable System candidate or Source-backed test intent. "
                "Source-fidelity, unsupported-generation, exact-trace, conflict, visual-unknown, and arbitrary test-value safeguards remain active."
            ),
        }

    @staticmethod
    def _render_situation_check(report: dict[str, Any]) -> str:
        ident = report.get("handoff_identity") or {}
        c = report.get("observed_counts") or {}
        sem = report.get("semantic_coverage") or {}
        policy = report.get("policy_interpretation") or {}
        lines = [
            "Requirement Studio Situation Check",
            "=" * 72,
            f"Requirement Studio Version: {ident.get('requirement_studio_version','')}",
            f"Evaluation Mode: {ident.get('evaluation_mode','')}",
            f"Source File: {ident.get('source_file','')}",
            f"Source SHA-256: {ident.get('source_sha256','')}",
            f"Run ID: {ident.get('run_id','')}",
            f"Evidence Package SHA-256: {ident.get('evidence_package_sha256','')}",
            f"Generated At: {report.get('generated_at','')}",
            "",
            "1. Question Being Checked",
            "-------------------------",
            "Does V0.88 preserve Source-backed engineering content through SYS.1/SWE.1 into E2E Requirements and through SYS.5/SWE.6/deferred/native verification into the E2E Evaluation specification, while preserving logic, exact fact scope, human-review states and Source fidelity?",
            "",
            "2. Observed Counts",
            "------------------",
            f"Canonical Engineering Requirements: {c.get('canonical_srs_count',0)}",
            f"Engineering Main Visible: {c.get('engineering_main_requirement_count',0)}",
            f"SYS.1 Eligible: {c.get('sys1_eligible_count',0)}",
            f"SYS.1 Review Needed: {c.get('sys1_review_needed_count',0)}",
            f"SYS.5 Candidate Requirements: {c.get('sys5_candidate_requirement_count',0)}",
            f"Human Decision Required: {c.get('human_decision_required_count',0)}",
            f"Main SWE.1 Eligible: {c.get('main_swe1_eligible_count',0)} ({c.get('main_swe1_ratio_percent',0)}%)",
            f"Allocation Review Annex / Review Needed: {c.get('pending_swe1_annex_count',0)}",
            f"SWE.1 Not Applicable: {c.get('not_applicable_swe1_count',0)}",
            f"SWE.6 Eligible SRS: {c.get('swe6_eligible_srs_count',0)}",
            f"SWE.6 Deferred SRS: {c.get('swe6_deferred_srs_count',0)}",
            f"Generated SWE.6 TC: {c.get('generated_tc_count',0)}",
            f"Integrated Test Objects: {c.get('integrated_test_object_count',0)}",
            f"  - SYS.5 Candidates: {c.get('integrated_sys5_candidate_count',0)}",
            f"  - SWE.6 TCs: {c.get('integrated_swe6_tc_count',0)}",
            f"  - Allocation Pending Intents: {c.get('integrated_allocation_pending_intent_count',0)}",
            f"  - External Dependency Intents: {c.get('integrated_external_dependency_intent_count',0)}",
            f"  - Integrated Single Truth Gate: {c.get('integrated_single_truth_gate','')}",
            f"  - Mode Transition Allocation Review Findings: {c.get('mode_transition_allocation_finding_count',0)}",
            f"Deferred Test Intents: {c.get('deferred_intent_count',0)}",
            f"  - Review Package summary count: {c.get('deferred_intent_count_review_summary',0)}",
            f"  - Actual SWE.6 XLSX count: {c.get('deferred_intent_count_actual_xlsx',0)}",
            f"  - Count consistency: {c.get('deferred_intent_count_consistent')}",
            f"SRS with explicit positive SW evidence: {c.get('positive_sw_evidence_srs_count',0)}",
            f"SRS promoted by V0.88 inferred controller-SW behavior: {c.get('inferred_sw_behavior_srs_count',0)}",
            f"Still held by no-positive-SW allocation policy: {c.get('policy_held_no_positive_sw_count',0)}",
            f"Historical positive-SW gate affected (held + inferred): {c.get('historical_positive_sw_gate_affected_count',0)} ({c.get('policy_held_ratio_percent',0)}%)",
            "",
            "3. Allocation Status Distribution",
            "---------------------------------",
        ]
        for key, value in sorted((report.get("allocation_status_counts") or {}).items()):
            lines.append(f"- {key}: {value}")
        lines += [
            "",
            "4. Source Extraction / Coverage Cross-check",
            "-------------------------------------------",
            f"Eligible Semantic Source Units: {sem.get('eligible_unit_count')}",
            f"Covered Semantic Source Units: {sem.get('covered_unit_count')}",
            f"OPEN REVIEW_NEEDED Semantic Units: {sem.get('open_review_needed_count')}",
            f"Potential Extraction Loss Units: {sem.get('potential_extraction_loss_count')}",
            f"Terminal Disposition Complete: {sem.get('terminal_disposition_complete')}",
            f"Source Conflict Count: {report.get('source_conflict_count',0)}",
            "",
            "5. Hypothesis Verdict",
            "---------------------",
            f"Verdict: {report.get('hypothesis_verdict','')}",
            f"Reason: {report.get('hypothesis_reason','')}",
            f"Current Policy: {policy.get('current_rule','')}",
            f"Primary Gate: {policy.get('primary_gate','')}",
            f"Hypothetical Main SWE.1 upper-bound if only allocation gate were relaxed: {policy.get('hypothetical_main_swe1_upper_bound_if_only_allocation_gate_relaxed',0)}",
            f"Warning: {policy.get('upper_bound_warning','')}",
            "",
            "6. Pending Requirement Evidence",
            "-------------------------------",
        ]
        pending = report.get("pending_requirements") or []
        if not pending:
            lines.append("- None")
        for item in pending:
            lines += [
                f"- {item.get('srs_id','')} | {item.get('allocation_status','')} | SWE.6={item.get('swe6_eligibility','')}",
                f"  Requirement: {item.get('requirement','')}",
                f"  Positive SW Evidence: {', '.join(item.get('positive_sw_evidence') or []) or '[NONE]'}",
                f"  Rationale: {item.get('allocation_rationale','')}",
                f"  Source Locations: {' | '.join(item.get('source_locations') or []) or '[NONE]'}",
            ]
        lines += [
            "",
            "7. OPEN Semantic Units",
            "----------------------",
        ]
        open_units = report.get("open_semantic_units") or []
        if not open_units:
            lines.append("- None")
        for item in open_units:
            lines += [
                f"- {item.get('source_semantic_unit_id','')} | {item.get('source_location','')}",
                f"  Disposition: {item.get('disposition_status','')}",
                f"  Reason: {item.get('disposition_reason','')}",
                f"  Required Resolution: {item.get('required_resolution','')}",
            ]
        lines += [
            "",
            "8. Strategy Boundary for Next Discussion",
            "----------------------------------------",
            str(report.get("strategy_guardrail") or ""),
            "Interpret inferred software allocation as REVIEW_REQUIRED until engineering approval; generated SWE.6 cases are candidate verification artifacts and execution readiness is still independently gated.",
        ]
        return "\n".join(lines).rstrip() + "\n"

    @staticmethod
    def split_situation_check(text: str, folder: Path, *, lines_per_file: int = 110, handoff_identity: dict[str, Any] | None = None) -> list[Path]:
        folder.mkdir(parents=True, exist_ok=True)
        for old in folder.glob("check_situation_*.txt"):
            old.unlink()
        lines = text.splitlines() or [""]
        line_limit = max(1, int(lines_per_file))
        ident = handoff_identity or {}

        def header(part_no: int, part_count: int) -> list[str]:
            return [
                "[Requirement Studio CHECK_SITUATION Handoff]",
                f"Version: {ident.get('requirement_studio_version','')}",
                f"Mode: {ident.get('evaluation_mode','')}",
                f"Source: {ident.get('source_file','')}",
                f"Source SHA-256: {ident.get('source_sha256','')}",
                f"Run ID: {ident.get('run_id','')}",
                f"Evidence SHA-256: {ident.get('evidence_package_sha256','')}",
                f"Part: {part_no}/{part_count}",
                "-" * 72,
            ] if ident else []

        header_len = len(header(1, 1))
        payload_limit = max(1, line_limit - header_len) if ident else line_limit
        chunks = [lines[i:i + payload_limit] for i in range(0, len(lines), payload_limit)] or [[""]]
        paths = []
        for idx, chunk in enumerate(chunks, 1):
            path = folder / f"check_situation_{idx:03d}.txt"
            _atomic_write_text(path, "\n".join(header(idx, len(chunks)) + chunk) + "\n")
            paths.append(path)
        manifest = {
            "line_limit_per_file": line_limit,
            "payload_line_limit_per_file": payload_limit,
            "total_payload_lines": len(lines),
            "part_count": len(paths),
            "handoff_identity": ident,
            "files": [{"name": p.name, "sha256": _hash_file(p), "line_count": len(p.read_text(encoding="utf-8").splitlines())} for p in paths],
        }
        _atomic_write_text(folder / "manifest.json", json.dumps(manifest, ensure_ascii=False, indent=2))
        return paths


    @staticmethod
    def _build_guide_person(review: dict[str, Any], evidence: dict[str, Any], situation: dict[str, Any]) -> dict[str, Any]:
        """Build cross-document human policy questions instead of per-SRS approval prompts.

        V0.80 still builds stable cross-document question families first.  A persistent
        HumanPolicyRegistry is applied after this deterministic candidate-question build so
        already-approved policy families disappear from the new-question list while their
        application remains auditable in an applied-policy snapshot.
        """
        canonical = review.get("canonical_requirement") if isinstance(review.get("canonical_requirement"), dict) else {}
        calc = review.get("calculated_review_evidence") if isinstance(review.get("calculated_review_evidence"), dict) else {}
        reqs = [x for x in (canonical.get("requirements") or []) if isinstance(x, dict)]
        sem_cov = canonical.get("semantic_source_unit_coverage") if isinstance(canonical.get("semantic_source_unit_coverage"), dict) else {}
        audit = calc.get("swe6_export_preservation_audit") if isinstance(calc.get("swe6_export_preservation_audit"), dict) else {}
        audit_rows = [x for x in (audit.get("rows") or []) if isinstance(x, dict)]
        inferred = [r for r in reqs if str(r.get("allocation_status") or "") == "SW_BEHAVIOR_INFERRED_REVIEW_REQUIRED"]
        cross_domain = [x for x in audit_rows if bool(x.get("cross_domain_bundle_review_required"))]
        open_records = [x for x in (sem_cov.get("open_review_needed_records") or []) if isinstance(x, dict)]
        contextual = [x for x in (canonical.get("shared_fragment_contextualization") or []) if isinstance(x, dict)]
        counts = situation.get("observed_counts") if isinstance(situation.get("observed_counts"), dict) else {}

        def examples(items: list[dict[str, Any]], key: str = "srs_id", limit: int = 8) -> list[str]:
            out = []
            for item in items:
                value = str(item.get(key) or "").strip()
                if value and value not in out:
                    out.append(value)
                if len(out) >= limit:
                    break
            return out

        questions: list[dict[str, Any]] = []
        if inferred:
            questions.append({
                "question_id": "GP-ALLOC-001",
                "policy_topic": "Inferred Controller Software Candidate Promotion",
                "question": "명시적 SW allocation 문구가 없더라도 ECU/HU/Controller가 Source에서 실행 가능한 저장·판단·송수신·제어·복구·처리 behavior를 수행하고 high-confidence non-SW 근거가 없다면, 이를 Main SWE.1/SWE.6 후보인 SW_BEHAVIOR_INFERRED_REVIEW_REQUIRED로 승격하는 일반 정책을 승인할까요?",
                "suggested_ai_position": "YES",
                "why_asked": f"이번 Run에서 {len(inferred)}개 SRS가 이 정책으로 복구되었습니다. 개별 SRS가 아니라 동일 구조의 향후 문서에도 재사용할 일반 정책인지 판단이 필요합니다.",
                "observed_examples": examples(inferred),
                "reuse_scope": "Cross-document: controller/ECU actor + literal executable behavior + no high-confidence non-SW allocation",
            })
            if int(counts.get("generated_tc_count") or 0) > 0:
                questions.append({
                    "question_id": "GP-TC-001",
                    "policy_topic": "Candidate SWE.6 Generation Before Allocation Sign-off",
                    "question": "Inferred SW 후보에 대해 사람의 최종 SW allocation 승인 전에도 Source-backed SWE.6 Candidate TC/Intent Draft 생성을 허용하되, Execution Readiness와 Official Release는 REVIEW_REQUIRED/HOLD로 유지하는 정책을 승인할까요?",
                    "suggested_ai_position": "YES",
                    "why_asked": f"이번 Run에서 {counts.get('generated_tc_count',0)}개 TC 후보가 생성되었으나 승인된 qualification evidence는 아닙니다.",
                    "observed_examples": examples(inferred),
                    "reuse_scope": "Cross-document: inferred SW candidate only; no synthetic test data; result fields blank",
                })
        if cross_domain:
            questions.append({
                "question_id": "GP-MIX-001",
                "policy_topic": "Mixed-domain Parent Requirement Handling",
                "question": "하나의 Canonical Parent에 SWE.1 Eligible fact와 SYS.5/Pending fact가 함께 있으면 Parent Canonical은 유지하고 Main SWE.1/TC에는 Eligible fact만 scope-isolation하여 표시하며, 나머지는 Allocation Annex/Review Context로 보존하는 정책을 기본으로 승인할까요? (NO는 parent SRS 자체 분할을 기본 정책으로 요구한다는 의미)",
                "suggested_ai_position": "YES",
                "why_asked": f"이번 Run에서 {len(cross_domain)}개 SRS가 cross-domain bundle review 대상으로 확인되었습니다.",
                "observed_examples": examples(cross_domain),
                "reuse_scope": "Cross-document: one Canonical parent with fact-level verification domains >=2",
            })
        if open_records:
            questions.append({
                "question_id": "GP-OPEN-001",
                "policy_topic": "OPEN Semantic Unit Disposition",
                "question": "기존 Canonical/SRS가 이미 동일 behavior를 충분히 보존하고 추가 Semantic Unit이 중복·배경·외부참조 성격이라면, 사람 승인 하에 Intentional Exclusion/External Dependency로 terminal disposition하는 일반 정책을 허용할까요? 단, 고유 normative behavior가 있으면 반드시 신규/기존 SRS 또는 Gap으로 연결해야 합니다.",
                "suggested_ai_position": "CONDITIONAL",
                "why_asked": f"이번 Run에서 {len(open_records)}개 Semantic Unit이 OPEN/non-terminal 상태입니다.",
                "observed_examples": [str(x.get("source_semantic_unit_id") or "") for x in open_records[:8]],
                "reuse_scope": "Cross-document: open Semantic Unit after duplicate/context/external-reference review",
            })
        if contextual:
            questions.append({
                "question_id": "GP-SHARED-001",
                "policy_topic": "Shared Literal Fact with Parent-context Allocation",
                "question": "동일 Source literal clause가 여러 Parent behavior에 합법적으로 적용되지만 Parent별 verification/allocation context가 다른 경우, 물리 Source Fragment는 하나로 보존하면서 parent별 contextual alias와 allocation rationale을 명시하는 방식을 허용할까요?",
                "suggested_ai_position": "YES",
                "why_asked": f"이번 Run에서 {len(contextual)}개 shared Fragment가 parent-context alias로 명시화되었습니다.",
                "observed_examples": [str(x.get("source_fact_fragment_id") or "") for x in contextual[:8]],
                "reuse_scope": "Cross-document: identical literal fact shared by multiple parents with divergent allocation/domain",
            })

        source_meta = evidence.get("source") if isinstance(evidence.get("source"), dict) else {}
        return {
            "schema_version": "1.0",
            "generated_at": datetime.now().isoformat(timespec="seconds"),
            "mode": "POLICY_CANDIDATE_BUILD",
            "auto_policy_application": False,
            "handoff_identity": {
                "requirement_studio_version": evidence.get("requirement_studio_version") or "",
                "evaluation_mode": evidence.get("evaluation_mode") or "",
                "source_file": source_meta.get("file") or "",
                "source_sha256": source_meta.get("sha256") or "",
                "run_id": evidence.get("run_dir") or "",
                "evidence_package_sha256": evidence.get("evidence_package_sha256") or "",
            },
            "question_count": len(questions),
            "questions": questions,
            "governance_note": "V0.88 candidate questions are resolved against the persistent Human Policy Registry after deterministic question-family generation.",
        }

    @staticmethod
    def _render_guide_person(report: dict[str, Any]) -> str:
        lines = [
            "Requirement Studio GUIDE_PERSON — Human Policy Questions",
            "=" * 72,
            f"Mode: {report.get('mode','')}",
            f"Candidate Policy Families: {report.get('candidate_question_count', report.get('question_count',0))}",
            f"Applied Existing Human Policies: {report.get('applied_policy_count',0)}",
            f"New Question Count: {report.get('question_count',0)}",
            "",
            "Purpose",
            "-------",
            "개별 SRS를 반복 승인하는 문서가 아니라, 여러 Source에 재사용할 수 있는 일반 Engineering Policy를 사람이 결정하기 위한 질문입니다.",
            "V0.88은 승인된 Human Policy Registry를 버전 외부 사용자 저장소에서 재사용하며, 이미 결정된 Policy family는 다시 묻지 않습니다.",
            "DEFER 또는 아직 없는 새로운 Policy family만 질문으로 남깁니다.",
            "",
            "Applied Existing Human Policies",
            "-------------------------------",
        ]
        applied = report.get("applied_policies") or []
        if not applied:
            lines.append("- None")
        for p in applied:
            lines += [
                f"- {p.get('policy_id','')} | Decision={p.get('decision','')} | {p.get('title','')}",
                f"  Scope: {p.get('reuse_scope','')}",
                f"  Notes: {p.get('notes','') or '[NONE]'}",
                f"  Examples this Run: {', '.join(p.get('observed_examples') or []) or '[NONE]'}",
            ]
        lines += ["", "New / Unresolved Policy Questions", "---------------------------------"]
        questions = report.get("questions") or []
        if not questions:
            lines += ["No new cross-document policy question was generated for this Run."]
        for idx, q in enumerate(questions, 1):
            lines += [
                f"{idx}. {q.get('question_id','')} — {q.get('policy_topic','')}",
                "-" * 72,
                f"Question: {q.get('question','')}",
                f"Why Asked: {q.get('why_asked','')}",
                f"AI Suggested Position: {q.get('suggested_ai_position','')}",
                f"Observed Examples: {', '.join(q.get('observed_examples') or []) or '[NONE]'}",
                f"Policy Reuse Scope: {q.get('reuse_scope','')}",
                "Human Decision:",
                "  [ ] YES",
                "  [ ] NO",
                "  [ ] CONDITIONAL",
                "  [ ] DEFER",
                "Conditional Rule / Notes:",
                "  ",
                "",
            ]
        lines += [
            "Governance Boundary",
            "-------------------",
            str(report.get("governance_note") or ""),
            "REVIEW_REQUIRED/HOLD 항목은 삭제하지 않고 Main Requirement Specification 및 Review View에 보존합니다.",
            "Deterministic Tool defects (parser/token/provenance/count bugs) are not delegated to GUIDE_PERSON.",
        ]
        return "\n".join(lines).rstrip() + "\n"

    @staticmethod
    def split_guide_person(text: str, folder: Path, *, lines_per_file: int = 110, handoff_identity: dict[str, Any] | None = None) -> list[Path]:
        folder.mkdir(parents=True, exist_ok=True)
        for old in folder.glob("guide_person_*.txt"):
            old.unlink()
        lines = text.splitlines() or [""]
        ident = handoff_identity or {}
        def header(part_no: int, part_count: int) -> list[str]:
            return [
                "[Requirement Studio GUIDE_PERSON Handoff]",
                f"Version: {ident.get('requirement_studio_version','')}",
                f"Mode: {ident.get('evaluation_mode','')}",
                f"Source: {ident.get('source_file','')}",
                f"Source SHA-256: {ident.get('source_sha256','')}",
                f"Run ID: {ident.get('run_id','')}",
                f"Evidence SHA-256: {ident.get('evidence_package_sha256','')}",
                f"Part: {part_no}/{part_count}",
                "-" * 72,
            ] if ident else []
        header_len = len(header(1, 1))
        payload_limit = max(1, int(lines_per_file) - header_len) if ident else max(1, int(lines_per_file))
        chunks = [lines[i:i+payload_limit] for i in range(0, len(lines), payload_limit)] or [[""]]
        paths = []
        for idx, chunk in enumerate(chunks, 1):
            path = folder / f"guide_person_{idx:03d}.txt"
            _atomic_write_text(path, "\n".join(header(idx, len(chunks)) + chunk) + "\n")
            paths.append(path)
        manifest = {
            "line_limit_per_file": int(lines_per_file),
            "payload_line_limit_per_file": payload_limit,
            "total_payload_lines": len(lines),
            "part_count": len(paths),
            "handoff_identity": ident,
            "files": [{"name": p.name, "sha256": _hash_file(p), "line_count": len(p.read_text(encoding="utf-8").splitlines())} for p in paths],
        }
        _atomic_write_text(folder / "manifest.json", json.dumps(manifest, ensure_ascii=False, indent=2))
        return paths

    @staticmethod
    def _system_prompt() -> str:
        return (
            "You are an independent automotive requirements and verification review model. "
            "Evaluate only the supplied Requirement Studio evidence. Do not infer missing DB values, signals, thresholds, "
            "test environments, allocation, or pass/fail results. Correct deferral is not a defect. Distinguish Tool Quality "
            "defects from Artifact Readiness limitations and Release Governance. In unified mode, review Source -> Canonical -> "
            "SWE.1 -> allocation -> Child Intent -> SWE.6 -> actual artifacts as one trace graph and report CROSS_OUTPUT issues when appropriate. "
            "Never treat an artifact that was not selected/generated as a defect merely because it is absent. "
            "Exact fragment ownership is evaluated over the union of all fragment IDs linked to the same SRS; when owned_fragment_fact_token_sources proves a token is owned, do not re-flag that token merely because it is split across multiple fragments. "
            "If an eligible Semantic Unit is explicitly surfaced as REVIEW_NEEDED with review_state=OPEN and terminal_disposition=false, treat that as an Artifact Readiness/Release Governance item rather than a Tool Quality defect unless the tool silently dropped or falsely closed it. "
            "V0.88 retains the practical promotion policy and persistent Human Policy Registry: Source-backed controller/ECU executable behavior may be classified as SW_BEHAVIOR_INFERRED_REVIEW_REQUIRED even without a literal SW keyword. Treat that as an approved review-required modeling policy, not explicit allocation evidence and not a Tool Quality defect by itself; Official Release still requires allocation review. "
            "V0.88 retains the rule that keeps every Canonical Engineering Requirement visible in the Engineering Main Specification and keeps REVIEW_REQUIRED/HOLD items visible in both Main and Review Views. System-level items may be SYS.1/SYS.5 candidates even when they are not SWE.1/SWE.6 items. Do not flag their mere Main-view presence as a Tool Quality defect; judge whether the review status is explicit and whether release approval is kept separate. Human Policy reuse governs processing policy, not per-SRS official allocation. "
            "Respect deterministic semantic_parent_containment_review and fragment_parent_containment_review records: a Source unit/fragment removed from a weaker unrelated parent by a dominant-ownership rule is not missing merely because another materially matching SRS still owns it. "
            "In V0.88 E2E Source/Fact Traceability, Fact-to-Test-Object coverage is exact-fragment bounded: a SWE.6 TC is linked only when that Source Fact Fragment is present in the finalized TC scope. EXCLUDED_FROM_SWE6_SCOPE is an explicit non-coverage state, not a missing TC defect. "
            "MODE_TRANSITION_ALLOCATION_INCONSISTENCY is review-only and precision-bounded by engineering transition family. If review_closed=true from the per-finding Review Decision Registry, do not re-open it as an unreviewed release issue unless contradictory new evidence is supplied. "
            "When swe6_scope_consistency_status=DEFERRED_NO_ELIGIBLE_FACT_SCOPE, treat the parent defer as an explicit consistency safeguard rather than demanding a TC from zero eligible fact scope. "
            "Do not see or assume any other reviewer result. Return only the required JSON object."
        )

    @staticmethod
    def _user_prompt(evidence: dict[str, Any]) -> str:
        mode = str(evidence.get("evaluation_mode") or "unified")
        if mode == "unified":
            scope = (
                "Perform a Unified End-to-End review. Check Source fidelity, Semantic Unit/Fact Fragment ownership, Canonical/SRS preservation, "
                "SWE.1 Word/Excel preservation, explicit/inferred SW allocation basis, Child Intent traceability, SWE.6 eligibility/TC/deferred governance, "
                "cross-output consistency, E2E Requirements preservation, all sheets/rows of the E2E Evaluation workbook, Structural vs Semantic gates, unsupported generation, regression governance, and execution readiness. "
                "Use category CROSS_OUTPUT for information that is correct in one output but lost/leaked/contradicted in another."
            )
        elif mode == "swe6":
            scope = "Review SWE.6 eligibility, test intent preservation, deferred governance, actual XLSX contract, traceability, and execution readiness."
        else:
            scope = "Review SWE.1 Source fidelity, allocation, atomicity, review-annex usability, Word/Excel preservation, and traceability."
        return (
            scope + " Findings must cite concrete Source/SRS/Fragment/TC/artifact references from the evidence. "
            "Do not create a finding solely because an unselected artifact is absent. Use stable finding_key values.\n"
            "Return exactly one JSON object matching this schema; do not wrap it in Markdown.\n"
            "OUTPUT JSON SCHEMA:\n" + json.dumps(EVALUATION_SCHEMA, ensure_ascii=False) +
            "\n\nEVIDENCE PACKAGE:\n" + json.dumps(evidence, ensure_ascii=False)
        )

    def _hchat_call(self, backend: str, evidence: dict[str, Any]) -> tuple[dict[str, Any], str]:
        cfg = copy.deepcopy(self.provider_config.get("hchat", {}))
        if backend == "claude":
            # Evaluation reports can be materially longer than generation snippets; keep the
            # evaluator response budget large enough without changing the user's normal model config.
            cfg["claude_max_tokens"] = max(int(cfg.get("claude_max_tokens") or 0), 12000)
        provider = HChatProvider(self.project_root, cfg, backend=backend)
        text = provider.generate(
            self._user_prompt(evidence),
            system_message=self._system_prompt() + " The JSON must match the published Requirement Studio Evaluation Schema.",
            timeout_seconds=int(self.config.get("timeout_seconds") or 240),
        )
        data = _safe_json_extract(text)
        data.setdefault("provider", f"H-Chat/{backend}")
        data.setdefault("model", provider.model)
        data.setdefault("schema_version", EVALUATION_SCHEMA_VERSION)
        _schema_validate_minimum(data)
        return data, provider.model

    def _direct_openai(self, evidence: dict[str, Any], cfg: dict[str, Any]) -> tuple[dict[str, Any], str]:
        key = os.environ.get(str(cfg.get("api_key_env") or "OPENAI_API_KEY"), "").strip()
        if not key:
            raise RuntimeError("OPENAI_API_KEY가 없습니다.")
        model = str(cfg.get("direct_model") or "gpt-6.1-sol")
        url = str(cfg.get("direct_base_url") or "https://api.openai.com/v1").rstrip("/") + "/responses"
        payload = {
            "model": model,
            "input": [
                {"role": "system", "content": [{"type": "input_text", "text": self._system_prompt()}]},
                {"role": "user", "content": [{"type": "input_text", "text": self._user_prompt(evidence)}]},
            ],
            "text": {"format": {"type": "json_schema", "name": "requirement_studio_evaluation", "strict": True, "schema": EVALUATION_SCHEMA}},
        }
        r = requests.post(url, headers={"Authorization": f"Bearer {key}", "Content-Type": "application/json"}, json=payload, timeout=int(self.config.get("timeout_seconds") or 240))
        r.raise_for_status()
        obj = r.json()
        text = obj.get("output_text")
        if not text:
            pieces = []
            for item in obj.get("output") or []:
                for part in (item or {}).get("content") or []:
                    if isinstance(part, dict) and part.get("text"):
                        pieces.append(str(part["text"]))
            text = "\n".join(pieces)
        data = _safe_json_extract(str(text or ""))
        data["provider"] = "OpenAI"
        data["model"] = model
        data["schema_version"] = EVALUATION_SCHEMA_VERSION
        _schema_validate_minimum(data)
        return data, model

    def _direct_gemini(self, evidence: dict[str, Any], cfg: dict[str, Any]) -> tuple[dict[str, Any], str]:
        key = os.environ.get(str(cfg.get("api_key_env") or "GEMINI_API_KEY"), "").strip()
        if not key:
            raise RuntimeError("GEMINI_API_KEY가 없습니다.")
        model = str(cfg.get("direct_model") or "gemini-3.8-flash")
        base = str(cfg.get("direct_base_url") or "https://generativelanguage.googleapis.com/v1beta").rstrip("/")
        # Current Gemini structured-output API uses the Interactions endpoint with a top-level
        # response_format. Keep the full evaluator instruction in one input so the direct adapter
        # and H-Chat adapter receive semantically identical evidence/instructions.
        url = f"{base}/interactions"
        payload = {
            "model": model,
            "input": self._system_prompt() + "\n\n" + self._user_prompt(evidence),
            "response_format": {
                "type": "text",
                "mime_type": "application/json",
                "schema": EVALUATION_SCHEMA,
            },
        }
        r = requests.post(url, headers={"x-goog-api-key": key, "Content-Type": "application/json"}, json=payload, timeout=int(self.config.get("timeout_seconds") or 240))
        r.raise_for_status()
        obj = r.json()
        text = str(obj.get("output_text") or "")
        if not text:
            pieces = []
            for item in obj.get("output") or []:
                if isinstance(item, dict) and item.get("text"):
                    pieces.append(str(item.get("text")))
                for part in (item or {}).get("content") or [] if isinstance(item, dict) else []:
                    if isinstance(part, dict) and part.get("text"):
                        pieces.append(str(part.get("text")))
            text = "\n".join(pieces)
        data = _safe_json_extract(text)
        data["provider"] = "Google Gemini"
        data["model"] = model
        data["schema_version"] = EVALUATION_SCHEMA_VERSION
        _schema_validate_minimum(data)
        return data, model

    def _direct_claude(self, evidence: dict[str, Any], cfg: dict[str, Any]) -> tuple[dict[str, Any], str]:
        key = os.environ.get(str(cfg.get("api_key_env") or "ANTHROPIC_API_KEY"), "").strip()
        if not key:
            raise RuntimeError("ANTHROPIC_API_KEY가 없습니다.")
        model = str(cfg.get("direct_model") or "claude-sonnet-5-5")
        url = str(cfg.get("direct_base_url") or "https://api.anthropic.com").rstrip("/") + "/v1/messages"
        payload = {
            "model": model,
            "max_tokens": 12000,
            "system": self._system_prompt(),
            "messages": [{"role": "user", "content": self._user_prompt(evidence)}],
            "output_config": {"format": {"type": "json_schema", "schema": EVALUATION_SCHEMA}},
        }
        headers = {
            "x-api-key": key,
            "anthropic-version": str(cfg.get("anthropic_version") or "2023-06-01"),
            "content-type": "application/json",
        }
        r = requests.post(url, headers=headers, json=payload, timeout=int(self.config.get("timeout_seconds") or 240))
        r.raise_for_status()
        obj = r.json()
        text = "\n".join(str(x.get("text")) for x in obj.get("content") or [] if isinstance(x, dict) and x.get("type") == "text" and x.get("text"))
        data = _safe_json_extract(text)
        data["provider"] = "Anthropic Claude"
        data["model"] = model
        data["schema_version"] = EVALUATION_SCHEMA_VERSION
        _schema_validate_minimum(data)
        return data, model

    def _call_one(self, name: str, evidence: dict[str, Any]) -> tuple[dict[str, Any], str]:
        transport = str(self.config.get("transport") or "hchat").lower()
        cfg = (self.config.get("providers") or {}).get(name) or {}
        if not bool(cfg.get("enabled", True)):
            raise RuntimeError(f"{name} evaluator disabled")
        if transport == "hchat":
            return self._hchat_call(name, evidence)
        if transport == "direct_api":
            if name == "gpt":
                return self._direct_openai(evidence, cfg)
            if name == "gemini":
                return self._direct_gemini(evidence, cfg)
            if name == "claude":
                return self._direct_claude(evidence, cfg)
        raise RuntimeError(f"지원하지 않는 evaluation transport: {transport}")

    @staticmethod
    def _render_provider_txt(data: dict[str, Any]) -> str:
        lines = [
            f"Requirement Studio Automated Evaluation - {data.get('provider','')}",
            "=" * 72,
            f"Model: {data.get('model','')}",
            f"Verdict: {data.get('verdict','')}",
            f"Tool Quality Gate: {data.get('tool_quality_gate','')}",
            f"Artifact Readiness Gate: {data.get('artifact_readiness_gate','')}",
            f"Official Release: {data.get('official_release','')}",
            "",
            "[Executive Summary]",
        ]
        lines += [f"- {x}" for x in data.get("executive_summary") or []]
        lines += ["", "[Positive Checks]"]
        lines += [f"- {x}" for x in data.get("positive_checks") or []]
        lines += ["", "[Findings]"]
        for i, f in enumerate(data.get("findings") or [], 1):
            lines += [
                f"{i}. {f.get('finding_key')} | {f.get('severity')} | {f.get('gate_scope')}",
                f"   Title: {f.get('title','')}",
                f"   Category: {f.get('category','')}",
                f"   Affected: {', '.join(f.get('affected_ids') or [])}",
                f"   Source: {' | '.join(f.get('source_references') or [])}",
                f"   Problem: {f.get('problem','')}",
                f"   Recommendation: {f.get('recommendation','')}",
                f"   Confidence: {f.get('confidence','')}",
            ]
            for e in f.get("evidence") or []:
                lines.append(f"   Evidence: {e}")
        lines += ["", "[Release Gates]"]
        for g in data.get("release_gates") or []:
            lines.append(f"- {g.get('gate')}: {g.get('result')} | {g.get('evidence')}")
        lines += ["", "[Next Version Recommendations]"]
        for r in data.get("next_version_recommendations") or []:
            lines.append(f"- {r.get('priority')} {r.get('title')}: {r.get('detail')}")
        if data.get("limitations"):
            lines += ["", "[Limitations]"] + [f"- {x}" for x in data.get("limitations") or []]
        return "\n".join(lines).rstrip() + "\n"

    @staticmethod
    def _tokens(value: str) -> set[str]:
        stop = {"requirement", "source", "review", "issue", "missing", "required", "swe", "fact", "fragment"}
        return {x.lower() for x in re.findall(r"[A-Za-z0-9가-힣_.+-]+", value or "") if len(x) > 1 and x.lower() not in stop}

    @classmethod
    def _finding_concepts(cls, finding: dict[str, Any]) -> set[str]:
        """Return stable failure concepts for cross-reviewer semantic clustering.

        V0.81 intentionally does not depend on provider-specific category wording.  Concepts are
        conservative, evidence-facing tags used only when findings already overlap on concrete IDs.
        """
        text = _normalize_space(" ".join([
            str(finding.get("title") or ""), str(finding.get("problem") or ""),
            str(finding.get("recommendation") or ""), " ".join(str(x) for x in (finding.get("evidence") or [])),
        ])).lower()
        concepts: set[str] = set()
        rules = {
            "ELIGIBILITY": r"eligible|eligibility|적격|적격성",
            "TC_GENERATION_GAP": r"(?:zero|no|missing|absent|미생성|생성되지|없(?:음|다)|누락).{0,45}(?:tc|test case|테스트 ?케이스)|(?:tc|test case|테스트 ?케이스).{0,45}(?:zero|no|missing|absent|미생성|생성되지|없(?:음|다)|누락)",
            "OWNERSHIP_LEAK": r"unrelated|leak|contamination|wrong parent|잘못.*(?:연결|부착)|오염|침범|소유권",
            "PROVENANCE": r"provenance|traceability|trace|출처|근거",
            "ALLOCATION": r"allocation|할당",
            "OPEN_SEMANTIC": r"open.{0,30}semantic|semantic.{0,30}open|미종결.{0,30}(?:semantic|시맨틱)|review_needed",
            "REGRESSION_BASELINE": r"gold source|baseline|regression|회귀",
            "MIXED_DOMAIN": r"mixed[- ]domain|cross[- ]domain|혼합.*도메인|fact[- ]level.*domain",
            "EXECUTION_READINESS": r"execution readiness|intent draft|independent coverage|실행.*준비|실행가능",
            "DEFERRED_SCOPE": r"defer|deferred|보류|pending",
            "GAP_VISIBILITY": r"document[- ]wide|open_issues|gap.*(?:excel|sheet)|엑셀.*gap",
        }
        for tag, pattern in rules.items():
            if re.search(pattern, text, re.I | re.S):
                concepts.add(tag)
        return concepts

    @classmethod
    def _same_finding(cls, a: dict[str, Any], b: dict[str, Any]) -> bool:
        if str(a.get("finding_key") or "").strip() and str(a.get("finding_key") or "").strip() == str(b.get("finding_key") or "").strip():
            return True
        aa, bb = set(a.get("affected_ids") or []), set(b.get("affected_ids") or [])
        shared_ids = aa & bb
        same_category = str(a.get("category") or "").lower() == str(b.get("category") or "").lower()
        if same_category and shared_ids:
            return True

        # V0.81 cross-category semantic dedup: reviewers often describe the same invariant failure
        # as SWE6, CROSS_OUTPUT, or TRACEABILITY.  Merge only when they share a concrete affected
        # ID *and* the same specific failure concept. This avoids merging unrelated review items
        # merely because both mention the same SRS.
        if shared_ids:
            ca, cb = cls._finding_concepts(a), cls._finding_concepts(b)
            shared_concepts = ca & cb
            specific = {
                "TC_GENERATION_GAP", "OWNERSHIP_LEAK", "OPEN_SEMANTIC", "REGRESSION_BASELINE",
                "MIXED_DOMAIN", "EXECUTION_READINESS", "GAP_VISIBILITY",
            }
            if shared_concepts & specific:
                # Eligibility+TC is the strongest signature for the V0.80 Eligible/zero-TC defect.
                if "TC_GENERATION_GAP" in shared_concepts and "ELIGIBILITY" in ca and "ELIGIBILITY" in cb:
                    return True
                if any(x in shared_concepts for x in specific - {"TC_GENERATION_GAP"}):
                    return True

        if not same_category:
            return False
        ta = cls._tokens(str(a.get("title") or "") + " " + str(a.get("problem") or ""))
        tb = cls._tokens(str(b.get("title") or "") + " " + str(b.get("problem") or ""))
        if not ta or not tb:
            return False
        return len(ta & tb) / max(1, min(len(ta), len(tb))) >= 0.45

    @classmethod
    def _consensus(cls, provider_results: dict[str, dict[str, Any]]) -> dict[str, Any]:
        clusters: list[dict[str, Any]] = []
        severity_rank = {"ADVISORY": 1, "RELEASE_REVIEW": 2, "RELEASE_BLOCKING": 3}
        for provider, result in provider_results.items():
            for finding in result.get("findings") or []:
                target = None
                for cluster in clusters:
                    if any(cls._same_finding(finding, member["finding"]) for member in cluster["members"]):
                        target = cluster
                        break
                if target is None:
                    target = {"members": []}
                    clusters.append(target)
                target["members"].append({"provider": provider, "finding": finding})

        consensus_findings = []
        for idx, cluster in enumerate(clusters, 1):
            members = cluster["members"]
            providers = list(dict.fromkeys(x["provider"] for x in members))
            primary = max((x["finding"] for x in members), key=lambda f: severity_rank.get(str(f.get("severity")), 0))
            count = len(providers)
            consensus_findings.append({
                "consensus_id": f"CF-{idx:03d}",
                "strength": "STRONG_CONSENSUS" if count >= 3 else ("CONSENSUS" if count == 2 else "SINGLE_REVIEWER"),
                "reviewer_count": count,
                "member_count": len(members),
                "reviewers": providers,
                "severity": primary.get("severity"),
                "gate_scope": primary.get("gate_scope"),
                "category": primary.get("category"),
                "title": primary.get("title"),
                "affected_ids": list(dict.fromkeys(y for x in members for y in (x["finding"].get("affected_ids") or []))),
                "source_references": list(dict.fromkeys(y for x in members for y in (x["finding"].get("source_references") or []))),
                "evidence": list(dict.fromkeys(y for x in members for y in (x["finding"].get("evidence") or [])))[:12],
                "problem": primary.get("problem"),
                "recommendation": primary.get("recommendation"),
            })

        tool_values = [str(r.get("tool_quality_gate") or "NOT_EVALUATED") for r in provider_results.values()]
        readiness_values = [str(r.get("artifact_readiness_gate") or "NOT_EVALUATED") for r in provider_results.values()]
        release_values = [str(r.get("official_release") or "NOT_EVALUATED") for r in provider_results.values()]
        if "FAIL" in tool_values:
            tool = "FAIL"
        elif "PASS_WITH_REVIEW_ITEMS" in tool_values:
            tool = "PASS_WITH_REVIEW_ITEMS"
        elif tool_values and all(x == "PASS" for x in tool_values):
            tool = "PASS"
        else:
            tool = "NOT_EVALUATED"
        if "FAIL" in readiness_values:
            readiness = "FAIL"
        elif "REVIEW_REQUIRED" in readiness_values:
            readiness = "REVIEW_REQUIRED"
        elif readiness_values and all(x == "PASS" for x in readiness_values):
            readiness = "PASS"
        else:
            readiness = "NOT_EVALUATED"
        reviewer_count = len(provider_results)
        completeness = "COMPLETE" if reviewer_count == 3 else ("PARTIAL" if reviewer_count else "NOT_EVALUATED")
        # Three-model evaluation is a review contract, not majority voting. A partial reviewer set may
        # still provide useful findings, but it cannot independently satisfy an Official Release gate.
        release = "HOLD" if reviewer_count < 3 or "HOLD" in release_values or tool != "PASS" or readiness != "PASS" else ("PASS" if release_values and all(x == "PASS" for x in release_values) else "NOT_EVALUATED")
        return {
            "consensus_schema_version": "1.0",
            "reviewer_count": reviewer_count,
            "evaluation_completeness": completeness,
            "tool_quality_gate": tool,
            "artifact_readiness_gate": readiness,
            "official_release": release,
            "consensus_findings": consensus_findings,
        }

    @staticmethod
    def _apply_unified_artifact_contract(consensus: dict[str, Any], evidence: dict[str, Any]) -> dict[str, Any]:
        contract = evidence.get("unified_artifact_contract") if isinstance(evidence.get("unified_artifact_contract"), dict) else {}
        if str(evidence.get("evaluation_mode") or "").lower() != "unified" or str(contract.get("status") or "") != "FAIL":
            return consensus
        missing = [str(x) for x in (contract.get("missing_artifacts") or []) if str(x)]
        finding = {
            "consensus_id": "CF-CONTRACT-001",
            "strength": "DETERMINISTIC",
            "reviewer_count": 0,
            "member_count": 0,
            "reviewers": ["deterministic_contract"],
            "severity": "RELEASE_BLOCKING",
            "gate_scope": "TOOL_QUALITY",
            "category": "UNIFIED_ARTIFACT_CONTRACT",
            "title": "Required legacy/integrated artifacts are missing from Unified AI evaluation evidence",
            "affected_ids": missing,
            "source_references": [],
            "evidence": [f"Missing required Unified artifacts: {', '.join(missing)}"],
            "problem": "Recommended Unified AI evaluation cannot claim end-to-end cross-output coverage when one or more required engineer-facing artifacts are absent.",
            "recommendation": "Regenerate/copy all required SWE.1 Word/Excel, SWE.6 Excel, Integrated Requirements Excel, and Integrated Tests Excel before release evaluation.",
        }
        out = dict(consensus)
        findings = [x for x in (out.get("consensus_findings") or []) if isinstance(x, dict)]
        if not any(str(x.get("category") or "") == "UNIFIED_ARTIFACT_CONTRACT" for x in findings):
            findings.insert(0, finding)
        out["consensus_findings"] = findings
        out["tool_quality_gate"] = "FAIL"
        out["official_release"] = "HOLD"
        return out

    @staticmethod
    def _render_consensus_txt(consensus: dict[str, Any]) -> str:
        lines = [
            "Requirement Studio Multi-Model Consensus",
            "=" * 72,
            f"Reviewer Count: {consensus.get('reviewer_count',0)}",
            f"Evaluation Completeness: {consensus.get('evaluation_completeness','')}",
            f"Tool Quality Gate: {consensus.get('tool_quality_gate','')}",
            f"Artifact Readiness Gate: {consensus.get('artifact_readiness_gate','')}",
            f"Official Release: {consensus.get('official_release','')}",
            "",
            "[Consensus Findings]",
        ]
        for f in consensus.get("consensus_findings") or []:
            lines += [
                f"{f.get('consensus_id')} | {f.get('strength')} | {f.get('severity')} | {f.get('gate_scope')}",
                f"Title: {f.get('title','')}",
                f"Reviewers: {', '.join(f.get('reviewers') or [])}",
                f"Affected: {', '.join(f.get('affected_ids') or [])}",
                f"Problem: {f.get('problem','')}",
                f"Recommendation: {f.get('recommendation','')}",
                "",
            ]
        return "\n".join(lines).rstrip() + "\n"

    @classmethod
    def _final_result(cls, provider_results: dict[str, dict[str, Any]], consensus: dict[str, Any], failures: dict[str, str], evidence_path: Path, evidence: dict[str, Any] | None = None) -> dict[str, Any]:
        priorities = []
        seen = set()
        for result in provider_results.values():
            for item in result.get("next_version_recommendations") or []:
                key = (_normalize_space(item.get("priority")), _normalize_space(item.get("title")).lower())
                if key not in seen:
                    seen.add(key)
                    priorities.append(item)
        priorities.sort(key=lambda x: {"P0": 0, "P1": 1, "P2": 2}.get(str(x.get("priority")), 9))
        provider_finding_counts = {name: len(result.get("findings") or []) for name, result in provider_results.items()}
        total_occurrences = sum(provider_finding_counts.values())
        mapped_occurrences = sum(int(f.get("member_count") or f.get("reviewer_count") or 0) for f in (consensus.get("consensus_findings") or []))
        completeness_audit = {
            "provider_finding_counts": provider_finding_counts,
            "provider_finding_occurrences": total_occurrences,
            "unique_consensus_findings": len(consensus.get("consensus_findings") or []),
            "mapped_finding_occurrences": mapped_occurrences,
            "unmapped_finding_occurrences": max(0, total_occurrences - mapped_occurrences),
            "passed": total_occurrences == mapped_occurrences,
        }
        evidence = evidence or {}
        source_meta = evidence.get("source") if isinstance(evidence.get("source"), dict) else {}
        actual_artifacts = evidence.get("actual_artifacts") if isinstance(evidence.get("actual_artifacts"), dict) else {}
        handoff_identity = {
            "requirement_studio_version": evidence.get("requirement_studio_version") or "",
            "evaluation_mode": evidence.get("evaluation_mode") or "",
            "source_file": source_meta.get("file") or "",
            "source_sha256": source_meta.get("sha256") or "",
            "run_id": evidence.get("run_dir") or "",
            "artifacts_evaluated": [f"{key}:{(value or {}).get('artifact','')}" for key, value in sorted(actual_artifacts.items()) if isinstance(value, dict)],
            "evidence_package_sha256": evidence.get("evidence_package_sha256") or "",
        }
        return {
            "final_result_schema_version": "1.2",
            "generated_at": datetime.now().isoformat(timespec="seconds"),
            "evidence_file": evidence_path.name,
            "handoff_identity": handoff_identity,
            "providers_completed": sorted(provider_results),
            "provider_failures": failures,
            "evaluation_completeness": consensus.get("evaluation_completeness"),
            "tool_quality_gate": consensus.get("tool_quality_gate"),
            "artifact_readiness_gate": consensus.get("artifact_readiness_gate"),
            "official_release": consensus.get("official_release"),
            "consensus_findings": consensus.get("consensus_findings") or [],
            "next_version_recommendations": priorities,
            "completeness_audit": completeness_audit,
            "provider_verdicts": {
                name: {
                    "verdict": result.get("verdict"),
                    "tool_quality_gate": result.get("tool_quality_gate"),
                    "artifact_readiness_gate": result.get("artifact_readiness_gate"),
                    "official_release": result.get("official_release"),
                }
                for name, result in provider_results.items()
            },
        }

    @staticmethod
    def _render_final_txt(final: dict[str, Any]) -> str:
        ident = final.get("handoff_identity") or {}
        lines = [
            "Requirement Studio Final Automated Evaluation",
            "=" * 72,
            f"Requirement Studio Version: {ident.get('requirement_studio_version','')}",
            f"Evaluation Mode: {ident.get('evaluation_mode','')}",
            f"Source File: {ident.get('source_file','')}",
            f"Source SHA-256: {ident.get('source_sha256','')}",
            f"Run ID: {ident.get('run_id','')}",
            f"Artifacts Evaluated: {', '.join(ident.get('artifacts_evaluated') or [])}",
            f"Evidence Package SHA-256: {ident.get('evidence_package_sha256','')}",
            f"Generated At: {final.get('generated_at','')}",
            f"Evidence: {final.get('evidence_file','')}",
            f"Providers Completed: {', '.join(final.get('providers_completed') or [])}",
            f"Evaluation Completeness: {final.get('evaluation_completeness','')}",
            f"Tool Quality Gate: {final.get('tool_quality_gate','')}",
            f"Artifact Readiness Gate: {final.get('artifact_readiness_gate','')}",
            f"Official Release: {final.get('official_release','')}",
            "",
            "1. Provider Verdicts",
            "--------------------",
        ]
        for name, v in (final.get("provider_verdicts") or {}).items():
            lines.append(f"- {name}: {v.get('verdict')} / Tool={v.get('tool_quality_gate')} / Artifact={v.get('artifact_readiness_gate')} / Release={v.get('official_release')}")
        if final.get("provider_failures"):
            lines += ["", "2. Provider Failures", "--------------------"]
            for name, err in final.get("provider_failures", {}).items():
                lines.append(f"- {name}: {err}")
        lines += ["", "3. Consensus Findings", "---------------------"]
        for i, f in enumerate(final.get("consensus_findings") or [], 1):
            lines += [
                f"{i}) {f.get('consensus_id')} / {f.get('strength')} / {f.get('severity')} / {f.get('gate_scope')}",
                f"   Title: {f.get('title','')}",
                f"   Reviewers: {', '.join(f.get('reviewers') or [])}",
                f"   Affected IDs: {', '.join(f.get('affected_ids') or [])}",
                f"   Source References: {' | '.join(f.get('source_references') or [])}",
                f"   Problem: {f.get('problem','')}",
                f"   Recommendation: {f.get('recommendation','')}",
            ]
            for ev in f.get("evidence") or []:
                lines.append(f"   Evidence: {ev}")
            lines.append("")
        lines += ["4. Next Version Recommendations", "--------------------------------"]
        for rec in final.get("next_version_recommendations") or []:
            lines.append(f"- {rec.get('priority')} {rec.get('title')}: {rec.get('detail')}")
        audit = final.get("completeness_audit") or {}
        lines += [
            "",
            "5. Final Result Completeness Audit",
            "---------------------------------",
            "Provider Findings: " + " / ".join(f"{k}={v}" for k, v in (audit.get("provider_finding_counts") or {}).items()),
            f"Provider Finding Occurrences: {audit.get('provider_finding_occurrences', 0)}",
            f"Unique Consensus Findings: {audit.get('unique_consensus_findings', 0)}",
            f"Mapped Finding Occurrences: {audit.get('mapped_finding_occurrences', 0)}",
            f"Unmapped Finding Occurrences: {audit.get('unmapped_finding_occurrences', 0)}",
            f"Completeness: {'PASS' if audit.get('passed') else 'FAIL'}",
            "",
            "6. Interpretation Rule",
            "----------------------",
            "- 3/3 reviewer agreement = STRONG_CONSENSUS.",
            "- 2/3 reviewer agreement = CONSENSUS.",
            "- 1/3 only = SINGLE_REVIEWER and remains a human-review candidate.",
            "- Final gates are conservative; a reviewer FAIL/HOLD is never hidden by majority voting.",
            "- Source-insufficient Deferred/External Dependency does not by itself mean Tool Quality failure.",
        ]
        return "\n".join(lines).rstrip() + "\n"

    @staticmethod
    def split_final_result(text: str, folder: Path, *, lines_per_file: int = 110, handoff_identity: dict[str, Any] | None = None) -> list[Path]:
        folder.mkdir(parents=True, exist_ok=True)
        for old in folder.glob("FINAL_RESULT_*.txt"):
            old.unlink()
        lines = text.splitlines()
        if not lines:
            lines = [""]
        line_limit = max(1, int(lines_per_file))
        ident = handoff_identity or {}

        def header(part_no: int, part_count: int) -> list[str]:
            if not ident:
                return []
            return [
                "[Requirement Studio FINAL_RESULT Handoff]",
                f"Version: {ident.get('requirement_studio_version','')}",
                f"Mode: {ident.get('evaluation_mode','')}",
                f"Source: {ident.get('source_file','')}",
                f"Source SHA-256: {ident.get('source_sha256','')}",
                f"Run ID: {ident.get('run_id','')}",
                f"Evidence SHA-256: {ident.get('evidence_package_sha256','')}",
                f"Part: {part_no}/{part_count}",
                "-" * 72,
            ]

        header_len = len(header(1, 1))
        payload_limit = max(1, line_limit - header_len) if ident else line_limit
        chunks = [lines[i:i + payload_limit] for i in range(0, len(lines), payload_limit)] or [[""]]
        paths = []
        for idx, chunk in enumerate(chunks, 1):
            part = header(idx, len(chunks)) + chunk
            path = folder / f"FINAL_RESULT_{idx:03d}.txt"
            _atomic_write_text(path, "\n".join(part) + "\n")
            paths.append(path)
        manifest = {
            "line_limit_per_file": line_limit,
            "payload_line_limit_per_file": payload_limit,
            "total_payload_lines": len(lines),
            "part_count": len(paths),
            "handoff_identity": ident,
            "files": [{"name": p.name, "sha256": _hash_file(p), "line_count": len(p.read_text(encoding='utf-8').splitlines())} for p in paths],
        }
        _atomic_write_text(folder / "manifest.json", json.dumps(manifest, ensure_ascii=False, indent=2))
        return paths

    def _write_status(self, output_dir: Path, payload: dict[str, Any]) -> None:
        body = dict(payload)
        body.setdefault("updated_at", datetime.now().isoformat(timespec="seconds"))
        _atomic_write_text(output_dir / "00_EVALUATION_STATUS.json", json.dumps(body, ensure_ascii=False, indent=2))
        lines = [
            "Requirement Studio Automated Evaluation Status",
            "=" * 72,
            f"Status: {body.get('status', '')}",
            f"Stage: {body.get('stage', '')}",
            f"Progress: {int(body.get('progress_percent') or 0)}%",
            f"Updated At: {body.get('updated_at', '')}",
            f"Providers Completed: {', '.join(body.get('providers_completed') or [])}",
        ]
        provider_status = body.get("provider_status") or {}
        if provider_status:
            lines.append("Provider Status: " + " / ".join(f"{k}={v}" for k, v in provider_status.items()))
        if body.get("tool_quality_gate"):
            lines.append(f"Tool Quality Gate: {body.get('tool_quality_gate')}")
        if body.get("artifact_readiness_gate"):
            lines.append(f"Artifact Readiness Gate: {body.get('artifact_readiness_gate')}")
        if body.get("official_release"):
            lines.append(f"Official Release: {body.get('official_release')}")
        failures = body.get("provider_failures") or {}
        if failures:
            lines.append("Provider Failures: " + " / ".join(f"{k}: {v}" for k, v in failures.items()))
        if body.get("message"):
            lines.append(f"Message: {body.get('message')}")
        _atomic_write_text(output_dir / "00_EVALUATION_STATUS.txt", "\n".join(lines).rstrip() + "\n")

    @staticmethod
    def _validate_output_artifacts(output_dir: Path, results: dict[str, dict[str, Any]]) -> dict[str, Any]:
        required = [
            "00_EVALUATION_STATUS.json",
            "00_EVALUATION_STATUS.txt",
            "01_GPT_EVALUATION.txt",
            "02_GEMINI_EVALUATION.txt",
            "03_CLAUDE_EVALUATION.txt",
            "04_EVALUATION_CONSENSUS.json",
            "04_EVALUATION_CONSENSUS.txt",
            "05_EVALUATION_RESULT.json",
            "06_EVALUATION_EVIDENCE.json",
            "AUTOMATIC_EVALUATION_MANIFEST.json",
            "final_result/manifest.json",
            "check_situation/manifest.json",
            "guide_person/manifest.json",
            "guide_person/guide_person_questions.json",
            "guide_person/applied_policy_snapshot.json",
            "guide_person/applied_policy_snapshot.txt",
        ]
        for name in results:
            prefix = {"gpt": "01_GPT_EVALUATION", "gemini": "02_GEMINI_EVALUATION", "claude": "03_CLAUDE_EVALUATION"}[name]
            required.append(prefix + ".json")
        final_parts = sorted((output_dir / "final_result").glob("FINAL_RESULT_*.txt"))
        if not final_parts:
            raise RuntimeError("final_result TXT가 생성되지 않았습니다.")
        required.extend(str(p.relative_to(output_dir)).replace("\\", "/") for p in final_parts)
        situation_parts = sorted((output_dir / "check_situation").glob("check_situation_*.txt"))
        if not situation_parts:
            raise RuntimeError("check_situation TXT가 생성되지 않았습니다.")
        required.extend(str(p.relative_to(output_dir)).replace("\\", "/") for p in situation_parts)
        guide_parts = sorted((output_dir / "guide_person").glob("guide_person_*.txt"))
        if not guide_parts:
            raise RuntimeError("guide_person TXT가 생성되지 않았습니다.")
        required.extend(str(p.relative_to(output_dir)).replace("\\", "/") for p in guide_parts)
        missing = []
        zero = []
        for rel in required:
            path = output_dir / rel
            if not path.is_file():
                missing.append(rel)
            elif path.stat().st_size <= 0:
                zero.append(rel)
        if missing or zero:
            raise RuntimeError(f"자동평가 결과 파일 검증 실패 · missing={missing} · zero_size={zero}")
        return {
            "verified": True,
            "file_count": sum(1 for p in output_dir.rglob("*") if p.is_file()),
            "required_file_count": len(required),
            "final_result_part_count": len(final_parts),
            "check_situation_part_count": len(situation_parts),
            "guide_person_part_count": len(guide_parts),
        }

    def run(self, run_dir: Path) -> dict[str, Any]:
        if not self.is_enabled():
            return {"status": "SKIPPED_DISABLED", "run_dir": str(run_dir)}
        run_dir = Path(run_dir).resolve()
        output_dir = run_dir / "automatic_evaluation"
        output_dir.mkdir(parents=True, exist_ok=True)
        started_at = datetime.now().isoformat(timespec="seconds")
        results: dict[str, dict[str, Any]] = {}
        failures: dict[str, str] = {}
        self._write_status(output_dir, {
            "status": "RUNNING",
            "stage": "BUILD_EVIDENCE",
            "progress_percent": 5,
            "started_at": started_at,
            "providers_completed": [],
            "provider_failures": {},
            "provider_status": {name: "WAITING" for name in ("gpt", "gemini", "claude")},
            "message": "자동평가 Evidence 생성 중",
        })
        try:
            evidence, evidence_path = self.build_evidence_package(run_dir)
            review_payload = json.loads((run_dir / "02_REQUIREMENT_STUDIO_REVIEW_PACKAGE.json").read_text(encoding="utf-8"))
            situation_report = self._build_situation_check(review_payload, evidence)
            situation_text = self._render_situation_check(situation_report)
            situation_parts = self.split_situation_check(
                situation_text,
                output_dir / "check_situation",
                lines_per_file=int(self.config.get("situation_check_chunk_lines") or 110),
                handoff_identity=situation_report.get("handoff_identity"),
            )
            raw_guide_report = self._build_guide_person(review_payload, evidence, situation_report)
            guide_dir = output_dir / "guide_person"
            human_policy_registry = HumanPolicyRegistry(self.project_root)
            human_policy_data = human_policy_registry.load()
            guide_report = human_policy_registry.resolve_guide_report(raw_guide_report, human_policy_data)
            _atomic_write_text(guide_dir / "guide_person_questions.json", json.dumps(guide_report, ensure_ascii=False, indent=2))
            human_policy_registry.write_run_snapshot(guide_dir, guide_report, human_policy_data)
            guide_text = self._render_guide_person(guide_report)
            guide_parts = self.split_guide_person(
                guide_text, guide_dir,
                lines_per_file=int(self.config.get("guide_person_chunk_lines") or 110),
                handoff_identity=guide_report.get("handoff_identity"),
            )
            contract = evidence.get("unified_artifact_contract") if isinstance(evidence.get("unified_artifact_contract"), dict) else {}
            preflight_blocked = (
                str(evidence.get("evaluation_mode") or "").lower() == "unified"
                and str(contract.get("status") or "").upper() == "FAIL"
            )
            requested_provider_names = [name for name in ("gpt", "gemini", "claude") if bool(((self.config.get("providers") or {}).get(name) or {}).get("enabled", True))]
            provider_names = [] if preflight_blocked else requested_provider_names
            file_map = {"gpt": "01_GPT_EVALUATION", "gemini": "02_GEMINI_EVALUATION", "claude": "03_CLAUDE_EVALUATION"}
            provider_status = {
                name: ("PRECHECK_SKIPPED" if preflight_blocked and name in requested_provider_names else ("RUNNING" if name in provider_names else "DISABLED"))
                for name in ("gpt", "gemini", "claude")
            }
            self._write_status(output_dir, {
                "status": "RUNNING",
                "stage": "PROVIDER_REVIEW",
                "progress_percent": 15,
                "started_at": started_at,
                "providers_requested": provider_names,
                "providers_completed": [],
                "provider_failures": {},
                "provider_status": provider_status,
                "message": (
                    "Unified Artifact 사전검증 실패 · 외부 3-AI 호출 생략 · deterministic Tool FAIL/Release HOLD 생성"
                    if preflight_blocked else
                    "GPT/Gemini/Claude 독립 평가 진행 중 · 완료되는 모델부터 즉시 파일 저장"
                ),
            })

            with ThreadPoolExecutor(max_workers=max(1, len(provider_names))) as pool:
                futures = {pool.submit(self._call_one, name, evidence): name for name in provider_names}
                for fut in as_completed(futures):
                    name = futures[fut]
                    base = file_map[name]
                    try:
                        data, _model = fut.result()
                        results[name] = data
                        _atomic_write_text(output_dir / f"{base}.json", json.dumps(data, ensure_ascii=False, indent=2))
                        _atomic_write_text(output_dir / f"{base}.txt", self._render_provider_txt(data))
                    except Exception as exc:
                        failures[name] = str(exc)
                        _atomic_write_text(
                            output_dir / f"{base}.txt",
                            f"Requirement Studio Automated Evaluation - {name}\nSTATUS: FAILED_OR_SKIPPED\nREASON: {failures[name]}\n",
                        )
                    done_count = len(results) + len(failures)
                    for p_name in provider_names:
                        if p_name in results:
                            provider_status[p_name] = "COMPLETED"
                        elif p_name in failures:
                            provider_status[p_name] = "FAILED"
                        else:
                            provider_status[p_name] = "RUNNING"
                    progress = 15 + int(60 * done_count / max(1, len(provider_names)))
                    self._write_status(output_dir, {
                        "status": "RUNNING",
                        "stage": "PROVIDER_REVIEW",
                        "progress_percent": min(75, progress),
                        "started_at": started_at,
                        "providers_requested": provider_names,
                        "providers_completed": sorted(results),
                        "provider_failures": failures,
                        "provider_status": provider_status,
                        "message": f"Reviewer 진행 {done_count}/{len(provider_names)}",
                    })

            # Disabled providers still receive an explicit human-readable status file so the folder
            # contract is stable and never looks mysteriously empty.
            for name in ("gpt", "gemini", "claude"):
                base = file_map[name]
                txt_path = output_dir / f"{base}.txt"
                if not txt_path.exists():
                    _atomic_write_text(
                        txt_path,
                        f"Requirement Studio Automated Evaluation - {name}\nSTATUS: FAILED_OR_SKIPPED\nREASON: "
                        + ("Unified Artifact Contract preflight failed; external reviewer call skipped" if preflight_blocked else "not enabled")
                        + "\n",
                    )

            if bool(self.config.get("require_all_providers")) and failures:
                raise RuntimeError("3-model evaluation requires all providers: " + json.dumps(failures, ensure_ascii=False))

            self._write_status(output_dir, {
                "status": "RUNNING",
                "stage": "CONSENSUS",
                "progress_percent": 82,
                "started_at": started_at,
                "providers_requested": provider_names,
                "providers_completed": sorted(results),
                "provider_failures": failures,
                "provider_status": provider_status,
                "message": "Reviewer 결과 Consensus 정리 중",
            })
            consensus = self._consensus(results)
            consensus = self._apply_unified_artifact_contract(consensus, evidence)
            _atomic_write_text(output_dir / "04_EVALUATION_CONSENSUS.json", json.dumps(consensus, ensure_ascii=False, indent=2))
            _atomic_write_text(output_dir / "04_EVALUATION_CONSENSUS.txt", self._render_consensus_txt(consensus))
            self._write_status(output_dir, {
                "status": "RUNNING",
                "stage": "FINAL_RESULT",
                "progress_percent": 90,
                "started_at": started_at,
                "providers_requested": provider_names,
                "providers_completed": sorted(results),
                "provider_failures": failures,
                "provider_status": provider_status,
                "tool_quality_gate": consensus.get("tool_quality_gate"),
                "artifact_readiness_gate": consensus.get("artifact_readiness_gate"),
                "official_release": consensus.get("official_release"),
                "message": "최종 종합 결과 및 110줄 분할 파일 생성 중",
            })
            final = self._final_result(results, consensus, failures, evidence_path, evidence)
            _atomic_write_text(output_dir / "05_EVALUATION_RESULT.json", json.dumps(final, ensure_ascii=False, indent=2))
            final_text = self._render_final_txt(final)
            parts = self.split_final_result(final_text, output_dir / "final_result", lines_per_file=int(self.config.get("final_result_chunk_lines") or 110), handoff_identity=final.get("handoff_identity"))
            manifest_payload = {
                "generated_at": datetime.now().isoformat(timespec="seconds"),
                "transport": self.config.get("transport"),
                "evidence": evidence_path.name,
                "provider_results": sorted(results),
                "provider_failures": failures,
                "final_result_parts": [p.name for p in parts],
                "check_situation_parts": [p.name for p in situation_parts],
                "guide_person_parts": [p.name for p in guide_parts],
                "applied_policy_snapshot": ["guide_person/applied_policy_snapshot.txt", "guide_person/applied_policy_snapshot.json"],
            }
            _atomic_write_text(output_dir / "AUTOMATIC_EVALUATION_MANIFEST.json", json.dumps(manifest_payload, ensure_ascii=False, indent=2))

            self._write_status(output_dir, {
                "status": "RUNNING",
                "stage": "OUTPUT_VERIFY",
                "progress_percent": 96,
                "started_at": started_at,
                "providers_requested": provider_names,
                "providers_completed": sorted(results),
                "provider_failures": failures,
                "provider_status": provider_status,
                "tool_quality_gate": consensus.get("tool_quality_gate"),
                "artifact_readiness_gate": consensus.get("artifact_readiness_gate"),
                "official_release": consensus.get("official_release"),
                "message": "자동평가 결과 파일 존재/크기/분할 계약 검증 중",
            })
            verification = self._validate_output_artifacts(output_dir, results)
            _atomic_write_text(output_dir / "99_EVALUATION_COMPLETE.ok", json.dumps({
                "completed_at": datetime.now().isoformat(timespec="seconds"),
                "verification": verification,
            }, ensure_ascii=False, indent=2))
            for p_name in provider_names:
                provider_status[p_name] = "COMPLETED" if p_name in results else ("FAILED" if p_name in failures else provider_status.get(p_name, "UNKNOWN"))
            self._write_status(output_dir, {
                "status": "COMPLETED",
                "stage": "OUTPUT_VERIFIED",
                "progress_percent": 100,
                "started_at": started_at,
                "providers_requested": provider_names,
                "providers_completed": sorted(results),
                "provider_failures": failures,
                "provider_status": provider_status,
                "tool_quality_gate": consensus.get("tool_quality_gate"),
                "artifact_readiness_gate": consensus.get("artifact_readiness_gate"),
                "official_release": consensus.get("official_release"),
                "final_result_part_count": verification.get("final_result_part_count", 0),
                "check_situation_part_count": verification.get("check_situation_part_count", 0),
                "guide_person_part_count": verification.get("guide_person_part_count", 0),
                "verification": verification,
                "message": "결과 파일 저장 및 재검증 완료",
            })
            # Status is rewritten after verification; re-check the completed marker and status too.
            if not (output_dir / "99_EVALUATION_COMPLETE.ok").is_file():
                raise RuntimeError("자동평가 완료 마커가 생성되지 않았습니다.")
            return {
                "status": ("PRECHECK_FAILED" if preflight_blocked else ("COMPLETED" if results else "NO_PROVIDER_COMPLETED")),
                "preflight_blocked": preflight_blocked,
                "unified_artifact_contract": contract,
                "output_dir": str(output_dir),
                "providers_completed": sorted(results),
                "provider_failures": failures,
                "final_result_parts": [str(p) for p in parts],
                "check_situation_parts": [str(p) for p in situation_parts],
                "guide_person_parts": [str(p) for p in guide_parts],
                "tool_quality_gate": consensus.get("tool_quality_gate"),
                "artifact_readiness_gate": consensus.get("artifact_readiness_gate"),
                "official_release": consensus.get("official_release"),
                "verified_file_count": verification.get("file_count", 0),
                "final_result_part_count": verification.get("final_result_part_count", 0),
                "check_situation_part_count": verification.get("check_situation_part_count", 0),
                "guide_person_part_count": verification.get("guide_person_part_count", 0),
            }
        except Exception as exc:
            failed_provider_status = locals().get("provider_status", {name: "UNKNOWN" for name in ("gpt", "gemini", "claude")})
            self._write_status(output_dir, {
                "status": "FAILED",
                "stage": "ERROR",
                "progress_percent": 0,
                "started_at": started_at,
                "providers_completed": sorted(results),
                "provider_failures": failures,
                "provider_status": failed_provider_status,
                "message": str(exc),
            })
            raise

