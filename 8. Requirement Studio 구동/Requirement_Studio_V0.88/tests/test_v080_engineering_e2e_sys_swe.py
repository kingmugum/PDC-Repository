import json
from pathlib import Path

from core.cross_document_semantics import (
    _e2e_engineering_projection,
    _source_location_keys,
    apply_allocation_gate,
)
from core.quality_audit import apply_quality_audits
from core.swe1_exporter import build_engineering_records
from core.swe6_exporter import build_swe6_cases, build_sys5_candidates


def _base_req(text: str, srs_id: str = "SRS_001"):
    return {
        "candidate_id": "REQ-CAND-001",
        "srs_id": srs_id,
        "scenario_candidate_id": "SCN-1",
        "category": "기능",
        "function_name": "F",
        "requirement": text,
        "user_input": "",
        "system_input_preconditions": "",
        "processing_action": text,
        "output": "",
        "acceptance_criteria": text,
        "failure_situations": [],
        "user_intervention_points": [],
        "derivation_type": "explicit",
        "derivation_reason": "source",
        "clarification_needed": [],
        "source_evidence": [{"document": "x.docx", "location": "Paragraph 1", "text": text}],
        "confidence": 0.95,
        "classification_basis": "source",
        "activation_trigger": "",
        "preconditions": "",
        "behavior_flows": [],
        "evaluation_method": "",
        "exception_conditions": [],
        "related_artifacts": [],
        "source_requirement_ids": [],
        "source_requirement_occurrence_ids": [],
        "source_chunk_ids": [],
        "source_backed_atomic_behaviors": [],
        "applicability": {},
        "external_dependencies": [],
        "tbd_items": [],
        "conflicts": [],
        "open_issue_ids": [],
        "requirement_status": "Review Needed",
        "verification_constraints": [],
        "knowledge_state": "KNOWN",
        "source_semantic_unit_ids": [],
        "source_backed_facts": [],
        "fact_level_allocations": [],
        "source_table_fact_matches": [],
        "source_fact_fragments": [],
    }


def _data(reqs):
    return {
        "schema_version": "REQ-STUDIO-CANONICAL-REQ-2.0",
        "source_document": "x.docx",
        "scenario_candidates": [{
            "scenario_candidate_id": "SCN-1", "scenario_name": "x", "user_goal_context": "x",
            "scenario_flow": [], "expected_outcome": "x", "source_evidence": [],
        }],
        "requirements": reqs,
        "gaps": [],
    }


def test_system_pending_is_first_class_sys1_sys5_candidate():
    # V0.80 projection test: start from a known system-allocation record rather than
    # exercising the V0.77+ inferred-SW classifier in this test.
    req = _base_req("차량 시스템은 외부 인터페이스 조건을 만족해야 한다.")
    req.update({
        "allocation_status": "SYSTEM_REQUIREMENT_PENDING_SW_ALLOCATION",
        "requirement_level": "System",
        "swe1_eligibility": "Review Needed",
        "swe6_eligibility": "Deferred pending SW allocation",
        "verification_domain": "System Integration / SYS.5",
    })
    _e2e_engineering_projection(req)
    assert req["allocation_status"] == "SYSTEM_REQUIREMENT_PENDING_SW_ALLOCATION"
    assert req["sys1_eligibility"] == "Eligible"
    assert req["sys5_eligibility"] == "Eligible"
    assert req["main_spec_visibility"] is True
    assert "System" in req["engineering_domains"]


def test_explicit_software_stays_swe_view_without_forcing_sys1():
    req = _base_req("Software Task는 checksum을 검사해야 한다.")
    apply_allocation_gate(req, [])
    _e2e_engineering_projection(req)
    assert req["allocation_status"] == "SW_IMPLEMENTATION_REQUIREMENT"
    assert req["swe1_eligibility"] == "Eligible"
    assert req["swe6_eligibility"] == "Eligible"
    assert req["sys1_eligibility"] == "Not Applicable"


def test_non_sw_normative_fact_recovered_as_canonical_engineering_requirement():
    req = _base_req("제어기는 DC 7V~18V에서 정상 동작해야 한다.")
    req["source_evidence"] = [{"document": "x.docx", "location": "Table 10 / Row 3", "text": req["requirement"]}]
    data = _data([req])
    compact = """[DOCUMENT] x.docx
[SRC C1 | Table 10 / Row 2 | table]
정격 전압 | DC 13.5 V
[SRC C2 | Table 10 / Row 3 | table]
정상 동작 전압 범위 | DC 7V~18V
"""
    out = apply_quality_audits(data, compact)
    texts = [str(r.get("requirement") or "") for r in out["requirements"]]
    assert any("13.5" in text for text in texts)
    recovered = [r for r in out["requirements"] if "13.5" in str(r.get("requirement") or "")]
    assert recovered
    assert recovered[0]["allocation_status"] == "ELECTRICAL_REQUIREMENT"
    assert recovered[0]["main_spec_visibility"] is True
    assert recovered[0]["swe1_eligibility"] == "Not Applicable"


def test_plural_comma_source_locations_preserve_every_explicit_paragraph():
    keys = _source_location_keys("Paragraphs 205, 208")
    assert "paragraph:205" in keys
    assert "paragraph:208" in keys


def test_deferred_sys5_fact_cannot_leak_into_swe6_qualifying_fields():
    req = _base_req("HU는 wake-up 후 software 상태를 판단하고 응답해야 한다.")
    req.update({
        "allocation_status": "SW_BEHAVIOR_INFERRED_REVIEW_REQUIRED",
        "requirement_level": "Software Candidate",
        "swe1_eligibility": "Eligible",
        "swe6_eligibility": "Eligible",
        "verification_domain": "SWE.6 Software Qualification (Allocation Review Required)",
        "semantic_provenance_status": "COMPLETE",
        "source_semantic_unit_ids": ["U1", "U2"],
        "source_backed_atomic_behaviors": [
            {"source_semantic_unit_id": "U1", "source_fact_fragment_id": "F1", "behavior_text": "HU는 software 상태를 판단해야 한다."},
            {"source_semantic_unit_id": "U2", "source_fact_fragment_id": "F2", "behavior_text": "HU는 C_PreCrankReq=0x1을 송신해야 한다."},
        ],
        "fact_level_allocations": [
            {"source_semantic_unit_id": "U1", "source_fact_fragment_id": "F1", "source_fact": "HU는 software 상태를 판단해야 한다.", "swe6_eligibility": "Eligible", "verification_domain": "SWE.6 Software Qualification"},
            {"source_semantic_unit_id": "U2", "source_fact_fragment_id": "F2", "source_fact": "HU는 C_PreCrankReq=0x1을 송신해야 한다.", "swe6_eligibility": "Deferred pending SW allocation", "verification_domain": "System Integration / SYS.5"},
        ],
        "activation_trigger": "System wake-up 완료 후 C_PreCrankReq=0x1 송신",
        "output": "C_PreCrankReq=0x1",
        "acceptance_criteria": "C_PreCrankReq=0x1",
    })
    cases = build_swe6_cases({"requirements": [req]})
    assert cases
    case = cases[0]
    qualifying = " ".join(str(case.get(k) or "") for k in (
        "prep_desc", "prep_var", "prep_compare", "prep_value",
        "exec_desc", "exec_var", "exec_compare", "exec_value",
        "expected_desc", "expected_var", "expected_compare", "expected_value",
    ))
    assert "C_PreCrankReq" not in qualifying
    assert "software 상태" in qualifying


def test_system_candidate_builder_keeps_result_fields_blank():
    req = _base_req("HU는 차량 상태 조합에 따라 sidemirrorOpen을 전송해야 한다.")
    apply_allocation_gate(req, [])
    _e2e_engineering_projection(req)
    candidates = build_sys5_candidates({"requirements": [req]})
    assert candidates
    assert candidates[0]["sys5_id"].startswith("SYS5_TC_")
    assert candidates[0]["status"] == "CANDIDATE_REVIEW_REQUIRED"
    assert candidates[0]["execution_result"] == ""
    assert candidates[0]["pass_fail"] == ""


def test_engineering_main_includes_native_domain_even_when_not_swe1():
    req = _base_req("커넥터는 DIP 타입을 적용해야 한다.")
    apply_allocation_gate(req, [])
    _e2e_engineering_projection(req)
    records = build_engineering_records({"requirements": [req]})
    assert len(records) == 1
    assert records[0]["SWE.1 Eligibility"] == "Not Applicable"
    assert "Hardware" in records[0]["Engineering Domain"]


def test_v080_seed_carries_e2e_core_policies():
    seed = json.loads((Path(__file__).parents[1] / "policy" / "human_policy_seed.json").read_text(encoding="utf-8"))
    pmap = {p["policy_id"]: p for p in seed["policies"]}
    assert pmap["HP-E2E-001"]["decision"] == "YES"
    assert pmap["HP-SYSVER-001"]["decision"] == "YES"


def test_v080_exporters_emit_integrated_sys_swe_views(tmp_path):
    from docx import Document
    from openpyxl import load_workbook
    from core.swe1_exporter import SWE1Exporter
    from core.swe6_exporter import SWE6Exporter

    src = tmp_path / "source.docx"
    Document().save(src)
    system_req = _base_req("차량 시스템은 외부 인터페이스 조건을 만족해야 한다.")
    system_req.update({
        "allocation_status": "SYSTEM_REQUIREMENT_PENDING_SW_ALLOCATION",
        "requirement_level": "System",
        "swe1_eligibility": "Review Needed",
        "swe6_eligibility": "Deferred pending SW allocation",
        "verification_domain": "System Integration / SYS.5",
    })
    _e2e_engineering_projection(system_req)
    data = _data([system_req])
    swe1_path = SWE1Exporter(tmp_path).export_excel(src, data)
    swe6_path = SWE6Exporter(tmp_path).export_excel(src, data)

    wb1 = load_workbook(swe1_path, read_only=True)
    try:
        assert {"08_Engineering_Main", "09_SYS1_System_View", "10_SWE1_Software_View", "11_Human_Decision_View"}.issubset(set(wb1.sheetnames))
    finally:
        wb1.close()
    wb6 = load_workbook(swe6_path, read_only=True)
    try:
        assert {"6_SYS5_Candidates", "7_Engineering_Verification"}.issubset(set(wb6.sheetnames))
        ws = wb6["6_SYS5_Candidates"]
        assert ws["J4"].value in (None, "")
        assert ws["K4"].value in (None, "")
    finally:
        wb6.close()
