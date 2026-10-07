from pathlib import Path

from openpyxl import load_workbook

from core.cross_document_semantics import (
    _source_fact_fragments_for_unit,
    _structured_source_facts,
    apply_allocation_gate,
    attach_semantic_traceability,
)
from core.quality_audit import (
    _source_backed_child_intents,
    build_swe6_export_preservation_audit,
    normalize_requirement_extensions,
)
from core.regression_engine import build_regression_report
from core.swe6_exporter import SWE6Exporter


def _req(text: str, source_text: str | None = None, *, srs: str = "SRS_001", location: str = "Paragraph 1") -> dict:
    row = {
        "candidate_id": f"C-{srs}",
        "srs_id": srs,
        "scenario_candidate_id": "SCN-CAND-001",
        "category": "기능",
        "function_name": "Feature",
        "requirement": text,
        "user_input": "",
        "system_input_preconditions": "",
        "processing_action": text,
        "output": text,
        "acceptance_criteria": text,
        "failure_situations": [],
        "user_intervention_points": [],
        "derivation_type": "explicit",
        "derivation_reason": "",
        "clarification_needed": [],
        "source_evidence": [{"document": "x.docx", "location": location, "text": source_text or text}],
        "confidence": 0.95,
        "classification_basis": "source",
        "activation_trigger": "",
        "preconditions": "",
        "behavior_flows": [],
        "evaluation_method": "",
        "exception_conditions": [],
        "related_artifacts": [],
        "source_requirement_ids": [],
    }
    normalize_requirement_extensions(row)
    return row


def _data(reqs):
    if not isinstance(reqs, list):
        reqs = [reqs]
    return {
        "schema_version": "REQ-STUDIO-CANONICAL-REQ-2.0",
        "source_document": "x.docx",
        "scenario_candidates": [],
        "requirements": reqs,
        "gaps": [],
    }


def test_table_decimal_version_and_heading_are_not_ordered_list_markers():
    table_unit = {
        "source_semantic_unit_id": "U-TABLE",
        "source_chunk_id": "C1",
        "source_location": "Table 10 / Row 4",
        "source_kind": "table",
        "source_excerpt": "작동 전압 범위 | 저전압 : DC 6.5V 이하 | 고전압 : DC 18.5V 이상 (통신 disable)",
    }
    fragments = _source_fact_fragments_for_unit(table_unit)
    assert len(fragments) == 1
    assert "6.5V" in fragments[0]["source_excerpt_raw"]
    assert "18.5V" in fragments[0]["source_excerpt_raw"]
    assert not fragments[0].get("source_list_marker")

    for literal in ("1.3 Revision History", "5.4.1.1 B-CAN 통신 요구 사항", "2019-07-19 Release"):
        unit = {
            "source_semantic_unit_id": literal,
            "source_chunk_id": "C2",
            "source_location": "Paragraph 2",
            "source_kind": "text",
            "source_excerpt": literal,
        }
        out = _source_fact_fragments_for_unit(unit)
        assert len(out) == 1
        assert not out[0].get("source_list_marker")


def test_true_numbered_list_is_still_split_into_four_members():
    unit = {
        "source_semantic_unit_id": "U-LIST",
        "source_chunk_id": "C3",
        "source_location": "Paragraph 136",
        "source_kind": "text",
        "source_excerpt": "High Power Mode 조건: 1) ON 명령 수신 시 2) B-CAN WAKE UP 상태 3) ERR PIN HIGH에서 LOW 상태 변환 시 4) 최초 B+ 인가 시",
    }
    out = _source_fact_fragments_for_unit(unit)
    members = [x for x in out if x.get("source_list_marker")]
    assert [x["source_list_marker"] for x in members] == ["1)", "2)", "3)", "4)"]
    assert any("최초 B+ 인가" in x["source_excerpt"] for x in members)


def test_composite_controller_clause_gets_behavior_level_fragment_ownership():
    source = (
        "무드램프 마스터 제어기는 AVN 또는 클러스터 등으로부터 사용자 설정 정보를 CAN 통신으로 전달 받아 "
        "마스터 제어기의 비휘발성 메모리에 저장하고 차량으로부터 오는 무드램프 On/Off 제어 정보를 받아 "
        "슬레이브 제어기를 통신으로 제어 한다."
    )
    compact = f"""[DOCUMENT] x.docx
[SRC C82 | Paragraph 82 | text]
{source}
"""
    r1 = _req("사용자 설정 정보를 비휘발성 메모리에 저장한다.", source, srs="SRS_001", location="Paragraph 82")
    r2 = _req("차량 On/Off 제어 정보를 수신하여 슬레이브 제어기를 통신으로 제어한다.", source, srs="SRS_002", location="Paragraph 82")
    data = _data([r1, r2])
    attach_semantic_traceability(data, compact)

    f1 = {x["source_fact_fragment_id"] for x in r1.get("source_fact_fragments") or []}
    f2 = {x["source_fact_fragment_id"] for x in r2.get("source_fact_fragments") or []}
    assert f1 and f2
    assert f1 != f2
    assert not (f1 & f2), "different behavior ownership should not reuse the whole composite clause"
    assert data.get("source_fact_multi_srs_allocation_conflicts") == []


def test_positive_sw_evidence_gate_prevents_system_behavior_overgeneration():
    system_req = _req("슬레이브 제어기를 통신으로 제어한다.", "슬레이브 제어기를 통신으로 제어한다.")
    apply_allocation_gate(system_req, [])
    assert system_req["allocation_status"] == "SYSTEM_REQUIREMENT_PENDING_SW_ALLOCATION"
    assert system_req["swe6_eligibility"].startswith("Deferred")
    assert system_req["positive_sw_allocation_evidence"] == []

    sw_req = _req("S/W 오류 발생 시 안전 상태로 전환한다.")
    apply_allocation_gate(sw_req, [])
    assert sw_req["allocation_status"] == "SW_IMPLEMENTATION_REQUIREMENT"
    assert sw_req["swe6_eligibility"] == "Eligible"
    assert sw_req["positive_sw_allocation_evidence"]

    structural = _req("Watchdog Timer는 가장 낮은 Task 이후 작동해야 하며 ISR 내부에 있을 수 없다.")
    apply_allocation_gate(structural, [])
    assert structural["swe6_eligibility"] == "Eligible"
    assert structural["software_allocation_evidence_type"] == "SOFTWARE_STRUCTURAL_CONSTRAINT"


def test_ms_and_general_reference_identifier_are_not_ranges():
    ms = _structured_source_facts("MS181-15 규격을 참조한다.", unit_id="U1", location="Table 1 / Row 1")[0]
    assert ms["fact_type"] == "EXTERNAL_SPEC_IDENTIFIER"
    assert "MS181-15" in ms["external_spec_identifiers"]
    assert ms["range_values"] == []
    assert ms["numeric_values"] == []

    iso = _structured_source_facts("ISO26262-6 reference", unit_id="U2", location="Table 1 / Row 2")[0]
    assert "ISO26262-6" in iso["external_spec_identifiers"]
    assert iso["range_values"] == []


def test_no_id_approved_gold_can_pass_by_stable_source_key_without_occurrence_ids():
    source_text = "차량 On/Off 제어 정보를 수신하여 슬레이브 제어기를 통신으로 제어한다."
    previous = {
        "run": {
            "requirement_studio_version": "v0.64",
            "source_sha256": "abc",
            "provider_id": "hchat_gpt",
            "provider": "HChat GPT",
            "model": "gpt",
            "extraction_core_profile": "v0.46-compatible-1.1",
        },
        "canonical_requirement": {
            "requirements": [{
                "candidate_id": "OLD",
                "srs_id": "OLD-SRS",
                "requirement": source_text,
                "source_requirement_ids": [],
                "source_evidence": [{"location": "Paragraph 82", "text": source_text}],
            }]
        },
        "gold_source_contract": {
            "baseline_type": "compact_gold_regression_contract",
            "baseline_origin_version": "v0.64",
            "is_full_review_package": False,
            "behavior_scope": "source-backed compact behavior contract",
            "approval_status": "APPROVED_FOR_REGRESSION",
        },
    }
    current_req = _req(source_text, source_text, srs="SRS_002", location="Paragraph 82")
    current_req["source_fact_fragments"] = [{
        "source_fact_fragment_id": "SRC-FRAG-X",
        "source_location": "Paragraph 82",
        "source_excerpt": source_text,
        "source_excerpt_raw": source_text,
    }]
    current = {"requirements": [current_req], "source_coverage": {"coverage_validity": "NOT_APPLICABLE", "source_requirement_occurrences": []}}
    report = build_regression_report(current, previous, "Gold Source Contract", current_run={
        "requirement_studio_version": "v0.65",
        "source_sha256": "abc",
        "provider_id": "hchat_gpt",
        "provider": "HChat GPT",
        "model": "gpt",
        "extraction_core_profile": "v0.46-compatible-1.1",
    })
    assert report["baseline_is_approved_gold_contract"] is True
    assert report["regression_gate"]["status"] == "PASS"
    assert report["regression_gate"]["passed"] is True
    assert report["findings"] == []
    assert report["comparison_coverage"]["content_hash_match_count"] >= 1


def test_tool_quality_and_artifact_readiness_are_separate_and_maturity_is_exported(tmp_path):
    req = _req("S/W는 Input_A를 수신하면 Output_B를 송신한다.")
    req.update({
        "requirement_level": "Software",
        "allocation_status": "SW_IMPLEMENTATION_REQUIREMENT",
        "swe1_eligibility": "Eligible",
        "swe6_eligibility": "Eligible",
        "verification_domain": "SWE.6 Software Qualification",
        "source_semantic_unit_ids": ["S1"],
        "source_fact_fragments": [{
            "source_fact_fragment_id": "F1",
            "parent_source_semantic_unit_id": "S1",
            "source_location": "Paragraph 1",
            "source_excerpt": "S/W는 Input_A를 수신하면 Output_B를 송신한다.",
            "knowledge_state": "KNOWN",
        }],
        "source_backed_atomic_behaviors": [{
            "source_semantic_unit_id": "S1",
            "source_fact_fragment_id": "F1",
            "source_location": "Paragraph 1",
            "behavior_text": "S/W는 Input_A를 수신하면 Output_B를 송신한다.",
            "knowledge_state": "KNOWN",
        }],
        "semantic_provenance_status": "COMPLETE",
        "mixed_external_dependency_review_required": True,
        "external_dependency_domains": ["External Standard Evidence Required"],
        "fact_level_allocations": [],
    })
    children = _source_backed_child_intents(req)
    data = _data(req)
    data["testability_and_decomposition_result"] = {
        "summary": {},
        "by_srs": [{
            "srs_id": "SRS_001",
            "swe6_eligibility": "Eligible",
            "verification_domain": "SWE.6 Software Qualification",
            "required_test_intents": ["Normal / Positive"],
            "not_generated_test_intents": [],
            "source_backed_child_intents": children,
        }],
    }
    case = {
        "srs_id": "SRS_001",
        "tc_id": "TC_001",
        "description": "Source-backed generic intent",
        "prep_desc": "",
        "prep_var": "",
        "prep_compare": "",
        "prep_value": "",
        "exec_desc": "S/W는 Input_A를 수신하면 Output_B를 송신한다.",
        "exec_var": "",
        "exec_compare": "",
        "exec_value": "",
        "expected_desc": "Output_B를 송신한다.",
        "expected_var": "",
        "expected_compare": "",
        "expected_value": "",
    }
    audit = build_swe6_export_preservation_audit(data, [case])
    assert audit["tool_quality_gate_passed"] is True
    assert audit["artifact_readiness_gate_passed"] is False
    assert audit["artifact_readiness_gate_status"] == "REVIEW_REQUIRED"
    assert case["tc_maturity"] == "INTENT_DRAFT"
    assert case["independent_coverage"] == "NOT_ESTABLISHED"
    assert case["execution_readiness"] == "REVIEW_REQUIRED"

    # Exporter carries maturity/readiness columns without moving the protected result columns R:U.
    data["testability_and_decomposition_result"]["by_srs"][0]["source_backed_child_intents"] = children
    out = SWE6Exporter(tmp_path).export_excel(tmp_path / "x.docx", data, source_text="S/W는 Input_A를 수신하면 Output_B를 송신한다.")
    wb = load_workbook(out, data_only=False)
    try:
        ws = wb["2_테스트 케이스"]
        assert ws["V3"].value == "TC Maturity"
        assert ws["W3"].value == "Independent Coverage"
        assert ws["X3"].value == "Execution Readiness"
        assert ws["V4"].value in {"INTENT_DRAFT", "PARTIALLY_TESTABLE", "EXECUTION_READY"}
        assert [ws.cell(4, c).value for c in (18, 19, 20, 21)] == [None, None, None, None]
    finally:
        wb.close()
