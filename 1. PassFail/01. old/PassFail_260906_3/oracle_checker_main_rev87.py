# -*- coding: utf-8 -*-
"""
oracle_checker_main_rev87.py

변경 사항:
- done_event 추가
  * stop_event: 진짜 중단
  * done_event: 사용자가 "수행 완료"를 눌러 현재 시점 기준으로 실시간 판정을 마감하고
                이후 로그 재검토 단계로 자연스럽게 넘어가게 하는 신호
- 실시간 입력/출력 판정 루프에서 done_event를 인식
- 수행 완료 시점의 관측 결과를 최대한 result에 반영
- 기존 export / Excel / DBC / ASC / BLF 로직은 유지
- rev20: CANoe COM 재연결/Measurement Stop-Start 안정화
- rev21: __pycache__ 생성 억제 및 종료 시 자동 삭제
- rev22: CANoe 미실행 상태에서 연결됨처럼 보이는 문제 완화
- rev23: 입력 다중 조건 지원
- rev24: 실행 중인 CANoe COM 객체 연결 및 유효성 검증
- rev25:
  * CANoe COM 연결 모드 분리 및 유효성 검증 강화
- rev26:
  * 누락된 CANoeClient 클래스 복구
  * 기본 연결은 GetActiveObject 우선이며, 이미 실행 중인 Vector 제품이 확인된 경우에만 guarded COM activation fallback 허용
- rev37:
  * rev36 기능 유지
  * 실행 중인 CANoe COM 객체의 Version/FullName 검증 강화
- rev40:
  * CANoe 15 고정 버전 검증 해제
  * 실행 중인 CANoe가 있으면 CANoe 10/15/기타 버전 모두 연결 가능
  * 버전 제한이 필요한 경우에만 connect(expected_version_keyword=...)로 별도 지정
- rev42:
  * GUI rev42의 다중 실행 UI/상태 유지 개선과 연동되는 core 계약 유지
  * CheckResult/FinalCheckResult 데이터 구조는 하위 호환을 위해 변경하지 않음
- rev44:
  * GUI 탭명/배치/열 폭 조정 요청 대응용 버전명 갱신
  * Core 판정 계약 및 CheckResult/FinalCheckResult 데이터 구조 변경 없음
- rev45:
  * 다중 TC 로그 검토용 review_conditions_in_log_batch 추가
  * 여러 TC/시그널을 ASC 1회 스캔으로 일괄 확인하여 대용량 로그 반복 스캔 부담 완화
  * Vector 연결은 GetActiveObject 우선 + 기존 프로세스 확인 기반 guarded Dispatch fallback
  * 실시간 PASS는 로그 재검토 결과와 무관하게 최종 PASS 유지
- rev87:
  * Vector CANoe/CANalyzer 공통 연결 계층 추가
    - 연결 대상: 자동 / CANoe / CANalyzer
    - CANoe.Application 및 CANalyzer.Application을 GetActiveObject 우선으로 탐지
    - ProgID 실패 시 Registry 32/64-bit view에서 read-only로 확인한 CLSID로 GetActiveObject 재시도
    - 두 제품이 동시에 서로 다른 활성 객체로 감지되면 자동 선택하지 않고 사용자가 제품을 지정
  * 구형 CANalyzer Runtime Kernel을 포함하도록 프로세스 진단을 제품명/설명/경로까지 확장
  * 기존 판정 Core는 관측(read) 역할 유지; 실제 Signal write API는 추가하지 않음
  * 실제 입력 PoC는 독립 oracle_checker_stimulus_rev87.py에서만 수행
  * DispatchEx/Registry 쓰기/무조건적 새 인스턴스 생성은 금지. Dispatch는 기존 제품 프로세스 확인 시 guarded fallback으로만 사용
  * 기존 DBC 자동매핑/COM 진단/판정 계약은 유지
  * rev87 TX/RX Stimulus 진단: per-run 분석 window, 전체/채널 Tx 기록 능력, 동일-ID Rx 지속, intended payload diff, TC 출력 Expected 반응을 BLF/ASC에서 분리 분석
  * rev87 추가 진단: intended baseline과 같은 실행창의 실제 RX raw frame byte diff를 표시하여 재구성 baseline 불일치를 검출
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Callable, Dict, Iterable, List, Optional, Tuple
import datetime as _dt
import math
import re
import threading
import time
import traceback
import subprocess
import sys
import atexit
import shutil
import os
import json
import platform
import tempfile


# =========================================================
# __pycache__ 생성 억제 및 기존 캐시 폴더 정리
# =========================================================
sys.dont_write_bytecode = True

def _cleanup_pycache(base_dir=None):
    try:
        root = Path(base_dir) if base_dir is not None else Path(__file__).resolve().parent
        for p in root.glob("__pycache__"):
            if p.is_dir():
                shutil.rmtree(p, ignore_errors=True)
    except Exception:
        pass

_cleanup_pycache()
atexit.register(_cleanup_pycache)


# =========================================================
# 공통 유틸
# =========================================================
_HEX_RE = re.compile(r"^[-+]?0x[0-9a-fA-F]+$")
_INT_RE = re.compile(r"^[-+]?\d+$")
_FLOAT_RE = re.compile(r"^[-+]?\d+\.\d+$")

DEBUG_ENABLED = False
_COM_THREAD_STATE = threading.local()


def debug_print(msg: str) -> None:
    if DEBUG_ENABLED:
        ts = _dt.datetime.now().strftime("%H:%M:%S.%f")[:-3]
        print(f"[DEBUG {ts}] {msg}", flush=True)


def _hidden_subprocess_kwargs() -> Dict[str, Any]:
    """Windows .pyw 실행 중 subprocess 호출 시 검은 콘솔창이 뜨지 않도록 한다."""
    if not sys.platform.startswith("win"):
        return {}

    kwargs: Dict[str, Any] = {}
    try:
        create_no_window = getattr(subprocess, "CREATE_NO_WINDOW", 0)
        if create_no_window:
            kwargs["creationflags"] = create_no_window

        startupinfo = subprocess.STARTUPINFO()
        startupinfo.dwFlags |= subprocess.STARTF_USESHOWWINDOW
        startupinfo.wShowWindow = 0
        kwargs["startupinfo"] = startupinfo
    except Exception:
        pass
    return kwargs



# =========================================================
# Vector CANoe/CANalyzer COM read-only diagnostics (rev87)
# =========================================================
def _exception_chain_details(error: Optional[BaseException]) -> List[Dict[str, Any]]:
    out: List[Dict[str, Any]] = []
    seen = set()
    cur = error
    while cur is not None and id(cur) not in seen and len(out) < 8:
        seen.add(id(cur))
        hresult = getattr(cur, "hresult", None)
        if hresult is None:
            try:
                if getattr(cur, "args", None) and isinstance(cur.args[0], int):
                    hresult = int(cur.args[0])
            except Exception:
                pass
        item: Dict[str, Any] = {
            "type": type(cur).__name__,
            "message": str(cur),
            "args": repr(getattr(cur, "args", ())),
        }
        if hresult is not None:
            try:
                hr = int(hresult)
                item["hresult"] = hr
                item["hresult_hex"] = f"0x{(hr & 0xFFFFFFFF):08X}"
            except Exception:
                item["hresult"] = str(hresult)
        out.append(item)
        cur = getattr(cur, "__cause__", None) or getattr(cur, "__context__", None)
    return out


def _python_is_admin() -> Optional[bool]:
    if not sys.platform.startswith("win"):
        return None
    try:
        import ctypes
        return bool(ctypes.windll.shell32.IsUserAnAdmin())
    except Exception:
        return None


VECTOR_PRODUCT_SPECS = (
    ("CANoe", "CANoe.Application"),
    ("CANalyzer", "CANalyzer.Application"),
)


def normalize_vector_product_preference(value: Any) -> str:
    raw = str(value or "").strip().lower()
    if raw in {"canoe", "can oe"}:
        return "CANoe"
    if raw in {"canalyzer", "can alyzer", "canalyser"}:
        return "CANalyzer"
    return "Auto"


def _vector_specs_for_preference(preference: Any = "Auto") -> List[Tuple[str, str]]:
    pref = normalize_vector_product_preference(preference)
    if pref == "CANoe":
        return [("CANoe", "CANoe.Application")]
    if pref == "CANalyzer":
        return [("CANalyzer", "CANalyzer.Application")]
    return list(VECTOR_PRODUCT_SPECS)


def _collect_vector_processes() -> Dict[str, Any]:
    """CANoe/CANalyzer 프로세스를 제품명/설명/경로까지 포함해 read-only 탐색한다."""
    result: Dict[str, Any] = {"method": "", "items": [], "error": ""}
    if not sys.platform.startswith("win"):
        result["error"] = "Windows가 아니므로 Vector 프로세스 조회를 수행하지 않음"
        return result

    ps_script = """
$ErrorActionPreference='SilentlyContinue'
$x = Get-CimInstance Win32_Process | ForEach-Object {
  $path = $_.ExecutablePath
  $fv=''; $pn=''; $fd=''
  if ($path) {
    try {
      $vi=(Get-Item -LiteralPath $path).VersionInfo
      $fv=$vi.FileVersion; $pn=$vi.ProductName; $fd=$vi.FileDescription
    } catch {}
  }
  $hay = (($_.Name)+' '+$path+' '+$pn+' '+$fd+' '+$_.CommandLine)
  if ($_.Name -notmatch '(?i)^powershell([.]exe)?$' -and $hay -match '(?i)CANoe|CANalyzer|CANw32|CANw64|CANoe32|CANoe64') {
    [PSCustomObject]@{
      ProcessId=$_.ProcessId; Name=$_.Name; ExecutablePath=$path; CommandLine=$_.CommandLine;
      FileVersion=$fv; ProductName=$pn; FileDescription=$fd
    }
  }
}
$x | ConvertTo-Json -Compress
"""
    ps_error = ""
    try:
        completed = subprocess.run(
            ["powershell.exe", "-NoProfile", "-Command", ps_script],
            capture_output=True,
            text=True,
            timeout=10,
            **_hidden_subprocess_kwargs(),
        )
        raw = (completed.stdout or "").strip()
        if completed.returncode == 0:
            result["method"] = "PowerShell Get-CimInstance + VersionInfo"
            if raw:
                parsed = json.loads(raw)
                if isinstance(parsed, dict):
                    parsed = [parsed]
                if isinstance(parsed, list):
                    result["items"] = parsed
        else:
            ps_error = (completed.stderr or "").strip() or f"PowerShell returncode={completed.returncode}"
    except Exception as e:
        ps_error = f"PowerShell 조회 실패: {type(e).__name__}: {e}"

    # PowerShell이 성공했더라도 0건이면 tasklist로 한 번 더 확인한다.
    if not result["items"]:
        try:
            completed = subprocess.run(
                ["tasklist.exe", "/V", "/FO", "CSV", "/NH"],
                capture_output=True,
                text=True,
                timeout=6,
                **_hidden_subprocess_kwargs(),
            )
            lines = []
            for line in (completed.stdout or "").splitlines():
                low = line.lower()
                if any(k in low for k in ("canoe", "canalyzer", "canw32", "vector")):
                    lines.append(line.strip())
            if lines:
                result["method"] = (result.get("method") + " + " if result.get("method") else "") + "tasklist /V fallback"
                result["items"] = [{"raw": x} for x in lines]
        except Exception as e:
            extra = f"tasklist 실패: {type(e).__name__}: {e}"
            ps_error = (ps_error + " | " if ps_error else "") + extra

    result["error"] = ps_error
    return result


# 이전 함수명은 하위 호환을 위해 유지한다.
def _collect_canoe_processes() -> Dict[str, Any]:
    return _collect_vector_processes()


def _vector_process_haystack(item: Dict[str, Any]) -> str:
    parts = [
        item.get("Name"), item.get("ExecutablePath"), item.get("ProductName"),
        item.get("FileDescription"), item.get("CommandLine"), item.get("raw"),
    ]
    return " ".join(str(x or "") for x in parts).lower()


def _vector_process_product(item: Dict[str, Any]) -> str:
    """Best-effort product classification for already-running Vector GUI processes.

    RuntimeKernel is intentionally treated as generic because it can belong to CANoe or CANalyzer.
    Guarded COM activation requires a product-specific main process, not RuntimeKernel alone.
    """
    hay = _vector_process_haystack(item)
    name = str(item.get("Name") or "").lower()
    if "runtimekernel" in name or "runtime kernel" in hay:
        return ""
    if name in {"canw32.exe", "canw64.exe"} or "vector canalyzer" in hay or "\\vector canalyzer" in hay:
        return "CANalyzer"
    if name in {"canoe32.exe", "canoe64.exe"} or "vector canoe" in hay or "\\vector canoe" in hay:
        # Avoid classifying the generic "Vector CANalyzer/CANoe" Runtime Kernel as CANoe.
        if "canalyzer/canoe" not in hay:
            return "CANoe"
    return ""


def _vector_running_products(process_snapshot: Optional[Dict[str, Any]] = None) -> List[str]:
    snap = process_snapshot if process_snapshot is not None else _collect_vector_processes()
    products = []
    for item in snap.get("items") or []:
        product = _vector_process_product(item)
        if product and product not in products:
            products.append(product)
    return products


def _vector_main_process_pids(process_snapshot: Dict[str, Any], product: str) -> set[int]:
    out: set[int] = set()
    target = str(product or "").strip()
    for item in process_snapshot.get("items") or []:
        if _vector_process_product(item) != target:
            continue
        pid = item.get("ProcessId") or item.get("PID") or item.get("pid")
        try:
            out.add(int(pid))
        except Exception:
            pass
    return out


def _registry_read_default(winreg, root, path: str, access: int) -> str:
    try:
        with winreg.OpenKey(root, path, 0, winreg.KEY_READ | access) as key:
            value, _ = winreg.QueryValueEx(key, "")
            return str(value)
    except Exception:
        return ""


def _registry_collect_view(
    view_name: str,
    root_name: str,
    classes_path: str,
    prog_id: str,
    product: str,
    wow_flag: int = 0,
) -> Dict[str, Any]:
    result: Dict[str, Any] = {
        "view": view_name,
        "root": root_name,
        "product": product,
        "prog_id": prog_id,
        "description": "",
        "clsid": "",
        "curver": "",
        "local_server32": "",
        "versioned_progids": [],
        "error": "",
    }
    try:
        import winreg
        root = winreg.HKEY_LOCAL_MACHINE if root_name == "HKLM" else winreg.HKEY_CURRENT_USER
        prefix = classes_path.rstrip("\\")
        base = prefix + "\\" + prog_id
        result["description"] = _registry_read_default(winreg, root, base, wow_flag)
        result["clsid"] = _registry_read_default(winreg, root, base + r"\CLSID", wow_flag)
        result["curver"] = _registry_read_default(winreg, root, base + r"\CurVer", wow_flag)
        if result["clsid"]:
            result["local_server32"] = _registry_read_default(
                winreg, root, prefix + "\\CLSID\\" + result["clsid"] + "\\LocalServer32", wow_flag
            )

        target_low = prog_id.lower()
        with winreg.OpenKey(root, prefix, 0, winreg.KEY_READ | wow_flag) as classes_key:
            i = 0
            while i < 50000 and len(result["versioned_progids"]) < 40:
                try:
                    name = winreg.EnumKey(classes_key, i)
                except OSError:
                    break
                i += 1
                low = name.lower()
                if not (low == target_low or low.startswith(target_low + ".")):
                    continue
                sub = prefix + "\\" + name
                clsid = _registry_read_default(winreg, root, sub + r"\CLSID", wow_flag)
                curver = _registry_read_default(winreg, root, sub + r"\CurVer", wow_flag)
                server = ""
                if clsid:
                    server = _registry_read_default(winreg, root, prefix + "\\CLSID\\" + clsid + "\\LocalServer32", wow_flag)
                result["versioned_progids"].append({
                    "progid": name,
                    "clsid": clsid,
                    "curver": curver,
                    "local_server32": server,
                })
    except Exception as e:
        result["error"] = f"{type(e).__name__}: {e}"
    return result


def _collect_vector_registry() -> List[Dict[str, Any]]:
    if not sys.platform.startswith("win"):
        return [{"view": "N/A", "error": "Windows가 아니므로 Registry 조회를 수행하지 않음"}]
    try:
        import winreg
        view64 = getattr(winreg, "KEY_WOW64_64KEY", 0)
        view32 = getattr(winreg, "KEY_WOW64_32KEY", 0)
    except Exception as e:
        return [{"view": "Registry", "error": f"winreg import 실패: {e}"}]

    out: List[Dict[str, Any]] = []
    for product, prog_id in VECTOR_PRODUCT_SPECS:
        out.extend([
            _registry_collect_view("HKLM 64-bit", "HKLM", r"SOFTWARE\Classes", prog_id, product, view64),
            _registry_collect_view("HKLM 32-bit", "HKLM", r"SOFTWARE\Classes", prog_id, product, view32),
            _registry_collect_view("HKCU user classes", "HKCU", r"Software\Classes", prog_id, product, 0),
        ])
    return out


def _collect_canoe_registry() -> List[Dict[str, Any]]:
    return _collect_vector_registry()


def _registry_clsids_for_product(registry_views: List[Dict[str, Any]], product: str) -> List[str]:
    out: List[str] = []
    for view in registry_views or []:
        if str(view.get("product") or "").lower() != product.lower():
            continue
        clsid = str(view.get("clsid") or "").strip()
        if clsid and clsid not in out:
            out.append(clsid)
        for item in view.get("versioned_progids", []) or []:
            clsid = str(item.get("clsid") or "").strip()
            if clsid and clsid not in out:
                out.append(clsid)
    return out


def _read_vector_app_info(app: Any, product_hint: str = "", prog_id: str = "", connected_by: str = "") -> Dict[str, Any]:
    out: Dict[str, Any] = {
        "product": product_hint,
        "prog_id": prog_id,
        "connected_by": connected_by,
        "version": "",
        "fullname": "",
        "configuration": "",
        "measurement_running": None,
    }
    try:
        out["version"] = str(app.Version)
    except Exception:
        pass
    try:
        out["fullname"] = str(app.FullName)
    except Exception:
        pass
    try:
        out["configuration"] = str(app.Configuration.FullName)
    except Exception:
        pass
    try:
        out["measurement_running"] = bool(app.Measurement.Running)
    except Exception as e:
        out["measurement_error"] = f"{type(e).__name__}: {e}"

    hay = (str(out.get("fullname") or "") + " " + str(out.get("version") or "")).lower()
    if "canalyzer" in hay or "canalyser" in hay:
        out["product"] = "CANalyzer"
    elif "canoe" in hay:
        out["product"] = "CANoe"
    return out


def _guarded_activate_vector_application(
    product: str,
    prog_id: str,
    before_processes: Optional[Dict[str, Any]] = None,
) -> Dict[str, Any]:
    """Activate a COM Application only when that Vector product is already running.

    Safety policy (rev87):
    - never call Dispatch when the product-specific main GUI process is absent;
    - use Dispatch, never DispatchEx;
    - snapshot product-specific main PIDs before/after activation;
    - if a new main Vector process appears, reject the returned object and block automation.

    Dispatch may activate a registered local-server COM class.  It is intentionally isolated here so
    normal diagnostics remain read-only.
    """
    result: Dict[str, Any] = {
        "product": product, "prog_id": prog_id, "attempted": False, "success": False,
        "unsafe_new_instance": False, "before_pids": [], "after_pids": [], "new_pids": [],
        "error": "", "connected_by": "", "process_snapshot_before": {}, "process_snapshot_after": {},
    }
    if not sys.platform.startswith("win"):
        result["error"] = "Windows가 아니므로 guarded COM activation을 수행하지 않음"
        return result

    before = before_processes if before_processes is not None else _collect_vector_processes()
    result["process_snapshot_before"] = before
    running_products = _vector_running_products(before)
    if product not in running_products:
        result["error"] = f"{product} 메인 프로세스가 사전에 확인되지 않아 Dispatch fallback을 차단했습니다."
        return result

    before_pids = _vector_main_process_pids(before, product)
    result["before_pids"] = sorted(before_pids)
    if not before_pids:
        result["error"] = f"{product} 제품별 메인 PID를 확인하지 못해 Dispatch fallback을 차단했습니다."
        return result

    _ensure_com_initialized()
    try:
        import win32com.client
    except Exception as e:
        result["error"] = f"win32com import 실패: {type(e).__name__}: {e}"
        return result

    result["attempted"] = True
    app = None
    try:
        # Guarded fallback only. DispatchEx is intentionally prohibited.
        app = win32com.client.Dispatch(prog_id)
        time.sleep(0.25)
        after = _collect_vector_processes()
        result["process_snapshot_after"] = after
        after_pids = _vector_main_process_pids(after, product)
        result["after_pids"] = sorted(after_pids)
        new_pids = sorted(after_pids - before_pids)
        result["new_pids"] = new_pids
        if new_pids:
            result["unsafe_new_instance"] = True
            result["error"] = (
                f"Dispatch 후 새로운 {product} 메인 프로세스가 생성되었습니다(PID={new_pids}). "
                "기존 실행 인스턴스 연결로 간주할 수 없어 자동 시험/Stimulus를 차단합니다."
            )
            app = None
            return result

        info = _read_vector_app_info(
            app, product_hint=product, prog_id=prog_id,
            connected_by="Guarded Dispatch(existing-process fallback)",
        )
        detected_product = str(info.get("product") or product)
        if detected_product and detected_product != product:
            result["error"] = f"요청 제품={product}, Dispatch 반환 제품={detected_product}로 불일치하여 차단했습니다."
            app = None
            return result
        result["success"] = True
        result["selected"] = info
        result["connected_by"] = info.get("connected_by")
        result["object"] = app
        return result
    except Exception as e:
        result["error"] = f"{type(e).__name__}: {e}"
        result["exception_chain"] = _exception_chain_details(e)
        return result


def _probe_active_vector_applications(
    preference: Any = "Auto",
    registry_views: Optional[List[Dict[str, Any]]] = None,
    include_objects: bool = False,
) -> Dict[str, Any]:
    """실행 중 CANoe/CANalyzer를 GetActiveObject only로 탐색한다.

    64-bit Python에서 32-bit 구형 제품의 ProgID 조회가 보이지 않는 경우를 위해
    Registry 32/64-bit view에서 읽은 CLSID를 사용한 GetActiveObject 재시도도 수행한다.
    새 인스턴스 생성/Dispatch/DispatchEx는 하지 않는다.
    """
    result: Dict[str, Any] = {
        "preference": normalize_vector_product_preference(preference),
        "success": False,
        "ambiguous": False,
        "selected": {},
        "successes": [],
        "attempts": [],
        "error": "",
    }
    if not sys.platform.startswith("win"):
        result["error"] = "Windows가 아니므로 GetActiveObject probe를 수행하지 않음"
        return result

    _ensure_com_initialized()
    try:
        import win32com.client
    except Exception as e:
        result["error"] = f"win32com import 실패: {type(e).__name__}: {e}"
        return result

    registry_views = registry_views if registry_views is not None else _collect_vector_registry()
    successes_with_obj: List[Tuple[Any, Dict[str, Any]]] = []

    for product, prog_id in _vector_specs_for_preference(preference):
        candidates: List[Tuple[str, str]] = [("ProgID", prog_id)]
        for clsid in _registry_clsids_for_product(registry_views, product):
            candidates.append(("Registry CLSID", clsid))

        seen_keys = set()
        for method, key in candidates:
            dedupe_key = (method, key.lower())
            if dedupe_key in seen_keys:
                continue
            seen_keys.add(dedupe_key)
            attempt = {"product": product, "prog_id": prog_id, "method": method, "key": key, "success": False}
            try:
                app = win32com.client.GetActiveObject(key)
                info = _read_vector_app_info(app, product_hint=product, prog_id=prog_id, connected_by=f"GetActiveObject({method})")
                attempt["success"] = True
                attempt.update({k: v for k, v in info.items() if k not in {"connected_by"}})
                result["attempts"].append(attempt)
                successes_with_obj.append((app, info))
                break
            except Exception as e:
                attempt["error"] = f"{type(e).__name__}: {e}"
                attempt["exception_chain"] = _exception_chain_details(e)
                result["attempts"].append(attempt)

    unique: List[Tuple[Any, Dict[str, Any]]] = []
    seen_fingerprints = set()
    for app, info in successes_with_obj:
        fp = (
            str(info.get("fullname") or "").strip().lower(),
            str(info.get("configuration") or "").strip().lower(),
            str(info.get("version") or "").strip().lower(),
        )
        if not any(fp):
            fp = (str(info.get("product") or "").lower(), str(info.get("prog_id") or "").lower(), id(app))
        if fp in seen_fingerprints:
            continue
        seen_fingerprints.add(fp)
        unique.append((app, info))

    result["successes"] = [dict(info) for _, info in unique]
    if len(unique) == 1:
        app, info = unique[0]
        result["success"] = True
        result["selected"] = dict(info)
        if include_objects:
            result["selected_object"] = app
        return result
    if len(unique) > 1:
        result["ambiguous"] = True
        result["error"] = (
            "실행 중인 CANoe와 CANalyzer COM 객체가 둘 이상 감지되었습니다. "
            "설정의 'Vector 연결 대상'을 CANoe 또는 CANalyzer로 지정하세요."
        )
        return result

    errors = []
    for a in result["attempts"]:
        if a.get("error"):
            errors.append(f"{a.get('product')} {a.get('method')}={a.get('key')}: {a.get('error')}")
    result["error"] = " | ".join(errors) if errors else "실행 중인 Vector COM 활성 객체를 찾지 못했습니다."
    return result


def _active_object_probe(preference: Any = "Auto", registry_views: Optional[List[Dict[str, Any]]] = None) -> Dict[str, Any]:
    try:
        return _probe_active_vector_applications(preference, registry_views=registry_views, include_objects=False)
    except Exception as e:
        return {"success": False, "error": f"{type(e).__name__}: {e}", "exception_chain": _exception_chain_details(e), "attempts": []}


def _extract_registered_server_paths(registry_views: List[Dict[str, Any]]) -> List[str]:
    paths: List[str] = []
    for view in registry_views or []:
        for raw in [view.get("local_server32", "")]:
            if raw and raw not in paths:
                paths.append(str(raw))
        for item in view.get("versioned_progids", []) or []:
            raw = item.get("local_server32", "")
            if raw and raw not in paths:
                paths.append(str(raw))
    return paths


def _build_vector_diag_interpretation(diag: Dict[str, Any]) -> List[str]:
    notes: List[str] = []
    processes = (diag.get("vector_processes") or {}).get("items") or []
    active = diag.get("active_object_probe") or {}
    registry = diag.get("registry") or []
    servers = _extract_registered_server_paths(registry)

    if active.get("success"):
        selected = active.get("selected") or {}
        notes.append(
            f"GetActiveObject 재확인 성공: {selected.get('product') or 'Vector'} / "
            f"{selected.get('prog_id') or '-'} 활성 객체를 찾았습니다."
        )
    elif active.get("ambiguous"):
        notes.append("CANoe/CANalyzer 활성 객체가 둘 이상 감지되어 자동 선택하지 않았습니다. 설정에서 연결 대상을 명시하세요.")
    elif processes:
        notes.append("Vector 프로세스는 감지되지만 GetActiveObject가 실패했습니다. rev87 실제 연결은 제품별 메인 프로세스가 사전 확인된 경우에만 guarded Dispatch fallback을 1회 시도합니다. 수동 진단 자체는 read-only이며 fallback을 실행하지 않습니다.")
    else:
        notes.append("진단 시점에 CANoe/CANalyzer 관련 프로세스를 찾지 못했습니다. 구형 Runtime Kernel은 프로세스 이름보다 ProductName/FileDescription/경로로도 탐색합니다.")

    products_registered = sorted({str(v.get("product")) for v in registry if v.get("clsid") or v.get("versioned_progids")})
    if not servers:
        notes.append("CANoe.Application/CANalyzer.Application의 LocalServer32 등록 경로를 찾지 못했습니다. COM 등록/Registry view를 확인하세요.")
    else:
        if products_registered:
            notes.append("Registry 감지 제품: " + ", ".join(products_registered))
        process_paths = []
        for item in processes:
            path = str(item.get("ExecutablePath") or item.get("executable_path") or "").strip()
            if path:
                process_paths.append(path.lower())
        if process_paths and not any(any(pp in server.lower() or server.lower().strip(' \"') in pp for server in servers) for pp in process_paths):
            notes.append("실행 중 Vector 프로그램 경로와 COM LocalServer32 경로가 일치하지 않아 보입니다. 다중 버전 설치/등록 상태를 확인하세요.")

    py_admin = (diag.get("python") or {}).get("is_admin")
    if py_admin is True:
        notes.append("Oracle Checker(Python)는 관리자 권한으로 실행 중입니다. Vector 프로그램도 동일한 권한 수준인지 확인하세요.")
    elif py_admin is False:
        notes.append("Oracle Checker(Python)는 일반 권한으로 실행 중입니다. Vector 프로그램도 일반 권한인지 확인하세요.")
    else:
        notes.append("Python 관리자 권한 상태를 자동 판별하지 못했습니다.")

    notes.append("이 진단은 read-only입니다. 진단 실행 자체는 Dispatch/DispatchEx/새 CANoe·CANalyzer 실행/Registry 수정/-regserver를 수행하지 않습니다. 실제 connect()에서만 기존 제품 메인 프로세스 확인 후 guarded Dispatch fallback을 제한적으로 사용할 수 있습니다.")
    return notes


def collect_vector_com_diagnostics(trigger_error: Optional[BaseException] = None, preference: Any = "Auto") -> Dict[str, Any]:
    registry = _collect_vector_registry()
    diag: Dict[str, Any] = {
        "diagnostic_version": "rev87",
        "created_at": _dt.datetime.now().isoformat(timespec="seconds"),
        "platform": platform.platform(),
        "preference": normalize_vector_product_preference(preference),
        "python": {
            "executable": sys.executable,
            "version": sys.version.replace("\n", " "),
            "bitness": 64 if sys.maxsize > 2**32 else 32,
            "pid": os.getpid(),
            "is_admin": _python_is_admin(),
        },
        "trigger_exception_chain": _exception_chain_details(trigger_error),
        "vector_processes": _collect_vector_processes(),
        "registry": registry,
        "active_object_probe": _active_object_probe(preference=preference, registry_views=registry),
    }
    diag["canoe_processes"] = diag["vector_processes"]
    diag["interpretation"] = _build_vector_diag_interpretation(diag)
    return diag


def collect_canoe_com_diagnostics(trigger_error: Optional[BaseException] = None, preference: Any = "Auto") -> Dict[str, Any]:
    return collect_vector_com_diagnostics(trigger_error=trigger_error, preference=preference)


def format_vector_com_diagnostics(diag: Dict[str, Any]) -> str:
    lines: List[str] = []
    lines.append("Oracle Checker - Vector CANoe/CANalyzer COM 연결 진단")
    lines.append("=" * 72)
    lines.append(f"진단 버전: {diag.get('diagnostic_version', '-')}")
    lines.append(f"생성 시각: {diag.get('created_at', '-')}")
    lines.append(f"연결 대상 설정: {diag.get('preference', 'Auto')}")
    lines.append(f"OS: {diag.get('platform', '-')}")
    py = diag.get("python") or {}
    lines.append(f"Python: {py.get('version', '-')}")
    lines.append(f"Python executable: {py.get('executable', '-')}")
    lines.append(f"Python bitness: {py.get('bitness', '-')} bit")
    lines.append(f"Python PID: {py.get('pid', '-')}")
    lines.append(f"Python 관리자 권한: {py.get('is_admin', '-')}")

    lines.append("")
    lines.append("[1] 최초 오류 / HRESULT")
    chain = diag.get("trigger_exception_chain") or []
    if not chain:
        lines.append("- 수동 진단 또는 최초 오류 정보 없음")
    for i, item in enumerate(chain, start=1):
        lines.append(f"- #{i} {item.get('type')}: {item.get('message')}")
        if item.get("hresult_hex"):
            lines.append(f"  HRESULT: {item.get('hresult')} ({item.get('hresult_hex')})")
        lines.append(f"  args: {item.get('args', '-')}")

    lines.append("")
    lines.append("[2] 실행 중 Vector CANoe/CANalyzer 관련 프로세스")
    proc = diag.get("vector_processes") or diag.get("canoe_processes") or {}
    lines.append(f"조회 방식: {proc.get('method', '-')}")
    if proc.get("error"):
        lines.append(f"조회 경고: {proc.get('error')}")
    items = proc.get("items") or []
    if not items:
        lines.append("- CANoe/CANalyzer 관련 프로세스 미검출")
    for item in items:
        if "raw" in item:
            lines.append(f"- {item.get('raw')}")
        else:
            lines.append(
                f"- PID={item.get('ProcessId', '-')} | Name={item.get('Name', '-')} | "
                f"Version={item.get('FileVersion', '-') } | Path={item.get('ExecutablePath', '-') }"
            )
            if item.get("ProductName") or item.get("FileDescription"):
                lines.append(f"  Product={item.get('ProductName') or '-'} | Description={item.get('FileDescription') or '-'}")
            if item.get("CommandLine"):
                lines.append(f"  CommandLine={item.get('CommandLine')}")

    lines.append("")
    lines.append("[3] CANoe.Application / CANalyzer.Application COM Registry")
    for view in diag.get("registry") or []:
        lines.append(
            f"- {view.get('product', '-')}: {view.get('prog_id', '-')} | "
            f"{view.get('view', '-')} / {view.get('root', '-')}"
        )
        if view.get("error"):
            lines.append(f"  error={view.get('error')}")
        lines.append(f"  CLSID={view.get('clsid') or '-'}")
        lines.append(f"  CurVer={view.get('curver') or '-'}")
        lines.append(f"  LocalServer32={view.get('local_server32') or '-'}")
        for item in view.get("versioned_progids", []) or []:
            lines.append(
                f"  * {item.get('progid')} | CLSID={item.get('clsid') or '-'} | "
                f"CurVer={item.get('curver') or '-'} | LocalServer32={item.get('local_server32') or '-'}"
            )

    lines.append("")
    lines.append("[4] GetActiveObject 제품별 재확인")
    active = diag.get("active_object_probe") or {}
    lines.append(f"성공 여부: {active.get('success', False)}")
    lines.append(f"다중 활성 객체 모호성: {active.get('ambiguous', False)}")
    selected = active.get("selected") or {}
    if active.get("success"):
        lines.append(f"선택 제품: {selected.get('product') or '-'}")
        lines.append(f"ProgID: {selected.get('prog_id') or '-'}")
        lines.append(f"연결 방식: {selected.get('connected_by') or '-'}")
        lines.append(f"Version: {selected.get('version') or '-'}")
        lines.append(f"FullName: {selected.get('fullname') or '-'}")
        lines.append(f"Configuration: {selected.get('configuration') or '-'}")
        lines.append(f"Measurement.Running: {selected.get('measurement_running', '-')}")
        if selected.get("measurement_error"):
            lines.append(f"Measurement 접근 오류: {selected.get('measurement_error')}")
    else:
        lines.append(f"오류: {active.get('error') or '-'}")

    for a in active.get("attempts", []) or []:
        status = "OK" if a.get("success") else "FAIL"
        lines.append(
            f"- [{status}] {a.get('product')} | {a.get('method')} | key={a.get('key')}"
        )
        if a.get("success"):
            lines.append(
                f"  Version={a.get('version') or '-'} | FullName={a.get('fullname') or '-'} | "
                f"Configuration={a.get('configuration') or '-'}"
            )
        elif a.get("error"):
            lines.append(f"  error={a.get('error')}")
            chain2 = a.get("exception_chain") or []
            if chain2:
                first = chain2[0]
                hr = f" / HRESULT={first.get('hresult_hex')}" if first.get("hresult_hex") else ""
                lines.append(f"  {first.get('type')}: {first.get('message')}{hr}")

    lines.append("")
    lines.append("[5] 자동 해석 / 현장 확인 포인트")
    for note in diag.get("interpretation") or []:
        lines.append(f"- {note}")

    lines.append("")
    lines.append("[주의] 본 진단은 read-only이며 COM 등록/Registry/Vector 프로세스를 변경하거나 새 CANoe/CANalyzer를 실행하지 않습니다.")
    return "\n".join(lines) + "\n"


def format_canoe_com_diagnostics(diag: Dict[str, Any]) -> str:
    return format_vector_com_diagnostics(diag)


def save_vector_com_diagnostic_report(
    output_dir: str | Path,
    diag: Dict[str, Any],
    prefix: str = "vector_com_diagnostic",
) -> Path:
    stamp = _dt.datetime.now().strftime("%Y%m%d_%H%M%S_%f")[:-3]
    report_text = format_vector_com_diagnostics(diag)
    primary = Path(output_dir)
    try:
        primary.mkdir(parents=True, exist_ok=True)
        path = primary / f"{prefix}_{stamp}.txt"
        path.write_text(report_text, encoding="utf-8-sig")
        return path
    except Exception as primary_error:
        fallback = Path(tempfile.gettempdir()) / "OracleCheckerDiagnostics"
        fallback.mkdir(parents=True, exist_ok=True)
        path = fallback / f"{prefix}_{stamp}.txt"
        path.write_text(
            report_text + f"\n[저장 참고] 프로그램 폴더 저장 실패: {type(primary_error).__name__}: {primary_error}\n",
            encoding="utf-8-sig",
        )
        return path


def save_canoe_com_diagnostic_report(
    output_dir: str | Path,
    diag: Dict[str, Any],
    prefix: str = "vector_com_diagnostic",
) -> Path:
    return save_vector_com_diagnostic_report(output_dir, diag, prefix=prefix)


# =========================================================
# CAN network name / DBC auto mapping utilities (rev87)
# =========================================================
_CAN_LOGICAL_RE = re.compile(r"(?<![A-Z0-9])CAN\s*([1-4])(?![0-9])", re.IGNORECASE)
_CH_LOGICAL_RE = re.compile(r"(?<![A-Z0-9])CH(?:ANNEL)?\s*([1-4])(?![0-9])", re.IGNORECASE)
_CHANNEL_NETWORK_PAIR_RE = re.compile(
    r"(?<![A-Z0-9])(?:CAN|CH(?:ANNEL)?)\s*([1-4])(?![0-9])"
    r"[\s:=><|/_-]{0,16}"
    r"([A-Z][A-Z0-9]{0,12}[\s_.-]*CAN)(?![A-Z0-9])",
    re.IGNORECASE,
)
_NETWORK_TOKEN_RE = re.compile(
    r"(?<![A-Z0-9])([A-Z][A-Z0-9]{0,12}[\s_.-]*CAN)(?![A-Z0-9])",
    re.IGNORECASE,
)


def normalize_can_network_name(raw: Any) -> str:
    """네트워크 표시명을 비교용 표준 이름으로 정규화한다.

    예: P CAN/P-CAN/P1_CAN -> P1-CAN, M_CAN -> M-CAN.
    사용자 운용상 P-CAN과 P1-CAN은 동일 네트워크 별칭으로 취급한다.
    """
    text = str(raw or "").strip().upper()
    if not text:
        return ""
    text = text.strip("'\"[](){}")
    compact = re.sub(r"[^A-Z0-9]", "", text)
    if not compact:
        return ""

    aliases = {
        "PCAN": "P1-CAN",
        "P1CAN": "P1-CAN",
        "BCAN": "B-CAN",
        "B1CAN": "B1-CAN",
        "B2CAN": "B2-CAN",
        "ECAN": "E-CAN",
        "MCAN": "M-CAN",
    }
    if compact in aliases:
        return aliases[compact]

    # 일반적인 XXX-CAN 명명도 지원한다. 단 CAN1/CAN2 같은 논리 채널명은 제외한다.
    if compact.startswith("CAN") and compact[3:].isdigit():
        return ""
    if compact.endswith("CAN") and len(compact) > 3:
        prefix = compact[:-3]
        if prefix and not prefix.isdigit():
            return f"{prefix}-CAN"
    return ""


def _network_match_keys(network_name: str) -> List[str]:
    norm = normalize_can_network_name(network_name)
    if not norm:
        return []
    compact = re.sub(r"[^A-Z0-9]", "", norm.upper())
    keys = [compact]
    # 사용자 운용 정책: P-CAN == P1-CAN
    if compact == "P1CAN":
        keys.append("PCAN")
    return keys


def _candidate_network_tokens(text: str) -> List[str]:
    out: List[str] = []
    for m in _NETWORK_TOKEN_RE.finditer(str(text or "")):
        norm = normalize_can_network_name(m.group(1))
        if norm and norm not in out:
            out.append(norm)
    return out


def extract_can_network_mappings_from_text(text: str, source_label: str = "") -> Dict[str, Any]:
    """텍스트/파일명에서 CAN1~CAN4와 네트워크명 동시 등장 근거를 추출한다.

    추론은 하지 않는다. 같은 줄/짧은 문맥에 CANx(또는 CHx)와 명시적 XXX-CAN 토큰이
    함께 있을 때만 evidence로 인정한다. 동일 채널에 서로 다른 이름이 잡히면 conflict로 남긴다.
    """
    evidence: Dict[str, List[Dict[str, str]]] = {f"CAN{i}": [] for i in range(1, 5)}
    lines = str(text or "").replace("\r\n", "\n").replace("\r", "\n").split("\n")

    for line_no, line in enumerate(lines, start=1):
        if not line.strip():
            continue

        # 1) CAN1=P1-CAN / CAN1_P1-CAN_CAN2_M-CAN 같은 명시적 pair를 우선 추출한다.
        paired = []
        for m in _CHANNEL_NETWORK_PAIR_RE.finditer(line):
            ch = int(m.group(1))
            network = normalize_can_network_name(m.group(2))
            if network:
                paired.append((ch, network))
                can_name = f"CAN{ch}"
                item = {
                    "network": network,
                    "source": source_label or "text",
                    "line": str(line_no),
                    "evidence": line.strip()[:500],
                }
                if item not in evidence[can_name]:
                    evidence[can_name].append(item)
        if paired:
            continue

        # 2) 명시 pair가 없으면 한 줄에 채널 1개 + 네트워크 1개일 때만 근거로 인정한다.
        channels = set()
        for m in _CAN_LOGICAL_RE.finditer(line):
            channels.add(int(m.group(1)))
        for m in _CH_LOGICAL_RE.finditer(line):
            channels.add(int(m.group(1)))
        networks = _candidate_network_tokens(line)
        if len(channels) != 1 or len(networks) != 1:
            continue
        ch = next(iter(channels))
        network = networks[0]
        can_name = f"CAN{ch}"
        item = {
            "network": network,
            "source": source_label or "text",
            "line": str(line_no),
            "evidence": line.strip()[:500],
        }
        if item not in evidence[can_name]:
            evidence[can_name].append(item)

    mapping: Dict[str, str] = {}
    conflicts: Dict[str, List[str]] = {}
    for can_name, items in evidence.items():
        names = []
        for item in items:
            n = item.get("network", "")
            if n and n not in names:
                names.append(n)
        if len(names) == 1:
            mapping[can_name] = names[0]
        elif len(names) > 1:
            conflicts[can_name] = names

    return {
        "source": source_label,
        "mapping": mapping,
        "conflicts": conflicts,
        "evidence": evidence,
    }


def read_text_probe(path: str | Path, max_bytes: int = 5_000_000, max_lines: int = 1200) -> str:
    """ASC/CANoe cfg 등에서 네트워크명 탐색에 필요한 범위만 best-effort로 읽는다."""
    p = Path(path)
    if not p.is_file():
        return ""
    try:
        raw = p.read_bytes()[: max(1, int(max_bytes))]
    except Exception:
        return ""

    candidates: List[str] = []
    for enc in ("utf-8-sig", "utf-16", "cp949", "cp1252", "latin1"):
        try:
            decoded = raw.decode(enc, errors="ignore")
            if decoded and decoded not in candidates:
                candidates.append(decoded)
        except Exception:
            pass
    if not candidates:
        return ""

    # CAN/Network 관련 문자가 가장 많이 살아있는 decode 결과를 선택한다.
    def score(t: str) -> int:
        up = t.upper()
        return up.count("CAN1") + up.count("CAN2") + up.count("CAN3") + up.count("CAN4") + up.count("-CAN")
    text = max(candidates, key=score)
    if max_lines > 0:
        text = "\n".join(text.splitlines()[:max_lines])
    return text


def find_dbc_candidates_for_network(dbc_dir: str | Path, network_name: str) -> List[Path]:
    """DBC 폴더에서 네트워크명과 파일명이 명시적으로 일치하는 후보만 반환한다."""
    folder = Path(dbc_dir)
    if not folder.is_dir():
        return []
    keys = _network_match_keys(network_name)
    if not keys:
        return []
    found: List[Path] = []
    try:
        dbcs = sorted((x for x in folder.rglob("*.dbc") if x.is_file()), key=lambda x: str(x).lower())
    except Exception:
        return []
    for dbc in dbcs:
        stem_compact = re.sub(r"[^A-Z0-9]", "", dbc.stem.upper())
        if any(key and key in stem_compact for key in keys):
            found.append(dbc)
    return found


def _active_canoe_configuration_path_readonly() -> Dict[str, Any]:
    """GetActiveObject only로 실행 중 CANoe/CANalyzer의 Configuration.FullName을 best-effort 조회한다."""
    out: Dict[str, Any] = {"success": False, "configuration": "", "product": "", "error": ""}
    if not sys.platform.startswith("win"):
        out["error"] = "Windows가 아님"
        return out
    try:
        probe = _probe_active_vector_applications("Auto", include_objects=True)
        if not probe.get("success"):
            out["error"] = str(probe.get("error") or "Vector 활성 객체 없음")
            return out
        app = probe.get("selected_object")
        info = probe.get("selected") or {}
        out["configuration"] = str(app.Configuration.FullName)
        out["product"] = str(info.get("product") or "")
        out["success"] = bool(out["configuration"])
    except Exception as e:
        out["error"] = f"{type(e).__name__}: {e}"
    return out


def auto_map_dbc_by_network_sources(
    dbc_dir: str | Path,
    source_paths: Optional[List[str | Path]] = None,
    source_texts: Optional[List[Tuple[str, str]]] = None,
    include_active_canoe_config: bool = True,
) -> Dict[str, Any]:
    """네트워크명 근거를 수집해 CAN1~CAN3 DBC 자동 매핑 후보를 만든다.

    안전 정책:
    - DBC 내용 유사도나 Message 겹침으로 네트워크를 추정하지 않는다.
    - 동일 CAN에 서로 다른 네트워크 근거가 있으면 conflict 처리한다.
    - DBC 파일 후보가 정확히 1개일 때만 selected에 넣는다.
    - 0개/복수 후보는 unresolved로 남긴다.
    """
    folder = Path(dbc_dir)
    result: Dict[str, Any] = {
        "dbc_dir": str(folder),
        "network_by_can": {},
        "evidence_by_can": {f"CAN{i}": [] for i in range(1, 4)},
        "conflicts": {},
        "dbc_candidates_by_can": {},
        "selected_dbc_by_can": {},
        "unresolved": {},
        "sources": [],
        "active_canoe_config": {},
    }
    if not folder.is_dir():
        result["unresolved"]["GLOBAL"] = f"DBC 자동 설정 폴더가 유효하지 않음: {folder}"
        return result

    all_source_texts: List[Tuple[str, str]] = list(source_texts or [])
    for raw_path in source_paths or []:
        p = Path(raw_path)
        if not p.is_file():
            continue
        label = str(p)
        # 파일명 자체에도 CAN1/P1-CAN 같은 명시 정보가 있을 수 있어 함께 검사한다.
        all_source_texts.append((f"filename:{p.name}", p.name))
        text = read_text_probe(p)
        if text:
            all_source_texts.append((label, text))

    if include_active_canoe_config:
        active = _active_canoe_configuration_path_readonly()
        result["active_canoe_config"] = active
        cfg = str(active.get("configuration") or "").strip()
        if cfg:
            p = Path(cfg)
            all_source_texts.append((f"configuration-filename:{p.name}", p.name))
            text = read_text_probe(p)
            if text:
                all_source_texts.append((f"configuration:{p}", text))

    aggregate: Dict[str, List[Dict[str, str]]] = {f"CAN{i}": [] for i in range(1, 4)}
    for label, text in all_source_texts:
        parsed = extract_can_network_mappings_from_text(text, source_label=label)
        result["sources"].append({
            "label": label,
            "mapping": parsed.get("mapping") or {},
            "conflicts": parsed.get("conflicts") or {},
        })
        for can_name in aggregate:
            for item in (parsed.get("evidence") or {}).get(can_name, []) or []:
                if item not in aggregate[can_name]:
                    aggregate[can_name].append(item)

    for can_name, items in aggregate.items():
        names: List[str] = []
        for item in items:
            name = normalize_can_network_name(item.get("network"))
            if name and name not in names:
                names.append(name)
        result["evidence_by_can"][can_name] = items
        if len(names) == 1:
            result["network_by_can"][can_name] = names[0]
        elif len(names) > 1:
            result["conflicts"][can_name] = names
            result["unresolved"][can_name] = "네트워크명 근거 충돌: " + ", ".join(names)
        else:
            result["unresolved"][can_name] = "CAN 네트워크명을 명시적으로 식별할 근거 없음"

    for can_name, network in result["network_by_can"].items():
        if can_name in result["conflicts"]:
            continue
        candidates = find_dbc_candidates_for_network(folder, network)
        result["dbc_candidates_by_can"][can_name] = [str(x) for x in candidates]
        if len(candidates) == 1:
            result["selected_dbc_by_can"][can_name] = str(candidates[0])
            result["unresolved"].pop(can_name, None)
        elif len(candidates) == 0:
            result["unresolved"][can_name] = f"{network} 일치 DBC 후보 없음"
        else:
            result["unresolved"][can_name] = f"{network} 일치 DBC 후보 {len(candidates)}개 - 자동 선택 안 함"

    return result


# =========================================================
# Vector CANoe/CANalyzer COM Client
# =========================================================
def _ensure_com_initialized() -> None:
    """현재 스레드에서 pywin32 COM 사용 준비."""
    if not sys.platform.startswith("win"):
        raise RuntimeError("Vector CANoe/CANalyzer COM 연결은 Windows 환경에서만 사용할 수 있습니다.")

    if getattr(_COM_THREAD_STATE, "initialized", False):
        return

    try:
        import pythoncom
    except Exception as e:
        raise ImportError("pywin32가 필요합니다. py -m pip install pywin32") from e

    pythoncom.CoInitialize()
    _COM_THREAD_STATE.initialized = True


class VectorApplicationClient:
    """CANoe/CANalyzer 공통 COM 연결 래퍼.

    rev87 연결 정책:
    - GetActiveObject를 항상 우선한다.
    - Auto/CANoe/CANalyzer 제품 선택을 지원한다.
    - ProgID 우선 탐색 후 Registry 32/64-bit view의 CLSID로 GetActiveObject를 재시도한다.
    - GetActiveObject 실패 시 해당 제품의 메인 프로세스가 이미 실행 중인 경우에만 guarded Dispatch fallback을 허용한다.
    - Dispatch 전후 제품별 메인 PID를 비교하고 새 인스턴스가 생기면 자동 시험/Stimulus를 차단한다.
    - CANoe/CANalyzer가 둘 다 실행 중이고 Auto라면 임의 선택하지 않는다.
    """

    DEFAULT_EXPECTED_VERSION_KEYWORD = None
    DEFAULT_EXPECTED_EXE_KEYWORD = None

    def __init__(self, bus_name: str = "CAN", product_preference: Any = "Auto"):
        self.bus_name = (bus_name or "CAN").strip() or "CAN"
        self.product_preference = normalize_vector_product_preference(product_preference)
        self.app = None
        self._connected_by = ""
        self._app_product = ""
        self._app_progid = ""
        self._app_version = ""
        self._app_fullname = ""
        self._cfg_fullname = ""
        # rev87: CAPLFunction objects must be acquired during Measurement.OnInit.
        self._measurement_event_sink = None

    def _read_application_info(self) -> Dict[str, str]:
        info = _read_vector_app_info(
            self.app,
            product_hint=self._app_product,
            prog_id=self._app_progid,
            connected_by=self._connected_by,
        ) if self.app is not None else {
            "product": self._app_product,
            "prog_id": self._app_progid,
            "connected_by": self._connected_by,
            "version": "",
            "fullname": "",
            "configuration": "",
        }
        self._app_product = str(info.get("product") or self._app_product or "")
        self._app_progid = str(info.get("prog_id") or self._app_progid or "")
        self._app_version = str(info.get("version") or "")
        self._app_fullname = str(info.get("fullname") or "")
        self._cfg_fullname = str(info.get("configuration") or "")
        return {
            "product": self._app_product,
            "prog_id": self._app_progid,
            "product_preference": self.product_preference,
            "connected_by": self._connected_by,
            "version": self._app_version,
            "fullname": self._app_fullname,
            "configuration": self._cfg_fullname,
        }

    def get_application_info(self) -> Dict[str, str]:
        return self._read_application_info()

    def connect(
        self,
        expected_version_keyword: Optional[str] = None,
        expected_exe_keyword: Optional[str] = None,
        expected_cfg_path: Optional[str | Path] = None,
        product_preference: Any = None,
    ) -> "VectorApplicationClient":
        """GetActiveObject 우선, 기존 프로세스 확인 기반 guarded Dispatch fallback으로 연결한다."""
        _ensure_com_initialized()
        if product_preference is not None:
            self.product_preference = normalize_vector_product_preference(product_preference)

        probe = _probe_active_vector_applications(self.product_preference, include_objects=True)
        selected = {}
        selected_object = None

        if probe.get("success"):
            selected_object = probe.get("selected_object")
            selected = probe.get("selected") or {}
        else:
            # rev87 guarded fallback: only when a product-specific main process already exists.
            processes_before = _collect_vector_processes()
            running_products = _vector_running_products(processes_before)
            pref = normalize_vector_product_preference(self.product_preference)
            activation_products: List[str] = []
            if pref == "Auto":
                if len(running_products) == 1:
                    activation_products = list(running_products)
                elif len(running_products) > 1:
                    self.app = None
                    raise RuntimeError(
                        "CANoe와 CANalyzer 메인 프로세스가 동시에 실행 중이며 GetActiveObject가 모두 실패했습니다.\n"
                        "안전상 Dispatch fallback 대상을 임의 선택하지 않습니다. [Vector 연결 대상]을 CANoe 또는 CANalyzer로 명시하세요."
                    )
            else:
                if pref in running_products:
                    activation_products = [pref]

            activation = None
            for product in activation_products:
                prog_id = dict(VECTOR_PRODUCT_SPECS).get(product, f"{product}.Application")
                activation = _guarded_activate_vector_application(
                    product, prog_id, before_processes=processes_before
                )
                if activation.get("success"):
                    selected_object = activation.get("object")
                    selected = activation.get("selected") or {}
                    break

            if selected_object is None:
                self.app = None
                attempts = probe.get("attempts") or []
                detail_lines = []
                for a in attempts:
                    if a.get("success"):
                        continue
                    detail_lines.append(
                        f"- {a.get('product')} {a.get('method')} ({a.get('key')}): {a.get('error') or '-'}"
                    )
                if activation is not None:
                    detail_lines.append(
                        f"- Guarded Dispatch {activation.get('product')}: {activation.get('error') or '실패'}"
                    )
                detail = "\n".join(detail_lines[-8:])
                extra = ""
                if not activation_products:
                    if pref == "Auto":
                        extra = "\n제품별 메인 Vector 프로세스를 하나로 특정하지 못해 Dispatch fallback은 실행하지 않았습니다."
                    else:
                        extra = f"\n{pref} 메인 프로세스가 사전 확인되지 않아 Dispatch fallback은 실행하지 않았습니다."
                raise RuntimeError(
                    "실행 중인 Vector CANoe/CANalyzer COM 객체에 연결하지 못했습니다.\n"
                    f"연결 대상 설정: {self.product_preference}\n"
                    "GetActiveObject를 우선 시도했고, 허용 조건을 만족한 경우에만 guarded Dispatch fallback을 시도했습니다.\n"
                    "사용하려는 Vector 프로그램을 직접 실행하고 Configuration을 연 상태인지 확인하세요."
                    + extra
                    + f"\n원인: {probe.get('error') or '-'}"
                    + (f"\n탐색 상세:\n{detail}" if detail else "")
                )

        self.app = selected_object
        self._connected_by = str(selected.get("connected_by") or "GetActiveObject")
        self._app_product = str(selected.get("product") or "")
        self._app_progid = str(selected.get("prog_id") or "")

        expected_version_keyword = (
            expected_version_keyword
            if expected_version_keyword is not None
            else self.DEFAULT_EXPECTED_VERSION_KEYWORD
        )
        expected_exe_keyword = (
            expected_exe_keyword
            if expected_exe_keyword is not None
            else self.DEFAULT_EXPECTED_EXE_KEYWORD
        )

        self._validate_application(
            expected_version_keyword=expected_version_keyword,
            expected_exe_keyword=expected_exe_keyword,
            expected_cfg_path=expected_cfg_path,
        )
        return self

    def _validate_application(
        self,
        expected_version_keyword: Optional[str] = None,
        expected_exe_keyword: Optional[str] = None,
        expected_cfg_path: Optional[str | Path] = None,
    ) -> None:
        if self.app is None:
            raise RuntimeError("Vector CANoe/CANalyzer COM 객체가 없습니다.")

        try:
            measurement = self.app.Measurement
            _ = bool(measurement.Running)
        except Exception as e:
            self.app = None
            raise RuntimeError(
                "Vector COM 객체는 찾았지만 Measurement 상태를 읽지 못했습니다.\n"
                "CANoe/CANalyzer가 정상 실행 중인지, Configuration이 열린 상태인지, "
                "Python과 Vector 프로그램의 관리자 권한이 서로 다른지 확인하세요.\n"
                f"원인: {type(e).__name__}: {e}"
            ) from e

        info = self._read_application_info()
        version = info.get("version", "")
        fullname = info.get("fullname", "")
        cfg_fullname = info.get("configuration", "")
        product = info.get("product", "Vector") or "Vector"

        if expected_version_keyword and expected_version_keyword not in version:
            self.app = None
            raise RuntimeError(
                f"{product} 버전이 예상과 다릅니다.\n"
                f"예상 버전 키워드: {expected_version_keyword}\n"
                f"실제 Version: {version or '-'}\n"
                f"실제 실행 파일: {fullname or '-'}\n"
                f"현재 Configuration: {cfg_fullname or '-'}"
            )

        if expected_exe_keyword and expected_exe_keyword.lower() not in (fullname or "").lower():
            self.app = None
            raise RuntimeError(
                f"{product} 실행 파일 경로가 예상과 다릅니다.\n"
                f"예상 경로 키워드: {expected_exe_keyword}\n"
                f"실제 실행 파일: {fullname or '-'}\n"
                f"실제 Version: {version or '-'}\n"
                f"현재 Configuration: {cfg_fullname or '-'}"
            )

        if expected_cfg_path:
            try:
                expected_cfg_norm = str(Path(expected_cfg_path).resolve()).lower()
                actual_cfg_norm = str(Path(cfg_fullname).resolve()).lower()
            except Exception:
                expected_cfg_norm = str(expected_cfg_path).strip().lower()
                actual_cfg_norm = str(cfg_fullname).strip().lower()
            if expected_cfg_norm != actual_cfg_norm:
                self.app = None
                raise RuntimeError(
                    f"{product} Configuration이 예상과 다릅니다.\n"
                    f"예상 cfg: {expected_cfg_path}\n"
                    f"실제 cfg: {cfg_fullname or '-'}"
                )

    def is_connected(self) -> bool:
        if self.app is None:
            return False
        try:
            _ = bool(self.app.Measurement.Running)
            return True
        except Exception:
            self.app = None
            return False

    def is_measurement_running(self) -> bool:
        if not self.is_connected():
            raise RuntimeError("Vector CANoe/CANalyzer 연결이 없습니다.")
        return bool(self.app.Measurement.Running)

    def start_measurement(self, wait_sec: float = 0.0, start_timeout_sec: float = 5.0) -> None:
        if not self.is_connected():
            raise RuntimeError("Vector CANoe/CANalyzer 연결이 없습니다.")
        measurement = self.app.Measurement
        if not bool(measurement.Running):
            measurement.Start()
            deadline = time.time() + max(0.5, float(start_timeout_sec))
            while time.time() <= deadline:
                try:
                    if bool(measurement.Running):
                        break
                except Exception:
                    pass
                time.sleep(0.1)
            if not bool(measurement.Running):
                raise RuntimeError("Vector Measurement 시작 요청 후에도 Running=True가 되지 않았습니다.")
        if wait_sec and float(wait_sec) > 0:
            time.sleep(float(wait_sec))

    def compile_existing_capl_nodes(self) -> None:
        """Compile CAPL programs already attached to the current Vector configuration.

        rev87 ACTIVE-bridge workflow intentionally does not create/insert Simulation Setup nodes.
        CANoe/CANalyzer COM exposes CAPL compilation for nodes that already exist in the loaded
        configuration; node creation/source-path assignment remains a one-time configuration step.
        """
        _ensure_com_initialized()
        if not self.is_connected():
            raise RuntimeError("Vector CANoe/CANalyzer 연결이 없습니다.")
        capl = self.app.CAPL
        errors = []
        for args in ((), (None,)):
            try:
                capl.Compile(*args)
                return
            except Exception as e:
                errors.append(f"Compile{args}: {type(e).__name__}: {e}")
        detail = " | ".join(errors)
        migration_hint = ""
        if "PassFail_" in detail and "PF_Stimulus_Bridge_" in detail and "ACTIVE.can" in detail:
            migration_hint = (
                "\n[rev87 경로 마이그레이션] 현재 CANoe/CANalyzer Network Node가 과거 PassFail 패키지 폴더의 "
                "PF_Stimulus_Bridge_CANx_ACTIVE.can을 참조하는 것으로 보입니다. "
                "rev87부터는 PassFail 상위 폴더의 PF_Stimulus_Runtime/PF_Stimulus_Bridge_CANx_ACTIVE.can을 "
                "최초 1회 다시 연결/Compile하세요. 이후 PassFail 패키지 버전이 바뀌어도 이 경로는 유지됩니다."
            )
        raise RuntimeError("기존 CAPL Node 자동 Compile 실패: " + detail + migration_hint)

    def start_measurement_with_capl_bindings(
        self,
        function_names: Iterable[str],
        start_timeout_sec: float = 5.0,
        restart_if_running: bool = True,
    ) -> Dict[str, Any]:
        """Start a fresh Measurement and acquire CAPLFunction objects inside Measurement.OnInit.

        Vector AN-AND-1-117 states that CAPL.GetFunction assignment must be done from
        Measurement.OnInit. rev87 uses this helper for isolated TX/RX Stimulus and the optional
        read-only CAPL Event Collector.
        """
        _ensure_com_initialized()
        if self.app is None:
            raise RuntimeError("Vector CANoe/CANalyzer COM 연결이 없습니다.")
        names = []
        for raw in function_names or []:
            name = str(raw or "").strip()
            if name and name not in names:
                names.append(name)
        if not names:
            raise RuntimeError("OnInit에서 획득할 CAPL 함수 이름이 없습니다.")
        try:
            import pythoncom
            import win32com.client
        except Exception as e:
            raise ImportError("CAPL OnInit event 연결에는 pywin32가 필요합니다.") from e

        measurement = self.app.Measurement
        if bool(measurement.Running):
            if not restart_if_running:
                raise RuntimeError("CAPL 함수는 Measurement.OnInit에서 획득해야 하므로 실행 중 Measurement를 재시작해야 합니다.")
            self.stop_measurement()

        state: Dict[str, Any] = {"done": False, "functions": {}, "errors": []}
        capl = self.app.CAPL

        class _MeasurementInitEvents:
            def OnInit(self):
                try:
                    funcs = {}
                    errors = []
                    for fname in list(getattr(self, "_pf_names", []) or []):
                        try:
                            fn = self._pf_capl.GetFunction(fname)
                            if fn is None:
                                raise RuntimeError("GetFunction returned None")
                            funcs[fname] = fn
                        except Exception as ex:
                            errors.append(f"{fname}: {type(ex).__name__}: {ex}")
                    self._pf_state["functions"] = funcs
                    self._pf_state["errors"] = errors
                except Exception as ex:
                    self._pf_state["errors"] = [f"OnInit handler: {type(ex).__name__}: {ex}"]
                finally:
                    self._pf_state["done"] = True

        try:
            sink = win32com.client.WithEvents(measurement, _MeasurementInitEvents)
            sink._pf_names = list(names)
            sink._pf_capl = capl
            sink._pf_state = state
            self._measurement_event_sink = sink
        except Exception as e:
            raise RuntimeError(
                "Measurement.OnInit COM event handler 연결에 실패했습니다. "
                "Vector COM type library/pywin32 event 지원 상태를 확인하세요. "
                f"원인={type(e).__name__}: {e}"
            ) from e

        measurement.Start()
        deadline = time.time() + max(0.5, float(start_timeout_sec))
        while time.time() < deadline:
            try:
                pythoncom.PumpWaitingMessages()
            except Exception:
                pass
            if state.get("done") and bool(measurement.Running):
                break
            time.sleep(0.01)

        if not state.get("done"):
            try:
                self.stop_measurement()
            except Exception:
                pass
            raise RuntimeError(
                "Measurement.OnInit 이벤트를 제한시간 내 받지 못했습니다. "
                "CAPL node가 Compile된 Configuration인지 확인하고 다시 시도하세요."
            )
        errors = list(state.get("errors") or [])
        if errors:
            try:
                self.stop_measurement()
            except Exception:
                pass
            raise RuntimeError(
                "Measurement.OnInit에서 요청한 CAPL 함수를 획득하지 못했습니다. "
                "사용 중인 CAPL Stimulus Bridge 또는 Event Collector .can 파일이 현재 Configuration에 "
                "삽입/Compile되어 있는지 확인하세요.\n- " + "\n- ".join(errors[:12])
            )
        funcs = dict(state.get("functions") or {})
        missing = [x for x in names if x not in funcs]
        if missing:
            raise RuntimeError("OnInit CAPL 함수 일부가 누락되었습니다: " + ", ".join(missing[:12]))
        return funcs

    def stop_measurement(self, stop_timeout_sec: float = 5.0) -> None:
        if not self.is_connected():
            raise RuntimeError("Vector CANoe/CANalyzer 연결이 없습니다.")
        measurement = self.app.Measurement
        if bool(measurement.Running):
            measurement.Stop()
            deadline = time.time() + max(0.5, float(stop_timeout_sec))
            while time.time() <= deadline:
                try:
                    if not bool(measurement.Running):
                        break
                except Exception:
                    break
                time.sleep(0.1)
        self._measurement_event_sink = None

    def _bus_name_candidates(self) -> List[str]:
        raw = (self.bus_name or "CAN").strip()
        candidates: List[str] = []
        for name in (raw, raw.replace(" ", ""), "CAN"):
            if name and name not in candidates:
                candidates.append(name)
        return candidates

    def get_signal_value(self, channel: int, message_name: str, signal_name: str) -> Any:
        if not self.is_connected():
            raise RuntimeError("Vector CANoe/CANalyzer 연결이 없습니다.")

        last_error = None
        ch = int(channel)
        msg = str(message_name).strip()
        sig = str(signal_name).strip()

        for bus_name in self._bus_name_candidates():
            try:
                bus = self.app.GetBus(bus_name)
                signal = bus.GetSignal(ch, msg, sig)
                return signal.Value
            except Exception as e:
                last_error = e
                continue

        raise RuntimeError(
            f"Signal 읽기 실패: bus 후보={self._bus_name_candidates()}, channel={ch}, "
            f"message={msg}, signal={sig}, 원인={last_error}"
        )


# 기존 모듈/타입 힌트 호환: 이름은 유지하되 구현은 제품 중립 Client를 사용한다.
CANoeClient = VectorApplicationClient

def normalize_value(value: Any) -> Any:
    if value is None:
        return None

    if isinstance(value, bool):
        return int(value)

    if isinstance(value, int):
        return value

    if isinstance(value, float):
        if math.isfinite(value) and value.is_integer():
            return int(value)
        return value

    s = str(value).strip()
    if not s:
        return None

    s = s.replace(" ", "")

    try:
        if _HEX_RE.match(s):
            return int(s, 16)
        if _INT_RE.match(s):
            return int(s, 10)
        if _FLOAT_RE.match(s):
            f = float(s)
            if f.is_integer():
                return int(f)
            return f
    except Exception:
        pass

    return s


def values_equal(expected: Any, observed: Any) -> bool:
    e = normalize_value(expected)
    o = normalize_value(observed)

    if isinstance(e, float) or isinstance(o, float):
        try:
            return abs(float(e) - float(o)) < 1e-9
        except Exception:
            return False

    return e == o


def split_multiline_cell(value: Any) -> List[str]:
    if value is None:
        return []
    text = str(value).replace("\r\n", "\n").replace("\r", "\n")
    return [line.strip() for line in text.split("\n") if line.strip()]


def join_expectation_field(expectations: List["SignalExpectation"], field_name: str) -> str:
    out = []
    for exp in expectations:
        out.append(str(getattr(exp, field_name, "") or ""))
    return "\n".join(out)


# =========================================================
# 신호 조건 데이터 구조
# =========================================================
@dataclass
class SignalExpectation:
    message: str
    signal: str
    expected_value_raw: str

    @property
    def expected_value_norm(self):
        return normalize_value(self.expected_value_raw)


@dataclass
class LogReviewSample:
    timestamp: float
    value: Any


@dataclass
class LogReviewOutcome:
    status: str
    message: str
    source: str = "log_review"
    observed_value_raw: Optional[Any] = None
    observed_value_norm: Optional[Any] = None
    observed_time: Optional[float] = None
    elapsed_sec: Optional[float] = None
    asc_path: Optional[str] = None
    matched_channel: Optional[int] = None
    matched_frame_id: Optional[int] = None
    samples: List[Tuple[float, Any]] = field(default_factory=list)


@dataclass
class SingleExpectationCheck:
    expectation: SignalExpectation
    status: str = "대기중"
    observed_value_raw: Optional[Any] = None
    observed_value_norm: Optional[Any] = None
    observed_time: Optional[float] = None
    observed_channel: Optional[int] = None
    message: str = ""
    samples: List[Tuple[float, Any]] = field(default_factory=list)
    observation_source: str = ""
    collector_hit_ms: Optional[float] = None


@dataclass
class OracleCondition:
    row_index: int
    tc_no: str
    subcategory: str
    tc_content: str
    tc_expected_result: str

    input_conditions: List[SignalExpectation] = field(default_factory=list)
    output_conditions: List[SignalExpectation] = field(default_factory=list)

    judge_mode: str = "once"
    timeout_sec: float = 30.0
    use_yn: str = "Y"
    note: str = ""
    condition_id: str = ""
    tc_logic: str = "all"
    revision: str = ""

    @property
    def display_id(self) -> str:
        if self.condition_id:
            return self.condition_id
        return f"{self.tc_no}-{self.row_index}"

    @property
    def enabled(self) -> bool:
        s = str(self.use_yn or "Y").strip().upper()
        return s not in ("N", "NO", "FALSE", "0", "미사용")


@dataclass
class OracleRowLoadResult:
    row_index: int
    tc_no: str = ""
    subcategory: str = ""
    tc_content: str = ""
    tc_expected_result: str = ""

    input_message: str = ""
    input_signal: str = ""
    input_expected_value_raw: str = ""

    output_message_raw: str = ""
    output_signal_raw: str = ""
    output_expected_value_raw: str = ""

    is_valid: bool = False
    status: str = "ERROR"
    error_reason: str = ""
    condition: Optional[OracleCondition] = None


@dataclass
class CheckResult:
    condition: OracleCondition
    status: str = "대기중"
    source: str = ""
    started_at: Optional[_dt.datetime] = None
    ended_at: Optional[_dt.datetime] = None
    elapsed_sec: Optional[float] = None
    message: str = ""

    input_results: List[SingleExpectationCheck] = field(default_factory=list)
    output_results: List[SingleExpectationCheck] = field(default_factory=list)

    observed_value_raw: Optional[Any] = None
    observed_value_norm: Optional[Any] = None
    observed_time: Optional[float] = None
    observed_channel: Optional[int] = None
    samples: List[Tuple[float, Any]] = field(default_factory=list)

    def to_row(self) -> List[Any]:
        c = self.condition

        return [
            c.tc_no,
            c.subcategory,
            c.tc_content,
            c.tc_expected_result,
            join_expectation_field(c.input_conditions, "message"),
            join_expectation_field(c.input_conditions, "signal"),
            join_expectation_field(c.input_conditions, "expected_value_raw"),
            join_expectation_field(c.output_conditions, "message"),
            join_expectation_field(c.output_conditions, "signal"),
            join_expectation_field(c.output_conditions, "expected_value_raw"),
            c.judge_mode,
            c.timeout_sec,
            c.use_yn,
            c.note,
            c.condition_id,
            c.tc_logic,
            c.revision,
            self.status,
            self.source,
            self.observed_value_raw,
            self.observed_time,
            self.observed_channel,
            self.elapsed_sec,
            self.message,
        ]


@dataclass
class FinalCheckResult:
    condition: OracleCondition
    realtime_result: Optional[CheckResult] = None
    log_result: Optional[CheckResult] = None
    final_status: str = "대기중"
    final_message: str = ""
    final_observed_value_raw: Optional[Any] = None
    final_observed_value_norm: Optional[Any] = None
    final_observed_time: Optional[float] = None
    final_observed_channel: Optional[int] = None
    used_log_path: Optional[str] = None
    started_at: Optional[_dt.datetime] = None
    ended_at: Optional[_dt.datetime] = None
    elapsed_sec: Optional[float] = None

    def to_row(self) -> List[Any]:
        c = self.condition

        return [
            c.tc_no,
            c.subcategory,
            c.tc_content,
            c.tc_expected_result,
            join_expectation_field(c.input_conditions, "message"),
            join_expectation_field(c.input_conditions, "signal"),
            join_expectation_field(c.input_conditions, "expected_value_raw"),
            join_expectation_field(c.output_conditions, "message"),
            join_expectation_field(c.output_conditions, "signal"),
            join_expectation_field(c.output_conditions, "expected_value_raw"),
            c.judge_mode,
            c.timeout_sec,
            c.use_yn,
            c.note,
            c.condition_id,
            c.tc_logic,
            c.revision,
            self.realtime_result.status if self.realtime_result else "",
            self.log_result.status if self.log_result else "",
            self.final_status,
            self.final_observed_value_raw,
            self.final_observed_time,
            self.final_observed_channel,
            self.used_log_path or "",
            self.final_message,
        ]


# =========================================================
# Excel Oracle 로딩
# =========================================================
DEFAULT_COLUMN_MAP = {
    "tc_no": 1,
    "subcategory": 2,
    "tc_content": 3,
    "tc_expected_result": 4,
    "input_message": 5,
    "input_signal": 6,
    "input_expected_value": 7,
    "output_message": 8,
    "output_signal": 9,
    "output_expected_value": 10,
    "judge_mode": 11,
    "timeout_sec": 12,
    "use_yn": 13,
    "note": 14,
    "condition_id": 15,
    "tc_logic": 16,
    "revision": 17,
}


def _cell_text(v: Any) -> str:
    if v is None:
        return ""
    return str(v).strip()


def _float_or_default(v: Any, default: float) -> float:
    s = _cell_text(v)
    if not s:
        return default
    try:
        return float(s)
    except Exception:
        return default


def list_excel_sheets(excel_path: str | Path) -> List[str]:
    from openpyxl import load_workbook

    wb = load_workbook(excel_path, read_only=True, data_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def load_oracle_conditions_with_diagnostics(
    excel_path: str | Path,
    sheet_name: str,
    include_disabled: bool = False,
    default_timeout_sec: float = 30.0,
) -> List[OracleRowLoadResult]:
    from openpyxl import load_workbook

    excel_path = Path(excel_path)
    if not excel_path.exists():
        raise FileNotFoundError(f"Oracle Excel 파일을 찾지 못했습니다: {excel_path}")

    wb = load_workbook(excel_path, read_only=True, data_only=True)
    try:
        if sheet_name not in wb.sheetnames:
            raise ValueError(f"시트를 찾지 못했습니다: {sheet_name}. 사용 가능: {wb.sheetnames}")

        ws = wb[sheet_name]
        start_row = 2

        results: List[OracleRowLoadResult] = []

        last_tc_no = ""
        last_subcategory = ""
        last_tc_content = ""
        last_tc_expected_result = ""

        for r in range(start_row, ws.max_row + 1):
            tc_no = _cell_text(ws.cell(r, DEFAULT_COLUMN_MAP["tc_no"]).value) or last_tc_no
            subcategory = _cell_text(ws.cell(r, DEFAULT_COLUMN_MAP["subcategory"]).value) or last_subcategory
            tc_content = _cell_text(ws.cell(r, DEFAULT_COLUMN_MAP["tc_content"]).value) or last_tc_content
            tc_expected_result = _cell_text(ws.cell(r, DEFAULT_COLUMN_MAP["tc_expected_result"]).value) or last_tc_expected_result

            input_message_raw = _cell_text(ws.cell(r, DEFAULT_COLUMN_MAP["input_message"]).value)
            input_signal_raw = _cell_text(ws.cell(r, DEFAULT_COLUMN_MAP["input_signal"]).value)
            input_expected_raw = _cell_text(ws.cell(r, DEFAULT_COLUMN_MAP["input_expected_value"]).value)

            output_message_raw = _cell_text(ws.cell(r, DEFAULT_COLUMN_MAP["output_message"]).value)
            output_signal_raw = _cell_text(ws.cell(r, DEFAULT_COLUMN_MAP["output_signal"]).value)
            output_expected_raw = _cell_text(ws.cell(r, DEFAULT_COLUMN_MAP["output_expected_value"]).value)

            judge_mode = _cell_text(ws.cell(r, DEFAULT_COLUMN_MAP["judge_mode"]).value) or "once"
            timeout_sec = _float_or_default(ws.cell(r, DEFAULT_COLUMN_MAP["timeout_sec"]).value, default_timeout_sec)
            use_yn = _cell_text(ws.cell(r, DEFAULT_COLUMN_MAP["use_yn"]).value) or "Y"
            note = _cell_text(ws.cell(r, DEFAULT_COLUMN_MAP["note"]).value)
            condition_id = _cell_text(ws.cell(r, DEFAULT_COLUMN_MAP["condition_id"]).value)
            tc_logic = _cell_text(ws.cell(r, DEFAULT_COLUMN_MAP["tc_logic"]).value) or "all"
            revision = _cell_text(ws.cell(r, DEFAULT_COLUMN_MAP["revision"]).value)

            if tc_no:
                last_tc_no = tc_no
            if subcategory:
                last_subcategory = subcategory
            if tc_content:
                last_tc_content = tc_content
            if tc_expected_result:
                last_tc_expected_result = tc_expected_result

            if not any([
                tc_no, subcategory, tc_content, tc_expected_result,
                input_message_raw, input_signal_raw, input_expected_raw,
                output_message_raw, output_signal_raw, output_expected_raw
            ]):
                continue

            errors: List[str] = []

            in_msgs = split_multiline_cell(input_message_raw)
            in_sigs = split_multiline_cell(input_signal_raw)
            in_exps = split_multiline_cell(input_expected_raw)

            if not in_msgs:
                errors.append("입력 Message 비어 있음")
            if not in_sigs:
                errors.append("입력 Signal 비어 있음")
            if not in_exps:
                errors.append("입력 Expected Value 비어 있음")

            if in_msgs or in_sigs or in_exps:
                if not (len(in_msgs) == len(in_sigs) == len(in_exps)):
                    errors.append("입력 Message/Signal/Expected Value 줄 수가 서로 다름")

            out_msgs = split_multiline_cell(output_message_raw)
            out_sigs = split_multiline_cell(output_signal_raw)
            out_exps = split_multiline_cell(output_expected_raw)

            if not out_msgs:
                errors.append("출력 Message 비어 있음")
            if not out_sigs:
                errors.append("출력 Signal 비어 있음")
            if not out_exps:
                errors.append("출력 Expected Value 비어 있음")

            if out_msgs or out_sigs or out_exps:
                if not (len(out_msgs) == len(out_sigs) == len(out_exps)):
                    errors.append("출력 Message/Signal/Expected Value 줄 수가 서로 다름")

            input_conditions: List[SignalExpectation] = []
            output_conditions: List[SignalExpectation] = []

            if not errors:
                for m, s, e in zip(in_msgs, in_sigs, in_exps):
                    input_conditions.append(
                        SignalExpectation(message=m, signal=s, expected_value_raw=e)
                    )

                for m, s, e in zip(out_msgs, out_sigs, out_exps):
                    output_conditions.append(
                        SignalExpectation(message=m, signal=s, expected_value_raw=e)
                    )

                if not input_conditions:
                    errors.append("입력 조건이 1개 이상 필요함")
                if not output_conditions:
                    errors.append("출력 조건이 1개 이상 필요함")

            if errors:
                results.append(
                    OracleRowLoadResult(
                        row_index=r,
                        tc_no=tc_no,
                        subcategory=subcategory,
                        tc_content=tc_content,
                        tc_expected_result=tc_expected_result,
                        input_message=input_message_raw,
                        input_signal=input_signal_raw,
                        input_expected_value_raw=input_expected_raw,
                        output_message_raw=output_message_raw,
                        output_signal_raw=output_signal_raw,
                        output_expected_value_raw=output_expected_raw,
                        is_valid=False,
                        status="ERROR",
                        error_reason=" / ".join(errors),
                        condition=None,
                    )
                )
                continue

            cond = OracleCondition(
                row_index=r,
                tc_no=tc_no,
                subcategory=subcategory,
                tc_content=tc_content,
                tc_expected_result=tc_expected_result,
                input_conditions=input_conditions,
                output_conditions=output_conditions,
                judge_mode=judge_mode,
                timeout_sec=timeout_sec,
                use_yn=use_yn,
                note=note,
                condition_id=condition_id,
                tc_logic=tc_logic,
                revision=revision,
            )

            if cond.enabled or include_disabled:
                results.append(
                    OracleRowLoadResult(
                        row_index=r,
                        tc_no=tc_no,
                        subcategory=subcategory,
                        tc_content=tc_content,
                        tc_expected_result=tc_expected_result,
                        input_message=input_message_raw,
                        input_signal=input_signal_raw,
                        input_expected_value_raw=input_expected_raw,
                        output_message_raw=output_message_raw,
                        output_signal_raw=output_signal_raw,
                        output_expected_value_raw=output_expected_raw,
                        is_valid=True,
                        status="Ready",
                        error_reason="",
                        condition=cond,
                    )
                )

        return results
    finally:
        wb.close()


def load_oracle_conditions(
    excel_path: str | Path,
    sheet_name: str,
    include_disabled: bool = False,
    default_timeout_sec: float = 30.0,
) -> List[OracleCondition]:
    rows = load_oracle_conditions_with_diagnostics(
        excel_path=excel_path,
        sheet_name=sheet_name,
        include_disabled=include_disabled,
        default_timeout_sec=default_timeout_sec,
    )
    return [x.condition for x in rows if x.is_valid and x.condition is not None]


# =========================================================
# 버스/채널 probe 유틸
# =========================================================
def build_default_bus_channel_candidates(channels: List[int]) -> List[Tuple[str, int]]:
    normalized = []
    seen = set()
    for ch in channels:
        try:
            ch_int = int(ch)
        except Exception:
            continue
        if ch_int not in seen:
            seen.add(ch_int)
            normalized.append(ch_int)

    # 실제 진단 결과 GetBus("CAN")만 유효했던 환경을 기준으로,
    # bus 이름은 CAN으로 고정하고 채널 번호만 바꿔 probe한다.
    return [("CAN", ch) for ch in normalized]


def debug_probe_signal(
    expectation: SignalExpectation,
    bus_channel_candidates: List[Tuple[str, int]],
) -> List[Tuple[str, int, bool, Optional[Any], Optional[str]]]:
    results = []
    client_cache: Dict[str, CANoeClient] = {}

    for bus_name, ch in bus_channel_candidates:
        client = client_cache.get(bus_name)
        if client is None:
            try:
                client = CANoeClient(bus_name=bus_name).connect()
                client_cache[bus_name] = client
            except Exception as e:
                msg = f"connect fail: {type(e).__name__}: {e}"
                results.append((bus_name, ch, False, None, msg))
                continue

        try:
            value = client.get_signal_value(ch, expectation.message, expectation.signal)
            results.append((bus_name, ch, True, value, None))
        except Exception as e:
            msg = f"{type(e).__name__}: {e}"
            results.append((bus_name, ch, False, None, msg))

    return results


# =========================================================
# 로그 자동 탐색 유틸
# =========================================================
def find_latest_log_file(
    log_dir: str | Path,
    after_ts: Optional[float] = None,
    wait_timeout: float = 5.0,
    poll_interval: float = 0.5,
    exts: Tuple[str, ...] = (".blf", ".asc"),
) -> Optional[Path]:
    folder = Path(log_dir)
    if not folder.exists():
        raise FileNotFoundError(f"로그 폴더를 찾지 못했습니다: {folder}")
    if not folder.is_dir():
        raise NotADirectoryError(f"로그 폴더가 아닙니다: {folder}")

    deadline = time.time() + max(0.1, float(wait_timeout))

    while time.time() <= deadline:
        candidates: List[Tuple[float, Path]] = []
        for ext in exts:
            for p in folder.glob(f"*{ext}"):
                try:
                    mtime = p.stat().st_mtime
                    if after_ts is None or mtime >= after_ts:
                        candidates.append((mtime, p))
                except FileNotFoundError:
                    pass

        if candidates:
            candidates.sort(key=lambda x: x[0], reverse=True)
            return candidates[0][1]

        time.sleep(max(0.1, float(poll_interval)))

    return None


def wait_until_file_stable(
    path: str | Path,
    stable_sec: float = 1.0,
    timeout: float = 5.0,
    poll_interval: float = 0.2,
) -> bool:
    p = Path(path)
    deadline = time.time() + max(0.1, float(timeout))
    last_size = None
    last_change = None

    while time.time() <= deadline:
        if not p.exists():
            time.sleep(poll_interval)
            continue

        try:
            size = p.stat().st_size
        except FileNotFoundError:
            time.sleep(poll_interval)
            continue

        if last_size is None or size != last_size:
            last_size = size
            last_change = time.time()
        else:
            if last_change is not None and (time.time() - last_change) >= stable_sec:
                return True

        time.sleep(poll_interval)

    return p.exists()


# =========================================================
# BLF -> ASC
# =========================================================
def convert_blf_to_asc(input_blf: str | Path, output_asc: Optional[str | Path] = None) -> Path:
    try:
        import can
    except Exception as e:
        raise ImportError("python-can이 필요합니다. py -m pip install python-can") from e

    input_blf = Path(input_blf)
    if not input_blf.exists():
        raise FileNotFoundError(f"BLF 파일을 찾지 못했습니다: {input_blf}")

    if output_asc is None:
        output_asc = input_blf.with_suffix(".asc")
    output_asc = Path(output_asc)

    with can.BLFReader(str(input_blf)) as reader:
        with can.ASCWriter(str(output_asc)) as writer:
            for msg in reader:
                writer.on_message_received(msg)

    return output_asc


def ensure_asc_path(log_path: str | Path) -> Path:
    p = Path(log_path)
    if not p.exists():
        raise FileNotFoundError(f"로그 파일을 찾지 못했습니다: {p}")

    ext = p.suffix.lower()

    if ext == ".asc":
        return p
    if ext == ".blf":
        return convert_blf_to_asc(p)

    raise ValueError(f"지원하지 않는 로그 확장자입니다: {ext} (허용: .asc, .blf)")


# =========================================================
# ASC Parser
# =========================================================
HEAD_RE = re.compile(
    r"^\s*(?P<time>\d+\.\d+)\s+(?P<bus>CANFD|CAN)\s+(?P<ch>\d+)\s+(?P<dir>Rx|Tx)\s+(?P<id>[0-9A-Fa-f]+x?)\s+",
    re.IGNORECASE,
)
HEXBYTE_RE = re.compile(r"^[0-9A-Fa-f]{2}$")
BASE_RE = re.compile(r"\bbase\s+(hex|dec)\b", re.IGNORECASE)
DATA_MARKERS = {"a", "b", "c", "d"}


def detect_asc_base(path: str | Path, max_lines: int = 200) -> int:
    with open(path, "r", encoding="utf-8", errors="ignore") as f:
        for i, line in enumerate(f):
            if i >= max_lines:
                break
            m = BASE_RE.search(line)
            if m:
                return 16 if m.group(1).lower() == "hex" else 10
    return 16


def parse_can_id(id_str: str, default_base: int) -> int:
    s = id_str.strip().lower()

    if s.endswith("x"):
        return int(s[:-1], 16)
    if s.startswith("0x"):
        return int(s, 16)
    if re.search(r"[a-f]", s):
        return int(s, 16)

    return int(s, default_base)


def _is_uint_token(tok: str) -> bool:
    return bool(re.fullmatch(r"\d+", tok))


def _is_hexbyte_token(tok: str) -> bool:
    return bool(HEXBYTE_RE.fullmatch(tok))


def _collect_hexbytes(tokens: List[str], start_idx: int) -> List[str]:
    out: List[str] = []
    for tok in tokens[start_idx:]:
        if _is_hexbyte_token(tok):
            out.append(tok)
        else:
            if out:
                break
    return out


def _try_marker_mode(tokens: List[str]) -> Tuple[bool, Optional[int], List[str]]:
    for i, tok in enumerate(tokens):
        if tok.lower() in DATA_MARKERS and i + 1 < len(tokens) and _is_uint_token(tokens[i + 1]):
            length = int(tokens[i + 1])
            if not (0 <= length <= 64):
                return False, None, []

            data_tokens = _collect_hexbytes(tokens, i + 2)
            if len(data_tokens) < length:
                return False, length, data_tokens

            return True, length, data_tokens[:length]

    return False, None, []


def _try_numeric_head_mode(tokens: List[str]) -> Tuple[bool, Optional[int], List[str]]:
    for i, tok in enumerate(tokens):
        if not _is_uint_token(tok):
            continue

        length = int(tok)
        if not (0 <= length <= 64):
            continue

        data_tokens = _collect_hexbytes(tokens, i + 1)
        if len(data_tokens) >= length:
            return True, length, data_tokens[:length]

    return False, None, []


def _parse_vector_payload_tokens(tokens: List[str]) -> Tuple[bool, Optional[int], List[str]]:
    ok, length, data_tokens = _try_marker_mode(tokens)
    if ok:
        return ok, length, data_tokens

    ok, length, data_tokens = _try_numeric_head_mode(tokens)
    if ok:
        return ok, length, data_tokens

    return False, None, []


def read_frames_from_asc_with_direction(path: str | Path):
    """Yield ASC CAN frames including direction for TX/RX diagnostics.

    This is intentionally separate from the legacy log-judgement iterator below.
    Existing Pass/Fail log review remains Rx-only; TX/RX Stimulus diagnostics may
    inspect both Rx and Tx so a CAPL output() call can be confirmed in the log.
    """
    default_base = detect_asc_base(path)

    with open(path, "r", encoding="utf-8", errors="ignore") as f:
        for line in f:
            m = HEAD_RE.match(line)
            if not m:
                continue

            direction = str(m.group("dir") or "").strip().upper()
            t = float(m.group("time"))
            ch = int(m.group("ch"))
            can_id = parse_can_id(m.group("id"), default_base)

            rest = line[m.end():].strip()
            tokens = rest.split()

            ok, length, data_tokens = _parse_vector_payload_tokens(tokens)
            if not ok or length is None:
                continue
            if len(data_tokens) < length:
                continue

            try:
                data = bytes(int(b, 16) for b in data_tokens[:length])
            except Exception:
                continue

            yield t, ch, direction, can_id, data


def read_frames_from_asc(path: str | Path):
    """Legacy Pass/Fail log iterator: keep Rx-only semantics unchanged."""
    for t, ch, direction, can_id, data in read_frames_from_asc_with_direction(path):
        if direction != "RX":
            continue
        yield t, ch, can_id, data


def _diag_compact_sequence(samples: Iterable[Tuple[float, str, Any]], max_items: int = 16) -> List[Tuple[float, str, Any]]:
    out: List[Tuple[float, str, Any]] = []
    last = object()
    for t, direction, value in samples or []:
        norm = normalize_value(value)
        if out and values_equal(last, norm):
            continue
        out.append((float(t), str(direction or "-"), value))
        last = norm
        if len(out) >= max(1, int(max_items)):
            break
    return out


def _diag_suspicious_signal_name(name: str) -> bool:
    text = re.sub(r"[^a-z0-9]", "", str(name or "").lower())
    return any(k in text for k in ("crc", "checksum", "alive", "alv", "counter", "cnt", "e2e"))



ASC_DATE_RE = re.compile(r"^\s*date\s+(?P<value>.+?)\s*$", re.IGNORECASE)


def detect_asc_start_datetime(path: str | Path, max_lines: int = 80) -> Optional[_dt.datetime]:
    """Best-effort Vector ASC wall-clock start parser.

    Typical Vector header:
      date Tue Aug 25 13:16:50.000 2026
    The helper is diagnostic only. If parsing fails, callers fall back to a
    measurement-relative or whole-log window rather than failing the test.
    """
    formats = (
        "%a %b %d %H:%M:%S.%f %Y",
        "%a %b %d %H:%M:%S %Y",
        "%a %b %d %I:%M:%S.%f %p %Y",
        "%a %b %d %I:%M:%S %p %Y",
    )
    try:
        with open(path, "r", encoding="utf-8", errors="ignore") as f:
            for idx, line in enumerate(f):
                if idx >= max_lines:
                    break
                m = ASC_DATE_RE.match(line)
                if not m:
                    continue
                raw = re.sub(r"\s+", " ", str(m.group("value") or "").strip())
                for fmt in formats:
                    try:
                        return _dt.datetime.strptime(raw, fmt)
                    except Exception:
                        continue
                return None
    except Exception:
        return None
    return None


def _diag_payload_bytes(payload_hex: Any) -> bytes:
    text = str(payload_hex or "").strip()
    if not text:
        return b""
    compact = re.sub(r"[^0-9A-Fa-f]", "", text)
    if not compact or len(compact) % 2:
        return b""
    try:
        return bytes.fromhex(compact)
    except Exception:
        return b""


def _diag_payload_diff(baseline_hex: Any, active_hex: Any) -> List[Dict[str, Any]]:
    before = _diag_payload_bytes(baseline_hex)
    after = _diag_payload_bytes(active_hex)
    size = max(len(before), len(after))
    out: List[Dict[str, Any]] = []
    for idx in range(size):
        b = before[idx] if idx < len(before) else None
        a = after[idx] if idx < len(after) else None
        if b != a:
            out.append({
                "index": idx,
                "baseline": None if b is None else int(b),
                "active": None if a is None else int(a),
            })
    return out


def _diag_stats_add(stats: Dict[str, Any], ch: int, direction: str) -> None:
    direction = str(direction or "").upper()
    if direction not in ("TX", "RX"):
        return
    stats["total"] = int(stats.get("total", 0) or 0) + 1
    stats[direction.lower()] = int(stats.get(direction.lower(), 0) or 0) + 1
    by_ch = stats.setdefault("by_channel", {})
    row = by_ch.setdefault(int(ch), {"total": 0, "tx": 0, "rx": 0})
    row["total"] += 1
    row[direction.lower()] += 1


def _diag_window_contains(t: float, start_sec: Optional[float], end_sec: Optional[float]) -> bool:
    if start_sec is not None and float(t) < float(start_sec):
        return False
    if end_sec is not None and float(t) > float(end_sec):
        return False
    return True


def _derive_stimulus_diag_window(
    asc_path: str | Path,
    stimulus_wall_start_ts: Optional[float],
    stimulus_wall_end_ts: Optional[float],
    measurement_restarted_for_stimulus: bool,
    hold_sec: Optional[float],
) -> Dict[str, Any]:
    """Choose a narrow per-Stimulus ASC analysis window when possible."""
    if measurement_restarted_for_stimulus:
        # CAPL Stimulus always restarts Measurement immediately before the first frame.
        # Allow generous OnInit/start/release margins while excluding older unrelated frames.
        hold = max(0.0, float(hold_sec or 0.0))
        return {
            "mode": "measurement_restart",
            "start_sec": 0.0,
            "end_sec": max(2.0, hold + 1.5),
            "asc_start_datetime": None,
        }

    asc_start = detect_asc_start_datetime(asc_path)
    if asc_start is not None and stimulus_wall_start_ts is not None:
        try:
            rel_start = float(stimulus_wall_start_ts) - float(asc_start.timestamp()) - 0.5
            wall_end = float(stimulus_wall_end_ts) if stimulus_wall_end_ts is not None else float(stimulus_wall_start_ts) + max(0.5, float(hold_sec or 0.0))
            rel_end = wall_end - float(asc_start.timestamp()) + 0.8
            # Reject obviously mismatched/local-time parse results.
            if rel_end >= -1.0 and rel_start <= 24 * 3600 and rel_end - rel_start <= 600:
                return {
                    "mode": "asc_wallclock",
                    "start_sec": max(0.0, rel_start),
                    "end_sec": max(0.0, rel_end),
                    "asc_start_datetime": asc_start.isoformat(sep=" "),
                }
        except Exception:
            pass

    return {
        "mode": "whole_log_fallback",
        "start_sec": None,
        "end_sec": None,
        "asc_start_datetime": asc_start.isoformat(sep=" ") if asc_start is not None else None,
    }


def inspect_stimulus_transitions_in_log(
    log_path: str | Path,
    dbc_path_by_channel: Dict[int, str | Path],
    targets: Iterable[Dict[str, Any]],
    max_samples_per_signal: int = 200,
    max_changed_signals: int = 12,
    response_groups: Optional[Iterable[Dict[str, Any]]] = None,
    stimulus_wall_start_ts: Optional[float] = None,
    stimulus_wall_end_ts: Optional[float] = None,
    measurement_restarted_for_stimulus: bool = False,
    hold_sec: Optional[float] = None,
) -> Dict[str, Any]:
    """Diagnostic evidence for a TX/RX Stimulus run.

    rev87 deliberately separates these layers:
      1) PassFail/CAPL call succeeded,
      2) logger actually contains TX/RX/message/signal evidence,
      3) ECU/vehicle functional reaction.

    The function never creates a FinalCheckResult and never changes existing
    single/multi P/F log semantics. It reads both Tx and Rx only for this
    diagnostic path.
    """
    asc_path = ensure_asc_path(log_path)
    db_by_ch = load_dbc_map(dbc_path_by_channel)
    window = _derive_stimulus_diag_window(
        asc_path,
        stimulus_wall_start_ts,
        stimulus_wall_end_ts,
        bool(measurement_restarted_for_stimulus),
        hold_sec,
    )
    window_start = window.get("start_sec")
    window_end = window.get("end_sec")

    target_rows: List[Dict[str, Any]] = []
    target_lookup: Dict[Tuple[int, int], List[int]] = {}

    for raw in list(targets or []):
        ch = int(raw.get("channel"))
        msg_name = str(raw.get("message") or "").strip()
        sig_name = str(raw.get("signal") or "").strip()
        row: Dict[str, Any] = {
            "channel": ch,
            "logical_can": str(raw.get("logical_can") or f"CAN{ch}"),
            "message": msg_name,
            "signal": sig_name,
            "active_value": raw.get("active_value"),
            "original_value": raw.get("original_value"),
            "baseline_payload_hex": str(raw.get("baseline_payload_hex") or ""),
            "active_payload_hex": str(raw.get("active_payload_hex") or ""),
            "cycle_ms": float(raw.get("cycle_ms", 0) or 0),
            "payload_diff_bytes": [],
            "intended_suspicious": {},
            "error": "",
            "frame_count": 0,
            "tx_count": 0,
            "rx_count": 0,
            "whole_frame_count": 0,
            "whole_tx_count": 0,
            "whole_rx_count": 0,
            "samples": [],
            "active_hits": [],
            "original_hits": [],
            "raw_active_frames": [],
            "raw_message_frames": [],
            "observed_baseline_raw_hex": "",
            "observed_baseline_direction": "",
            "baseline_vs_observed_diff_bytes": [],
            "signal_sequences": {},
            "suspicious_signals": [],
            "changed_signals": [],
            "restore_seen_after_active": False,
            "active_seen": False,
            "source_rx_during_window": False,
        }
        db = db_by_ch.get(ch)
        if db is None:
            row["error"] = f"CAN{ch} DBC를 로드하지 못했습니다."
            target_rows.append(row)
            continue
        try:
            msg_obj = db.get_message_by_name(msg_name)
        except Exception:
            msg_obj = None
        if msg_obj is None:
            row["error"] = f"DBC에서 Message를 찾지 못했습니다: {msg_name}"
            target_rows.append(row)
            continue
        try:
            msg_obj.get_signal_by_name(sig_name)
        except Exception:
            row["error"] = f"DBC에서 Signal을 찾지 못했습니다: {msg_name}.{sig_name}"
            target_rows.append(row)
            continue

        row["_msg_obj"] = msg_obj
        row["frame_id"] = int(getattr(msg_obj, "frame_id"))
        row["suspicious_signals"] = [
            str(getattr(sig, "name", "")) for sig in (getattr(msg_obj, "signals", []) or [])
            if _diag_suspicious_signal_name(str(getattr(sig, "name", "")))
        ]
        row["payload_diff_bytes"] = _diag_payload_diff(row["baseline_payload_hex"], row["active_payload_hex"])

        baseline_payload = _diag_payload_bytes(row["baseline_payload_hex"])
        active_payload = _diag_payload_bytes(row["active_payload_hex"])
        if baseline_payload and active_payload:
            try:
                baseline_decoded = msg_obj.decode(baseline_payload, decode_choices=False)
                active_decoded = msg_obj.decode(active_payload, decode_choices=False)
                for suspicious_name in row["suspicious_signals"]:
                    bval = baseline_decoded.get(suspicious_name)
                    aval = active_decoded.get(suspicious_name)
                    row["intended_suspicious"][suspicious_name] = {
                        "baseline": bval,
                        "active": aval,
                        "changed": not values_equal(bval, aval),
                    }
            except Exception:
                pass

        idx = len(target_rows)
        target_rows.append(row)
        target_lookup.setdefault((ch, row["frame_id"]), []).append(idx)

    # Prepare optional TC output-response expectations. Multiple logical CAN locations
    # are kept as OR candidates; observing any candidate may satisfy the diagnostic group.
    response_rows: List[Dict[str, Any]] = []
    response_lookup: Dict[Tuple[int, int], List[Tuple[int, int]]] = {}
    for raw_group in list(response_groups or []):
        group: Dict[str, Any] = {
            "role": str(raw_group.get("role") or "output"),
            "message": str(raw_group.get("message") or "").strip(),
            "signal": str(raw_group.get("signal") or "").strip(),
            "expected_value": raw_group.get("expected_value"),
            "candidates": [],
            "seen": False,
            "seen_after_active": None,
            "first_hit_time": None,
            "error": "",
        }
        for raw_candidate in list(raw_group.get("candidates") or []):
            ch = int(raw_candidate.get("channel", 0) or 0)
            logical = str(raw_candidate.get("logical_can") or f"CAN{ch}")
            candidate = {
                "channel": ch,
                "logical_can": logical,
                "frame_count": 0,
                "tx_count": 0,
                "rx_count": 0,
                "samples": [],
                "expected_hits": [],
                "compact_sequence": [],
                "error": "",
            }
            db = db_by_ch.get(ch)
            if db is None:
                candidate["error"] = f"CAN{ch} DBC 없음"
                group["candidates"].append(candidate)
                continue
            try:
                msg_obj = db.get_message_by_name(group["message"])
            except Exception:
                msg_obj = None
            if msg_obj is None:
                candidate["error"] = f"Message 없음: {group['message']}"
                group["candidates"].append(candidate)
                continue
            try:
                msg_obj.get_signal_by_name(group["signal"])
            except Exception:
                candidate["error"] = f"Signal 없음: {group['message']}.{group['signal']}"
                group["candidates"].append(candidate)
                continue
            candidate["_msg_obj"] = msg_obj
            candidate["frame_id"] = int(getattr(msg_obj, "frame_id"))
            cidx = len(group["candidates"])
            group["candidates"].append(candidate)
            response_lookup.setdefault((ch, candidate["frame_id"]), []).append((len(response_rows), cidx))
        response_rows.append(group)

    whole_stats: Dict[str, Any] = {"total": 0, "tx": 0, "rx": 0, "by_channel": {}}
    window_stats: Dict[str, Any] = {"total": 0, "tx": 0, "rx": 0, "by_channel": {}}

    for t, ch, direction, can_id, data in read_frames_from_asc_with_direction(asc_path):
        _diag_stats_add(whole_stats, ch, direction)
        in_window = _diag_window_contains(t, window_start, window_end)
        if in_window:
            _diag_stats_add(window_stats, ch, direction)

        indexes = target_lookup.get((int(ch), int(can_id))) or []
        for idx in indexes:
            row = target_rows[idx]
            if direction == "TX":
                row["whole_tx_count"] += 1
            elif direction == "RX":
                row["whole_rx_count"] += 1
            row["whole_frame_count"] += 1
            if not in_window:
                continue
            msg_obj = row.get("_msg_obj")
            if msg_obj is None:
                continue
            try:
                decoded = msg_obj.decode(data, decode_choices=False)
            except Exception:
                continue

            row["frame_count"] += 1
            if direction == "TX":
                row["tx_count"] += 1
            elif direction == "RX":
                row["rx_count"] += 1
            if len(row["raw_message_frames"]) < 8:
                row["raw_message_frames"].append({
                    "time": float(t), "direction": direction, "data_hex": data.hex(" ").upper()
                })

            value = decoded.get(row["signal"])
            if len(row["samples"]) < max_samples_per_signal:
                row["samples"].append((float(t), direction, value))
            if values_equal(row.get("active_value"), value):
                row["active_hits"].append((float(t), direction, value))
                if len(row["raw_active_frames"]) < 5:
                    row["raw_active_frames"].append({
                        "time": float(t), "direction": direction, "data_hex": data.hex(" ").upper()
                    })
            original = row.get("original_value")
            if original is not None and values_equal(original, value):
                row["original_hits"].append((float(t), direction, value))

            for sig_name, sig_value in decoded.items():
                seq = row["signal_sequences"].setdefault(str(sig_name), [])
                if len(seq) < max_samples_per_signal:
                    seq.append((float(t), direction, sig_value))

        for group_idx, candidate_idx in response_lookup.get((int(ch), int(can_id))) or []:
            if not in_window:
                continue
            group = response_rows[group_idx]
            candidate = group["candidates"][candidate_idx]
            msg_obj = candidate.get("_msg_obj")
            if msg_obj is None:
                continue
            try:
                decoded = msg_obj.decode(data, decode_choices=False)
            except Exception:
                continue
            candidate["frame_count"] += 1
            if direction == "TX":
                candidate["tx_count"] += 1
            elif direction == "RX":
                candidate["rx_count"] += 1
            value = decoded.get(group["signal"])
            if len(candidate["samples"]) < max_samples_per_signal:
                candidate["samples"].append((float(t), direction, value))
            if values_equal(group.get("expected_value"), value):
                candidate["expected_hits"].append((float(t), direction, value))

    first_active_time: Optional[float] = None
    for row in target_rows:
        row.pop("_msg_obj", None)
        compact = _diag_compact_sequence(row.get("samples") or [])
        row["compact_sequence"] = compact
        row["active_seen"] = bool(row.get("active_hits"))
        row["source_rx_during_window"] = int(row.get("rx_count", 0) or 0) > 0

        active_hits = list(row.get("active_hits") or [])
        original_hits = list(row.get("original_hits") or [])
        if active_hits:
            this_first = min(float(x[0]) for x in active_hits)
            first_active_time = this_first if first_active_time is None else min(first_active_time, this_first)
        if active_hits and original_hits:
            last_active_t = max(float(x[0]) for x in active_hits)
            row["restore_seen_after_active"] = any(float(x[0]) > last_active_t for x in original_hits)

        raw_candidates = list(row.get("raw_message_frames") or [])
        observed = next((x for x in raw_candidates if str(x.get("direction") or "").upper() == "RX"), None)
        if observed is None and raw_candidates:
            observed = raw_candidates[0]
        if observed is not None:
            row["observed_baseline_raw_hex"] = str(observed.get("data_hex") or "")
            row["observed_baseline_direction"] = str(observed.get("direction") or "")
            row["baseline_vs_observed_diff_bytes"] = _diag_payload_diff(
                row.get("baseline_payload_hex"), row.get("observed_baseline_raw_hex")
            )

        changed = []
        for sig_name, samples in (row.get("signal_sequences") or {}).items():
            cseq = _diag_compact_sequence(samples, max_items=16)
            if len(cseq) >= 2:
                changed.append({"signal": sig_name, "sequence": cseq})
        target_sig = row.get("signal")
        suspicious = set(row.get("suspicious_signals") or [])
        changed.sort(key=lambda x: (0 if x["signal"] == target_sig else 1 if x["signal"] in suspicious else 2, x["signal"]))
        row["changed_signals"] = changed[:max(1, int(max_changed_signals))]
        row["suspicious_sequences"] = {
            sig_name: _diag_compact_sequence(
                (row.get("signal_sequences") or {}).get(sig_name, []), max_items=16
            )
            for sig_name in row.get("suspicious_signals") or []
        }
        row.pop("signal_sequences", None)

    for group in response_rows:
        all_hits: List[Tuple[float, str, Any]] = []
        for candidate in group.get("candidates") or []:
            candidate.pop("_msg_obj", None)
            candidate["compact_sequence"] = _diag_compact_sequence(candidate.get("samples") or [])
            all_hits.extend(list(candidate.get("expected_hits") or []))
        group["seen"] = bool(all_hits)
        if all_hits:
            group["first_hit_time"] = min(float(x[0]) for x in all_hits)
        if first_active_time is None:
            group["seen_after_active"] = None
        else:
            group["seen_after_active"] = any(float(x[0]) >= float(first_active_time) for x in all_hits)

    return {
        "log_path": str(Path(log_path)),
        "asc_path": str(asc_path),
        "analysis_window": window,
        "direction_stats": {
            "whole": whole_stats,
            "window": window_stats,
        },
        "targets": target_rows,
        "response_groups": response_rows,
        "first_active_time": first_active_time,
    }


# =========================================================
# DBC 로드
# =========================================================
def load_dbc_map(dbc_path_by_channel: Dict[int, str | Path]):
    try:
        import cantools
    except Exception as e:
        raise ImportError("cantools가 필요합니다. py -m pip install cantools") from e

    db_by_ch = {}
    for ch, dbc_path in dbc_path_by_channel.items():
        p = Path(dbc_path)
        if not p.exists():
            continue
        db_by_ch[int(ch)] = cantools.database.load_file(str(p))

    return db_by_ch


def lookup_dbc_choice_text(
    db_by_ch: Dict[int, Any],
    message_name: str,
    signal_name: str,
    raw_value: Any,
    max_len: int = 24,
) -> str:
    """DBC choice/value table에서 Expected Value의 표시명을 찾는다.

    예: expected=0x01, DBC choice={1: "up"}이면 "up" 반환.
    찾지 못하거나 cantools/DBC 구조가 맞지 않으면 "-"를 반환한다.
    """
    try:
        msg_name = str(message_name or "").strip()
        sig_name = str(signal_name or "").strip()
        if not msg_name or not sig_name or not db_by_ch:
            return "-"

        norm = normalize_value(raw_value)
        candidates: List[Any] = []
        if norm is not None:
            candidates.append(norm)
            try:
                f = float(norm)
                if math.isfinite(f) and f.is_integer():
                    candidates.append(int(f))
            except Exception:
                pass
        raw_text = str(raw_value or "").strip()
        if raw_text:
            candidates.append(raw_text)

        deduped: List[Any] = []
        for item in candidates:
            if item not in deduped:
                deduped.append(item)

        for _ch, db in sorted(db_by_ch.items(), key=lambda x: int(x[0])):
            try:
                msg = db.get_message_by_name(msg_name)
            except Exception:
                continue

            signal_obj = None
            try:
                signal_obj = msg.get_signal_by_name(sig_name)
            except Exception:
                try:
                    signal_obj = next((s for s in getattr(msg, "signals", []) if getattr(s, "name", "") == sig_name), None)
                except Exception:
                    signal_obj = None

            choices = getattr(signal_obj, "choices", None) if signal_obj is not None else None
            if not choices:
                continue

            for key in deduped:
                try:
                    if key in choices:
                        label = str(choices[key])
                        return (label[: max_len - 3] + "...") if len(label) > max_len else label
                except Exception:
                    pass

            # 일부 DBC/라이브러리 버전에서 key 타입이 다르게 들어오는 경우를 대비한 완화 비교
            for choice_key, choice_val in choices.items():
                try:
                    if normalize_value(choice_key) == norm:
                        label = str(choice_val)
                        return (label[: max_len - 3] + "...") if len(label) > max_len else label
                except Exception:
                    continue

        return "-"
    except Exception:
        return "-"


def build_target_message_map(
    db_by_ch: Dict[int, Any],
    expectations: List[SignalExpectation],
) -> Tuple[Dict[Tuple[int, str], int], Dict[Tuple[int, str], Any]]:
    target_frame_ids: Dict[Tuple[int, str], int] = {}
    target_msgs: Dict[Tuple[int, str], Any] = {}

    unique_msg_names = sorted(set(exp.message for exp in expectations if exp.message))

    for ch, db in db_by_ch.items():
        for msg_name in unique_msg_names:
            try:
                msg = db.get_message_by_name(msg_name)
                target_frame_ids[(int(ch), msg_name)] = int(msg.frame_id)
                target_msgs[(int(ch), msg_name)] = msg
            except Exception:
                continue

    return target_frame_ids, target_msgs


# =========================================================
# 단일 기대조건 검사
# =========================================================
ProgressCallback = Callable[[str, Optional["CheckResult"]], None]


def _normalize_channels(channel) -> List[int]:
    if isinstance(channel, (list, tuple, set)):
        out = []
        for ch in channel:
            try:
                ch_int = int(ch)
            except Exception:
                continue
            if ch_int not in out:
                out.append(ch_int)
        return out
    try:
        return [int(channel)]
    except Exception:
        return []


def _read_expectation_with_existing_client(
    client: CANoeClient,
    expectation: SignalExpectation,
    channels: List[int],
) -> Tuple[List[Tuple[str, int, Any]], List[str]]:
    """
    이미 연결된 CANoeClient 하나만 사용해 선택 채널에서 신호를 읽는다.

    과거 로직은 이 함수 안에서 CANoeClient(...).connect()를 다시 호출했기 때문에
    PC의 COM 등록 상태에 따라 CANoe 10 같은 다른 버전으로 빠질 수 있었다.
    이 함수는 TC Start/Worker에서 확보한 Active CANoe 인스턴스만 사용한다.
    """
    readings: List[Tuple[str, int, Any]] = []
    errors: List[str] = []

    if client is None or not client.is_connected():
        return [], ["Vector CANoe/CANalyzer 연결이 없습니다."]

    bus_label = getattr(client, "bus_name", "CAN") or "CAN"

    for ch in channels:
        try:
            ch_int = int(ch)
        except Exception:
            errors.append(f"CAN{ch}: 잘못된 논리 CAN 번호")
            continue

        try:
            raw = client.get_signal_value(ch_int, expectation.message, expectation.signal)
            readings.append((bus_label, ch_int, raw))
        except Exception as e:
            errors.append(f"{bus_label}/CAN{ch_int}: {type(e).__name__}: {e}")

    return readings, errors


def wait_for_expectation_realtime(
    client: CANoeClient,
    expectation: SignalExpectation,
    channels: List[int],
    timeout_sec: float,
    stop_event: Optional[threading.Event] = None,
    done_event: Optional[threading.Event] = None,
    poll_interval_sec: float = 0.05,
    collector_probe: Optional[Callable[..., Optional[Dict[str, Any]]]] = None,
    collector_min_hit_ms: float = 0.0,
) -> SingleExpectationCheck:
    stop_event = stop_event or threading.Event()
    done_event = done_event or threading.Event()

    res = SingleExpectationCheck(
        expectation=expectation,
        status="검토중",
        message=f"실시간 확인 시작: {expectation.message}:{expectation.signal} == {expectation.expected_value_raw}",
    )

    # 채널 후보별로 새 COM 연결을 만들지 않고, 인자로 받은 client 하나만 사용한다.

    start = time.perf_counter()
    first_successful_read = False
    last_errors: List[str] = []

    while True:
        if stop_event.is_set():
            res.status = "N/A"
            res.message = "사용자 중지"
            return res

        elapsed = time.perf_counter() - start

        if timeout_sec > 0 and elapsed > timeout_sec:
            break

        # rev87 optional CAPL Event Collector: CAPL latches a short expected-value edge inside
        # CANoe/CANalyzer. Python only reads the thread-safe cache populated by the GUI thread.
        # If the collector is absent/disabled/stale, collector_probe returns None and the exact
        # existing COM polling path below remains unchanged.
        if collector_probe is not None:
            try:
                hit = collector_probe(
                    expectation=expectation,
                    channels=channels,
                    min_hit_ms=float(collector_min_hit_ms or 0.0),
                )
            except Exception:
                hit = None
            if hit:
                raw = hit.get("value", expectation.expected_value_raw)
                ch = hit.get("channel")
                hit_ms = float(hit.get("hit_ms") or 0.0)
                res.status = "PASS"
                res.observed_value_raw = raw
                res.observed_value_norm = normalize_value(raw)
                res.observed_time = hit_ms / 1000.0 if hit_ms > 0 else elapsed
                res.observed_channel = int(ch) if ch is not None else None
                res.observation_source = "capl_collector"
                res.collector_hit_ms = hit_ms if hit_ms > 0 else None
                res.message = f"Expected Value 관측 | CAPL Event Collector / CAN{ch}" if ch is not None else "Expected Value 관측 | CAPL Event Collector"
                if len(res.samples) < 100:
                    res.samples.append((res.observed_time or elapsed, f"CAPL/CAN{ch}={raw}" if ch is not None else f"CAPL={raw}"))
                return res

        readings, errors = _read_expectation_with_existing_client(
            client=client,
            expectation=expectation,
            channels=channels,
        )

        if readings:
            first_successful_read = True
            for bus_name, ch, raw in readings:
                res.observed_value_raw = raw
                res.observed_value_norm = normalize_value(raw)
                res.observed_time = elapsed
                res.observed_channel = ch
                if len(res.samples) < 100:
                    res.samples.append((elapsed, f"{bus_name}/CAN{ch}={raw}"))

                if values_equal(expectation.expected_value_raw, raw):
                    res.status = "PASS"
                    res.observation_source = "com_polling"
                    res.message = f"Expected Value 관측 | {bus_name}/CAN{ch}"
                    return res
        else:
            last_errors = errors
            if elapsed >= min(1.0, timeout_sec if timeout_sec > 0 else 1.0):
                res.status = "ERROR"
                res.message = "선택된 모든 버스/채널 후보에서 Signal 읽기 실패: " + " | ".join(last_errors[:6])
                return res

        if done_event.is_set():
            if first_successful_read:
                res.status = "FAIL"
                res.message = "수행 완료 시점까지 예상 결과값 미관측"
            else:
                res.status = "N/A"
                res.message = "수행 완료 시점까지 신호 미관측"
            return res

        time.sleep(max(0.005, float(poll_interval_sec)))

    if first_successful_read:
        res.status = "FAIL"
        res.message = f"제한시간 {timeout_sec:.1f}초 내 예상 결과값 미관측"
    else:
        res.status = "N/A"
        res.message = "신호 읽기 성공 없이 제한시간 종료"

    return res


# =========================================================
# 로그에서 단일 기대조건 재검토
# =========================================================
def review_single_expectation_in_log_robust(
    expectation: SignalExpectation,
    log_path: str | Path,
    dbc_path_by_channel: Dict[int, str | Path],
    drop_first_seconds: float = 0.0,
    max_samples: int = 100,
) -> LogReviewOutcome:
    started = time.perf_counter()

    result = LogReviewOutcome(
        status="ERROR",
        message="초기화됨",
    )

    try:
        asc_path = ensure_asc_path(log_path)
        result.asc_path = str(asc_path)

        db_by_ch = load_dbc_map(dbc_path_by_channel)
        if not db_by_ch:
            result.status = "ERROR"
            result.message = "유효한 DBC를 1개도 로드하지 못했습니다."
            return result

        target_frame_ids, target_msgs = build_target_message_map(db_by_ch, [expectation])

        if not target_frame_ids:
            result.status = "ERROR"
            result.message = f"DBC에서 Message를 찾지 못함: {expectation.message}"
            return result

        last_observed = None
        last_observed_time = None
        last_ch = None
        last_frame_id = None

        for t, ch, can_id, data in read_frames_from_asc(asc_path):
            if t < float(drop_first_seconds):
                continue

            key = (ch, expectation.message)
            if key not in target_frame_ids:
                continue

            if can_id != target_frame_ids[key]:
                continue

            msg = target_msgs[key]
            try:
                decoded = msg.decode(data, decode_choices=False)
            except Exception:
                continue

            if expectation.signal not in decoded:
                continue

            val = decoded[expectation.signal]
            last_observed = val
            last_observed_time = t
            last_ch = ch
            last_frame_id = can_id

            if len(result.samples) < max_samples:
                result.samples.append((t, val))

            result.observed_value_raw = val
            result.observed_value_norm = normalize_value(val)
            result.observed_time = t
            result.matched_channel = ch
            result.matched_frame_id = can_id

            if values_equal(expectation.expected_value_raw, val):
                result.status = "PASS"
                result.message = f"로그에서 예상 결과값 관측: {Path(log_path).name}"
                return result

        result.observed_value_raw = last_observed
        result.observed_value_norm = normalize_value(last_observed)
        result.observed_time = last_observed_time
        result.matched_channel = last_ch
        result.matched_frame_id = last_frame_id
        result.status = "FAIL"
        result.message = f"로그 재검토에서도 예상 결과값 미관측: {Path(log_path).name}"
        return result

    except Exception as e:
        result.status = "ERROR"
        result.message = f"로그 재검토 오류: {e}\n{traceback.format_exc()}"
        return result

    finally:
        result.elapsed_sec = time.perf_counter() - started




# =========================================================
# 대용량 로그 일괄 재검토: 여러 TC/시그널을 ASC 1회 스캔으로 확인
# =========================================================
def _make_error_log_check_result(condition: OracleCondition, message: str) -> CheckResult:
    check = CheckResult(
        condition=condition,
        status="ERROR",
        source="multi_log_batch_review",
        started_at=_dt.datetime.now(),
        ended_at=_dt.datetime.now(),
        message=message,
        input_results=[],
        output_results=[],
    )
    check.input_results = [
        SingleExpectationCheck(expectation=exp, status="ERROR", message=message)
        for exp in condition.input_conditions
    ]
    check.output_results = [
        SingleExpectationCheck(expectation=exp, status="ERROR", message=message)
        for exp in condition.output_conditions
    ]
    return check


def _set_check_result_last_observed_from_items(check: CheckResult) -> None:
    for one in list(check.output_results or [])[::-1] + list(check.input_results or [])[::-1]:
        if one.observed_value_raw is not None:
            check.observed_value_raw = one.observed_value_raw
            check.observed_value_norm = one.observed_value_norm
            check.observed_time = one.observed_time
            check.observed_channel = one.observed_channel
            check.samples = list(one.samples or [])
            return


def review_conditions_in_log_batch(
    conditions: Iterable[OracleCondition],
    log_path: str | Path,
    dbc_path_by_channel: Dict[int, str | Path],
    drop_first_seconds: float = 0.0,
    max_samples_per_expectation: int = 100,
    stop_event: Optional[threading.Event] = None,
    progress_callback: Optional[Callable[[str], None]] = None,
) -> List[CheckResult]:
    """여러 TC의 입력/출력 기대조건을 로그 파일 1회 스캔으로 재검토한다.

    기존 review_single_expectation_in_log_robust()는 조건 1개마다 ASC를 처음부터 읽는다.
    다중 TC 수행 후 로그 증적 확인에서는 조건 수가 많아질 수 있으므로, 이 함수는
    모든 대상 Message/Signal/Expected를 먼저 모은 뒤 ASC를 한 번만 순회한다.
    """
    stop_event = stop_event or threading.Event()
    conditions = [c for c in (conditions or []) if c is not None]
    if not conditions:
        return []

    started_at = _dt.datetime.now()
    started_perf = time.perf_counter()
    log_name = Path(log_path).name if log_path else "-"

    try:
        asc_path = ensure_asc_path(log_path)
        db_by_ch = load_dbc_map(dbc_path_by_channel)
        if not db_by_ch:
            return [_make_error_log_check_result(c, "유효한 DBC를 1개도 로드하지 못했습니다.") for c in conditions]

        all_expectations: List[SignalExpectation] = []
        for c in conditions:
            all_expectations.extend(list(c.input_conditions or []))
            all_expectations.extend(list(c.output_conditions or []))

        target_frame_ids, target_msgs = build_target_message_map(db_by_ch, all_expectations)

        # (channel, frame_id) -> [(message_name, cantools_message)]
        frame_targets: Dict[Tuple[int, int], List[Tuple[str, Any]]] = {}
        for (ch, msg_name), frame_id in target_frame_ids.items():
            msg_obj = target_msgs.get((ch, msg_name))
            if msg_obj is None:
                continue
            frame_targets.setdefault((int(ch), int(frame_id)), []).append((msg_name, msg_obj))

        checks: List[CheckResult] = []
        watch_by_sig: Dict[Tuple[int, str, str], List[SingleExpectationCheck]] = {}
        all_single_checks: List[SingleExpectationCheck] = []

        for c in conditions:
            check = CheckResult(
                condition=c,
                status="검토중",
                source="multi_log_batch_review",
                started_at=started_at,
                message=f"다중 TC 로그 일괄 재검토 시작: {log_name}",
            )

            def _make_one(exp: SignalExpectation) -> SingleExpectationCheck:
                one = SingleExpectationCheck(
                    expectation=exp,
                    status="검토중",
                    message="로그 일괄 재검토 대기",
                )
                matched_channels = [
                    int(ch) for (ch, msg_name), _frame_id in target_frame_ids.items()
                    if msg_name == exp.message
                ]
                if not matched_channels:
                    one.status = "ERROR"
                    one.message = f"DBC에서 Message를 찾지 못함: {exp.message}"
                else:
                    for ch in matched_channels:
                        watch_by_sig.setdefault((ch, exp.message, exp.signal), []).append(one)
                all_single_checks.append(one)
                return one

            check.input_results = [_make_one(exp) for exp in c.input_conditions]
            check.output_results = [_make_one(exp) for exp in c.output_conditions]
            checks.append(check)

        if progress_callback:
            progress_callback(f"로그 일괄 재검토 시작: {log_name}")

        for t, ch, can_id, data in read_frames_from_asc(asc_path):
            if stop_event.is_set():
                break
            if t < float(drop_first_seconds):
                continue

            targets = frame_targets.get((int(ch), int(can_id)))
            if not targets:
                continue

            for msg_name, msg_obj in targets:
                try:
                    decoded = msg_obj.decode(data, decode_choices=False)
                except Exception:
                    continue

                for sig_name, val in decoded.items():
                    watchers = watch_by_sig.get((int(ch), msg_name, sig_name))
                    if not watchers:
                        continue
                    for one in watchers:
                        if one.status == "PASS":
                            continue
                        one.observed_value_raw = val
                        one.observed_value_norm = normalize_value(val)
                        one.observed_time = t
                        one.observed_channel = int(ch)
                        if len(one.samples) < max_samples_per_expectation:
                            one.samples.append((t, val))
                        if values_equal(one.expectation.expected_value_raw, val):
                            one.status = "PASS"
                            one.message = f"로그에서 예상 결과값 관측: {log_name}"

        if stop_event.is_set():
            for one in all_single_checks:
                if one.status == "검토중":
                    one.status = "N/A"
                    one.message = "사용자 중단으로 로그 재검토 미완료"

        for one in all_single_checks:
            if one.status == "검토중":
                one.status = "FAIL"
                one.message = f"로그 재검토에서도 예상 결과값 미관측: {log_name}"

        ended_at = _dt.datetime.now()
        elapsed_sec = time.perf_counter() - started_perf

        for check in checks:
            check.ended_at = ended_at
            check.elapsed_sec = elapsed_sec
            inputs = list(check.input_results or [])
            outputs = list(check.output_results or [])
            all_items = inputs + outputs
            any_error = any(x.status == "ERROR" for x in all_items)
            inputs_ok = bool(inputs) and all(x.status == "PASS" for x in inputs)
            outputs_ok = bool(outputs) and all(x.status == "PASS" for x in outputs)

            if stop_event.is_set():
                check.status = "N/A"
                check.message = "사용자 중단으로 로그 재검토 미완료"
            elif any_error:
                check.status = "ERROR"
                check.message = "로그 재검토 중 DBC/Message 매핑 오류 발생"
            elif inputs_ok and outputs_ok:
                check.status = "PASS"
                check.message = "로그 일괄 재검토에서 모든 입력 및 모든 출력 조건 PASS"
            elif not inputs_ok:
                check.status = "FAIL"
                check.message = "로그 일괄 재검토에서도 입력 조건 일부 또는 전체 미관측"
            else:
                check.status = "FAIL"
                check.message = "로그 일괄 재검토에서도 출력 조건 일부 미관측"
            _set_check_result_last_observed_from_items(check)

        if progress_callback:
            progress_callback(f"로그 일괄 재검토 완료: {len(checks)}개 TC")
        return checks

    except Exception as e:
        msg = f"로그 일괄 재검토 오류: {e}\n{traceback.format_exc()}"
        return [_make_error_log_check_result(c, msg) for c in conditions]


# =========================================================
# 실시간 판정: 입력 다중 + 출력 다중(AND)
# =========================================================
def check_condition_realtime(
    client: CANoeClient,
    condition: OracleCondition,
    channel,
    stop_event: Optional[threading.Event] = None,
    done_event: Optional[threading.Event] = None,
    poll_interval_sec: float = 0.05,
    progress_callback: Optional[ProgressCallback] = None,
    realtime_timeout_override_sec: Optional[float] = None,
    collector_probe: Optional[Callable[..., Optional[Dict[str, Any]]]] = None,
    collector_measurement_start_perf: Optional[float] = None,
) -> CheckResult:
    stop_event = stop_event or threading.Event()
    done_event = done_event or threading.Event()
    channels = _normalize_channels(channel)

    result = CheckResult(
        condition=condition,
        status="검토중",
        source="realtime",
        started_at=_dt.datetime.now(),
        message=f"실시간 감시 시작 | channels={channels}",
    )

    if not channels:
        result.status = "ERROR"
        result.message = "감시할 CANoe 논리 CAN이 지정되지 않았습니다."
        result.ended_at = _dt.datetime.now()
        result.elapsed_sec = 0.0
        return result

    if progress_callback:
        progress_callback("검토중", result)

    started_perf = time.perf_counter()

    try:
        timeout_sec = (
            max(0.0, float(realtime_timeout_override_sec))
            if realtime_timeout_override_sec is not None
            else max(0.0, float(condition.timeout_sec or 0.0))
        )

        if not condition.input_conditions:
            result.status = "ERROR"
            result.message = "입력 조건이 없습니다."
            return result

        input_results: List[SingleExpectationCheck] = []
        latest_input_measurement_ms = 0.0
        for exp in condition.input_conditions:
            one = wait_for_expectation_realtime(
                client=client,
                expectation=exp,
                channels=channels,
                timeout_sec=timeout_sec,
                stop_event=stop_event,
                done_event=done_event,
                poll_interval_sec=poll_interval_sec,
                collector_probe=collector_probe,
                collector_min_hit_ms=0.0,
            )
            input_results.append(one)
            if one.status == "PASS":
                if one.collector_hit_ms is not None:
                    latest_input_measurement_ms = max(latest_input_measurement_ms, float(one.collector_hit_ms))
                elif collector_measurement_start_perf is not None:
                    latest_input_measurement_ms = max(
                        latest_input_measurement_ms,
                        max(0.0, (time.perf_counter() - float(collector_measurement_start_perf)) * 1000.0),
                    )
            result.input_results = list(input_results)
            result.observed_value_raw = one.observed_value_raw
            result.observed_value_norm = one.observed_value_norm
            result.observed_time = one.observed_time
            result.observed_channel = one.observed_channel
            result.samples.extend(one.samples[:20])

            if progress_callback:
                progress_callback("검토중", result)

            if one.status != "PASS":
                result.status = "ERROR" if one.status == "ERROR" else one.status
                result.message = f"입력 조건 미충족: {exp.message}:{exp.signal} / {one.message}"
                return result

        if not condition.output_conditions:
            result.status = "ERROR"
            result.message = "출력 조건이 없습니다."
            return result

        output_results: List[SingleExpectationCheck] = []
        for exp in condition.output_conditions:
            if stop_event.is_set():
                result.status = "N/A"
                result.message = "사용자 중지"
                result.output_results = output_results
                return result

            one = wait_for_expectation_realtime(
                client=client,
                expectation=exp,
                channels=channels,
                timeout_sec=timeout_sec,
                stop_event=stop_event,
                done_event=done_event,
                poll_interval_sec=poll_interval_sec,
                collector_probe=collector_probe,
                collector_min_hit_ms=latest_input_measurement_ms,
            )
            output_results.append(one)
            result.output_results = list(output_results)
            result.observed_value_raw = one.observed_value_raw
            result.observed_value_norm = one.observed_value_norm
            result.observed_time = one.observed_time
            result.observed_channel = one.observed_channel
            result.samples.extend(one.samples[:20])

            if progress_callback:
                progress_callback("검토중", result)

            if one.status != "PASS":
                result.status = "ERROR" if one.status == "ERROR" else one.status
                result.message = f"출력 조건 미충족: {exp.message}:{exp.signal} / {one.message}"
                return result

        result.status = "PASS"
        result.message = "모든 입력 조건 및 모든 출력 조건 PASS"
        return result

    except Exception as e:
        result.status = "ERROR"
        result.source = "error"
        result.message = f"실시간 감시 오류: {e}"
        return result

    finally:
        result.ended_at = _dt.datetime.now()
        result.elapsed_sec = time.perf_counter() - started_perf
        if progress_callback:
            progress_callback(result.status, result)


# =========================================================
# 로그 재검토 조립
# =========================================================
def _log_outcome_to_single_check(exp: SignalExpectation, lr: LogReviewOutcome) -> SingleExpectationCheck:
    return SingleExpectationCheck(
        expectation=exp,
        status=lr.status,
        observed_value_raw=lr.observed_value_raw,
        observed_value_norm=lr.observed_value_norm,
        observed_time=lr.observed_time,
        observed_channel=lr.matched_channel,
        message=lr.message,
        samples=list(lr.samples),
    )


# =========================================================
# 공통 최종 판정 보정: 실시간 + 로그 결과 결합
# =========================================================
def apply_log_result_to_final(
    final: FinalCheckResult,
    log_result: Optional[CheckResult],
    strict_log_reflect: bool = False,
    context_label: str = "로그 재검토",
) -> FinalCheckResult:
    """GUI 종류와 무관하게 동일한 실시간/로그 최종 판정 정책을 적용한다."""
    final.log_result = log_result
    if log_result is None:
        return final

    rt = final.realtime_result
    rt_status = getattr(rt, "status", "")
    log_name = Path(final.used_log_path).name if final.used_log_path else "-"

    # 실시간 PASS는 가장 우선하는 판정 근거다. 로그는 보조 증적으로만 기록한다.
    if rt_status == "PASS":
        final.final_status = "PASS"
        if rt is not None:
            final.final_observed_value_raw = rt.observed_value_raw
            final.final_observed_value_norm = rt.observed_value_norm
            final.final_observed_time = rt.observed_time
            final.final_observed_channel = rt.observed_channel

        if log_result.status == "PASS":
            final.final_message = f"실시간 PASS 우선 유지 / {context_label} PASS: {log_name}"
        else:
            final.final_message = (
                f"실시간 PASS 우선 유지 / {context_label} {log_result.status}: {log_result.message}"
            )
        return final

    # 실시간 미통과 시에는 로그의 마지막 관측값을 최종 표시값으로 사용할 수 있다.
    for seq in (getattr(log_result, "output_results", None), getattr(log_result, "input_results", None)):
        for one in list(seq or [])[::-1]:
            if getattr(one, "observed_value_raw", None) is not None:
                final.final_observed_value_raw = one.observed_value_raw
                final.final_observed_value_norm = one.observed_value_norm
                final.final_observed_time = one.observed_time
                final.final_observed_channel = one.observed_channel
                break
        else:
            continue
        break

    if log_result.status == "PASS":
        final.final_status = "PASS"
        final.final_message = f"실시간 미관측 조건을 {context_label}에서 PASS: {log_name}"
        return final

    if strict_log_reflect and log_result.status in ("FAIL", "N/A", "ERROR"):
        final.final_status = log_result.status
        final.final_message = f"{context_label} 결과 {log_result.status}: {log_result.message}"
        return final

    if final.final_status == "검토중":
        final.final_status = log_result.status
        final.final_message = log_result.message
    return final


# =========================================================
# 최종 판정: 실시간 + 자동 로그 재검토(입력 다중 + 출력 다중)
# =========================================================
def check_condition_with_log_fallback(
    client: CANoeClient,
    condition: OracleCondition,
    channel,
    stop_event: Optional[threading.Event] = None,
    done_event: Optional[threading.Event] = None,
    poll_interval_sec: float = 0.05,
    progress_callback: Optional[ProgressCallback] = None,
    log_progress_callback: Optional[Callable[[str], None]] = None,
    log_path: Optional[str] = None,
    log_dir: Optional[str] = None,
    tc_start_ts: Optional[float] = None,
    dbc_path_by_channel: Optional[Dict[int, str]] = None,
    auto_log_review: bool = True,
    auto_find_latest_log: bool = True,
    auto_stop_measurement: bool = True,
    skip_log_review_on_pass: bool = True,
    log_wait_timeout: float = 5.0,
    drop_first_seconds: float = 0.0,
    realtime_timeout_override_sec: Optional[float] = None,
    collector_probe: Optional[Callable[..., Optional[Dict[str, Any]]]] = None,
    collector_measurement_start_perf: Optional[float] = None,
) -> FinalCheckResult:
    started = _dt.datetime.now()

    final = FinalCheckResult(
        condition=condition,
        started_at=started,
    )

    rt = check_condition_realtime(
        client=client,
        condition=condition,
        channel=channel,
        stop_event=stop_event,
        done_event=done_event,
        poll_interval_sec=poll_interval_sec,
        progress_callback=progress_callback,
        realtime_timeout_override_sec=realtime_timeout_override_sec,
        collector_probe=collector_probe,
        collector_measurement_start_perf=collector_measurement_start_perf,
    )
    final.realtime_result = rt

    if rt.status == "PASS" and skip_log_review_on_pass:
        if auto_stop_measurement:
            try:
                client.stop_measurement()
            except Exception:
                pass
        final.final_status = "PASS"
        final.final_message = "실시간 감시 PASS - Measurement Stop 후 로그 재검토 생략"
        final.final_observed_value_raw = rt.observed_value_raw
        final.final_observed_value_norm = rt.observed_value_norm
        final.final_observed_time = rt.observed_time
        final.final_observed_channel = rt.observed_channel
        final.ended_at = _dt.datetime.now()
        final.elapsed_sec = (final.ended_at - started).total_seconds()
        return final

    if rt.status == "PASS" and not skip_log_review_on_pass:
        if auto_stop_measurement:
            try:
                client.stop_measurement()
            except Exception:
                pass
        # 사용자가 PASS 시 로그 검토 생략 옵션을 해제한 경우에는 아래 로그 재검토 흐름을 그대로 탄다.

    if stop_event is not None and stop_event.is_set():
        final.final_status = "N/A"
        final.final_message = "사용자 중지"
        final.final_observed_value_raw = rt.observed_value_raw
        final.final_observed_value_norm = rt.observed_value_norm
        final.final_observed_time = rt.observed_time
        final.final_observed_channel = rt.observed_channel
        final.ended_at = _dt.datetime.now()
        final.elapsed_sec = (final.ended_at - started).total_seconds()
        return final

    if rt.status == "ERROR" and not auto_log_review:
        final.final_status = "ERROR"
        final.final_message = rt.message
        final.final_observed_value_raw = rt.observed_value_raw
        final.final_observed_value_norm = rt.observed_value_norm
        final.final_observed_time = rt.observed_time
        final.final_observed_channel = rt.observed_channel
        final.ended_at = _dt.datetime.now()
        final.elapsed_sec = (final.ended_at - started).total_seconds()
        return final

    if not auto_log_review:
        final.final_status = rt.status
        final.final_message = f"로그 재검토 미사용: {rt.message}"
        final.final_observed_value_raw = rt.observed_value_raw
        final.final_observed_value_norm = rt.observed_value_norm
        final.final_observed_time = rt.observed_time
        final.final_observed_channel = rt.observed_channel
        final.ended_at = _dt.datetime.now()
        final.elapsed_sec = (final.ended_at - started).total_seconds()
        return final

    if not dbc_path_by_channel:
        final.final_status = rt.status
        final.final_message = f"로그 재검토용 DBC 미지정: {rt.message}"
        final.final_observed_value_raw = rt.observed_value_raw
        final.final_observed_value_norm = rt.observed_value_norm
        final.final_observed_time = rt.observed_time
        final.final_observed_channel = rt.observed_channel
        final.ended_at = _dt.datetime.now()
        final.elapsed_sec = (final.ended_at - started).total_seconds()
        return final

    selected_log_path: Optional[Path] = None

    if log_progress_callback:
        log_progress_callback("로그 파일 탐색/안정화 단계 시작")

    if log_path:
        p = Path(log_path)
        if p.exists():
            selected_log_path = p

    if selected_log_path is None and auto_find_latest_log and log_dir:
        if auto_stop_measurement:
            try:
                client.stop_measurement()
            except Exception:
                pass

        latest = find_latest_log_file(
            log_dir=log_dir,
            after_ts=tc_start_ts,
            wait_timeout=log_wait_timeout,
            poll_interval=0.5,
            exts=(".blf", ".asc"),
        )
        if latest is not None:
            wait_until_file_stable(
                latest,
                stable_sec=1.0,
                timeout=min(max(2.0, log_wait_timeout), 10.0),
                poll_interval=0.2,
            )
            selected_log_path = latest

    if selected_log_path is None:
        final.final_status = rt.status
        final.final_message = f"로그 파일을 찾지 못함: {rt.message}"
        final.final_observed_value_raw = rt.observed_value_raw
        final.final_observed_value_norm = rt.observed_value_norm
        final.final_observed_time = rt.observed_time
        final.final_observed_channel = rt.observed_channel
        final.ended_at = _dt.datetime.now()
        final.elapsed_sec = (final.ended_at - started).total_seconds()
        return final

    final.used_log_path = str(selected_log_path)

    if log_progress_callback:
        log_progress_callback(f"로그 재검토 시작: {selected_log_path.name} (단일 TC 1회 스캔)")

    batch_results = review_conditions_in_log_batch(
        conditions=[condition],
        log_path=str(selected_log_path),
        dbc_path_by_channel=dbc_path_by_channel,
        drop_first_seconds=drop_first_seconds,
        stop_event=stop_event,
        progress_callback=log_progress_callback,
    )
    if not batch_results:
        log_check = _make_error_log_check_result(condition, "단일 TC 로그 일괄 재검토 결과가 비어 있습니다.")
    else:
        log_check = batch_results[0]

    input_log_results = list(log_check.input_results or [])
    output_log_results = list(log_check.output_results or [])

    inputs_pass = (
        input_log_results
        and len(input_log_results) == len(condition.input_conditions)
        and all(x.status == "PASS" for x in input_log_results)
    )
    outputs_pass = (
        output_log_results
        and len(output_log_results) == len(condition.output_conditions)
        and all(x.status == "PASS" for x in output_log_results)
    )

    # 최우선 판정 정책: 실시간 PASS는 로그 결과가 FAIL/N/A/ERROR여도 최종 PASS를 유지한다.
    # 로그 재검토는 보조 증적이며, 최종 Observed 값도 실시간 PASS 관측값을 보존한다.
    if rt.status == "PASS":
        failed_input = next((x for x in input_log_results if x.status != "PASS"), None)
        failed_output = next((x for x in output_log_results if x.status != "PASS"), None)

        if inputs_pass and outputs_pass:
            log_check.status = "PASS"
            log_check.message = "로그 재검토에서도 모든 입력 및 모든 출력 조건 PASS"
            final.final_message = f"실시간 PASS 우선 유지 / 로그 재검토 PASS: {selected_log_path.name}"
        elif failed_input is not None:
            log_check.status = "ERROR" if failed_input.status == "ERROR" else failed_input.status
            log_check.message = failed_input.message
            final.final_message = (
                f"실시간 PASS 우선 유지 / 입력 로그 재검토 {log_check.status}: {failed_input.message}"
            )
        elif failed_output is not None:
            log_check.status = "ERROR" if failed_output.status == "ERROR" else failed_output.status
            log_check.message = failed_output.message
            final.final_message = (
                f"실시간 PASS 우선 유지 / 출력 로그 재검토 {log_check.status}: {failed_output.message}"
            )
        else:
            log_check.status = "FAIL"
            log_check.message = "로그 재검토 실패"
            final.final_message = f"실시간 PASS 우선 유지 / 로그 재검토 FAIL: {selected_log_path.name}"

        final.log_result = log_check
        final.final_status = "PASS"
        final.final_observed_value_raw = rt.observed_value_raw
        final.final_observed_value_norm = rt.observed_value_norm
        final.final_observed_time = rt.observed_time
        final.final_observed_channel = rt.observed_channel
        final.ended_at = _dt.datetime.now()
        final.elapsed_sec = (final.ended_at - started).total_seconds()
        return final

    if inputs_pass and outputs_pass:
        log_check.status = "PASS"
        log_check.message = "로그 재검토에서 모든 입력 및 모든 출력 조건 PASS"
        final.final_status = "PASS"
        final.final_message = f"실시간 미관측이었지만 로그 재검토에서 PASS: {selected_log_path.name}"
        if output_log_results:
            last = output_log_results[-1]
        else:
            last = input_log_results[-1]
        final.final_observed_value_raw = last.observed_value_raw
        final.final_observed_value_norm = last.observed_value_norm
        final.final_observed_time = last.observed_time
        final.final_observed_channel = last.observed_channel
    else:
        failed_input = next((x for x in input_log_results if x.status != "PASS"), None)
        failed_output = next((x for x in output_log_results if x.status != "PASS"), None)

        if failed_input is not None:
            if failed_input.status == "ERROR":
                log_check.status = "ERROR"
                log_check.message = failed_input.message
                final.final_status = "ERROR"
                final.final_message = f"입력 조건 로그 재검토 오류: {failed_input.message} | log={selected_log_path}"
            else:
                log_check.status = "FAIL"
                log_check.message = failed_input.message
                final.final_status = "FAIL"
                final.final_message = f"실시간/로그 재검토 모두 입력 예상 결과값 미관측: {selected_log_path.name}"

            final.final_observed_value_raw = failed_input.observed_value_raw
            final.final_observed_value_norm = failed_input.observed_value_norm
            final.final_observed_time = failed_input.observed_time
            final.final_observed_channel = failed_input.observed_channel

        elif failed_output is not None:
            if failed_output.status == "ERROR":
                log_check.status = "ERROR"
                log_check.message = failed_output.message
                final.final_status = "ERROR"
                final.final_message = f"{failed_output.message} | log={selected_log_path}"
            else:
                log_check.status = "FAIL"
                log_check.message = failed_output.message
                final.final_status = "FAIL"
                final.final_message = f"실시간/로그 재검토 모두 예상 결과값 미관측: {selected_log_path.name}"

            final.final_observed_value_raw = failed_output.observed_value_raw
            final.final_observed_value_norm = failed_output.observed_value_norm
            final.final_observed_time = failed_output.observed_time
            final.final_observed_channel = failed_output.observed_channel
        else:
            log_check.status = "FAIL"
            log_check.message = "로그 재검토 실패"
            final.final_status = "FAIL"
            final.final_message = f"실시간/로그 재검토 모두 실패: {selected_log_path.name}"

    final.log_result = log_check
    final.ended_at = _dt.datetime.now()
    final.elapsed_sec = (final.ended_at - started).total_seconds()
    return final


# =========================================================
# 결과 저장
# =========================================================
def export_results_to_excel(results: Iterable[CheckResult], output_path: str | Path) -> Path:
    from openpyxl import Workbook

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Result"

    headers = [
        "TC 번호", "소분류", "TC 내용", "TC 예상 결과",
        "입력 Message", "입력 Signal", "입력 Expected Value",
        "출력 Message", "출력 Signal", "출력 Expected Value",
        "판정 방식", "대기 시간", "사용 여부", "비고", "조건 ID", "TC Logic", "Revision",
        "최종 상태", "출처", "Observed Value", "Observed Time", "Observed Channel", "Elapsed Sec", "비고(결과)",
    ]
    ws.append(headers)

    for r in results:
        ws.append(r.to_row())

    wb.save(output_path)
    return output_path


def export_final_results_to_excel(results: Iterable[FinalCheckResult], output_path: str | Path) -> Path:
    from openpyxl import Workbook

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "FinalResult"

    headers = [
        "TC 번호", "소분류", "TC 내용", "TC 예상 결과",
        "입력 Message", "입력 Signal", "입력 Expected Value",
        "출력 Message", "출력 Signal", "출력 Expected Value",
        "판정 방식", "대기 시간", "사용 여부", "비고", "조건 ID", "TC Logic", "Revision",
        "실시간 결과", "로그 재검토 결과", "최종 결과",
        "최종 Observed Value", "최종 Observed Time", "최종 Observed Channel", "사용 로그 파일", "최종 메시지",
    ]
    ws.append(headers)

    for r in results:
        ws.append(r.to_row())

    wb.save(output_path)
    return output_path


# =========================================================
# 최신 GUI 실행
# =========================================================
_COUNTERPART_REV_RE = re.compile(
    r"^(?P<base>{base})"
    r"(?:(?P<sep>[_\-.])rev(?P<sep2>[_\-.])?(?P<num>\d+))?$",
    re.IGNORECASE,
)


def find_latest_rev_file(base_dir: str | Path, base_name: str) -> Path:
    base_dir = Path(base_dir)
    pattern = _COUNTERPART_REV_RE.pattern.format(base=re.escape(base_name))
    rx = re.compile(pattern, re.IGNORECASE)

    candidates: List[Tuple[int, Path]] = []
    seen_paths = set()
    for glob_pat in (f"{base_name}*.py", f"{base_name}*.pyw"):
        for p in base_dir.glob(glob_pat):
            if not p.is_file() or p in seen_paths:
                continue
            seen_paths.add(p)
            m = rx.match(p.stem)
            if not m:
                continue
            num = m.group("num")
            rev = int(num) if num is not None else -1
            candidates.append((rev, p))

    if not candidates:
        raise FileNotFoundError(
            f"{base_name} 계열 파일을 찾지 못했습니다. 예: {base_name}.py / {base_name}_rev02.py"
        )

    candidates.sort(key=lambda x: x[0], reverse=True)
    return candidates[0][1]


def launch_latest_gui() -> int:
    base_dir = Path(__file__).resolve().parent
    gui_path = find_latest_rev_file(base_dir, "oracle_checker_gui")
    return subprocess.call([sys.executable, str(gui_path)], cwd=str(base_dir))


def _demo_print_excel(excel_path: str, sheet_name: str):
    rows = load_oracle_conditions_with_diagnostics(excel_path, sheet_name)
    print(f"Loaded rows: {len(rows)}")
    for r in rows[:10]:
        print(r)


if __name__ == "__main__":
    import argparse

    parser = argparse.ArgumentParser(description="Oracle Checker main utility")
    parser.add_argument("--excel", help="TestCase_Oracle.xlsx path")
    parser.add_argument("--sheet", help="Sheet name")
    parser.add_argument(
        "--no-gui",
        action="store_true",
        help="인자 없이 실행했을 때 최신 GUI를 자동 실행하지 않고 안내만 출력",
    )
    args = parser.parse_args()

    if args.excel and args.sheet:
        _demo_print_excel(args.excel, args.sheet)
    elif args.no_gui:
        print("oracle_checker_main*.py loaded. GUI에서 사용하는 모듈입니다.")
    else:
        raise SystemExit(launch_latest_gui())

